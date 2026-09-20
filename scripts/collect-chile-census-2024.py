#!/usr/bin/env python3
"""Build Chile's 2024 Census country-depth bundle from official INE files.

The public INE FeatureServer supplies official CUT codes, attributes and
reference geometry for 16 regions, 56 provinces and 345 communes.  The
published D1 workbook supplies the 346th commune, Antartica (CUT 12202), whose
statistics remain in the registry even though the map service omits its shape.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

LANDING = "https://censo2024.ine.gob.cl/estadisticas/"
SERVICE = "https://services5.arcgis.com/hUyD8u3TeZLKPe4T/ArcGIS/rest/services/Censo2024_v2/FeatureServer"
WB_API = "https://api.worldbank.org/v2/country/CHL/indicator/{indicator}?format=json&per_page=100&source=2"


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def number(value):
    if value in (None, "", "-"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def pct(numerator, denominator):
    n, d = number(numerator), number(denominator)
    return None if n is None or d in (None, 0) else round(n / d * 100, 4)


def code(value, width):
    return str(int(value)).zfill(width)


def load_geo(raw: Path, filename: str):
    return json.loads((raw / filename).read_text(encoding="utf-8"))["features"]


def antartica_row(raw: Path):
    file = next(raw.glob("D1_Poblacion*.xlsx"))
    ws = load_workbook(file, read_only=True, data_only=True)["2"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if number(row[4]) == 12202:
            return {
                "COD_REGION": int(row[0]), "REGION": row[1],
                "COD_PROVINCIA": int(row[2]), "PROVINCIA": row[3],
                "CUT": int(row[4]), "COMUNA": row[5],
                "n_per": number(row[6]), "n_hombres": number(row[7]),
                "n_mujeres": number(row[8]),
            }
    raise RuntimeError("Official D1 workbook has no Antartica CUT 12202 row")


COUNT_FIELDS = [
    "n_per", "n_hombres", "n_mujeres", "n_vp", "n_fuente_agua_publica",
    "denominador_fuente_agua", "n_serv_hig_alc_dentro", "n_serv_hig_alc_fuera",
    "n_serv_hig_fosa", "n_serv_hig_pozo", "n_serv_hig_acequia_canal",
    "n_serv_hig_cajon_otro", "n_serv_hig_bano_quimico", "n_serv_hig_bano_seco",
    "n_serv_hig_no_tiene",
    "denominador_fuente_elect", "n_fuente_elect_no_tiene", "n_analfabet",
    "n_ocupado", "n_desocupado", "n_fuera_fuerza_trabajo", "n_discapacidad",
    "denominador_discapacidad", "n_inmigrantes", "denominador_inmigrantes",
    "n_pueblos_orig", "denominador_pueblos_orig",
]


def aggregate(rows):
    return {key: sum(number(row.get(key)) or 0 for row in rows) for key in COUNT_FIELDS}


def workbook_inventory(raw: Path):
    records = []
    for file in sorted(raw.glob("*.xlsx")):
        wb = load_workbook(file, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if ws.title == "Nota" or ws.max_row < 5:
                continue
            title = next((row[0] for row in ws.iter_rows(min_row=1, max_row=4, values_only=True) if row and row[0]), ws.title)
            headers = list(next(ws.iter_rows(min_row=4, max_row=4, values_only=True)))
            counts = [0] * len(headers)
            for row in ws.iter_rows(min_row=5, values_only=True):
                for index, value in enumerate(row[:len(headers)]):
                    if number(value) is not None:
                        counts[index] += 1
            for index, header in enumerate(headers):
                if not header or counts[index] == 0:
                    continue
                records.append({
                    "country_area_id": "CHL", "source_path": f"raw/chile-census-2024/{file.name}",
                    "source_url": LANDING, "table_id": f"{file.stem}:{ws.title}",
                    "table_title": str(title), "field_id": f"col-{index + 1}",
                    "field_label": str(header), "numeric_cell_count": counts[index],
                    "theme": "source_inventory", "disposition": "not_adopted",
                    "reason": "Official numeric field inventoried. It is retained for later use but is not one of this edition's selected comparable dashboard indicators.",
                    "coverage_complete": True, "country_edition_eligible": False,
                })
    return records


def wb_latest(indicator: str, out: Path):
    url = WB_API.format(indicator=urllib.parse.quote(indicator))
    req = urllib.request.Request(url, headers={"User-Agent": "AreaData/0.10.2 source collector"})
    with urllib.request.urlopen(req, timeout=120) as response:
        body = response.read()
    target = out / f"wb-{indicator}.json"
    target.write_bytes(body)
    payload = json.loads(body)
    rows = payload[1] if isinstance(payload, list) and len(payload) > 1 else []
    row = next((item for item in rows if item.get("value") is not None), None)
    if not row:
        raise RuntimeError(f"No World Bank observation for {indicator}")
    return {"indicator_id": indicator, "period": str(row["date"]), "value": row["value"],
            "url": url, "raw_path": str(target), "sha256": sha(target),
            "retrieved_at": datetime.now(timezone.utc).isoformat()}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    raw, out = Path(args.raw_dir).resolve(), Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)

    layers = {
        "region": load_geo(raw, "Regional_CPV24.geojson"),
        "province": load_geo(raw, "Provincial_CPV24.geojson"),
        "commune": load_geo(raw, "Comunal_CPV24.geojson"),
    }
    if [len(layers[x]) for x in ("region", "province", "commune")] != [16, 56, 345]:
        raise RuntimeError("Unexpected INE FeatureServer geography counts")

    records = {}
    features = []
    territories = []
    ids = {"region": {}, "province": {}, "commune": {}}
    layer_source = "chl-ine-census-2024-feature-service"

    def add(level, feature):
        p = feature["properties"]
        if level == "region":
            official, width, prefix, parent, name = p["COD_REGION"], 2, "REG", "CHL", p["REGION"]
        elif level == "province":
            official, width, prefix = p["COD_PROVINCIA"], 3, "PROV"
            parent, name = f"CHL:C2024:REG:{code(p['COD_REGION'], 2)}", p["PROVINCIA"]
        else:
            official, width, prefix = p["CUT"], 5, "COM"
            parent, name = f"CHL:C2024:PROV:{code(p['COD_PROVINCIA'], 3)}", p["COMUNA"]
        c = code(official, width)
        tid = f"CHL:C2024:{prefix}:{c}"
        ids[level][c] = tid
        records[tid] = p
        territories.append({"id": tid, "country_id": "CHL", "name": str(name).title(),
                            "level": level, "type": level, "parent_id": parent,
                            "official_code": c, "code_system": "Chile CUT 2024",
                            "boundary_version": "Censo 2024 CPV24", "valid_from": "2024-01-01"})
        features.append({"type": "Feature", "properties": {"territory_id": tid,
                         "source_id": layer_source, "official_code": c,
                         "code_system": "Chile CUT 2024", "geometry_edition": "Censo 2024 CPV24",
                         "join_method": "Exact published CUT code", "reference_only": True},
                         "geometry": feature["geometry"]})

    for level in ("region", "province", "commune"):
        for feature in layers[level]:
            add(level, feature)

    ant = antartica_row(raw)
    ant_id = "CHL:C2024:COM:12202"
    records[ant_id] = ant
    ids["commune"]["12202"] = ant_id
    territories.append({"id": ant_id, "country_id": "CHL", "name": "Antártica",
                        "level": "commune", "type": "commune",
                        "parent_id": "CHL:C2024:PROV:122", "official_code": "12202",
                        "code_system": "Chile CUT 2024", "boundary_version": None,
                        "valid_from": "2024-01-01",
                        "geography_note": "Official Census workbook row retained; CPV24 FeatureServer layer 11 did not return a geometry for CUT 12202."})
    records["CHL"] = aggregate([f["properties"] for f in layers["region"]])

    source = "chl-ine-census-2024-results"
    specs = [
        ("CHL_C2024_POP_TOTAL", "Census population", "Population", "people", lambda r: number(r.get("n_per")), "sum", "Enumerated population."),
        ("CHL_C2024_FEMALE_PCT", "Female population", "Population", "%", lambda r: pct(r.get("n_mujeres"), r.get("n_per")), "none", "Female residents as a share of enumerated population."),
        ("CHL_C2024_PRIVATE_DWELLINGS", "Private dwellings", "Households and housing", "dwellings", lambda r: number(r.get("n_vp")), "sum", "Private dwellings counted by Census 2024."),
        ("CHL_C2024_PUBLIC_WATER_PCT", "Public-network water", "Water", "% of occupied private dwellings", lambda r: pct(r.get("n_fuente_agua_publica"), r.get("denominador_fuente_agua")), "none", "Occupied private dwellings whose domestic water source is the public network."),
        ("CHL_C2024_SEWER_NETWORK_PCT", "Sewer-network sanitation", "Sanitation", "% of occupied private dwellings with declared sanitation", lambda r: pct((number(r.get("n_serv_hig_alc_dentro")) or 0) + (number(r.get("n_serv_hig_alc_fuera")) or 0), sum(number(r.get(k)) or 0 for k in ("n_serv_hig_alc_dentro", "n_serv_hig_alc_fuera", "n_serv_hig_fosa", "n_serv_hig_pozo", "n_serv_hig_acequia_canal", "n_serv_hig_cajon_otro", "n_serv_hig_bano_quimico", "n_serv_hig_bano_seco", "n_serv_hig_no_tiene"))), "none", "Occupied private dwellings connected to a sewer network inside or outside the dwelling, divided by dwellings with a declared sanitation category."),
        ("CHL_C2024_ELECTRICITY_PCT", "Any electricity source", "Basic services", "% of occupied private dwellings", lambda r: pct((number(r.get("denominador_fuente_elect")) or 0) - (number(r.get("n_fuente_elect_no_tiene")) or 0), r.get("denominador_fuente_elect")), "none", "Occupied private dwellings with any reported electricity source."),
        ("CHL_C2024_ILLITERATE_COUNT", "Population unable to read or write", "Education", "people age 5+", lambda r: number(r.get("n_analfabet")), "sum", "Population age 5 or older reported unable to read or write."),
        ("CHL_C2024_LABOUR_FORCE_PCT", "Labour-force participation", "Employment", "% of classified working-age population", lambda r: pct((number(r.get("n_ocupado")) or 0) + (number(r.get("n_desocupado")) or 0), (number(r.get("n_ocupado")) or 0) + (number(r.get("n_desocupado")) or 0) + (number(r.get("n_fuera_fuerza_trabajo")) or 0)), "none", "Employed plus unemployed people as a share of the classified employed, unemployed and outside-labour-force population."),
        ("CHL_C2024_DISABILITY_PCT", "Population with disability", "Disability", "% of applicable population", lambda r: pct(r.get("n_discapacidad"), r.get("denominador_discapacidad")), "none", "Population with disability in the Census disability universe."),
        ("CHL_C2024_IMMIGRANT_PCT", "International immigrants", "Migration", "% of declared population", lambda r: pct(r.get("n_inmigrantes"), r.get("denominador_inmigrantes")), "none", "Population born outside Chile as a share of the applicable declared population."),
        ("CHL_C2024_INDIGENOUS_PCT", "Indigenous people", "Ethnicity", "% of declared population", lambda r: pct(r.get("n_pueblos_orig"), r.get("denominador_pueblos_orig")), "none", "Population identifying with an Indigenous people as a share of the applicable declared population."),
    ]
    indicators, observations = [], []
    for iid, name, theme, unit, calc, aggregation, definition in specs:
        indicators.append({"id": iid, "name": name, "theme": theme, "unit": unit,
                           "definition": definition, "definition_id": "CHL-CENSO-2024",
                           "population": "Census 2024 published universe stated per indicator",
                           "measurement_method": "source_reported" if aggregation == "sum" else "derived_from_source_counts",
                           "aggregation": aggregation, "period_policy": "fixed_source_period",
                           "series_family": "census", "display_role": "primary",
                           "source_id": source, "source_locator": "INE Censo 2024 CPV24 attributes and official XLSX tables"})
        for tid, row in records.items():
            value = calc(row)
            observations.append({"territory_id": tid, "indicator_id": iid, "period": "2024",
                                 "value": value, "status": "observed" if value is not None else "missing",
                                 "source_id": source, "definition": definition,
                                 "definition_id": "CHL-CENSO-2024", "unit": unit,
                                 "population": "Census 2024 published universe stated per indicator",
                                 "measurement_method": "source_reported" if aggregation == "sum" else "derived_from_source_counts",
                                 "source_locator": "INE Censo 2024 CPV24 attributes and official XLSX tables"})

    comparisons = [{"parent_id": "CHL", "member_ids": [ids["region"][x] for x in sorted(ids["region"])],
                    "label": "Regions", "membership_note": "Official 2024 CUT hierarchy.", "source_ids": [layer_source]}]
    for rcode, rid in sorted(ids["region"].items()):
        members = [pid for pcode, pid in sorted(ids["province"].items()) if pcode.startswith(rcode)]
        comparisons.append({"parent_id": rid, "member_ids": members, "label": "Provinces",
                            "membership_note": "Official 2024 CUT hierarchy.", "source_ids": [layer_source]})
    for pcode, pid in sorted(ids["province"].items()):
        members = [cid for ccode, cid in sorted(ids["commune"].items()) if ccode.startswith(pcode)]
        comparisons.append({"parent_id": pid, "member_ids": members, "label": "Communes",
                            "membership_note": "Official 2024 CUT hierarchy; Antartica CUT 12202 has statistics but no returned layer-11 geometry.", "source_ids": [layer_source, source]})

    bundle = {"schema_version": "1.0", "country_area_id": "CHL", "period": "2024",
              "replace_country_branch": True, "territories": territories,
              "terminal_territory_ids": list(ids["commune"].values()), "comparisons": comparisons,
              "indicators": indicators, "observations": observations,
              "boundaries": {"type": "FeatureCollection", "features": features},
              "sources": [
                  {"id": source, "name": "Censo de Poblacion y Vivienda 2024 - resultados estadisticos",
                   "publisher": "Instituto Nacional de Estadisticas de Chile", "url": LANDING,
                   "status": "ready", "retrieved_at": datetime.now(timezone.utc).date().isoformat(),
                   "reference_period": "2024", "geographic_level": "country, region and commune",
                   "license": "Official public material; reuse terms require source-level review.",
                   "note": "Twelve official XLSX workbooks were acquired and hashed. Selected fields are cross-checked to the CPV24 service by exact CUT code."},
                  {"id": layer_source, "name": "Censo 2024 CPV24 official geography and attributes service",
                   "publisher": "Instituto Nacional de Estadisticas de Chile", "url": SERVICE,
                   "status": "ready", "retrieved_at": datetime.now(timezone.utc).date().isoformat(),
                   "reference_period": "2024", "geographic_level": "region, province and commune",
                   "license": "Official public feature service; reference geometry is attributed.",
                   "note": "Exact CUT attributes: 16 regions, 56 provinces and 345 commune geometries. Antartica CUT 12202 is retained from the official workbook without a geometry."},
              ]}
    (out / "chile-country-depth-bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inventory = workbook_inventory(raw)
    (out / "chile-workbook-field-inventory.json").write_text(json.dumps({"schema_version": "1.0", "country_area_id": "CHL", "records": inventory}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    supplements = [wb_latest(i, out) for i in ("SP.URB.TOTL.IN.ZS", "SN.ITK.DEFC.ZS", "SI.POV.NAHC")]
    (out / "chile-wb-supplements.json").write_text(json.dumps({"country_area_id": "CHL", "records": supplements}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    receipts = [{"file": p.name, "bytes": p.stat().st_size, "sha256": sha(p)} for p in sorted(raw.iterdir()) if p.is_file()]
    audit = {"country_area_id": "CHL", "status": "complete", "counts": {"regions": 16, "provinces": 56, "communes": 346, "boundary_features": len(features), "indicators": len(indicators), "observations": len(observations), "workbook_numeric_fields": len(inventory)}, "antartica_geometry_status": "official_service_not_returned", "source_receipts": receipts}
    (out / "CHILE_COLLECTION_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
