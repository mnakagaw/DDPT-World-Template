"""Extract the first audited census population layer without editing source XLSX files."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from census_pdf_extract import (
    extract_honduras_municipal_reports,
    extract_honduras_tome1,
    extract_nicaragua_table5,
)


INDICATOR_ID = "CENSUS_POP_TOTAL"
BELIZE_DISTRICTS = ["Corozal", "Orange Walk", "Belize", "Cayo", "Stann Creek", "Toledo"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_entry(receipt: dict, source_id: str) -> dict:
    entry = next(item for item in receipt["entries"] if item["source_id"] == source_id)
    if entry["status"] != "acquired":
        raise ValueError(f"Required source is not acquired: {source_id}")
    return entry


def source_path(root: Path, entry: dict) -> Path:
    path = root.joinpath(*entry["filename"].split("/"))
    if sha256(path) != entry["sha256"]:
        raise ValueError(f"Receipt hash mismatch: {entry['source_id']}")
    return path


def finite_number(value, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite number at {label}")
    return value


def difference(left, right):
    return float(left) - float(right)


def source_record(entry: dict, note: str, license_value=None, license_url=None) -> dict:
    if license_value is None:
        license_value = "Reuse terms not stated in the source publication; raw files are excluded from public redistribution pending terms review."
    return {
        "id": entry["source_id"],
        "name": entry["title"],
        "publisher": entry["publisher"],
        "url": entry["final_url"],
        "catalog_url": entry["catalog_url"],
        "status": "ready",
        "retrieved_at": entry["retrieved_at"] if "retrieved_at" in entry else None,
        "reference_period": entry["census_year"],
        "sha256": entry["sha256"],
        "bytes": entry["bytes"],
        "license": license_value,
        "license_url": license_url,
        "license_detail": "AreaData publishes source-linked factual observations and provenance only. This statement does not assert an open-data license." if license_url is None else None,
        "terms_review_status": "terms_not_stated_in_source" if license_url is None else "license_recorded",
        "raw_redistribution_status": "withheld_pending_terms_review" if license_url is None else "follow_recorded_license",
        "normalized_observation_publication": "source_attributed_factual_observation_only",
        "geographic_level": "subnational",
        "note": note,
    }


def observation(territory_id, value, source_id, period, footnote, precision="source_decimal"):
    return {
        "territory_id": territory_id,
        "indicator_id": INDICATOR_ID,
        "period": str(period),
        "value": finite_number(value, f"{territory_id} population"),
        "status": "observed",
        "source_id": source_id,
        "footnote": footnote,
        "source_precision": precision,
    }


def territory(
    territory_id,
    country_id,
    name,
    level,
    territory_type,
    parent_id,
    source_id,
    official_code=None,
    code_system=None,
    geography_note=None,
    administrative_subtype=None,
):
    row = {
        "id": territory_id,
        "country_id": country_id,
        "name": name,
        "level": level,
        "type": territory_type,
        "parent_id": parent_id,
        "official_code": official_code,
        "code_system": code_system,
        "boundary_version": None,
        "source_id": source_id,
        "geography_note": geography_note or "Name and hierarchy copied from the adopted official census table; polygon boundaries are not yet joined.",
    }
    if administrative_subtype:
        row["administrative_subtype"] = administrative_subtype
    return row


def comparison(parent_id, member_ids, label, source_ids, note):
    return {
        "parent_id": parent_id,
        "member_ids": member_ids,
        "label": label,
        "membership_note": note,
        "source_ids": source_ids,
    }


def coded_label(value: str) -> tuple[str, str]:
    match = re.match(r"^\s*(\d+)\s*-\s*(.+?)\s*$", str(value))
    if not match:
        raise ValueError(f"Expected coded geography label: {value}")
    return match.group(1).zfill(2), match.group(2).strip()


def clean_panama_label(value: str) -> str:
    return re.sub(r"\s+\(\d+\)\s*$", "", str(value).strip())


def extract_belize(root: Path, receipt: dict):
    general_entry = source_entry(receipt, "BLZ_C2022_GENERAL_CHARACTERISTICS")
    locality_entry = source_entry(receipt, "BLZ_C2022_POP_SETTLEMENT_SEX")
    general_path, locality_path = source_path(root, general_entry), source_path(root, locality_entry)

    general_book = load_workbook(general_path, read_only=True, data_only=True)
    locality_book = load_workbook(locality_path, read_only=True, data_only=True)
    try:
        admin = general_book["Admin_Area"]
        national = finite_number(admin["D4"].value, "BLZ Admin_Area!D4")
        published_admin = {
            str(admin.cell(row, 2).value).strip(): finite_number(admin.cell(row, 4).value, f"BLZ Admin_Area!D{row}")
            for row in range(4, admin.max_row + 1)
            if admin.cell(row, 2).value is not None and isinstance(admin.cell(row, 4).value, (int, float))
        }
        district_values = {name: published_admin[name] for name in BELIZE_DISTRICTS}
        territories, observations, comparisons, terminal_ids = [], [], [], []
        district_ids = []
        general_note = "SIB 2022 calibrated-weight census population estimate. The workbook retains fractional weighted values; SIB's report presents rounded persons."
        observations.append(observation("BLZ", national, general_entry["source_id"], 2022, general_note))
        for index, name in enumerate(BELIZE_DISTRICTS, start=1):
            territory_id = f"BLZ:C2022:DIST:{index}"
            district_ids.append(territory_id)
            territories.append({
                "id": territory_id,
                "country_id": "BLZ",
                "name": name,
                "level": "district",
                "type": "district",
                "parent_id": "BLZ",
                "official_code": None,
                "code_system": "No district code column in adopted SIB 2022 workbook",
                "boundary_version": None,
                "source_id": general_entry["source_id"],
                "geography_note": "District name as published by SIB; official code and polygon edition remain unverified.",
            })
            observations.append(observation(territory_id, district_values[name], general_entry["source_id"], 2022, general_note))
        comparisons.append({
            "parent_id": "BLZ",
            "member_ids": district_ids,
            "label": "Six Belize districts",
            "membership_note": "All six district totals published in SIB's 2022 general characteristics workbook. Codes and polygon boundaries remain unverified.",
            "source_ids": [general_entry["source_id"]],
        })

        sheet = locality_book["2022"]
        current_district = None
        locality_row_counts = {name: 0 for name in BELIZE_DISTRICTS}
        rejected = []
        published_totals = {}
        for row in range(5, sheet.max_row + 1):
            district_cell, name_cell = sheet.cell(row, 1).value, sheet.cell(row, 2).value
            if district_cell is not None:
                current_district = " ".join(str(district_cell).split())
            if current_district not in locality_row_counts or name_cell is None:
                continue
            name = str(name_cell).strip()
            value = sheet.cell(row, 5).value
            if name == "Total":
                published_totals[current_district] = finite_number(value, f"BLZ locality total row {row}")
                continue
            if name.lower().startswith("other "):
                rejected.append({"row": row, "district": current_district, "name": name, "reason": "Residual group is not one discrete locality."})
                continue
            finite_number(value, f"BLZ locality row {row}")
            locality_row_counts[current_district] += 1

        district_reconciliation = []
        for district in BELIZE_DISTRICTS:
            listed = published_totals[district]
            official = district_values[district]
            district_reconciliation.append({"district": district, "general_characteristics_total": official, "locality_table_total": listed, "difference": difference(listed, official)})
        audit = {
            "national_population": national,
            "district_count": len(district_ids),
            "district_sum": sum(district_values.values()),
            "district_sum_difference": difference(sum(district_values.values()), national),
            "published_discrete_localities_audited": sum(locality_row_counts.values()),
            "published_discrete_localities_by_district": locality_row_counts,
            "locality_adoption_status": "audited_not_loaded; no official code or statutory-type column, no verified polygons, and two district totals conflict with the adopted district table",
            "residual_rows_not_adopted_as_territories": rejected,
            "district_reconciliation": district_reconciliation,
            "precision_note": general_note,
        }
    finally:
        general_book.close()
        locality_book.close()
    return {
        "country_id": "BLZ",
        "period": "2022",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [
            source_record(general_entry, "National and district population source. Values are calibrated-weight estimates and retain source decimals."),
            source_record(locality_entry, "City/town/village/community population and sex table. Audited but not loaded as territories: residual rows, missing geographic codes and two district-total conflicts require reconciliation."),
        ],
        "audit": audit,
    }


def extract_guatemala(root: Path, receipt: dict):
    population_entry = source_entry(receipt, "GTM_C2018_POP_SEX_AGE_AREA")
    centroid_entry = source_entry(receipt, "GTM_C2018_POPULATED_PLACE_CENTROIDS")
    population_path, centroid_path = source_path(root, population_entry), source_path(root, centroid_entry)
    population_book = load_workbook(population_path, read_only=True, data_only=True)
    centroid_book = load_workbook(centroid_path, read_only=True, data_only=True)
    try:
        national_sheet = population_book["A1_1"]
        municipality_sheet = population_book["A1_2"]
        place_sheet = population_book["A1_3"]
        centroid_sheet = centroid_book["A_3"]
        national = finite_number(national_sheet["C8"].value, "GTM A1_1!C8")
        territories, observations, comparisons, terminal_ids = [], [], [], []
        department_ids, departments = [], {}
        note = "INE XII Censo 2018 enumerated population (poblacion total censada) in the published A1 table."
        observations.append(observation("GTM", national, population_entry["source_id"], 2018, note, "integer"))
        for row, values in enumerate(national_sheet.iter_rows(min_row=10, max_row=31, min_col=1, max_col=3, values_only=True), start=10):
            code = str(values[0])
            name = str(values[1]).strip()
            value = finite_number(values[2], f"GTM A1_1!C{row}")
            territory_id = f"GTM:C2018:DEP:{code}"
            department_ids.append(territory_id)
            departments[code] = {"id": territory_id, "name": name, "value": value, "municipality_ids": [], "municipality_sum": 0}
            territories.append({
                "id": territory_id,
                "country_id": "GTM",
                "name": name,
                "level": "department",
                "type": "department",
                "parent_id": "GTM",
                "official_code": code,
                "code_system": "INE Guatemala Censo 2018 published department code",
                "boundary_version": None,
                "source_id": population_entry["source_id"],
                "geography_note": "Code and name copied from the 2018 census table; polygon edition remains unverified.",
            })
            observations.append(observation(territory_id, value, population_entry["source_id"], 2018, note, "integer"))
        comparisons.append({
            "parent_id": "GTM",
            "member_ids": department_ids,
            "label": "22 Guatemala departments",
            "membership_note": "Complete department register and totals as published in INE Censo 2018 table A1.1; polygon boundaries are not yet joined.",
            "source_ids": [population_entry["source_id"]],
        })
        for row, values in enumerate(municipality_sheet.iter_rows(min_row=7, max_row=346, min_col=1, max_col=5, values_only=True), start=7):
            department_code = str(values[0])
            municipality_code = str(values[2])
            name = str(values[3]).strip()
            value = finite_number(values[4], f"GTM A1_2!E{row}")
            if department_code not in departments:
                raise ValueError(f"Unknown department code for municipality row {row}: {department_code}")
            territory_id = f"GTM:C2018:MUN:{municipality_code}"
            terminal_ids.append(territory_id)
            departments[department_code]["municipality_ids"].append(territory_id)
            departments[department_code]["municipality_sum"] += value
            territories.append({
                "id": territory_id,
                "country_id": "GTM",
                "name": name,
                "level": "municipality",
                "type": "municipality",
                "parent_id": departments[department_code]["id"],
                "official_code": municipality_code,
                "code_system": "INE Guatemala Censo 2018 published municipality code",
                "boundary_version": None,
                "source_id": population_entry["source_id"],
                "geography_note": "Municipality code and name copied from table A1.2; polygon edition remains unverified.",
            })
            observations.append(observation(territory_id, value, population_entry["source_id"], 2018, note, "integer"))
        for department in departments.values():
            comparisons.append({
                "parent_id": department["id"],
                "member_ids": department["municipality_ids"],
                "label": f"Municipalities within {department['name']}",
                "membership_note": "Complete municipality membership as published in INE Censo 2018 table A1.2; polygon boundaries are not yet joined.",
                "source_ids": [population_entry["source_id"]],
            })

        population_place_keys = []
        for values in place_sheet.iter_rows(min_row=7, max_row=20042, min_col=5, max_col=6, values_only=True):
            population_place_keys.append((str(values[0]), str(values[1]).strip()))
        centroid_place_keys, missing_coordinates = [], 0
        for values in centroid_sheet.iter_rows(min_row=7, max_row=20042, min_col=5, max_col=8, values_only=True):
            centroid_place_keys.append((str(values[0]), str(values[1]).strip()))
            if values[2] is None or values[3] is None:
                missing_coordinates += 1
        if population_place_keys != centroid_place_keys:
            raise ValueError("Guatemala population-place and centroid workbook identities do not match row-for-row")
        department_reconciliation = [
            {
                "department_code": code,
                "department": item["name"],
                "department_total": item["value"],
                "municipality_count": len(item["municipality_ids"]),
                "municipality_sum": item["municipality_sum"],
                "difference": difference(item["municipality_sum"], item["value"]),
            }
            for code, item in departments.items()
        ]
        audit = {
            "national_population": national,
            "department_count": len(department_ids),
            "department_sum": sum(item["value"] for item in departments.values()),
            "department_sum_difference": difference(sum(item["value"] for item in departments.values()), national),
            "municipality_count": len(terminal_ids),
            "municipality_sum": sum(item["municipality_sum"] for item in departments.values()),
            "municipality_sum_difference": difference(sum(item["municipality_sum"] for item in departments.values()), national),
            "department_reconciliation": department_reconciliation,
            "published_place_rows": len(population_place_keys),
            "place_rows_matching_centroid_register": len(population_place_keys),
            "place_rows_without_complete_coordinates": missing_coordinates,
            "place_adoption_status": "audited_not_loaded; municipalities are the first terminal analysis level",
        }
    finally:
        population_book.close()
        centroid_book.close()
    return {
        "country_id": "GTM",
        "period": "2018",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [
            source_record(population_entry, "Adopted national, department and municipality population source; table also contains sex, five-year age and urban/rural columns.", "Creative Commons Attribution (official catalog record)", population_entry["catalog_url"]),
            source_record(centroid_entry, "Acquired and row-identity checked against 20,036 published populated-place rows. Not yet used as municipality polygons or as a terminal dashboard layer.", "Creative Commons Attribution (official catalog record)", centroid_entry["catalog_url"]),
        ],
        "audit": audit,
    }


def extract_el_salvador(root: Path, receipt: dict):
    entry = source_entry(receipt, "SLV_C2024_POP_GEO_AGE_SEX")
    path = source_path(root, entry)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["TAB_POB_1"]
        total_rows = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=7, min_col=2, max_col=6, values_only=True), start=7
        ):
            department, municipality, district, age, value = values
            if age == "TOTAL":
                total_rows.append((row_number, department, municipality, district, finite_number(value, f"SLV TAB_POB_1!F{row_number}")))
    finally:
        book.close()

    national_row = next(row for row in total_rows if row[1] == "00 - República de El Salvador")
    national = national_row[4]
    unallocated_rows = [row for row in total_rows if str(row[1]).startswith("15 - ")]
    unallocated = next(row[4] for row in unallocated_rows if row[2] == "TOTAL")
    administrative_rows = [row for row in total_rows if not str(row[1]).startswith(("00 - ", "15 - "))]
    territories, observations, comparisons, terminal_ids = [], [], [], []
    departments = {}
    municipalities = {}
    note = "VII Censo de Población 2024 under the post-2023 structure of 14 departments, 44 municipalities and 262 districts. The official national total includes 107,055 persons with unspecified residence; that residual is not represented as an administrative territory."
    observations.append(observation("SLV", national, entry["source_id"], 2024, note, "integer"))

    for row_number, department_label, municipality_label, district_label, value in administrative_rows:
        department_code, department_name = coded_label(department_label)
        department_id = f"SLV:C2024:DEP:{department_code}"
        if municipality_label == "TOTAL":
            departments[department_code] = {"id": department_id, "name": department_name, "value": value, "municipality_ids": [], "district_ids": []}
            territories.append(territory(department_id, "SLV", department_name, "department", "department", "SLV", entry["source_id"], department_code, "Censo 2024 published department code"))
            observations.append(observation(department_id, value, entry["source_id"], 2024, note, "integer"))
            continue
        municipality_code, municipality_name = coded_label(municipality_label)
        municipality_key = f"{department_code}-{municipality_code}"
        municipality_id = f"SLV:C2024:MUN:{department_code}{municipality_code}"
        if district_label == "TOTAL":
            if department_code not in departments:
                raise ValueError(f"SLV municipality before department at row {row_number}")
            municipalities[municipality_key] = {"id": municipality_id, "name": municipality_name, "value": value, "district_ids": []}
            departments[department_code]["municipality_ids"].append(municipality_id)
            territories.append(territory(municipality_id, "SLV", municipality_name, "municipality", "municipality", department_id, entry["source_id"], municipality_key, "Censo 2024 composite department-municipality code"))
            observations.append(observation(municipality_id, value, entry["source_id"], 2024, note, "integer"))
            continue
        district_code, district_name = coded_label(district_label)
        district_id = f"SLV:C2024:DIST:{department_code}{district_code}"
        if municipality_key not in municipalities:
            raise ValueError(f"SLV district before municipality at row {row_number}")
        municipalities[municipality_key]["district_ids"].append(district_id)
        departments[department_code]["district_ids"].append(district_id)
        terminal_ids.append(district_id)
        territories.append(territory(district_id, "SLV", district_name, "district", "district", municipality_id, entry["source_id"], f"{department_code}-{district_code}", "Censo 2024 composite department-district code"))
        observations.append(observation(district_id, value, entry["source_id"], 2024, note, "integer"))

    if (len(departments), len(municipalities), len(terminal_ids)) != (14, 44, 262):
        raise ValueError(f"Unexpected El Salvador geography counts: {len(departments)}, {len(municipalities)}, {len(terminal_ids)}")
    administrative_total = national - unallocated
    department_sum = sum(item["value"] for item in departments.values())
    municipality_sum = sum(item["value"] for item in municipalities.values())
    district_sum = sum(row[4] for row in administrative_rows if row[3] != "TOTAL")
    if (department_sum, municipality_sum, district_sum) != (administrative_total,) * 3:
        raise ValueError("El Salvador administrative levels do not reconcile after excluding unspecified residence")
    comparisons.append(comparison("SLV", [item["id"] for item in departments.values()], "14 El Salvador departments", [entry["source_id"]], "All 14 named departments; 107,055 persons with unspecified residence remain in the exact national observation and outside the administrative comparison."))
    for department in departments.values():
        comparisons.append(comparison(department["id"], department["municipality_ids"], f"Municipalities within {department['name']}", [entry["source_id"]], "Complete post-2023 municipality membership as published in TAB_POB_1."))
    for municipality in municipalities.values():
        comparisons.append(comparison(municipality["id"], municipality["district_ids"], f"Districts within {municipality['name']}", [entry["source_id"]], "Complete district membership as published in TAB_POB_1."))
    return {
        "country_id": "SLV",
        "period": "2024",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [source_record(entry, "Adopted national, department, post-2023 municipality and district population source. The unspecified-residence residual is retained in the national total and excluded from administrative territory rows.")],
        "audit": {
            "national_population": national,
            "department_count": len(departments),
            "department_sum": department_sum,
            "department_sum_difference": difference(department_sum, national),
            "municipality_count": len(municipalities),
            "municipality_sum": municipality_sum,
            "municipality_sum_difference": difference(municipality_sum, national),
            "district_count": len(terminal_ids),
            "district_sum": district_sum,
            "district_sum_difference": difference(district_sum, national),
            "unallocated_residence_population": unallocated,
            "unallocated_rows_not_adopted_as_territories": len(unallocated_rows),
            "geography_note": "The census uses the post-2023 44-municipality and 262-district structure; older 262-municipality datasets need an explicit crosswalk.",
        },
    }


def extract_honduras(root: Path, receipt: dict):
    entry = source_entry(receipt, "HND_C2013_GENERAL_POPULATION_TOME1")
    tome = extract_honduras_tome1(source_path(root, entry))
    municipal_receipt_path = root / "HND" / "2013" / "municipal-receipt.json"
    municipal_receipt = json.loads(municipal_receipt_path.read_text(encoding="utf-8"))
    municipalities = extract_honduras_municipal_reports(root, municipal_receipt, tome["departments"])
    municipal_source_id = "HND_C2013_MUNICIPAL_REPORT_COLLECTION"
    municipal_source = {
        "id": municipal_source_id,
        "name": "XVII Censo 2013: 298 municipal characteristics reports",
        "publisher": "Instituto Nacional de Estadistica Honduras",
        "url": entry["catalog_url"],
        "catalog_url": entry["catalog_url"],
        "status": "ready",
        "retrieved_at": municipal_receipt.get("retrieved_at", receipt["retrieved_at"]),
        "reference_period": "2013",
        "sha256": sha256(municipal_receipt_path),
        "bytes": sum(item["bytes"] for item in municipal_receipt["entries"]),
        "license": "Reuse terms not stated in the source publication; raw files are excluded from public redistribution pending terms review.",
        "license_url": None,
        "license_detail": "AreaData publishes source-linked factual observations and provenance only. This statement does not assert an open-data license.",
        "terms_review_status": "terms_not_stated_in_source",
        "raw_redistribution_status": "withheld_pending_terms_review",
        "normalized_observation_publication": "source_attributed_factual_observation_only",
        "geographic_level": "municipality",
        "note": "Collection receipt for 298 official municipal PDFs. Every observation retains its individual report URL and verified source hash.",
    }
    territories, observations, comparisons, terminal_ids = [], [], [], []
    departments = {}
    note = "INE Honduras XVII Censo 2013 published population. The official national figure is one person higher than the sum of the 18 published department totals; both values are retained without alteration."
    observations.append(observation("HND", tome["national"], entry["source_id"], 2013, note, "integer"))
    for row in tome["departments"]:
        department_id = f"HND:C2013:DEP:{row['code']}"
        departments[row["code"]] = {**row, "id": department_id, "municipality_ids": [], "municipality_sum": 0}
        territories.append(territory(department_id, "HND", row["name"], "department", "department", "HND", entry["source_id"], row["code"], "INE Honduras Censo 2013 department code"))
        observations.append(observation(department_id, row["population"], entry["source_id"], 2013, note, "integer"))
    for row in municipalities:
        department = departments[row["department_code"]]
        municipality_id = f"HND:C2013:MUN:{row['code'].replace('-', '')}"
        department["municipality_ids"].append(municipality_id)
        if row["population"] is not None:
            department["municipality_sum"] += row["population"]
        terminal_ids.append(municipality_id)
        territories.append(territory(municipality_id, "HND", row["name"], "municipality", "municipality", department["id"], municipal_source_id, row["code"], "INE Honduras Censo 2013 department-municipality code", "Name and code copied from the official municipality report cover; polygon edition remains unverified."))
        if row["population"] is not None:
            item = observation(municipality_id, row["population"], municipal_source_id, 2013, "Population total copied from the municipality's official Censo 2013 report; the observation retains its report URL and SHA-256.", "integer")
            item["source_detail_url"] = row["pdf_url"]
            item["source_detail_sha256"] = row["source_sha256"]
            observations.append(item)
    comparisons.append(comparison("HND", [item["id"] for item in departments.values()], "18 Honduras departments", [entry["source_id"]], "All 18 department totals in Tome 1. Their published sum is one person below the separately published national total."))
    for department in departments.values():
        comparisons.append(comparison(department["id"], department["municipality_ids"], f"Municipalities within {department['name']}", [municipal_source_id], "Complete membership from the 298 official Censo 2013 municipal reports."))
    accepted_municipalities = [row for row in municipalities if row["population"] is not None]
    municipal_sum = sum(row["population"] for row in accepted_municipalities)
    department_reconciliation = [
        {
            "department_code": code,
            "department": department["name"],
            "department_total": department["population"],
            "municipality_count": len(department["municipality_ids"]),
            "municipality_observations_accepted": sum(
                row["department_code"] == code and row["population"] is not None for row in municipalities
            ),
            "municipality_sum": department["municipality_sum"],
            "difference": difference(department["municipality_sum"], department["population"]),
        }
        for code, department in departments.items()
    ]
    withheld_municipalities = [
        {
            "code": row["code"],
            "name": row["name"],
            "published_population": row["published_population"],
            "reason": row["adoption_status"],
            "source_url": row["pdf_url"],
            "source_sha256": row["source_sha256"],
        }
        for row in municipalities
        if row["population"] is None
    ]
    return {
        "country_id": "HND",
        "period": "2013",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [source_record(entry, "Adopted national and department source."), municipal_source],
        "audit": {
            "national_population": tome["national"],
            "department_count": len(tome["departments"]),
            "department_sum": tome["department_sum"],
            "department_sum_difference": tome["department_sum_difference"],
            "municipality_count": len(municipalities),
            "municipality_observations_accepted": len(accepted_municipalities),
            "municipality_sum": municipal_sum,
            "municipality_sum_difference": difference(municipal_sum, tome["national"]),
            "municipal_reports_acquired_and_hashed": len(municipalities),
            "municipal_values_withheld": withheld_municipalities,
            "unresolved_atlantida_residual": 18_449,
            "department_reconciliation": department_reconciliation,
            "reconciliation_note": "Esparta's official PDF repeats San Francisco's 14,559 value and is withheld; no residual value is imputed. For nine other departments, published municipality sums differ from their published parent by one or two persons. National and department observations remain the published figures rather than being replaced by child sums.",
        },
    }


def extract_nicaragua(root: Path, receipt: dict):
    entry = source_entry(receipt, "NIC_C2005_OFFICIAL_FIGURES_TABLES")
    extracted = extract_nicaragua_table5(source_path(root, entry))
    territories, observations, comparisons, terminal_ids = [], [], [], []
    note = "INIDE VIII Censo de Población 2005 official population from Table 5. The two autonomous regions are retained as published first-order units."
    observations.append(observation("NIC", extracted["national"], entry["source_id"], 2005, note, "integer"))
    municipality_index = 0
    department_ids = []
    for department_index, department in enumerate(extracted["departments"], start=1):
        department_id = f"NIC:C2005:FIRST:{department_index:02d}"
        department_ids.append(department_id)
        subtype = "autonomous_region" if department["name"].startswith("R.A.A.") else "department"
        territories.append(territory(department_id, "NIC", department["name"], "first_order_division", "first_order_division", "NIC", entry["source_id"], None, None, "First-order unit name and order copied from Table 5; the concise table provides no code or boundary edition.", subtype))
        observations.append(observation(department_id, department["population"], entry["source_id"], 2005, note, "integer"))
        member_ids = []
        for municipality in department["municipalities"]:
            municipality_index += 1
            municipality_id = f"NIC:C2005:MUN:{municipality_index:03d}"
            member_ids.append(municipality_id)
            terminal_ids.append(municipality_id)
            territories.append(territory(municipality_id, "NIC", municipality["name"], "municipality", "municipality", department_id, entry["source_id"], None, None, "Municipality name and first-order membership copied from Table 5; the concise table provides no code or boundary edition."))
            observations.append(observation(municipality_id, municipality["population"], entry["source_id"], 2005, note, "integer"))
        comparisons.append(comparison(department_id, member_ids, f"Municipalities within {department['name']}", [entry["source_id"]], "Complete municipality membership and totals as published in Table 5."))
    comparisons.insert(0, comparison("NIC", department_ids, "17 Nicaragua first-order divisions", [entry["source_id"]], "15 departments and two autonomous regions as published in Table 5."))
    return {
        "country_id": "NIC",
        "period": "2005",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [source_record(entry, "Adopted national, first-order and municipality population source from Table 5. Both 1995 and 2005 columns were audited; the active series uses 2005 only.")],
        "audit": {
            "national_population": extracted["national"],
            "department_count": len(extracted["departments"]),
            "department_sum": sum(item["population"] for item in extracted["departments"]),
            "department_sum_difference": 0.0,
            "municipality_count": len(extracted["municipalities"]),
            "municipality_sum": sum(item["population"] for item in extracted["municipalities"]),
            "municipality_sum_difference": 0.0,
            "geography_note": "The source table has no official code field. Internal sequential IDs preserve source order and must not be treated as national statistical codes.",
        },
    }


def extract_costa_rica(root: Path, receipt: dict):
    report_entry = source_entry(receipt, "CRI_C2022_POP_HOUSING_ESTIMATED_RESULTS")
    canton_entry = source_entry(receipt, "CRI_C2022_CANTON_POPULATION_ARCGIS")
    province_entry = source_entry(receipt, "CRI_C2022_PROVINCE_POPULATION_ARCGIS")
    report_path = source_path(root, report_entry)
    report_text = "\n".join((page.extract_text() or "") for page in PdfReader(report_path).pages[14:22])
    if "5 044 197" not in report_text:
        raise ValueError("Costa Rica official 2022 estimated national population was not found in the adopted report")
    national = 5_044_197
    province_json = json.loads(source_path(root, province_entry).read_text(encoding="utf-8"))
    canton_json = json.loads(source_path(root, canton_entry).read_text(encoding="utf-8"))
    province_rows = [item["attributes"] for item in province_json.get("features", [])]
    canton_rows = [item["attributes"] for item in canton_json.get("features", [])]
    if len(province_rows) != 7 or len(canton_rows) != 82:
        raise ValueError(f"Unexpected Costa Rica ArcGIS geography counts: {len(province_rows)}, {len(canton_rows)}")
    territories, observations, comparisons, terminal_ids = [], [], [], []
    provinces = {}
    note = "INEC 2022 official population and housing estimate produced after partial census coverage; it is an adjusted estimate associated with the 2022 census operation, not a complete enumeration count."
    observations.append(observation("CRI", national, report_entry["source_id"], 2022, note, "integer"))
    for row in province_rows:
        code = str(row["COD_UGEP"])
        province_id = f"CRI:C2022:PROV:{code}"
        provinces[code] = {"id": province_id, "name": row["NOM_UGEP"], "value": finite_number(row["POBLACION"], f"CRI province {code}"), "canton_ids": [], "canton_sum": 0}
        territories.append(territory(province_id, "CRI", row["NOM_UGEP"], "province", "province", "CRI", province_entry["source_id"], code, "INEC 2022 UGEP province code", "Official INEC ArcGIS result attribute; service response in this acquisition excludes geometry, so polygon joining remains pending."))
        observations.append(observation(province_id, row["POBLACION"], province_entry["source_id"], 2022, note, "integer"))
    for row in canton_rows:
        province_code = str(row["COD_UGEP"])
        code = str(row["COD_UGEC"])
        if province_code not in provinces:
            raise ValueError(f"Unknown Costa Rica province for canton {code}")
        canton_id = f"CRI:C2022:CANTON:{code}"
        value = finite_number(row["POB_TOTAL"], f"CRI canton {code}")
        provinces[province_code]["canton_ids"].append(canton_id)
        provinces[province_code]["canton_sum"] += value
        terminal_ids.append(canton_id)
        territories.append(territory(canton_id, "CRI", row["NOM_UGEC"], "canton", "canton", provinces[province_code]["id"], canton_entry["source_id"], code, "INEC 2022 UGEC canton code", "Official INEC ArcGIS result attribute; service response in this acquisition excludes geometry, so polygon joining remains pending."))
        observations.append(observation(canton_id, value, canton_entry["source_id"], 2022, note, "integer"))
    province_sum = sum(item["value"] for item in provinces.values())
    canton_sum = sum(item["canton_sum"] for item in provinces.values())
    if province_sum != national or canton_sum != national:
        raise ValueError("Costa Rica province or canton totals do not reconcile to the official estimate")
    comparisons.append(comparison("CRI", [item["id"] for item in provinces.values()], "7 Costa Rica provinces", [province_entry["source_id"]], "Complete province register and population from the official INEC 2022 ArcGIS result layer."))
    for province in provinces.values():
        if province["canton_sum"] != province["value"]:
            raise ValueError(f"Costa Rica canton subtotal does not reconcile for {province['name']}")
        comparisons.append(comparison(province["id"], province["canton_ids"], f"Cantons within {province['name']}", [canton_entry["source_id"]], "Complete canton membership and population from the official INEC 2022 ArcGIS result layer."))
    return {
        "country_id": "CRI",
        "period": "2022",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [
            source_record(report_entry, "Official methodological and national-results report; establishes that the 2022 values are adjusted estimates following partial census coverage."),
            source_record(province_entry, "Adopted official 2022 province population result service."),
            source_record(canton_entry, "Adopted official 2022 canton population result service."),
        ],
        "audit": {
            "national_population": national,
            "province_count": len(provinces),
            "province_sum": province_sum,
            "province_sum_difference": difference(province_sum, national),
            "canton_count": len(terminal_ids),
            "canton_sum": canton_sum,
            "canton_sum_difference": difference(canton_sum, national),
            "method_status": "official_corrected_estimate_due_partial_census_coverage",
            "rejected_series": "The revised 2000-2050 projection series was inspected but not mixed into this census-associated 2022 estimate because its 2022 canton total and statistical method differ.",
        },
    }


def extract_panama(root: Path, receipt: dict):
    entry = source_entry(receipt, "PAN_C2023_VOLUME5_TABLE1")
    path = source_path(root, entry)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Cuadro01"]
        rows = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=9, min_col=1, max_col=16, values_only=True), start=9):
            first_order, district, corregimiento = values[:3]
            population = values[15]
            if not isinstance(population, (int, float)):
                continue
            nonblank = [(index, value) for index, value in enumerate((first_order, district, corregimiento)) if value not in (None, "")]
            if len(nonblank) != 1:
                raise ValueError(f"Ambiguous Panama geography row {row_number}: {nonblank}")
            rows.append((row_number, nonblank[0][0], str(nonblank[0][1]).strip(), finite_number(population, f"PAN Cuadro01!P{row_number}")))
    finally:
        book.close()
    national_row = next(row for row in rows if row[1] == 0 and row[2] == "TOTAL")
    national = national_row[3]
    territories, observations, comparisons, terminal_ids = [], [], [], []
    first_orders, districts = [], []
    current_first = current_district = None
    note = "INEC Panama Censos Nacionales 2023 population from Volume V, Table 1. Province, indigenous comarca, district and corregimiento labels follow the published hierarchy."
    observations.append(observation("PAN", national, entry["source_id"], 2023, note, "integer"))
    for _, level_index, published_name, value in rows:
        if level_index == 0:
            if published_name == "TOTAL":
                continue
            name = clean_panama_label(published_name)
            index = len(first_orders) + 1
            territory_id = f"PAN:C2023:FIRST:{index:02d}"
            subtype = "indigenous_comarca" if name.startswith("Comarca ") else "province"
            current_first = {"id": territory_id, "name": name, "published_name": published_name, "value": value, "district_ids": [], "district_sum": 0, "subtype": subtype}
            first_orders.append(current_first)
            current_district = None
            territories.append(territory(territory_id, "PAN", name, "first_order_division", "first_order_division", "PAN", entry["source_id"], None, None, "Name, subtype and source order copied from Table 1; the table provides no geography code or polygon edition.", subtype))
            observations.append(observation(territory_id, value, entry["source_id"], 2023, note, "integer"))
        elif level_index == 1:
            if current_first is None:
                raise ValueError("Panama district encountered before a first-order unit")
            name = clean_panama_label(published_name)
            index = len(districts) + 1
            territory_id = f"PAN:C2023:DIST:{index:03d}"
            current_district = {"id": territory_id, "name": name, "published_name": published_name, "value": value, "corregimiento_ids": [], "corregimiento_sum": 0, "parent": current_first}
            districts.append(current_district)
            current_first["district_ids"].append(territory_id)
            current_first["district_sum"] += value
            territories.append(territory(territory_id, "PAN", name, "district", "district", current_first["id"], entry["source_id"], None, None, "District name and source order copied from Table 1; the table provides no geography code or polygon edition."))
            observations.append(observation(territory_id, value, entry["source_id"], 2023, note, "integer"))
        else:
            if current_district is None:
                raise ValueError("Panama corregimiento encountered before a district")
            name = clean_panama_label(published_name)
            index = len(terminal_ids) + 1
            territory_id = f"PAN:C2023:CORR:{index:03d}"
            current_district["corregimiento_ids"].append(territory_id)
            current_district["corregimiento_sum"] += value
            terminal_ids.append(territory_id)
            territories.append(territory(territory_id, "PAN", name, "corregimiento", "corregimiento", current_district["id"], entry["source_id"], None, None, "Corregimiento name and source order copied from Table 1; the table provides no geography code or polygon edition."))
            observations.append(observation(territory_id, value, entry["source_id"], 2023, note, "integer"))
    if (len(first_orders), len(districts), len(terminal_ids)) != (13, 82, 699):
        raise ValueError(f"Unexpected Panama geography counts: {len(first_orders)}, {len(districts)}, {len(terminal_ids)}")
    first_order_sum = sum(row["value"] for row in first_orders)
    district_sum = sum(row["value"] for row in districts)
    corregimiento_sum = sum(row["corregimiento_sum"] for row in districts)
    if (first_order_sum, district_sum, corregimiento_sum) != (national,) * 3:
        raise ValueError("Panama geography levels do not reconcile to the national total")
    comparisons.append(comparison("PAN", [row["id"] for row in first_orders], "13 Panama first-order divisions", [entry["source_id"]], "Ten provinces and three indigenous comarcas as published in Table 1."))
    for first_order in first_orders:
        if first_order["district_sum"] != first_order["value"]:
            raise ValueError(f"Panama district subtotal does not reconcile for {first_order['name']}")
        comparisons.append(comparison(first_order["id"], first_order["district_ids"], f"Districts within {first_order['name']}", [entry["source_id"]], "Complete district membership and population as published in Table 1."))
    for district in districts:
        if district["corregimiento_sum"] != district["value"]:
            raise ValueError(f"Panama corregimiento subtotal does not reconcile for {district['name']}")
        comparisons.append(comparison(district["id"], district["corregimiento_ids"], f"Corregimientos within {district['name']}", [entry["source_id"]], "Complete corregimiento membership and population as published in Table 1."))
    return {
        "country_id": "PAN",
        "period": "2023",
        "territories": territories,
        "observations": observations,
        "comparisons": comparisons,
        "terminal_territory_ids": terminal_ids,
        "sources": [source_record(entry, "Adopted national, first-order, district and corregimiento population source.", "CC BY 4.0", entry["catalog_url"])],
        "audit": {
            "national_population": national,
            "first_order_count": len(first_orders),
            "first_order_sum": first_order_sum,
            "first_order_sum_difference": 0.0,
            "province_count": sum(row["subtype"] == "province" for row in first_orders),
            "indigenous_comarca_count": sum(row["subtype"] == "indigenous_comarca" for row in first_orders),
            "district_count": len(districts),
            "district_sum": district_sum,
            "district_sum_difference": 0.0,
            "corregimiento_count": len(terminal_ids),
            "corregimiento_sum": corregimiento_sum,
            "corregimiento_sum_difference": 0.0,
            "geography_note": "The published table has no geography-code columns. Internal sequential IDs preserve source order and must not be treated as official codes.",
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True, help="Acquisition directory containing receipt.json and XLSX files")
    parser.add_argument("--out", help="Normalized JSON; defaults to ROOT/normalized-census.json")
    args = parser.parse_args()
    root = Path(args.root).resolve()
    receipt = json.loads((root / "receipt.json").read_text(encoding="utf-8"))
    for entry in receipt["entries"]:
        entry["retrieved_at"] = receipt["retrieved_at"]
    countries = [
        extract_belize(root, receipt),
        extract_guatemala(root, receipt),
        extract_el_salvador(root, receipt),
        extract_honduras(root, receipt),
        extract_nicaragua(root, receipt),
        extract_costa_rica(root, receipt),
        extract_panama(root, receipt),
    ]
    country_totals = {
        country["country_id"]: next(
            row["value"] for row in country["observations"] if row["territory_id"] == country["country_id"]
        )
        for country in countries
    }
    regional_total = sum(country_totals.values())
    component_periods = {country["country_id"]: country["period"] for country in countries}
    output = {
        "schema_version": "1.0",
        "scope_id": receipt["scope_id"],
        "generated_at": receipt["retrieved_at"],
        "status": "available_complete_country_coverage",
        "indicator": {
            "id": INDICATOR_ID,
            "name": "Official census population",
            "theme": "Population",
            "unit": "people",
            "definition": "Population published in each country's adopted official census product. Country year and country-specific weighting or enumeration method remain attached to every observation.",
            "definition_id": "official-census-population-country-product-v1",
            "population": "Population under each national statistical office's adopted census definition; see country observation footnote",
            "measurement_method": "country-specific official census product",
            "aggregation": "sum",
            "series_family": "census",
            "display_role": "primary",
            "period_policy": "latest_available_by_component",
            "display_decimals": 0
        },
        "series_source": {
            "id": "areadata-ca7-census-series",
            "name": "AreaData Central America census series contract",
            "publisher": "AreaData project",
            "url": "https://areadata.net/",
            "status": "ready",
            "retrieved_at": receipt["retrieved_at"],
            "reference_period": ", ".join(f"{country_id} {period}" for country_id, period in component_periods.items()),
            "license": "Project metadata; official observations retain source-specific terms",
            "geographic_level": "multi-country and subnational",
            "note": "This source defines only the cross-country indicator contract. Each observation cites its official national source. The regional value sums the seven exact national observations after coverage review and keeps every component year and method visible; it is not reconstructed from municipalities."
        },
        "countries": countries,
        "summary": {
            "pilot_country_count": 7,
            "countries_acquired": [country["country_id"] for country in countries],
            "countries_pending": [],
            "territories_added": sum(len(country["territories"]) for country in countries),
            "observations_added": sum(len(country["observations"]) for country in countries),
            "regional_total_status": "available_complete_country_coverage_mixed_reference_years",
            "regional_total": regional_total,
            "regional_total_source": "sum_of_exact_country_observations",
            "regional_total_component_periods": component_periods,
            "regional_total_country_values": country_totals,
        }
    }
    output["indicator"]["source_id"] = output["series_source"]["id"]
    target = Path(args.out).resolve() if args.out else root / "normalized-census.json"
    target.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Extracted {output['summary']['observations_added']} census observations across {output['summary']['territories_added']} added territories.")
    print(f"Regional total: {regional_total} ({output['summary']['regional_total_status']}). Output: {target}")


if __name__ == "__main__":
    main()
