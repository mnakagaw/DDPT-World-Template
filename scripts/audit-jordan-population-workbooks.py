"""Read-only sheet and field inventory for Jordan DoS population workbooks."""

import argparse
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


ADOPTED_CELLS = {
    ("PopulationEstimatesbyLocality.xlsx", "الملخص "): "B5:E122, excluding B118:E121 pending Aqaba hierarchy review",
    ("PopulationEstimates.xlsx", "2.3"): "B27:C39; B6:D18 used only to reconcile the summary workbook",
    ("PopulationEstimates.xlsx", "2.7"): "E29:E41 and G29:G41; D29:D41 used only to reconcile population",
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True)
    parser.add_argument("--detail", help="Print nonempty row/cell locations for one workbook:sheet")
    parser.add_argument("--public-output", help="Write a small metadata-only sheet and numeric-column ledger")
    args = parser.parse_args()
    raw = Path(args.project) / "raw" / "official-jordan"
    result = {}
    for filename in ("PopulationEstimates.xlsx", "PopulationEstimatesbyLocality.xlsx", "Municipalities.xlsx"):
        path = raw / filename
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet in book:
            preview = []
            numeric_columns = {}
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number <= 6:
                    compact = [str(value)[:120] if value is not None else None for value in row[:8]]
                    if any(value is not None for value in compact):
                        preview.append(compact)
                for column_number, value in enumerate(row, 1):
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        column = get_column_letter(column_number)
                        item = numeric_columns.setdefault(column, {"count": 0, "first_cell": None, "last_cell": None,
                                                                    "first_value": None, "last_value": None})
                        item["count"] += 1
                        item["last_cell"] = f"{column}{row_number}"
                        item["last_value"] = value
                        if item["first_cell"] is None:
                            item["first_cell"] = item["last_cell"]
                            item["first_value"] = value
            adopted = ADOPTED_CELLS.get((filename, sheet.title))
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                           "preview": preview, "numeric_columns": numeric_columns,
                           "disposition": "partially_adopted" if adopted else "not_adopted_pending_semantic_review",
                           "adopted_cells_or_checks": adopted,
                           "warning": "Numeric-column counts are a mechanical inventory, not a semantic source audit."})
        result[filename] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size, "sheets": sheets}
    sys.stdout.reconfigure(encoding="utf-8")
    output = Path(args.project) / "evidence" / "JOR_SOURCE_WORKBOOK_INVENTORY.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.public_output:
        public = {filename: {"sha256": item["sha256"], "bytes": item["bytes"],
                             "sheets": [{"name": sheet["name"], "rows": sheet["rows"],
                                         "columns": sheet["columns"],
                                         "numeric_columns": {column: {"count": entry["count"],
                                                                      "first_cell": entry["first_cell"],
                                                                      "last_cell": entry["last_cell"]}
                                                             for column, entry in sheet["numeric_columns"].items()},
                                         "disposition": sheet["disposition"],
                                         "adopted_cells_or_checks": sheet["adopted_cells_or_checks"]}
                                        for sheet in item["sheets"]]}
                  for filename, item in result.items()}
        Path(args.public_output).write_text(json.dumps(public, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.detail:
        filename, sheet_name = args.detail.split(":", 1)
        book = openpyxl.load_workbook(raw / filename, read_only=True, data_only=True)
        matches = [name for name in book.sheetnames if name.strip() == sheet_name.strip()]
        if len(matches) != 1:
            raise ValueError(f"Sheet not found unambiguously: {sheet_name!r}: {matches}")
        sheet = book[matches[0]]
        for row in sheet:
            cells = [f"{cell.coordinate}={str(cell.value)[:100]}" for cell in row[:8] if cell.value is not None]
            if cells:
                print(" | ".join(cells))
        return
    print(json.dumps({filename: {"sha256": item["sha256"], "bytes": item["bytes"],
                                  "sheets": [{"name": sheet["name"], "rows": sheet["rows"],
                                              "columns": sheet["columns"]} for sheet in item["sheets"]]}
                      for filename, item in result.items()}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
