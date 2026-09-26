"""Adopt a pinned, directly reported subset of Qatar Census 2020 municipal cells.

The remaining workbook columns and all unselected cells stay unassessed. This
collector never links 2015 provider polygons or newer GIS editions to 2020 cells.
"""

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


AREAS = [
    ("Doha", "الدوحة", "DOHA"),
    ("Al Rayyan", "الريان", "AL-RAYYAN"),
    ("Al Wakra", "الوكرة", "AL-WAKRA"),
    ("Umm Slal", "ام صلال", "UMM-SLAL"),
    ("Al Khor and Al Thakhira", "الخور والذخيرة", "AL-KHOR-AND-AL-THAKHIRA"),
    ("Al Shamal", "الشمال", "AL-SHAMAL"),
    ("Al Daayen", "الظعاين", "AL-DAAYEN"),
    ("Al Sheehaniya", "الشيحانية", "AL-SHEEHANIYA"),
]

# Table, sheet title, coordinate rule, column, row; only the listed cells are adopted.
TABLE_TITLES = {
    "1": "Population by Sex and Municipality",
    "31": "Population (10+) by Municipality, Sex and Educational Attainment",
    "108": "Population with Difficulties (Disabilities can be more than 1) by Type of Difficulty & Degree and Municipality",
    "131": "Buildings by Building Status and Municipality",
    "133": "Occupied Housing Units, by Type of Unit and Municipality",
    "137": "Households, by Type of Housing Unit and Municipality",
    "139": "Completed Residential Buildings, by Connection to Public Facilities and Municipality",
    "149": "Operating Establishments by Sector and Municipality",
    "152": "Business Establishments & Persons Engaged (P.E.) by Number of Employees and Municipality",
}

SPECS = {
    "POP_TOTAL": ("1", "row", "D", 8, "2020 census population", "Population", "people", "All persons in the December 2020 census municipality table"),
    "POP_MALE": ("1", "row", "C", 8, "2020 census male population", "Population", "people", "Male persons in the same table"),
    "POP_FEMALE": ("1", "row", "B", 8, "2020 census female population", "Population", "people", "Female persons in the same table"),
    "POP_10PLUS": ("31", "column", None, 7, "2020 census population aged 10+", "Education", "people", "Persons aged 10+ in the educational-attainment census table"),
    "ILLITERATE_10PLUS": ("31", "column", None, 10, "2020 census illiterate persons aged 10+", "Education", "people", "Persons aged 10+ in the source's illiterate educational-attainment category; a count, not a rate"),
    "UNIVERSITY_ABOVE_10PLUS": ("31", "column", None, 25, "2020 census university-and-above attainment, age 10+", "Education", "people", "Persons aged 10+ in the source's university-and-above educational-attainment category; a count"),
    "PERSONS_WITH_DIFFICULTIES": ("108", "column", None, 8, "2020 census persons with difficulties", "Disability", "people", "Distinct persons in the source's Total Persons row; types of difficulty may overlap"),
    "BUILDINGS_TOTAL": ("131", "row", "E", 7, "2020 census buildings", "Housing", "buildings", "All building statuses in the December 2020 census table"),
    "BUILDINGS_COMPLETE": ("131", "row", "D", 7, "2020 census completed buildings", "Housing", "buildings", "Buildings with Complete status; a count, not a completion rate"),
    "BUILDINGS_UNDER_CONSTRUCTION": ("131", "row", "C", 7, "2020 census buildings under construction", "Housing", "buildings", "Buildings with Under Construction status"),
    "OCCUPIED_HOUSING_UNITS": ("133", "row", "G", 7, "2020 census occupied housing units", "Housing", "housing units", "Occupied units of all listed housing-unit types"),
    "HOUSEHOLDS_TOTAL": ("137", "row", "G", 7, "2020 census households", "Households", "households", "Households of all listed housing-unit types; distinct from occupied housing units"),
    "RESIDENTIAL_BUILDINGS_SEWAGE_CONNECTED": ("139", "row", "C", 8, "2020 census completed residential buildings connected to sewage", "Infrastructure", "buildings", "Completed residential buildings listed as connected to a sewage network"),
    "RESIDENTIAL_BUILDINGS_SEWAGE_NOT_CONNECTED": ("139", "row", "B", 8, "2020 census completed residential buildings not connected to sewage", "Infrastructure", "buildings", "Completed residential buildings listed as not connected to a sewage network"),
    "OPERATING_ESTABLISHMENTS": ("149", "row", "I", 10, "2020 census operating establishments", "Economy", "establishments", "Operating establishments across all listed sectors"),
    "BUSINESS_ESTABLISHMENTS": ("149", "row", "F", 10, "2020 census business-sector establishments", "Economy", "establishments", "Operating establishments in the source's business-establishment sector subtotal"),
    "PERSONS_ENGAGED_BUSINESS": ("152", "pair", "J", 9, "2020 census persons engaged in business establishments", "Economy", "people", "Persons engaged in business establishments; not all employed persons and not a workforce rate"),
}


def count(value, locator):
    if isinstance(value, bool) or (not isinstance(value, int) and not (isinstance(value, str) and value.isdigit())):
        raise ValueError(f"Expected nonnegative integer at {locator}: {value!r}")
    result = int(value)
    if result < 0:
        raise ValueError(f"Negative count at {locator}")
    return result


def source_cell(sheet, row, col):
    locator = f"Census_Final_Results.xlsx!'{sheet.title}'!{col}{row}"
    return count(sheet[f"{col}{row}"].value, locator), locator


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    manifest_path = args.manifest.resolve(strict=True)
    if not manifest_path.is_relative_to(project):
        raise ValueError("Manifest must be inside the country project")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    receipt = next(r for r in manifest["receipts"] if r["file"] == "Census_Final_Results.xlsx")
    if receipt["status"] != "acquired":
        raise ValueError("Official census workbook was not acquired")
    original = manifest_path.parent / receipt["file"]
    actual_sha = hashlib.sha256(original.read_bytes()).hexdigest()
    if actual_sha != receipt["sha256"]:
        raise ValueError("Official workbook differs from its acquisition receipt")
    workbook = load_workbook(original, read_only=True, data_only=True)
    for table, expected in TABLE_TITLES.items():
        if workbook[table]["A3"].value != expected or "2020" not in str(workbook[table]["A2"].value):
            raise ValueError(f"Table {table} title or census period changed")
    dashboard_path = project / "data/dashboard.json"
    dataset = json.loads(dashboard_path.read_text(encoding="utf-8"))
    if dataset["country"]["id"] != "QAT" or any(item["id"].startswith("QAT:CENSUS2020:") for item in dataset["territories"]):
        raise ValueError("Expected an unmodified Qatar bootstrap")
    inventory_path = project / "evidence/QAT_CENSUS2020_SHEET_INVENTORY.json"
    inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    if inventory["source_sha256"] != actual_sha or inventory["sheet_count"] != 209 or inventory["numbered_table_count"] != 156:
        raise ValueError("Complete source workbook inventory is required")
    column_path = project / "evidence/QAT_CENSUS2020_NUMERIC_COLUMNS.csv"
    with column_path.open(encoding="utf-8-sig", newline="") as handle:
        columns = list(csv.DictReader(handle))
    if len(columns) != inventory["numeric_column_count"]:
        raise ValueError("Numeric-column register incomplete")

    # Table 1 supplies the complete 2020 reporting-area list and bilingual labels.
    for index, (name, arabic, _) in enumerate(AREAS):
        row = 9 + index
        if workbook["1"][f"A{row}"].value != name or workbook["1"][f"E{row}"].value != arabic:
            raise ValueError(f"2020 census area row {row} changed")
    dataset["territories"] = [item for item in dataset["territories"] if item["id"] == "QAT"]
    dataset["boundaries"] = {"type": "FeatureCollection", "features": []}
    ids = {}
    for name, arabic, slug in AREAS:
        area_id = f"QAT:CENSUS2020:MUNICIPALITY:{slug}"
        ids[name] = area_id
        dataset["territories"].append({"id": area_id, "name": name, "local_name": arabic,
            "level": "adm1", "type": "2020 census municipality", "parent_id": "QAT",
            "official_code": None, "code_system": "2020 census table 1 labels; official code not yet verified",
            "boundary_version": None, "source_id": "qat-npc-census2020-table-1"})
    now = datetime.now(timezone.utc).isoformat()
    raw_path = str(original.relative_to(project)).replace("\\", "/")
    for table, title in TABLE_TITLES.items():
        dataset["sources"].append({"id": f"qat-npc-census2020-table-{table}", "name": f"Census 2020 Table {table}: {title}",
            "url": receipt["url"], "publisher": "Qatar National Planning Council, Census 2020 results",
            "reference_period": "December 2020 census", "geographic_level": "national and eight census municipality reporting areas",
            "status": "ready", "retrieved_at": receipt["retrieved_at"], "raw_path": raw_path,
            "sha256": actual_sha, "license": "Official NPC website reuse terms require review before external release",
            "note": "Only the listed cells were semantically adopted; the remaining workbook columns/cells remain unassessed. Official 2020 municipality codes and polygon edition are unverified."})

    controls = {}
    adopted_cells = []
    for key, (table, layout, column, base_row, name, theme, unit, population) in SPECS.items():
        source_id = f"qat-npc-census2020-table-{table}"
        indicator_id = f"QAT_CENSUS2020_{key}"
        dataset["indicators"].append({"id": indicator_id, "name": name, "theme": theme, "unit": unit,
            "source_id": source_id, "definition": f"{population}. Directly reported in Census 2020 Table {table}; December 2020, not an annual estimate.",
            "population": population, "aggregation": "sum", "measurement_method": "census_2020_published_table_cell",
            "series_family": "census", "display_role": "primary", "period_policy": "latest_available_per_indicator", "display_decimals": 0})
        values = []
        for area_index in range(9):
            national = area_index == 0
            if layout == "row":
                row = base_row if national else base_row + area_index
                col = column
            elif layout == "column":
                row = base_row
                col = "K" if national and table == "31" else "J" if national else get_column_letter((11 if table == "31" else 10) - area_index)
            else:
                row = base_row if national else base_row + 2 * area_index
                col = column
            value, locator = source_cell(workbook[table], row, col)
            adopted_cells.append({"indicator_id": indicator_id, "table": table, "cell": f"{col}{row}", "value": value})
            territory_id = "QAT" if national else ids[AREAS[area_index - 1][0]]
            dataset["observations"].append({"territory_id": territory_id, "indicator_id": indicator_id,
                "period": "2020", "value": value, "status": "observed", "source_id": source_id,
                "measurement_method": "census_2020_published_table_cell", "source_locator": locator,
                "footnote": "Published December 2020 census table cell; current legal boundary and official code not inferred."})
            values.append(value)
        if sum(values[1:]) != values[0]:
            raise ValueError(f"Eight municipalities do not reconcile with the published national cell: {key}")
        controls[key] = values[0]

    if controls["POP_TOTAL"] != 2846118 or controls["POP_MALE"] + controls["POP_FEMALE"] != controls["POP_TOTAL"]:
        raise ValueError("Census 2020 population control changed")
    if controls["BUILDINGS_TOTAL"] != 222701 or controls["BUILDINGS_COMPLETE"] != 204056:
        raise ValueError("Building-status controls changed")
    if controls["HOUSEHOLDS_TOTAL"] != 281136 or controls["OCCUPIED_HOUSING_UNITS"] != 302119:
        raise ValueError("Households and occupied units must remain separate source measures")
    if controls["RESIDENTIAL_BUILDINGS_SEWAGE_CONNECTED"] + controls["RESIDENTIAL_BUILDINGS_SEWAGE_NOT_CONNECTED"] != 165619:
        raise ValueError("Sewage connection categories do not cover completed residential buildings")
    if controls["BUSINESS_ESTABLISHMENTS"] != count(workbook["152"]["J8"].value, "152!J8"):
        raise ValueError("Business-establishment crosscheck failed")
    if controls["HOUSEHOLDS_TOTAL"] != count(workbook["7"]["D6"].value, "7!D6"):
        raise ValueError("Independent household crosscheck failed")
    for index, (name, _, _) in enumerate(AREAS, 1):
        row = 10 + index
        if workbook["149"][f"A{row}"].value != name or workbook["131"][f"A{7+index}"].value != name:
            raise ValueError(f"Source municipality row alignment changed: {name}")
        for table, first_area_row in (("133", 8), ("137", 8), ("139", 9)):
            if workbook[table][f"A{first_area_row+index-1}"].value != name:
                raise ValueError(f"Table {table} municipality row alignment changed: {name}")
        if workbook["152"][f"A{8+2*index}"].value != name:
            raise ValueError(f"Table 152 municipality pair alignment changed: {name}")
        reverse_col31 = get_column_letter(11-index)
        if name not in str(workbook["31"][f"{reverse_col31}6"].value):
            raise ValueError(f"Education table municipality column alignment changed: {name}")
        reverse_col108 = get_column_letter(10-index)
        difficulty_alias = {"Al Shamal": "Madinat Al Shamal", "Al Khor and Al Thakhira": "Al Khor &. Al Thak"}.get(name, name)
        if difficulty_alias not in str(workbook["108"][f"{reverse_col108}7"].value):
            raise ValueError(f"Difficulty table municipality column alignment changed: {name}")
        table1row = 8 + index
        female = count(workbook["1"][f"B{table1row}"].value, "table1 female")
        male = count(workbook["1"][f"C{table1row}"].value, "table1 male")
        total = count(workbook["1"][f"D{table1row}"].value, "table1 total")
        if male + female != total or total != count(workbook["3"][f"{get_column_letter(10-index)}7"].value, "table3 population"):
            raise ValueError(f"Population partition/age-table crosscheck failed: {name}")
    dataset["analysis"]["terminal_territory_ids"] = list(ids.values())
    dataset["analysis"]["comparisons"] = [{"parent_id": "QAT", "member_ids": list(ids.values()),
        "label": "Eight December 2020 census municipalities",
        "membership_note": "Table 1 reports exactly these eight bilingual municipality rows and a national control. The statistical 2020 reporting areas are not joined to newer GIS polygons or legal plan jurisdictions.",
        "source_ids": ["qat-npc-census2020-table-1"]}]
    dataset["country"]["geography_note"] = ("Census 2020 directly reports eight municipalities in December 2020. Official 2020 codes and dated polygons remain unverified. The 2015 geoBoundaries provider shapes and newer Qatar GIS records are withheld; national WDI series remain separate.")
    for gap in dataset["gaps"]:
        if gap["category"] == "boundary_reconciliation":
            gap.update(status="partial", detail="The eight Census 2020 municipality labels are retained without polygons or verified official codes. The current official GIS layer has later STARTDATE values, so it is not a 2020 geometry crosswalk.", next_action="Acquire an official December 2020 municipality code and polygon edition, confirm area changes through 2026, then evaluate a dated join.")
        elif gap["category"] == "subnational_statistics":
            gap.update(status="partial", detail=f"The 209-sheet official Census 2020 workbook has {inventory['numeric_column_count']} numeric columns; {len(SPECS)} carefully selected direct count indicators have eight municipality cells plus national controls. Remaining cells are priority_unassessed.", next_action="Semantically assess all remaining municipal and zone tables, with population scope/denominators and official zone-to-municipality linkage before expansion.")
        elif gap["category"] == "planning_documents":
            gap.update(status="not_collected", detail="The official Qatar National Master Plan site locates municipality spatial-development strategies; adopted current local plans, budgets, execution and evaluations are not yet verified.", next_action="Acquire official plans and legal approval records by municipality, then budget, implementation and evaluation originals with period and page mapping.")
    dataset["collection"]["adapters"] = sorted(set(dataset["collection"]["adapters"] + ["qatar-npc-census2020-selected-direct-cells"]))
    dataset["collection"]["notes"].append(f"NPC Census 2020 workbook SHA-256 {actual_sha}; 17 direct indicators use 153 published national/municipality cells. Remaining workbook cells are unassessed; no historical boundary join.")
    dataset["generated_at"] = now
    dashboard_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    selected_columns = {(item["table"], "".join(c for c in item["cell"] if c.isalpha())) for item in adopted_cells}
    for sheet in inventory["sheets"]:
        if sheet["sheet"] in TABLE_TITLES:
            sheet["adoption_decision"] = "partially_adopted_selected_cells"
    inventory["adoption_note"] = "Only the coordinates in QAT_CENSUS2020_IMPORT_AUDIT.json are adopted; a selected column still contains unassessed cells."
    inventory_path.write_text(json.dumps(inventory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for item in columns:
        if (item["sheet"], item["column"]) in selected_columns:
            item["adoption_decision"] = "partially_adopted_selected_cells"
    with column_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns[0]))
        writer.writeheader()
        writer.writerows(columns)
    audit = {"status": "partial_candidate_not_accepted", "source_sha256": actual_sha,
        "source_manifest": str(manifest_path.relative_to(project)).replace("\\", "/"),
        "source_worksheets": inventory["worksheet_count"], "source_numbered_tables": inventory["numbered_table_count"],
        "numeric_columns": inventory["numeric_column_count"], "adopted_tables": list(TABLE_TITLES),
        "adopted_indicators": len(SPECS), "adopted_direct_cells": adopted_cells,
        "national_controls": controls, "census_municipality_ids": ids, "source_geometry_adopted": 0,
        "unresolved": ["Remaining workbook cells and zone tables", "Official 2020 codes and dated polygons",
                       "Current planning-law and plan approval status", "Individual budget/execution/evaluation originals",
                       "Representative browser/download/print QA and independent country acceptance"]}
    (project / "evidence/QAT_CENSUS2020_IMPORT_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"municipalities": 8, "indicators": len(SPECS), "direct_cells": len(adopted_cells),
                      "population": controls["POP_TOTAL"], "dataset": str(dashboard_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
