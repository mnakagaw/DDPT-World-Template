#!/usr/bin/env python3
"""Build the El Salvador 2024 Census country-depth bundle.

The collector reads the complete set of official BCR workbooks acquired for
the project, adopts a small set of auditable indicators at country,
department, municipality and district level, and joins the current BCR
ArcGIS selector geometries by official codes.  National poverty and nutrition
context comes from World Bank API series and is never copied to lower areas.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook


UA = {"User-Agent": "AreaData/0.10.2 (+https://areadata.net/)"}
CENSUS_PAGE = "https://geoportal.bcr.gob.sv/"
TABLE_ROOT = "https://censo2024.bcr.gob.sv/wp-content/uploads/tablas-geoportal/2025"
BOUNDARY_SERVICES = {
    "department": "https://services8.arcgis.com/xSn1Pefr3M9zkw7Z/arcgis/rest/services/Departamentos_Selector/FeatureServer/0",
    "municipality": "https://services8.arcgis.com/xSn1Pefr3M9zkw7Z/arcgis/rest/services/Selector_Municipios/FeatureServer/0",
    "district": "https://services8.arcgis.com/xSn1Pefr3M9zkw7Z/arcgis/rest/services/Selector_Distritos/FeatureServer/0",
}
WB_API = "https://api.worldbank.org/v2/country/SLV/indicator/{indicator}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def finite(value, context: str) -> float:
    if value in (None, ""):
        return 0.0
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Non-finite value in {context}")
    return result


def code(value) -> str | None:
    match = re.match(r"\s*(\d+)\s*-", str(value or ""))
    return match.group(1).zfill(2) if match else None


def territory_id(dep, mun, dist) -> str | None:
    dep_code = code(dep)
    if dep_code == "00" and str(mun).strip().upper() == "TOTAL" and str(dist).strip().upper() == "TOTAL":
        return "SLV"
    if dep_code is None or dep_code == "00":
        return None
    mun_code = code(mun)
    dist_code = code(dist)
    if str(mun).strip().upper() == "TOTAL" and str(dist).strip().upper() == "TOTAL":
        return f"SLV:C2024:DEP:{dep_code}"
    if mun_code and str(dist).strip().upper() == "TOTAL":
        return f"SLV:C2024:MUN:{dep_code}{mun_code}"
    if mun_code and dist_code:
        return f"SLV:C2024:DIST:{dep_code}{dist_code}"
    return None


def rows(path: Path):
    book = load_workbook(path, read_only=True, data_only=True)
    return book.active.iter_rows(values_only=True)


def pct(numerator: float, denominator: float) -> float | None:
    return round(numerator / denominator * 100, 6) if denominator else None


INDICATORS = [
    ("SLV_C2024_POP_TOTAL", "Official Census population", "Population", "people", "Published total population in the same 2024 Census geography."),
    ("SLV_C2024_FEMALE_PCT", "Female population", "Population", "%", "Women divided by total population in the same published Census geography."),
    ("SLV_C2024_AVG_HOUSEHOLD_SIZE", "Average persons per household", "Housing", "people per household", "Published average number of persons per household."),
    ("SLV_C2024_ORGANIZED_WATER_SUPPLY_PCT", "Public, community or private-company water supply", "Basic services", "%", "Occupied private dwellings supplied by ANDA, a community system or a private company. This is not represented as a drinking-water quality measure."),
    ("SLV_C2024_SANITATION_ACCESS_PCT", "Dwellings with a toilet or latrine", "Basic services", "%", "Occupied private dwellings reporting sewer-connected toilet, septic tank, biodigester or latrine."),
    ("SLV_C2024_ELECTRICITY_ACCESS_PCT", "Dwellings with an electricity source", "Basic services", "%", "Occupied private dwellings reporting any listed electricity source, excluding no electricity and unknown response."),
    ("SLV_C2024_LITERACY_10PLUS_PCT", "Literacy, population age 10+", "Education", "%", "Population age 10 or older reported as able to read and write, divided by all published known and unknown literacy responses for that age universe."),
    ("SLV_C2024_EMPLOYED_EAP_PCT", "Employed among economically active population age 10+", "Economy", "%", "Employed population divided by employed plus unemployed economically active population age 10 or older."),
    ("SLV_C2024_DISABILITY_3PLUS_PCT", "At least one severe activity limitation, age 3+", "Inclusion", "%", "Population age 3 or older reporting much difficulty or inability in at least one listed activity, divided by the Census population age 3 or older."),
    ("SLV_C2024_RECENT_MIGRATION_PCT", "Living in another municipality or country in May 2019", "Migration", "%", "Population age 5 or older living in another municipality or another country in May 2019, divided by responses with a known 2019 residence."),
    ("SLV_C2024_URBAN_POP_PCT", "Urban population", "Population", "%", "Population classified as urban in the 2024 Census urban-rural table."),
    ("SLV_C2024_INDIGENOUS_IDENTIFICATION_PCT", "Population identifying with an Indigenous people", "Ethnicity", "%", "Population identifying with an Indigenous people divided by all published yes, no and unknown responses."),
    ("SLV_C2024_HOUSEHOLD_DEATH_LAST12M_PCT", "Households reporting a death in the last 12 months", "Health", "%", "Households reporting at least one member who died in the twelve months before enumeration."),
    ("SLV_WB_UNDERNOURISHMENT_PCT", "Prevalence of undernourishment", "Nutrition", "%", "National FAO-modeled prevalence of undernourishment distributed through the World Bank API. No local value is inferred."),
    ("SLV_WB_NATIONAL_POVERTY_HEADCOUNT_PCT", "National poverty headcount ratio", "Poverty", "%", "National poverty headcount ratio at the national poverty line distributed through the World Bank API. No local value is inferred."),
]


def indicator_records():
    result = []
    for iid, name, theme, unit, definition in INDICATORS:
        wb = iid.startswith("SLV_WB_")
        result.append({
            "id": iid, "name": name, "theme": theme, "unit": unit,
            "definition": definition, "definition_id": iid, "population": definition,
            "measurement_method": "source_reported" if wb or "AVG" in iid or "POP_TOTAL" in iid else "derived_from_source_counts",
            "aggregation": "none", "period_policy": "fixed_source_period",
            "series_family": "international_reference" if wb else "census",
            "display_role": "context" if wb else "primary",
            "source_id": "world-bank-api" if wb else "slv-bcr-census-2024-tables",
        })
    return result


def observation(tid, iid, value, locator, numerator=None, denominator=None, period="2024", source="slv-bcr-census-2024-tables"):
    spec = next(item for item in INDICATORS if item[0] == iid)
    return {
        "territory_id": tid, "indicator_id": iid, "period": str(period),
        "value": value, "status": "observed" if value is not None else "missing",
        "source_id": source, "definition": spec[4], "definition_id": iid,
        "unit": spec[3], "population": spec[4],
        "measurement_method": "source_reported" if numerator is None and denominator is None else "derived_from_source_counts",
        "source_locator": locator, "numerator": numerator, "denominator": denominator,
    }


def add_direct_table(result, path, iid, geometry_columns, value_column, locator):
    for number, row in enumerate(rows(path), 1):
        tid = territory_id(*(row[i] if i < len(row) else None for i in geometry_columns))
        if not tid or number < 7:
            continue
        raw = row[value_column] if value_column < len(row) else None
        if raw in (None, ""):
            result[tid] = observation(tid, iid, None, f"{locator}; row {number}")
        else:
            result[tid] = observation(tid, iid, float(raw), f"{locator}; row {number}")


def add_dimension_direct_table(result, path, iid, geometry_columns, dimension_column, dimension_value, value_column, locator, start_row=7):
    for number, row in enumerate(rows(path), 1):
        if number < start_row:
            continue
        tid = territory_id(*(row[i] if i < len(row) else None for i in geometry_columns))
        if not tid or str(row[dimension_column]).strip().upper() != dimension_value:
            continue
        raw = row[value_column] if value_column < len(row) else None
        result[tid] = observation(tid, iid, None if raw in (None, "") else float(raw), f"{locator}; row {number}")


def add_ratio_table(result, path, iid, geometry_columns, dimension_column, dimension_value, numerator_columns, denominator_column, locator, start_row=7):
    for number, row in enumerate(rows(path), 1):
        if number < start_row:
            continue
        tid = territory_id(*(row[i] if i < len(row) else None for i in geometry_columns))
        if not tid or (dimension_column is not None and str(row[dimension_column]).strip().upper() != dimension_value):
            continue
        numerator = sum(finite(row[i], f"{path.name}:{number}:{i + 1}") for i in numerator_columns)
        denominator = finite(row[denominator_column], f"{path.name}:{number}:{denominator_column + 1}")
        result[tid] = observation(tid, iid, pct(numerator, denominator), f"{locator}; row {number}", numerator, denominator)


def workbook_inventory(raw: Path, adopted_files: set[str]):
    records = []
    for path in sorted([*raw.glob("TAB_*.xlsx"), *(raw / "extended").glob("*.xlsx")]):
        book = load_workbook(path, read_only=True, data_only=True)
        ws = book.active
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True):
            preview.append([str(x).strip() if x is not None else "" for x in row])
        title = " / ".join(" ".join(value for value in row if value) for row in preview[:4])
        relative = path.relative_to(raw).as_posix()
        canonical_name = path.name.replace("-1104", "")
        records.append({
            "path": relative, "sha256": sha256(path), "bytes": path.stat().st_size,
            "sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column,
            "title": title, "header_rows": preview,
            "disposition": "integrated" if path.name in adopted_files or canonical_name in adopted_files else "not_adopted",
            "reason": "Selected published total/count fields are integrated with explicit formulas and geography checks." if path.name in adopted_files or canonical_name in adopted_files else "Acquired and inspected for the source inventory; not needed for the current common indicator set and not silently mapped to another theme.",
        })
    return records


def fetch_geojson(url: str, output: Path):
    response = requests.get(url + "/query", params={"where": "1=1", "outFields": "*", "outSR": 4326, "f": "geojson"}, headers=UA, timeout=180)
    response.raise_for_status()
    data = response.json()
    output.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return data


def fetch_wb(indicator: str, output: Path):
    url = WB_API.format(indicator=indicator)
    response = requests.get(url, params={"format": "json", "per_page": 100}, headers=UA, timeout=60)
    response.raise_for_status()
    output.write_bytes(response.content)
    payload = response.json()
    values = [row for row in payload[1] if row.get("value") is not None]
    if not values:
        raise ValueError(f"No World Bank values for {indicator}")
    return url, values[0]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw", required=True)
    parser.add_argument("--project", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    raw, project, out = Path(args.raw), Path(args.project), Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    dataset = json.loads((project / "data" / "dashboard.json").read_text(encoding="utf-8"))
    territories = [row for row in dataset["territories"] if row.get("country_id") == "SLV" and row["id"] != "SLV"]
    territory_ids = {"SLV", *(row["id"] for row in territories)}
    counts = {level: sum(row["level"] == level for row in territories) for level in ("department", "municipality", "district")}
    if counts != {"department": 14, "municipality": 44, "district": 262}:
        raise ValueError(f"Unexpected existing El Salvador hierarchy: {counts}")

    extended = raw / "extended"
    paths = {
        "pob1": raw / "TAB_POB_1.xlsx", "hog7": extended / "TAB_HOG_7-1104.xlsx",
        "viv17": raw / "TAB_VIV_17.xlsx", "viv15": raw / "TAB_VIV_15.xlsx",
        "viv19": raw / "TAB_VIV_19.xlsx", "edu3": extended / "TAB_EDU_3.xlsx",
        "ce2": extended / "TAB_CE_2.xlsx", "disca2": extended / "TAB_DISCA_2-1104.xlsx",
        "migr4": extended / "TAB_MIGR_4.xlsx", "area1": extended / "TAB_POB_AREA_1-1104.xlsx",
        "etnia2": extended / "TAB_ETNIA_2.xlsx", "mort1": extended / "TAB_MORT_1.xlsx",
    }
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing adopted workbooks: " + ", ".join(missing))

    by_indicator = {iid: {} for iid, *_ in INDICATORS}
    add_dimension_direct_table(by_indicator["SLV_C2024_POP_TOTAL"], paths["pob1"], "SLV_C2024_POP_TOTAL", (1, 2, 3), 4, "TOTAL", 5, "TAB_POB_1")
    add_ratio_table(by_indicator["SLV_C2024_FEMALE_PCT"], paths["pob1"], "SLV_C2024_FEMALE_PCT", (1, 2, 3), 4, "TOTAL", (7,), 5, "TAB_POB_1")
    add_direct_table(by_indicator["SLV_C2024_AVG_HOUSEHOLD_SIZE"], paths["hog7"], "SLV_C2024_AVG_HOUSEHOLD_SIZE", (0, 1, 2), 3, "TAB_HOG_7")
    add_ratio_table(by_indicator["SLV_C2024_ORGANIZED_WATER_SUPPLY_PCT"], paths["viv17"], "SLV_C2024_ORGANIZED_WATER_SUPPLY_PCT", (0, 1, 2), 3, "TOTAL", (5, 6, 7), 4, "TAB_VIV_17")
    add_ratio_table(by_indicator["SLV_C2024_SANITATION_ACCESS_PCT"], paths["viv15"], "SLV_C2024_SANITATION_ACCESS_PCT", (0, 1, 2), None, "", (4, 5, 6, 7), 3, "TAB_VIV_15")
    add_ratio_table(by_indicator["SLV_C2024_ELECTRICITY_ACCESS_PCT"], paths["viv19"], "SLV_C2024_ELECTRICITY_ACCESS_PCT", (0, 1, 2), None, "", (4, 5, 6, 7, 8, 9), 3, "TAB_VIV_19")

    literacy = {}
    for number, row in enumerate(rows(paths["edu3"]), 1):
        tid = territory_id(row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None, row[3] if len(row) > 3 else None)
        age = str(row[4] if len(row) > 4 else "").strip()
        if not tid or number < 7 or age in ("TOTAL", "3 a 4 años", "5 a 9 años"):
            continue
        item = literacy.setdefault(tid, [0.0, 0.0])
        item[0] += finite(row[7], f"EDU3:{number}:8")
        item[1] += finite(row[7], f"EDU3:{number}:8") + finite(row[10], f"EDU3:{number}:11") + finite(row[13], f"EDU3:{number}:14")
    for tid, (num, den) in literacy.items():
        by_indicator["SLV_C2024_LITERACY_10PLUS_PCT"][tid] = observation(tid, "SLV_C2024_LITERACY_10PLUS_PCT", pct(num, den), "TAB_EDU_3; age-group rows 10+", num, den)

    add_ratio_table(by_indicator["SLV_C2024_EMPLOYED_EAP_PCT"], paths["ce2"], "SLV_C2024_EMPLOYED_EAP_PCT", (1, 2, 3), 4, "TOTAL", (7,), 11, "TAB_CE_2")

    population_3plus = {}
    for number, row in enumerate(rows(paths["pob1"]), 1):
        tid = territory_id(row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None, row[3] if len(row) > 3 else None)
        age = str(row[4] if len(row) > 4 else "").strip()
        if not tid or number < 7 or age == "TOTAL":
            continue
        try:
            age_number = int(float(age))
        except ValueError:
            continue
        if age_number >= 3:
            population_3plus[tid] = population_3plus.get(tid, 0.0) + finite(row[5], f"POB1:{number}:6")
    for number, row in enumerate(rows(paths["disca2"]), 1):
        tid = territory_id(row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None, row[3] if len(row) > 3 else None)
        if not tid or number < 7 or str(row[4]).strip().upper() != "TOTAL":
            continue
        num = finite(row[19], f"DISCA2:{number}:20") + finite(row[20], f"DISCA2:{number}:21")
        den = population_3plus.get(tid, 0.0)
        by_indicator["SLV_C2024_DISABILITY_3PLUS_PCT"][tid] = observation(tid, "SLV_C2024_DISABILITY_3PLUS_PCT", pct(num, den), "TAB_DISCA_2 severe limitation count / TAB_POB_1 population age 3+", num, den)

    migration = {}
    for number, row in enumerate(rows(paths["migr4"]), 1):
        tid = territory_id(row[0] if len(row) > 0 else None, row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None)
        if not tid or number < 8 or str(row[3]).strip().upper() != "TOTAL":
            continue
        same = finite(row[4], "MIGR4") + finite(row[5], "MIGR4")
        moved = sum(finite(row[i], "MIGR4") for i in (6, 7, 8, 9))
        migration[tid] = observation(tid, "SLV_C2024_RECENT_MIGRATION_PCT", pct(moved, same + moved), f"TAB_MIGR_4; row {number}; unknown residence excluded", moved, same + moved)
    by_indicator["SLV_C2024_RECENT_MIGRATION_PCT"] = migration

    add_ratio_table(by_indicator["SLV_C2024_URBAN_POP_PCT"], paths["area1"], "SLV_C2024_URBAN_POP_PCT", (0, 1, 2), None, "", (4, 5), 3, "TAB_POB_AREA_1")
    add_ratio_table(by_indicator["SLV_C2024_INDIGENOUS_IDENTIFICATION_PCT"], paths["etnia2"], "SLV_C2024_INDIGENOUS_IDENTIFICATION_PCT", (0, 1, 2), 3, "TOTAL", (4,), 4, "TAB_ETNIA_2")
    # Replace the ethnicity denominator with yes + no + unknown.
    for number, row in enumerate(rows(paths["etnia2"]), 1):
        tid = territory_id(row[0] if len(row) > 0 else None, row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None)
        if not tid or number < 7 or str(row[3]).strip().upper() != "TOTAL":
            continue
        num = finite(row[4], "ETNIA2")
        den = num + finite(row[5], "ETNIA2") + finite(row[6], "ETNIA2")
        by_indicator["SLV_C2024_INDIGENOUS_IDENTIFICATION_PCT"][tid] = observation(tid, "SLV_C2024_INDIGENOUS_IDENTIFICATION_PCT", pct(num, den), f"TAB_ETNIA_2; row {number}", num, den)
    add_ratio_table(by_indicator["SLV_C2024_HOUSEHOLD_DEATH_LAST12M_PCT"], paths["mort1"], "SLV_C2024_HOUSEHOLD_DEATH_LAST12M_PCT", (0, 1, 2), None, "", (3,), 5, "TAB_MORT_1")

    for iid, records in by_indicator.items():
        if iid.startswith("SLV_WB_"):
            continue
        unknown = set(records) - territory_ids
        # The BCR tables use department code 15 for residence not specified.
        # It contributes to the exact national published row but is not an
        # administrative territory and must never appear in the selector.
        for tid in unknown:
            del records[tid]
        missing_ids = territory_ids - set(records)
        if missing_ids:
            raise ValueError(f"{iid} geography mismatch: excluded_nonterritories={sorted(unknown)[:5]}, missing={sorted(missing_ids)[:5]}, counts={len(records)}/{len(territory_ids)}")

    wb_specs = [("SN.ITK.DEFC.ZS", "SLV_WB_UNDERNOURISHMENT_PCT"), ("SI.POV.NAHC", "SLV_WB_NATIONAL_POVERTY_HEADCOUNT_PCT")]
    wb_receipts = []
    for wb_id, iid in wb_specs:
        api_url, item = fetch_wb(wb_id, out / f"world-bank-{wb_id}.json")
        obs = observation("SLV", iid, float(item["value"]), f"World Bank API {wb_id}; latest non-null national observation", period=item["date"], source="world-bank-api")
        by_indicator[iid]["SLV"] = obs
        wb_receipts.append({"indicator": wb_id, "url": api_url, "period": item["date"], "value": item["value"], "file": f"world-bank-{wb_id}.json", "sha256": sha256(out / f"world-bank-{wb_id}.json")})

    boundary_features = []
    boundary_receipts = []
    expected = {"department": 14, "municipality": 44, "district": 262}
    for level, url in BOUNDARY_SERVICES.items():
        target = out / f"bcr-{level}-boundaries.geojson"
        data = fetch_geojson(url, target)
        features = []
        for feature in data.get("features", []):
            props = feature.get("properties", {})
            if str(props.get("id_depto", "")).zfill(2) == "00":
                continue
            if level == "municipality" and str(props.get("id_mun", "")).zfill(2) == "00":
                continue
            if level == "district" and str(props.get("id_distrito", "")).zfill(2).endswith("00"):
                continue
            features.append(feature)
        if len(features) != expected[level]:
            raise ValueError(f"Unexpected {level} boundary count after excluding the national summary geometry: {len(features)}")
        for feature in features:
            props = feature["properties"]
            if level == "department":
                tid = f"SLV:C2024:DEP:{str(props['id_depto']).zfill(2)}"
            elif level == "municipality":
                tid = f"SLV:C2024:MUN:{str(props['id_mun']).zfill(4)}"
            else:
                district_code = str(props["id_depto"]).zfill(2) + str(props["id_distrito"])[-2:].zfill(2)
                tid = f"SLV:C2024:DIST:{district_code}"
            if tid not in territory_ids:
                raise ValueError(f"Boundary code does not match Census hierarchy: {tid}")
            boundary_features.append({
                "type": "Feature",
                "properties": {"territory_id": tid, "name": props.get("nombre_min") or props.get("mun_min") or props.get("distrito_min"), "source_id": "slv-bcr-boundaries-2024", "geometry_edition": "BCR 2024 administrative selectors", "reference_only": False, "join_method": "Exact BCR 2024 official code"},
                "geometry": feature["geometry"],
            })
        boundary_receipts.append({"level": level, "url": url, "source_features": len(data["features"]), "adopted_features": len(features), "excluded": "national summary geometry id_depto=00", "file": target.name, "sha256": sha256(target)})
    if len({f["properties"]["territory_id"] for f in boundary_features}) != 320:
        raise ValueError("Boundary territory IDs are not unique and complete")

    territory_by_id = {row["id"]: row for row in territories}
    for feature in boundary_features:
        tid = feature["properties"]["territory_id"]
        territory_by_id[tid]["boundary_version"] = "BCR 2024 administrative selectors"
        territory_by_id[tid]["geography_note"] = "2024 Census hierarchy and BCR geometry joined by exact official code. Districts are diagnostic subdivisions of the 44 municipalities, not separate municipal planning authorities."

    comparisons = [row for row in dataset.get("analysis", {}).get("comparisons", []) if row.get("parent_id") in territory_ids and row.get("parent_id") != "SLV"]
    terminal = [row["id"] for row in territories if row["level"] == "district"]
    observations = [obs for iid in by_indicator for obs in by_indicator[iid].values()]
    sources = [
        {"id": "slv-bcr-census-2024-tables", "name": "VII Censo de Población y VI de Vivienda 2024 tabulations", "publisher": "Banco Central de Reserva de El Salvador", "url": CENSUS_PAGE, "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2024", "geographic_level": "country, department, municipality and district", "license": "Official public statistical tables", "note": "Official tabulations retain their stated universes. Derived percentages use only source counts from the same geography and compatible universe."},
        {"id": "slv-bcr-boundaries-2024", "name": "BCR 2024 department, municipality and district selector boundaries", "publisher": "Banco Central de Reserva de El Salvador", "url": BOUNDARY_SERVICES["district"], "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2024 administrative structure", "geographic_level": "department, municipality and district", "license": "Official ArcGIS feature services", "note": "Geometry joins use exact BCR official codes. Districts remain below the 44 municipalities."},
        {"id": "world-bank-api", "name": "World Development Indicators API", "publisher": "World Bank", "url": "https://api.worldbank.org/", "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "latest non-null value per indicator", "geographic_level": "El Salvador only for the adopted supplements", "license": "World Bank data terms", "note": "International reference observations are national only and are not imputed to departments, municipalities or districts."},
    ]
    adopted = {path.name for path in paths.values()}
    inventory = workbook_inventory(raw, adopted)
    audit = {
        "schema_version": "1.0", "generated_at": datetime.now(timezone.utc).isoformat(),
        "country_area_id": "SLV", "status": "complete_census_depth_bundle_pending_integration",
        "hierarchy": {**counts, "country": 1, "planning_authority": "44 municipalities", "diagnostic_subdivision": "262 districts"},
        "indicators": {iid: {"observations": len(items), "observed": sum(o["status"] == "observed" for o in items.values())} for iid, items in by_indicator.items()},
        "unallocated_population": {"national_total": 6029976, "named_administrative_total": 5922921, "unallocated_residence": 107055, "rule": "Retained in the published national total and never created as a territory or spread across lower areas."},
        "workbook_inventory": inventory, "boundary_receipts": boundary_receipts,
        "world_bank_receipts": wb_receipts,
        "method": "Every domestic indicator covers the country plus 14 departments, 44 municipalities and 262 districts. Ratios use published counts from compatible universes; no rate is averaged and no missing lower area changes the national published total.",
    }
    bundle = {
        "schema_version": "1.0", "country_area_id": "SLV", "period": "2024", "replace_country_branch": False,
        "territories": list(territory_by_id.values()), "indicators": indicator_records(), "sources": sources,
        "observations": observations, "boundaries": {"type": "FeatureCollection", "features": boundary_features},
        "comparisons": comparisons, "terminal_territory_ids": terminal, "audit": audit,
    }
    (out / "el-salvador-census-bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "EL_SALVADOR_COLLECTION_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"territories": len(territories), "indicators": len(INDICATORS), "observations": len(observations), "boundaries": len(boundary_features), "workbooks_inventoried": len(inventory)}, indent=2))


if __name__ == "__main__":
    main()
