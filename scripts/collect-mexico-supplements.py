#!/usr/bin/env python3
"""Collect Mexico municipal poverty and food-access indicators from CONEVAL."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


SOURCE_URL = "https://www.coneval.org.mx/Medicion/Documents/Pobreza_municipal/2020/Concentrado_indicadores_de_pobreza_2020.zip"
SOURCE_PAGE = "https://www.coneval.org.mx/Medicion/Paginas/Pobreza-municipio-2010-2020.aspx"
WORKBOOK_NAME = "Concentrado_indicadores_de_pobreza_2020.xlsx"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def numeric_rows(sheet, key_index: int):
    rows = []
    for row in sheet.iter_rows(min_row=9, values_only=True):
        key = row[key_index]
        if key not in (None, "") and str(key).strip().isdigit():
            rows.append(row)
    return rows


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def observation(territory_id: str, indicator_id: str, value, locator: str, *, calculated=False):
    numeric = number(value)
    return {
        "territory_id": territory_id,
        "indicator_id": indicator_id,
        "period": "2020",
        "value": numeric,
        "status": "observed" if numeric is not None else "missing",
        "source_id": "mexico-coneval-municipal-poverty-2020",
        "definition_id": indicator_id,
        "unit": "%",
        "population": "CONEVAL source-defined estimated population for municipal poverty measurement",
        "measurement_method": "weighted_from_complete_state_source_counts" if calculated else "source_reported_estimate",
        "source_locator": locator,
        **({"provenance": "calculated"} if calculated else {}),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    raw, out = Path(args.raw).resolve(), Path(args.out).resolve()
    source_zip = raw / "coneval-poverty-2010-2020.zip"
    with zipfile.ZipFile(source_zip) as archive:
        workbook_bytes = archive.read(WORKBOOK_NAME)
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    municipality_rows = numeric_rows(workbook["Concentrado municipal"], 1)
    state_rows = numeric_rows(workbook["Concentrado estatal"], 1)
    if len(municipality_rows) != 2469 or len(state_rows) != 32:
        raise RuntimeError(f"Unexpected CONEVAL geography counts: municipalities={len(municipality_rows)}, states={len(state_rows)}")
    municipality_codes = {str(row[3]).zfill(5) for row in municipality_rows}
    if len(municipality_codes) != 2469:
        raise RuntimeError("CONEVAL municipality codes are not unique.")

    poverty_id = "MEX_CONEVAL_POVERTY_PCT_2020"
    food_id = "MEX_CONEVAL_FOOD_ACCESS_DEPRIVATION_PCT_2020"
    observations = []
    for row in municipality_rows:
        code = str(row[3]).zfill(5)
        observations.append(observation(f"MEX:MUN:{code}", poverty_id, row[10], f"Concentrado municipal; {code}; Pobreza; Porcentaje 2020"))
        observations.append(observation(f"MEX:MUN:{code}", food_id, row[103], f"Concentrado municipal; {code}; Carencia por acceso a la alimentación; Porcentaje 2020"))
    for row in state_rows:
        code = str(row[1]).zfill(2)
        observations.append(observation(f"MEX:ENT:{code}", poverty_id, row[8], f"Concentrado estatal; {code}; Pobreza; Porcentaje 2020"))
        observations.append(observation(f"MEX:ENT:{code}", food_id, row[101], f"Concentrado estatal; {code}; Carencia por acceso a la alimentación; Porcentaje 2020"))

    state_population = sum(float(row[5]) for row in state_rows)
    poverty_people = sum(float(row[11]) for row in state_rows)
    food_people = sum(float(row[104]) for row in state_rows)
    observations.append(observation("MEX", poverty_id, 100 * poverty_people / state_population, "AreaData calculation from all 32 CONEVAL state population and poverty-person counts", calculated=True))
    observations.append(observation("MEX", food_id, 100 * food_people / state_population, "AreaData calculation from all 32 CONEVAL state population and food-access-deprivation person counts", calculated=True))

    common = {
        "unit": "%",
        "population": "CONEVAL source-defined estimated population for municipal poverty measurement",
        "aggregation": "none",
        "period_policy": "fixed_source_period",
        "series_family": "administrative",
        "display_role": "primary",
        "source_id": "mexico-coneval-municipal-poverty-2020",
    }
    indicators = [
        {
            "id": poverty_id,
            "name": "Population in multidimensional poverty",
            "theme": "Poverty",
            "definition": "Percentage of the population in multidimensional poverty under CONEVAL's municipal poverty measurement methodology.",
            "definition_id": poverty_id,
            "measurement_method": "source_reported_estimate; national value weighted from complete official state numerator and denominator counts",
            "source_locator": "Concentrado municipal/estatal, Pobreza, Porcentaje 2020",
            **common,
        },
        {
            "id": food_id,
            "name": "Food-access deprivation",
            "theme": "Health and nutrition",
            "definition": "Percentage of the population with deprivation due to lack of access to food, as defined in CONEVAL's 2020 municipal poverty measurement.",
            "definition_id": food_id,
            "measurement_method": "source_reported_estimate; national value weighted from complete official state numerator and denominator counts",
            "source_locator": "Concentrado municipal/estatal, Carencia por acceso a la alimentación, Porcentaje 2020",
            **common,
        },
    ]
    source = {
        "id": "mexico-coneval-municipal-poverty-2020",
        "name": "Medición de la pobreza municipal 2010–2020 — statistical annex",
        "publisher": "Consejo Nacional de Evaluación de la Política de Desarrollo Social (CONEVAL)",
        "url": SOURCE_PAGE,
        "download_url": SOURCE_URL,
        "status": "ready",
        "retrieved_at": datetime.now(timezone.utc).date().isoformat(),
        "reference_period": "2020",
        "geographic_level": "Mexico, federal entities, municipalities and territorial demarcations",
        "license": "Official public statistical annex; source terms apply",
        "raw_path": "raw/mexico-required-themes-planning/coneval-poverty-2010-2020.zip",
        "sha256": sha256(source_zip),
        "note": "Published percentages are retained for every state and municipality. National percentages are calculated only from the complete 32-state official population and person counts; municipal percentages are never averaged.",
    }
    audit = {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "country_area_id": "MEX",
        "status": "official_coneval_poverty_and_food_access_integrated",
        "counts": {
            "states": len(state_rows),
            "municipalities": len(municipality_rows),
            "indicators": len(indicators),
            "observations": len(observations),
            "observations_per_indicator": 1 + len(state_rows) + len(municipality_rows),
        },
        "national_calculation": {
            "state_population_denominator": state_population,
            "poverty_people_numerator": poverty_people,
            "food_access_deprivation_people_numerator": food_people,
            "complete_non_overlapping_state_cover": True,
        },
        "source_receipt": {"path": source_zip.name, "sha256": sha256(source_zip), "bytes": source_zip.stat().st_size},
    }
    bundle = {"schema_version": "1.0", "country_area_id": "MEX", "sources": [source], "indicators": indicators, "observations": observations, "audit": audit}
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    out.with_name("mexico-supplements-audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
