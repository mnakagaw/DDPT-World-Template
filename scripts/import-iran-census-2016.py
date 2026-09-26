"""Import the SCI-authored 2016 census workbooks distributed by Iran Data Portal.

This importer deliberately does not invent Iranian administrative codes or attach
unreconciled reference geometry. It preserves the exact source row/cell for each
adopted observation and fails if the expected source bytes or census totals drift.
"""

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


SOURCES = {
    "population": {
        "file": "Population-and-Households.xlsx",
        "sha256": "beaa06950520e65892959c6dabe9013284907b4a49fc9ac89485c8800b0f60b3",
        "url": "https://irandataportal.syr.edu/wp-content/uploads/Population-and-Households.xlsx",
        "sheet": "Population, household, Province",
        "id": "irn-sci-census-2016-population-copy",
    },
    "age_sex": {
        "file": "Population-by-Age-and-Sex-provincial-level.xlsx",
        "sha256": "b3154c61e73c58ff61b3452644a49d9d2534dfc7ce11a4eb3a2773bdacc07556",
        "url": "https://irandataportal.syr.edu/wp-content/uploads/Population-by-Age-and-Sex-provincial-level.xlsx",
        "sheet": "Age group",
        "id": "irn-sci-census-2016-age-sex-copy",
    },
}

BASE_FIELDS = (("POP_TOTAL", "Population", "people", 1, "B"),
               ("POP_MALE", "Male population", "people", 2, "C"),
               ("POP_FEMALE", "Female population", "people", 3, "D"),
               ("HOUSEHOLDS", "Households", "households", 4, "E"))
SETTLEMENTS = (("URBAN", "Urban", 1), ("RURAL", "Rural", 2),
               ("UNSETTLED", "Unsettled", 3))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_sheet(project, key):
    spec = SOURCES[key]
    path = project / "raw" / spec["file"]
    if digest(path) != spec["sha256"]:
        raise ValueError(f"Source hash mismatch: {path}")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if book.sheetnames != [spec["sheet"]]:
        raise ValueError(f"Unexpected sheets: {path}: {book.sheetnames}")
    return list(book.active.iter_rows(values_only=True))


def exact_int(value, label):
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"Expected nonnegative integer at {label}: {value!r}")
    return value


def observation(territory_id, indicator_id, value, source_id, cell):
    return {
        "territory_id": territory_id,
        "indicator_id": indicator_id,
        "period": "2016",
        "value": value,
        "status": "observed",
        "measurement_method": "source_reported",
        "source_id": source_id,
        "source_locator": cell,
    }


def indicator(indicator_id, name, unit, source_id, population):
    return {
        "id": indicator_id,
        "name": name,
        "theme": "Census",
        "unit": unit,
        "definition": f"2016 Population and Housing Census count: {name.lower()}.",
        "population": population,
        "source_id": source_id,
        "aggregation": "sum",
        "measurement_method": "source_reported",
        "series_family": "administrative",
        "period_policy": "latest_available_per_indicator",
        "display_decimals": 0,
        "metadata_status": "ready",
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve()
    pop = load_sheet(project, "population")
    age = load_sheet(project, "age_sex")
    if len(pop) != 1846 or len(age) != 35:
        raise ValueError(f"Source row count changed: {len(pop)}, {len(age)}")
    if tuple(pop[1][:5]) != ("Description", "Population", "Male", "Female", "Household"):
        raise ValueError("Population header changed")
    if tuple(age[2][1:4]) != ("Population", "Male", "Female"):
        raise ValueError("Age-sex header changed")

    age_rows = [(i, row) for i, row in enumerate(age[4:], 5)]
    if len(age_rows) != 31 or len({row[1] for _, row in age_rows}) != 31:
        raise ValueError("Expected 31 provinces with unique populations in age-sex sheet")
    age_by_population = {row[1]: (i, row) for i, row in age_rows}
    blocks = [(i, pop[i - 1:i + 3]) for i in range(3, len(pop) + 1, 4)]
    if len(blocks) != 461 or blocks[-1][0] != 1843:
        raise ValueError("Expected country + 31 province + 429 county four-row blocks")

    national_id = "IRN"
    territories = []
    observations = []
    provinces = []
    counties_by_province = []
    source_label_conflicts = []
    current_province = None
    seen_province_values = set()
    for block_index, (row_number, block) in enumerate(blocks):
        total = block[0]
        if not isinstance(total[0], str):
            raise ValueError(f"Missing territory name at row {row_number}")
        values = [exact_int(total[col], f"B{row_number}+{col}") for col in range(1, 5)]
        if values[0] != values[1] + values[2]:
            raise ValueError(f"Male/female total mismatch at row {row_number}")
        for offset, (key, label, _) in enumerate(SETTLEMENTS, 1):
            if str(block[offset][0]).strip().lower() not in {
                "setteled in urban areas", "settled in rural areas", "unsettled"
            }:
                raise ValueError(f"Unexpected settlement label at row {row_number + offset}")
            if any(not isinstance(block[offset][col], int) for col in range(1, 5)):
                raise ValueError(f"Noninteger settlement values at row {row_number + offset}")
        for col in range(1, 5):
            if total[col] != sum(block[offset][col] for offset in (1, 2, 3)):
                raise ValueError(f"Settlement subtotal mismatch at row {row_number}, column {col}")

        if block_index == 0:
            if total[0] != "Total country":
                raise ValueError("First block is not the national census row")
            territory_id = national_id
        elif values[0] in age_by_population and values[0] not in seen_province_values:
            seen_province_values.add(values[0])
            age_row_number, age_row = age_by_population[values[0]]
            source_name = total[0].strip()
            canonical_name = source_name
            if values[0] == 768898:
                canonical_name = "South Khorasan"
            if canonical_name != age_row[0].strip():
                source_label_conflicts.append({
                    "population_sheet_row": row_number,
                    "population_sheet_label": source_name,
                    "age_sheet_row": age_row_number,
                    "age_sheet_label": age_row[0].strip(),
                    "population": values[0],
                    "adopted_label": canonical_name,
                })
            territory_id = f"IRN-SCI2016-OSTAN-R{row_number}"
            current_province = territory_id
            provinces.append((row_number, territory_id, values))
            counties_by_province.append([])
            territories.append({
                "id": territory_id, "name": canonical_name,
                "level": "adm1", "type": "ostan", "parent_id": national_id,
                "official_code": None, "code_system": "SCI 2016 source row (provisional ID)",
                "boundary_version": None, "source_id": SOURCES["population"]["id"],
                "source_locator": f"{SOURCES['population']['sheet']}!A{row_number}:E{row_number}",
            })
        else:
            if current_province is None:
                raise ValueError(f"County before province at row {row_number}")
            territory_id = f"IRN-SCI2016-SHAHRESTAN-R{row_number}"
            counties_by_province[-1].append((row_number, territory_id, values))
            territories.append({
                "id": territory_id, "name": total[0].strip(),
                "level": "adm2", "type": "shahrestan", "parent_id": current_province,
                "official_code": None, "code_system": "SCI 2016 source row (provisional ID)",
                "boundary_version": None, "source_id": SOURCES["population"]["id"],
                "source_locator": f"{SOURCES['population']['sheet']}!A{row_number}:E{row_number}",
            })
        for suffix, _, _, col, letter in BASE_FIELDS:
            observations.append(observation(territory_id, f"IRN_CENSUS_2016_{suffix}",
                                            total[col], SOURCES["population"]["id"],
                                            f"{SOURCES['population']['sheet']}!{letter}{row_number}"))
        for settlement_key, _, offset in SETTLEMENTS:
            observations.append(observation(territory_id,
                                            f"IRN_CENSUS_2016_POP_{settlement_key}",
                                            block[offset][1], SOURCES["population"]["id"],
                                            f"{SOURCES['population']['sheet']}!B{row_number + offset}"))

    if len(provinces) != 31 or sum(map(len, counties_by_province)) != 429:
        raise ValueError("Province/county count changed")
    national_values = [blocks[0][1][0][col] for col in range(1, 5)]
    for index, national_value in enumerate(national_values):
        if sum(item[2][index] for item in provinces) != national_value:
            raise ValueError(f"Province national subtotal mismatch at column {index + 2}")
        for province, counties in zip(provinces, counties_by_province):
            if sum(item[2][index] for item in counties) != province[2][index]:
                raise ValueError(f"County subtotal mismatch in {province[1]}")

    territory_by_population = {values[0]: territory_id for _, territory_id, values in provinces}
    age_group_labels = list(age[2][4:25])
    if age_group_labels != list(age[2][25:46]) or len(age_group_labels) != 21:
        raise ValueError("Age-group headers changed")
    for row_number, row in [(4, age[3]), *age_rows]:
        territory_id = national_id if row_number == 4 else territory_by_population[row[1]]
        expected_values = national_values if row_number == 4 else next(
            values for _, tid, values in provinces if tid == territory_id
        )
        if tuple(row[1:4]) != tuple(expected_values[:3]):
            raise ValueError(f"Age-sex population mismatch at row {row_number}")
        if sum(row[4:25]) != row[2] or sum(row[25:46]) != row[3]:
            raise ValueError(f"Age bands do not sum to sex total at row {row_number}")
        for index, label in enumerate(age_group_labels):
            normalized = str(label).upper().replace("-", "_").replace(" ", "_")
            for sex, column in (("MALE", 4 + index), ("FEMALE", 25 + index)):
                indicator_id = f"IRN_CENSUS_2016_AGE_{normalized}_{sex}"
                observations.append(observation(
                    territory_id, indicator_id,
                    exact_int(row[column], f"{row_number}:{column + 1}"),
                    SOURCES["age_sex"]["id"],
                    f"{SOURCES['age_sex']['sheet']}!{openpyxl.utils.get_column_letter(column + 1)}{row_number}",
                ))

    indicators = [indicator(f"IRN_CENSUS_2016_{suffix}", name, unit,
                            SOURCES["population"]["id"], "Enumerated residents" if unit == "people" else "Private and other reported households")
                  for suffix, name, unit, _, _ in BASE_FIELDS]
    indicators.extend(indicator(f"IRN_CENSUS_2016_POP_{key}", f"{name} population",
                                "people", SOURCES["population"]["id"],
                                "Enumerated residents by settlement classification")
                      for key, name, _ in SETTLEMENTS)
    for label in age_group_labels:
        normalized = str(label).upper().replace("-", "_").replace(" ", "_")
        for sex in ("MALE", "FEMALE"):
            indicators.append(indicator(f"IRN_CENSUS_2016_AGE_{normalized}_{sex}",
                                        f"{sex.title()} population age {label}", "people",
                                        SOURCES["age_sex"]["id"], "Enumerated residents"))

    dataset_path = project / "data" / "dashboard.json"
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    prefix = "IRN-SCI2016-"
    indicator_prefix = "IRN_CENSUS_2016_"
    source_ids = {spec["id"] for spec in SOURCES.values()}
    dataset["territories"] = [x for x in dataset["territories"] if not x["id"].startswith(prefix)] + territories
    dataset["indicators"] = [x for x in dataset["indicators"] if not x["id"].startswith(indicator_prefix)] + indicators
    dataset["observations"] = [x for x in dataset["observations"] if not x["indicator_id"].startswith(indicator_prefix)] + observations
    dataset["sources"] = [x for x in dataset["sources"] if x["id"] not in source_ids]
    timestamp = datetime.now(timezone.utc).isoformat()
    dataset["generated_at"] = timestamp
    for key, spec in SOURCES.items():
        path = project / "raw" / spec["file"]
        receipt = {
            "url": spec["url"], "distributor": "Iran Data Portal, Syracuse University",
            "original_publisher": "Statistical Centre of Iran",
            "recorded_at_utc": timestamp, "size_bytes": path.stat().st_size,
            "sha256": spec["sha256"], "acquisition": "Downloaded public SCI-authored workbook copy",
            "official_host_status": "2016 SCI catalogue URL not retrievable in this session",
            "redistribution_terms": "review_required",
        }
        (project / "raw" / (spec["file"] + ".receipt.json")).write_text(
            json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        dataset["sources"].append({
            "id": spec["id"], "name": f"Statistical Centre of Iran 2016 census — {spec['sheet']}",
            "url": spec["url"], "publisher": "Statistical Centre of Iran",
            "distributor": "Iran Data Portal, Syracuse University",
            "reference_period": "2016", "geographic_level": "national, ostan, shahrestan" if key == "population" else "national, ostan",
            "status": "ready", "retrieved_at": timestamp, "sha256": spec["sha256"],
            "raw_path": f"raw/{spec['file']}", "license": "terms_review_required",
            "note": "SCI-authored workbook distributed as a copy. Original SCI catalogue not reachable; source labels with conflicts and provisional row IDs are documented in the import audit. No boundary/code correspondence adopted.",
        })
    dataset["country"]["geography_note"] = (
        "2016 SCI census tables cover 31 ostans and 429 shahrestans; provisional source-row IDs "
        "retain the reporting hierarchy. Official local codes and versioned boundaries are not yet matched. "
        "Country WDI estimates and 2016 census counts remain separate series."
    )
    dataset["gaps"] = [x for x in dataset["gaps"] if x["category"] != "subnational_statistics"]
    dataset["gaps"].append({
        "category": "subnational_statistics", "status": "partial",
        "detail": "2016 SCI census population/household totals are imported for 31 ostans and 429 shahrestans; provincial sex-age table is imported. Other census domains and official local codes/boundaries remain unreviewed.",
        "next_action": "Acquire official code and boundary version, inspect all remaining census tables, and join only with verified correspondence.",
    })
    dataset["collection"]["adapters"] = sorted(set(dataset["collection"]["adapters"] + ["iran-sci-census-2016-copy"]))
    census_note = "SCI-authored 2016 census workbook copies include 31 ostans and 429 shahrestans; official codes, boundaries, and redistribution terms remain unresolved. This is an incomplete country edition."
    dataset["collection"]["notes"] = [x for x in dataset["collection"]["notes"] if x != census_note] + [census_note]
    dataset["analysis"]["terminal_territory_ids"] = sorted(x["id"] for x in territories if x["level"] == "adm2")
    dataset["analysis"]["latest_values_only"] = True
    dataset_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    evidence = project / "evidence"
    evidence.mkdir(exist_ok=True)
    audit = {
        "country": "IRN", "census_year": 2016, "status": "partial_import_not_accepted",
        "raw_sources": {k: {"url": spec["url"], "sha256": spec["sha256"], "sheet": spec["sheet"]} for k, spec in SOURCES.items()},
        "counts": {"ostans": len(provinces), "shahrestans": sum(map(len, counties_by_province)),
                   "indicators": len(indicators), "observations": len(observations)},
        "national_values": dict(zip(("population", "male", "female", "households"), national_values)),
        "reconciliation": "Every ostan equals the sum of its shahrestans in all four base columns; all ostans equal national totals; every settlement component and sex-age component balances.",
        "source_label_conflicts": source_label_conflicts,
        "unresolved": ["Official administrative codes", "2016-compatible boundaries", "Original SCI server replay", "Redistribution terms", "Planning evidence", "Other census domains"],
    }
    (evidence / "IRN_CENSUS_IMPORT_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inventory = {
        "country": "IRN", "census_year": 2016,
        "resources": [
            {"source_id": spec["id"], "url": spec["url"], "raw_path": f"raw/{spec['file']}",
             "sha256": spec["sha256"], "sheet": spec["sheet"],
             "acquisition": "downloaded_SCI_authored_copy",
             "official_host_replay": "failed_TLS_handshake_or_unavailable",
             "redistribution": "terms_review_required"}
            for spec in SOURCES.values()
        ],
        "expected_but_unresolved": [
            "Official SCI catalogue and original-host copy of census tables",
            "Official administrative code list and 2016-compatible boundary files",
            "Other 2016 census domains beyond the two acquired workbooks",
            "Planning law, guidance, plans, budgets, implementation and evaluations",
        ],
    }
    (evidence / "SOURCE_RESOURCE_INVENTORY.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = []
    for column, name in (("A", "Description"), ("B", "Population"),
                         ("C", "Male"), ("D", "Female"), ("E", "Household")):
        fields.append({"sheet": SOURCES["population"]["sheet"], "column": column,
                       "source_field": name, "row_roles": "national, ostan, shahrestan total",
                       "adoption": "identifier" if column == "A" else "adopted"})
    for row_role in ("urban", "rural", "unsettled"):
        for column, name in (("B", "Population"), ("C", "Male"),
                             ("D", "Female"), ("E", "Household")):
            fields.append({"sheet": SOURCES["population"]["sheet"], "column": column,
                           "source_field": f"{row_role} {name}", "row_roles": row_role,
                           "adoption": "adopted" if column == "B" else "not_adopted_yet",
                           "reason": None if column == "B" else "Available numeric field; review presentation and definitions before adoption"})
    fields.extend({"sheet": SOURCES["age_sex"]["sheet"],
                   "column": openpyxl.utils.get_column_letter(index),
                   "source_field": label, "row_roles": "national, ostan",
                   "adoption": "identifier" if index == 1 else "verified_against_population_sheet"}
                  for index, label in enumerate(("Province", "Population", "Male", "Female"), 1))
    for index, label in enumerate(age_group_labels):
        for sex, column_number in (("male", 5 + index), ("female", 26 + index)):
            fields.append({"sheet": SOURCES["age_sex"]["sheet"],
                           "column": openpyxl.utils.get_column_letter(column_number),
                           "source_field": f"{sex} age {label}", "row_roles": "national, ostan",
                           "adoption": "adopted"})
    (evidence / "SOURCE_TABLE_INVENTORY.json").write_text(
        json.dumps({"country": "IRN", "census_year": 2016, "fields": fields},
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (evidence / "CODE_CROSSWALK.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("source_row_id", "source_label", "level", "parent_source_row_id", "official_code", "boundary_version", "match_status"))
        for territory in territories:
            writer.writerow((territory["id"], territory["name"], territory["level"], territory["parent_id"], "", "", "not_matched"))
    print(json.dumps(audit["counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
