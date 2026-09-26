"""Audit selected DCS CPH-2024 count tables and source-coded GN geography.

All acquired workbook sheet dimensions are inventoried. Detailed numeric-cell
decisions here cover final A5/A14/A16 and the GN population/code crosswalk only;
the other acquired table families remain semantically unassessed.
"""

import argparse
import hashlib
import json
import re
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


FIELDS = {
    "populationa-a5": ["population", "male", "female", "under_15", "age_15_59", "age_60_64", "age_65_plus"],
    "housinga-a14": ["households", "protected_well", "semi_protected_well", "unprotected_well",
                     "tube_well", "spring", "water_board_piped", "local_authority_piped",
                     "community_piped", "private_piped", "tank_river_stream", "rainwater",
                     "bottled_water", "reverse_osmosis", "bowser", "other_water"],
    "housinga-a16": ["households", "within_unit_exclusive", "within_unit_shared",
                     "outside_unit_exclusive", "outside_unit_shared", "other_unit_shared",
                     "public_toilet", "no_toilet"],
}


def require(ok, message):
    if not ok:
        raise ValueError(message)


def digest(body):
    return hashlib.sha256(body).hexdigest()


def english(value):
    return next((part.strip() for part in reversed(str(value).replace("_x000D_", "").split("\n"))
                 if part.strip()), "") if value is not None else ""


def normal(value):
    return re.sub(r"[^a-z0-9]", "", english(value).lower().replace("district", ""))


def numeric(value, locator):
    require(isinstance(value, (int, float)) and not isinstance(value, bool) and
            value >= 0 and int(value) == value, f"Invalid count {locator}: {value}")
    return int(value)


def read_selected(project, entry, fields):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(project / entry["raw_path"], read_only=True, data_only=True)
    require(len(book.sheetnames) == 1, f"Unexpected multiple sheets: {entry['id']}")
    sheet = book.active
    start = 2 if entry["id"] == "populationa-a5" else 1
    label = 1 if entry["id"] == "populationa-a5" else 0
    records = []
    district = None
    for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if not isinstance(row[start], (int, float)):
            continue
        name = english(row[label])
        values = {field: numeric(row[start + n], f"{entry['id']}!{rownum}:{field}")
                  for n, field in enumerate(fields)}
        if "District" in name:
            district = name.replace("District", "").strip()
            kind = "district"
        elif "Sri Lanka" in name and district is None:
            kind = "country"
        else:
            kind = "division"
        if entry["id"] == "populationa-a5":
            require(values["male"] + values["female"] == values["population"],
                    f"Sex sum {rownum}")
            require(sum(values[f] for f in fields[3:]) == values["population"],
                    f"Age sum {rownum}")
        else:
            require(sum(values[f] for f in fields[1:]) == values["households"],
                    f"Category sum {entry['id']} row {rownum}")
        records.append({"row": rownum, "name": name, "kind": kind,
                        "district": district, "values": values})
    require(len(records) == 366 and Counter(x["kind"] for x in records) ==
            {"country": 1, "district": 25, "division": 340},
            f"Unexpected selected table hierarchy: {entry['id']}")
    # The national direct count must equal all district direct counts and each
    # district direct count must equal all of its DS divisions for every field.
    for field in fields:
        require(records[0]["values"][field] == sum(x["values"][field] for x in records if x["kind"] == "district"),
                f"District coverage {entry['id']} {field}")
        for parent in (x for x in records if x["kind"] == "district"):
            children = [x for x in records if x["kind"] == "division" and
                        normal(x["district"]) == normal(parent["name"])]
            require(parent["values"][field] == sum(x["values"][field] for x in children),
                    f"DS coverage {entry['id']} {field} {parent['name']}")
    return {"sheet": sheet.title, "records": records}


def main(project):
    receipt = json.loads((project / "evidence/LKA_CPH2024_ACQUISITION.json").read_text(encoding="utf-8"))
    manifest_path = Path(__file__).resolve().parents[1] / "config/sri-lanka-cph2024-source-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    require(len(manifest["source_files"]) == len(receipt["records"]) == 34 and
            all({k: row[k] for k in ("id", "url", "kind", "raw_path", "bytes", "sha256")} == pinned
                for row, pinned in zip(receipt["records"], manifest["source_files"])),
            "Acquisition receipt differs from pinned source manifest")
    catalogue = project / manifest["catalogue_raw_path"]
    require(catalogue.exists() and catalogue.stat().st_size == manifest["catalogue_bytes"] and
            digest(catalogue.read_bytes()) == manifest["catalogue_sha256"],
            "Official DCS catalogue snapshot changed")
    original = {}
    sheet_inventory = []
    for entry in receipt["records"]:
        file = project / entry["raw_path"]
        require(file.exists() and file.stat().st_size == entry["bytes"] and
                digest(file.read_bytes()) == entry["sha256"], f"Original changed: {entry['id']}")
        original[entry["id"]] = entry
        if entry["kind"] == "xlsx":
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                book = openpyxl.load_workbook(file, read_only=True, data_only=True)
            for sheet in book:
                sheet_inventory.append({"source_id": entry["id"], "sheet": sheet.title,
                    "declared_max_row": sheet.max_row, "declared_max_column": sheet.max_column,
                    "disposition": "selected_count_fields_audited" if entry["id"] in FIELDS else
                                   "acquired_priority_unassessed",
                    "note": "Declared extent may include formatted empty rows; numeric-cell audit pending"
                    if entry["id"] not in FIELDS else "All selected count columns and reporting rows audited"})
    selected = {key: read_selected(project, original[key], fields)
                for key, fields in FIELDS.items()}
    a5 = selected["populationa-a5"]["records"]
    table_name_differences = {}
    for key in ("housinga-a14", "housinga-a16"):
        rows = selected[key]["records"]
        differences = [(n, x["name"], y["name"], x["district"], y["district"])
                       for n, (x, y) in enumerate(zip(a5, rows))
                       if x["kind"] != y["kind"] or normal(x["name"]) != normal(y["name"]) or
                          normal(x["district"]) != normal(y["district"])]
        table_name_differences[key] = differences
        require(len(differences) == 32 and
                {(x[1], x[2]) for x in differences if x[1] != x[2]} == {
                    ("Padavisripura", "Padavisiripura"),
                    ("Moneragala District", "Monaragala District"),
                    ("Ratnapura District", "Rathnapura District")},
                f"Selected housing table row order/name alignment changed: {key}")
    require(all(x["values"]["households"] == y["values"]["households"]
                for x, y in zip(selected["housinga-a14"]["records"],
                                selected["housinga-a16"]["records"])),
            "Water and toilet household universes differ")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        codes = openpyxl.load_workbook(project / original["admin-codes"]["raw_path"],
                                       read_only=True, data_only=True).active
        gn = openpyxl.load_workbook(project / original["gn-population-provisional"]["raw_path"],
                                    read_only=True, data_only=True).active
    code_rows = [row for row in list(codes.iter_rows(values_only=True))[1:] if isinstance(row[1], int)]
    gn_rows = [row for row in list(gn.iter_rows(values_only=True))[5:] if isinstance(row[0], int)]
    require(len(code_rows) == len(gn_rows) == 14008, "GN source/code row count changed")
    code_by_key = {(r[2], r[4], r[6], r[8]): r for r in code_rows}
    gn_by_key = {(r[0], r[2] % 10, r[4], r[6]): r for r in gn_rows}
    require(len(code_by_key) == len(gn_by_key) == 14008 and code_by_key.keys() == gn_by_key.keys(),
            "GN source-to-code key mismatch")
    gn_name_differences = [{"key": list(key), "code_name": code_by_key[key][10],
                            "population_name": gn_by_key[key][7]}
                           for key in code_by_key if normal(code_by_key[key][10]) != normal(gn_by_key[key][7])]
    require(len(gn_name_differences) == 1, f"Unexpected GN name changes: {gn_name_differences[:5]}")
    require(all(normal(code_by_key[key][5]) == normal(gn_by_key[key][3]) and
                normal(code_by_key[key][7]) == normal(gn_by_key[key][5])
                for key in code_by_key), "District/DS name mismatch in GN code crosswalk")
    for rownum, row in enumerate(gn_rows, 6):
        values = [numeric(row[n], f"GN population {rownum} col {n+1}") for n in range(9, 17)]
        require(values[0] == values[1] + values[2] == values[3] == sum(values[4:]),
                f"GN sex/age mismatch row {rownum}")
    require(sum(row[9] for row in gn_rows) == a5[0]["values"]["population"] == 21781800,
            "GN total and final A5 national total differ")
    gn_ds_counts = defaultdict(lambda: defaultdict(int))
    gn_district_counts = defaultdict(int)
    for row in gn_rows:
        dkey = (row[0], row[2])
        skey = (row[0], row[2], row[4])
        gn_district_counts[dkey] += row[9]
        for n, field in enumerate(FIELDS["populationa-a5"]):
            gn_ds_counts[skey][field] += row[9 + (n if n < 3 else n + 1)]
    district_crosswalk = {}
    for district in (x for x in a5 if x["kind"] == "district"):
        matches = [key for key, total in gn_district_counts.items()
                   if total == district["values"]["population"]]
        require(len(matches) == 1, f"District count is not a unique GN match: {district['name']}")
        district_crosswalk[normal(district["name"])] = matches[0]
    code_ds = defaultdict(list)
    for row in code_rows:
        key = (normal(row[5]), normal(row[7]))
        code_ds[key].append(row)
    ds_unmatched = []
    ds_ambiguous = []
    ds_crosswalk = []
    gn_age_differences = []
    for row in (x for x in a5 if x["kind"] == "division"):
        matches = code_ds[(normal(row["district"]), normal(row["name"]))]
        if not matches:
            ds_unmatched.append([row["district"], row["name"], row["row"]])
        elif len({(x[2], x[4], x[6]) for x in matches}) != 1:
            ds_ambiguous.append([row["district"], row["name"], row["row"]])
        dkey = district_crosswalk[normal(row["district"])]
        by_counts = [key for key, values in gn_ds_counts.items()
                     if key[:2] == dkey and
                     all(values[field] == row["values"][field]
                         for field in ("population", "male", "female"))]
        require(len(by_counts) == 1, f"Final A5 DS did not match GN population/sex uniquely: {row['name']}")
        key = by_counts[0]
        code_key = (key[0], key[1] % 10, key[2])
        ages = {field: gn_ds_counts[key][field] - row["values"][field]
                for field in FIELDS["populationa-a5"][3:]}
        if any(ages.values()):
            gn_age_differences.append({"district": row["district"], "division": row["name"],
                                       "source_row": row["row"], "gn_minus_final": ages})
        if matches:
            require({(x[2], x[4], x[6]) for x in matches} == {code_key},
                    f"Final A5 DS name/count code conflict: {row['name']}")
        ds_crosswalk.append({"district": row["district"], "final_name": row["name"],
                             "final_row": row["row"], "official_province_code": code_key[0],
                             "official_district_code": code_key[1], "official_ds_code": code_key[2],
                             "code_register_name": code_by_key[next(k for k in code_by_key if k[:3] == code_key)][7],
                             "match_method": "normalized_name_and_population_sex_counts" if matches else
                                             "unique_population_sex_counts_within_district_name_variant"})
    require(len(ds_crosswalk) == 340 and len(ds_unmatched) == 16 and not ds_ambiguous,
            "Unexpected A5-to-code division alignment")
    result = {"status": "selected_tables_audited_other_acquired_tables_unassessed",
              "manifest_sha256": digest(manifest_path.read_bytes()),
              "acquired_originals": len(receipt["records"]),
              "sheet_inventory": sheet_inventory,
              "selected_fields": FIELDS,
              "selected": selected,
              "housing_row_name_differences_from_population_a5": table_name_differences,
              "division_code_crosswalk": ds_crosswalk,
              "gn_age_differences_from_final_a5": gn_age_differences,
              "gn_code_reconciliation": {"source_rows": len(gn_rows),
                  "official_code_rows": len(code_rows), "exact_code_keys": len(code_by_key),
                  "gn_name_differences": gn_name_differences,
                  "national_gn_sum": sum(row[9] for row in gn_rows),
                  "ds_name_unmatched": ds_unmatched, "ds_name_ambiguous": ds_ambiguous},
              "cautions": ["GN population workbook says Provisional although its national sum equals final A5; source status remains distinct",
                  "GN age-category counts differ from final A5 in some DS rows; keep the provisional GN and final series distinct",
                  "The administrative code register's edition is not proved to be 2024 census boundary geometry",
                  "Local-government authorities are not equivalent to Divisional Secretary statistical units",
                  "Other acquired workbooks have dimensions inventoried but not every numeric field semantically decided"]}
    out = project / "evidence/LKA_CPH2024_AUDIT.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"sources": len(receipt["records"]), "sheets": len(sheet_inventory),
        "final_table_records": len(a5), "GN_rows": len(gn_rows),
        "GN_code_name_differences": len(gn_name_differences),
        "DS_name_unmatched": len(ds_unmatched), "DS_name_ambiguous": len(ds_ambiguous),
        "GN_age_difference_divisions": len(gn_age_differences)},
        ensure_ascii=False))
    if ds_unmatched or ds_ambiguous:
        print("DS unmatched", ds_unmatched[:20], "ambiguous", ds_ambiguous[:20])


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
