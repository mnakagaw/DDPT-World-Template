#!/usr/bin/env python3
"""Inventory every sheet and numeric column in Qatar's official Census 2020 Excel."""

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def numeric_kind(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return "typed"
    if isinstance(value, str) and re.fullmatch(r"[+-]?(?:\d+(?:,\d{3})*|\d+)(?:\.\d+)?", value.strip()):
        return "text"
    return None


def compact(value):
    return " ".join(str(value or "").split())[:240]


parser = argparse.ArgumentParser()
parser.add_argument("--project", required=True)
parser.add_argument("--manifest", required=True)
args = parser.parse_args()
project = Path(args.project).resolve()
manifest_path = Path(args.manifest).resolve()
if not manifest_path.is_relative_to(project):
    raise SystemExit("Manifest must be inside the country project")
manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
receipt = next((r for r in manifest["receipts"] if r["file"] == "Census_Final_Results.xlsx"), None)
if not receipt or receipt["status"] != "acquired":
    raise SystemExit("Acquired official Excel receipt is required")
workbook_path = manifest_path.parent / receipt["file"]
actual_sha = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
if actual_sha != receipt["sha256"]:
    raise SystemExit("Official Excel SHA-256 does not match receipt")

workbook = load_workbook(workbook_path, read_only=True, data_only=True)
sheet_inventory = []
column_inventory = []
for sheet in workbook:
    rows = list(sheet.iter_rows(values_only=True))
    title_en = compact(next((v for v in rows[2] if isinstance(v, str) and re.search(r"[A-Za-z]", v)), "")) if len(rows) > 2 else ""
    title_ar = compact(next((v for v in rows[0] if isinstance(v, str)), "")) if rows else ""
    first_numeric = None
    last_numeric = None
    typed_count = 0
    text_count = 0
    for col_idx in range(sheet.max_column):
        typed = []
        as_text = []
        for row_idx, row in enumerate(rows, 1):
            value = row[col_idx] if col_idx < len(row) else None
            kind = numeric_kind(value)
            if kind == "typed":
                typed.append((row_idx, value))
            elif kind == "text":
                as_text.append((row_idx, value))
        typed_count += len(typed)
        text_count += len(as_text)
        values = sorted(typed + as_text, key=lambda item: item[0])
        if not values:
            continue
        first_numeric = min(first_numeric or values[0][0], values[0][0])
        last_numeric = max(last_numeric or values[-1][0], values[-1][0])
        header = " | ".join(dict.fromkeys(compact(row[col_idx]) for row in rows[:9] if col_idx < len(row) and row[col_idx] is not None))
        column_inventory.append({
            "sheet": sheet.title,
            "column": get_column_letter(col_idx + 1),
            "title_en": title_en,
            "header_preview": header[:360],
            "typed_numeric_cells": len(typed),
            "numeric_text_cells": len(as_text),
            "first_numeric_row": values[0][0],
            "first_numeric_value": values[0][1],
            "last_numeric_row": values[-1][0],
            "last_numeric_value": values[-1][1],
            "adoption_decision": "priority_unassessed",
        })
    sheet_inventory.append({
        "sheet": sheet.title,
        "kind": "numbered_table" if re.fullmatch(r"[1-9]\d*", sheet.title) else "supporting_sheet",
        "title_en": title_en,
        "title_ar": title_ar,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "typed_numeric_cells": typed_count,
        "numeric_text_cells": text_count,
        "numeric_columns": sum(1 for item in column_inventory if item["sheet"] == sheet.title),
        "first_numeric_row": first_numeric,
        "last_numeric_row": last_numeric,
        "adoption_decision": "priority_unassessed",
    })
worksheet_names = {item["sheet"] for item in sheet_inventory}
for sheet_name in workbook.sheetnames:
    if sheet_name not in worksheet_names:
        sheet_inventory.append({
            "sheet": sheet_name,
            "kind": "chart_sheet",
            "title_en": "",
            "title_ar": "",
            "max_row": None,
            "max_column": None,
            "typed_numeric_cells": 0,
            "numeric_text_cells": 0,
            "numeric_columns": 0,
            "first_numeric_row": None,
            "last_numeric_row": None,
            "adoption_decision": "priority_unassessed",
        })
sheet_inventory.sort(key=lambda item: workbook.sheetnames.index(item["sheet"]))

evidence = project / "evidence"
evidence.mkdir(exist_ok=True)
(evidence / "QAT_CENSUS2020_SHEET_INVENTORY.json").write_text(json.dumps({
    "source_url": receipt["url"],
    "source_sha256": actual_sha,
    "source_manifest": str(manifest_path.relative_to(project)).replace("\\", "/"),
    "sheet_count": len(sheet_inventory),
    "worksheet_count": len(worksheet_names),
    "chart_sheet_count": len(sheet_inventory) - len(worksheet_names),
    "numbered_table_count": sum(item["kind"] == "numbered_table" for item in sheet_inventory),
    "numeric_column_count": len(column_inventory),
    "sheets": sheet_inventory,
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
with (evidence / "QAT_CENSUS2020_NUMERIC_COLUMNS.csv").open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(column_inventory[0]))
    writer.writeheader()
    writer.writerows(column_inventory)
print(json.dumps({"sheets": len(sheet_inventory), "worksheets": len(worksheet_names), "chart_sheets": len(sheet_inventory) - len(worksheet_names),
                  "numbered_tables": sum(item["kind"] == "numbered_table" for item in sheet_inventory),
                  "numeric_columns": len(column_inventory), "typed_numeric_cells": sum(item["typed_numeric_cells"] for item in sheet_inventory),
                  "numeric_text_cells": sum(item["numeric_text_cells"] for item in sheet_inventory)}, ensure_ascii=False))
