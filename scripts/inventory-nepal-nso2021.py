"""Audit selected NSO NPHC 2021 province workbooks and 2023 code register.

The three adopted table families have independent original-cell and parent
reconciliation checks. Literacy Table 17 is acquired and column-inventoried,
but no literacy cells are adopted until its age/sex hierarchy is audited.
"""

import argparse
import csv
import difflib
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import xlrd


MANIFEST = Path(__file__).resolve().parents[1] / "config/nepal-nso2021-source-manifest.json"
FIELDS = {
    "Indv01": ["households_all", "population", "male", "female"],
    "Hhld06": ["households_private", "piped_inside", "piped_outside", "tube_well",
               "covered_well", "uncovered_well", "spout", "river_stream",
               "jar_bottle", "other_water"],
    "Hhld09": ["households_private", "flush_sewer", "flush_septic", "pit_toilet",
               "public_toilet", "no_toilet"],
}
START_COLUMN = {"Indv01": 2, "Hhld06": 2, "Hhld09": 2}
UNRESOLVED_CODE_NAMES = {
    30304: ("Kalika Gaunpalika", "Shuva Kalika Rural Municipality"),
    70905: ("Dodhara Chandani Municipality", "Mahakali Urban Municipality"),
}


def require(ok, message):
    if not ok:
        raise ValueError(message)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def normal_name(value):
    value = str(value).lower().strip()
    value = re.sub(r"rural municipality|urban municipality|sub.metropolitan city|"
                   r"metropolitan city|municipality|gaunpalika|nagarpalika|"
                   r"mahanagarpalika", "", value)
    return re.sub(r"[^a-z0-9]", "", value)


def pin_sources(project, manifest):
    for entry in manifest["source_files"]:
        target = project / entry["raw_path"]
        require(target.exists(), f"Missing {target}")
        body = target.read_bytes()
        require(len(body) == entry["bytes"] and digest(body) == entry["sha256"],
                f"Original changed: {target}")
    code = project / manifest["administrative_code_raw_path"]
    require(code.exists() and digest(code.read_bytes()) ==
            manifest["administrative_code_sha256"], "NSO code workbook changed")
    for name, key in (("nso-provincial-koshi-data.html", "catalogue_page_sha256"),
                      ("nso-provincial-page.js", "catalogue_script_sha256")):
        file = project / "raw" / name
        require(file.exists() and digest(file.read_bytes()) == manifest[key],
                f"Official catalogue original changed: {name}")


def numeric_row(row, start, fields, locator):
    vals = row[start:start + len(fields)]
    require(len(vals) == len(fields) and
            all(isinstance(x, (int, float)) and not isinstance(x, bool) and
                int(x) == x and x >= 0 for x in vals),
            f"Non-integer or missing source cell: {locator}")
    return dict(zip(fields, map(int, vals)))


def read_table(project, manifest_map, province, table):
    entry = manifest_map[(province, table)]
    book = openpyxl.load_workbook(project / entry["raw_path"],
                                  read_only=True, data_only=True)
    require(len(book.sheetnames) == 1, f"Unexpected sheets: {entry['raw_path']}")
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    first = START_COLUMN[table]
    fields = FIELDS[table]
    national = numeric_row(rows[5], first, fields, f"{sheet.title}!{6}")
    province_value = numeric_row(rows[7], first, fields, f"{sheet.title}!{8}")
    for scope, values in (("national", national), ("province", province_value)):
        if table == "Indv01":
            require(values["male"] + values["female"] == values["population"],
                    f"Sex sum: {entry['raw_path']} {scope}")
        else:
            require(sum(list(values.values())[1:]) == values["households_private"],
                    f"Category sum: {entry['raw_path']} {scope}")
    require(str(rows[5][1 if table == "Indv01" else 0]).strip() == "Nepal" and
            int(rows[7][0]) == 0, f"Unexpected national/province rows: {entry['raw_path']}")
    area_rows = []
    district = 0
    for rownum, row in enumerate(rows[9:], start=10):
        serial, name = row[:2]
        if not isinstance(serial, (int, float)) or not name:
            continue
        serial = int(serial)
        if serial == 0:
            district += 1
            kind = "district"
        elif serial == 99:
            require(table == "Indv01" and district > 0,
                    f"Unexpected institutional row {entry['raw_path']}!{rownum}")
            kind = "institutional"
        else:
            require(district > 0 and serial > 0, "Local row without district")
            kind = "local"
        values = numeric_row(row, first, fields, f"{sheet.title}!{rownum}")
        if table == "Indv01":
            require(values["male"] + values["female"] == values["population"],
                    f"Sex sum: {entry['raw_path']}!{rownum}")
        else:
            require(sum(list(values.values())[1:]) == values["households_private"],
                    f"Category sum: {entry['raw_path']}!{rownum}")
        area_rows.append({"kind": kind, "province": province, "district_order": district,
                          "source_serial": serial, "name": str(name).strip(),
                          "source_cell_row": rownum, "sheet": sheet.title,
                          "source_id": f"npl-nso2021-{table.lower()}-p{province}",
                          "values": values})
    return {"province": province, "table": table, "national": national,
            "province_value": province_value, "rows": area_rows,
            "sheet": sheet.title, "xlsx_rows": sheet.max_row,
            "xlsx_columns": sheet.max_column, "hash": entry["sha256"]}


def read_codes(project, manifest):
    book = xlrd.open_workbook(str(project / manifest["administrative_code_raw_path"]))
    require(book.sheet_names() == ["pradesh_code", "district_code", "localunit_code"],
            "NSO code workbook layout changed")
    province = {int(book.sheet_by_name("pradesh_code").cell_value(i, 4)):
                str(book.sheet_by_name("pradesh_code").cell_value(i, 2)).strip()
                for i in range(5, 12)}
    sheet = book.sheet_by_name("district_code")
    district = {int(sheet.cell_value(i, 4)): str(sheet.cell_value(i, 2)).strip()
                for i in range(3, sheet.nrows)}
    sheet = book.sheet_by_name("localunit_code")
    local = {int(sheet.cell_value(i, 4)): {
        "district": str(sheet.cell_value(i, 1)).strip(),
        "name": str(sheet.cell_value(i, 2)).strip(),
    } for i in range(3, sheet.nrows)}
    require(len(province) == 7 and len(district) == 77 and len(local) == 753,
            "Expected official NSO 2023 code-register cardinalities")
    return province, district, local


def check_rows(tables, codes):
    provinces, district_codes, local_codes = codes
    for table in FIELDS:
        group = [tables[(p, table)] for p in range(1, 8)]
        require(len({json.dumps(x["national"], sort_keys=True) for x in group}) == 1,
                f"{table} national row differs by province file")
        fields = FIELDS[table]
        for field in fields:
            require(sum(x["province_value"][field] for x in group) ==
                    group[0]["national"][field], f"{table} national reconciliation {field}")
        for p, source in enumerate(group, start=1):
            rows = source["rows"]
            districts = [r for r in rows if r["kind"] == "district"]
            locals_ = [r for r in rows if r["kind"] == "local"]
            institutionals = [r for r in rows if r["kind"] == "institutional"]
            require(len(districts) == sum(1 for c in district_codes if c // 100 == p),
                    f"District count changed P{p} {table}")
            for field in fields:
                require(sum(r["values"][field] for r in districts) ==
                        source["province_value"][field],
                        f"Province/district mismatch P{p} {table} {field}")
            for r in districts:
                children = [x for x in locals_ + institutionals
                            if x["district_order"] == r["district_order"]]
                require(children, f"No children P{p} {r['name']}")
                for field in fields:
                    require(sum(x["values"][field] for x in children) ==
                            r["values"][field],
                            f"District/child mismatch P{p} {r['name']} {table} {field}")
            if table == "Indv01":
                require(len(institutionals) == len(districts),
                        f"Missing institutional rows P{p}")
            else:
                require(not institutionals, f"Household table has institutional rows P{p}")
    pop = [r for p in range(1, 8) for r in tables[(p, "Indv01")]["rows"]]
    water = [r for p in range(1, 8) for r in tables[(p, "Hhld06")]["rows"]]
    toilet = [r for p in range(1, 8) for r in tables[(p, "Hhld09")]["rows"]]
    counts = Counter(r["kind"] for r in pop)
    require(counts == {"district": 77, "local": 753, "institutional": 77},
            f"Population geography changed: {counts}")
    def keyed(rows):
        return {(r["province"], r["district_order"], r["source_serial"]): r
                for r in rows if r["kind"] != "institutional"}
    base = keyed(pop)
    for table, other in (("Hhld06", water), ("Hhld09", toilet)):
        current = keyed(other)
        require(set(current) == set(base), f"{table} geographic rows differ")
        for key, value in current.items():
            require(normal_name(value["name"]) == normal_name(base[key]["name"]),
                    f"{table} name differs: {key}")
    district_matches = []
    local_matches = []
    used_codes = set()
    for row in pop:
        p, order, serial = row["province"], row["district_order"], row["source_serial"]
        if row["kind"] == "district":
            code = p * 100 + order
            require(code in district_codes and normal_name(row["name"]) ==
                    normal_name(district_codes[code]),
                    f"District/code name differs: {code} {row['name']}")
            row["official_code"] = str(code)
            district_matches.append(code)
        elif row["kind"] == "local":
            code = (p * 100 + order) * 100 + serial
            require(code in local_codes and code not in used_codes,
                    f"Local code missing/duplicate: {code}")
            used_codes.add(code)
            official = local_codes[code]["name"]
            ratio = difflib.SequenceMatcher(None, normal_name(row["name"]),
                                            normal_name(official)).ratio()
            if code in UNRESOLVED_CODE_NAMES:
                require((row["name"], official) == UNRESOLVED_CODE_NAMES[code],
                        f"Unresolved code row changed: {code}")
                status, row["official_code"] = "name_conflict_unresolved", None
            else:
                require(ratio >= 0.85,
                        f"New code/name mismatch: {code} {row['name']} {official}")
                status = "exact_normalized" if ratio == 1 else "minor_transliteration"
                row["official_code"] = str(code)
            row["code_name_status"] = status
            row["code_register_name"] = official
            row["code_candidate"] = str(code)
            local_matches.append(status)
        else:
            row["official_code"] = None
            row["code_name_status"] = "non_territorial_institutional_row"
    require(len(used_codes) == 753 and len(district_matches) == 77,
            "Code-sheet row cover incomplete")
    require(Counter(local_matches) == {
        "exact_normalized": 749, "minor_transliteration": 2,
        "name_conflict_unresolved": 2}, f"Local code correspondence changed: {Counter(local_matches)}")
    institutional_households = sum(r["values"]["households_all"] for r in pop
                                   if r["kind"] == "institutional")
    institutional_population = sum(r["values"]["population"] for r in pop
                                   if r["kind"] == "institutional")
    full_households = tables[(1, "Indv01")]["national"]["households_all"]
    private_households = tables[(1, "Hhld06")]["national"]["households_private"]
    require(tables[(1, "Hhld09")]["national"]["households_private"] == private_households
            and full_households - private_households == institutional_households == 6096,
            "Population-vs-amenity household universe changed")
    return {"population_row_types": dict(counts),
            "local_code_matches": dict(Counter(local_matches)),
            "district_code_matches": len(district_matches),
            "institutional_households": institutional_households,
            "institutional_population": institutional_population,
            "all_households": full_households,
            "private_amenity_households": private_households}


def inventory_literacy(project, manifest_map):
    rows = []
    for p in range(1, 8):
        entry = manifest_map[(p, "Indv17")]
        sheet = openpyxl.load_workbook(project / entry["raw_path"],
                                       read_only=True, data_only=True).active
        counts = {openpyxl.utils.get_column_letter(i): 0 for i in range(6, 11)}
        for row in sheet.iter_rows(min_row=6, values_only=True):
            for i in range(6, 11):
                value = row[i - 1] if len(row) >= i else None
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    counts[openpyxl.utils.get_column_letter(i)] += 1
        rows.append({"province": p, "table": "Indv17", "sheet": sheet.title,
                     "xlsx_rows": sheet.max_row, "xlsx_columns": sheet.max_column,
                     "numeric_columns": ["F: aged 5+ total", "G: can read/write",
                                         "H: can read only", "I: cannot read/write",
                                         "J: not stated"],
                     "numeric_cell_counts": counts,
                     "status": "priority_unassessed_age_sex_hierarchy"})
    return rows


def main(project):
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    pin_sources(project, manifest)
    manifest_map = {(int(x["province_directory"][1:]), x["table"]): x
                    for x in manifest["source_files"]}
    require(len(manifest_map) == 28, "Expected 7 provinces x 4 selected originals")
    tables = {(p, t): read_table(project, manifest_map, p, t)
              for p in range(1, 8) for t in FIELDS}
    reconciliation = check_rows(tables, read_codes(project, manifest))
    literacy = inventory_literacy(project, manifest_map)
    evidence = project / "evidence"
    inventory = evidence / "NPL_NSO2021_ACQUIRED_TABLE_INVENTORY.csv"
    with inventory.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "province", "table", "sheet", "xlsx_rows", "xlsx_columns",
            "numeric_columns", "numeric_cell_counts", "adoption"])
        writer.writeheader()
        for p in range(1, 8):
            for table in FIELDS:
                item = tables[(p, table)]
                writer.writerow({"province": p, "table": table, "sheet": item["sheet"],
                                 "xlsx_rows": item["xlsx_rows"],
                                 "xlsx_columns": item["xlsx_columns"],
                                 "numeric_columns": "; ".join(FIELDS[table]),
                                 "numeric_cell_counts": len(item["rows"]) * len(FIELDS[table]) +
                                 2 * len(FIELDS[table]),
                                 "adoption": "selected_all_count_columns_pending_site_import"})
            item = literacy[p - 1]
            writer.writerow({"province": p, "table": "Indv17", "sheet": item["sheet"],
                             "xlsx_rows": item["xlsx_rows"],
                             "xlsx_columns": item["xlsx_columns"],
                             "numeric_columns": "; ".join(item["numeric_columns"]),
                             "numeric_cell_counts": json.dumps(item["numeric_cell_counts"]),
                             "adoption": item["status"]})
    audit = {
        "audited_at_utc": datetime.now(timezone.utc).isoformat(),
        "manifest_sha256": digest(MANIFEST.read_bytes()),
        "code_xls_sha256": manifest["administrative_code_sha256"],
        "source_catalogue": {"entries": manifest["catalogue_entries"],
                             "xlsx_tables_per_province": manifest["catalogue_xlsx_entries"],
                             "acquired_originals": len(manifest_map)},
        "reconciliation": reconciliation,
        "national": {t: tables[(1, t)]["national"] for t in FIELDS},
        "adopted_tables": {f"P{p}-{t}": tables[(p, t)]
                           for p in range(1, 8) for t in FIELDS},
        "unassessed_literacy": literacy,
        "unacquired_catalogue_tables_per_province": 85,
        "source_limits": [
            "2021 census versus 2023 NSO code sheet: two local names conflict; hold official_code on those rows",
            "2021 census population Table 1 includes institutional rows; water/toilet tables cover private households only",
            "No official same-edition polygon or ward statistics adopted",
            "Indv17 literacy age/sex hierarchy inventoried but not semantically adopted",
        ],
    }
    (evidence / "NPL_NSO2021_AUDIT.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"sources": 28, "adopted_tables": 21,
                      "population": audit["national"]["Indv01"]["population"],
                      "households_private": audit["national"]["Hhld06"]["households_private"],
                      **reconciliation}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
