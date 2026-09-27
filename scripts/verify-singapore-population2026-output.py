"""Compare real AreaData exports with selected cells in the pinned SingStat ZIP."""

import argparse
import csv
import hashlib
import json
import subprocess
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CASES = {
    "SGP": ("Total", "Total"),
    "Tampines": ("Tampines", "Total"),
    "TampinesEast": ("Tampines", "Tampines East"),
    "AngMoKio": ("Ang Mo Kio", "Total"),
    "ChangiBay": ("Changi Bay", "Total"),
}
INDICATOR_COLUMNS = {
    "RESIDENT_TOTAL": ("2026", "Total", "Total"),
    "RESIDENT_MALE": ("2026", "Total", "Males"),
    "RESIDENT_FEMALE": ("2026", "Total", "Females"),
    "RESIDENT_HDB": ("2026(Total)", "Total", "Total HDB^"),
    "RESIDENT_CONDO": ("2026(Total)", "Total", "Condominiums and Other Apartments"),
    "RESIDENT_LANDED": ("2026(Total)", "Total", "Landed Properties"),
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def source_values(zip_path):
    result = {}
    with zipfile.ZipFile(zip_path) as archive:
        for member in archive.namelist():
            if not member.endswith(".xlsx") or not ("Single Year" in member or "Type of Dwelling" in member):
                continue
            with archive.open(member) as stream:
                workbook = load_workbook(stream, read_only=True, data_only=True)
                for sheet in workbook:
                    if sheet.title not in {"2026", "2026(Total)"}:
                        continue
                    for row_num, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
                        if len(row) != 5:
                            continue
                        area, subzone, third, fourth, value = row
                        if (area, subzone) not in CASES.values():
                            continue
                        for indicator, (selected_sheet, selected_third, selected_fourth) in INDICATOR_COLUMNS.items():
                            if sheet.title == selected_sheet and third == selected_third and fourth == selected_fourth:
                                key = (area, subzone, indicator)
                                require(key not in result, f"Duplicated original {key}")
                                result[key] = (value, f"{member}:{sheet.title}:E{row_num}")
                workbook.close()
    require(len(result) == len(CASES) * len(INDICATOR_COLUMNS), "Original selected cells incomplete")
    return result


def overall_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = [r for r in csv.DictReader(handle) if r["Record scope"] == "overall"]
    return {r["Indicator ID"]: r for r in rows}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve()
    dashboard = (project / "data/dashboard.json").read_bytes()
    data = json.loads(dashboard)
    require(data["country"]["id"] == "SGP" and len(data["territories"]) == 388 and
            len(data["boundaries"]["features"]) == 55, "Candidate scope changed")
    audit = json.loads((project / "evidence/SGP_POPULATION2026_AUDIT.json").read_text(encoding="utf-8"))
    source_observed = sum(row["status"] == "observed" for row in data["observations"]
                          if row["indicator_id"].startswith("SGP_SINGSTAT_"))
    source_missing = sum(row["status"] == "missing" for row in data["observations"]
                         if row["indicator_id"].startswith("SGP_SINGSTAT_"))
    require((source_observed, source_missing) ==
            (audit["observed_slots"], audit["missing_nil_or_negligible_slots"]),
            "Imported observation statuses differ from source audit")
    exported = subprocess.run(["node", str(ROOT / "scripts/export-singapore-population2026-cases.mjs"), str(project)],
                              cwd=ROOT, text=True, capture_output=True, check=True)
    export_state = json.loads(exported.stdout)
    export_cases = export_state["cases"]
    require(export_state["comparison"] == {"national_total": 55, "national_observed": 48,
                                           "tampines_total": 5, "tampines_observed": 5},
            f"Comparable rows differ from direct source geography: {export_state['comparison']}")
    original = source_values(project / "raw/singstat-population-trends-2026-geospatial.zip")
    out = project / "evidence/output-verification"
    checks = []
    for stem, pair in CASES.items():
        area, subzone = pair
        tid = export_cases[stem]
        rows = overall_rows(out / f"{stem}-diagnostic.csv")
        for indicator in INDICATOR_COLUMNS:
            oid = "SGP_SINGSTAT_" + indicator
            require(oid in rows, f"Missing export indicator {stem}/{oid}")
            row = rows[oid]
            value, locator = original[(area, subzone, indicator)]
            require(row["Territory ID"] == tid and row["Period"] == "2026-06" and
                    row["Source URL"].startswith("https://www.singstat.gov.sg/") and
                    locator in row["Source locator"], f"Export provenance differs {stem}/{indicator}")
            if value == "-":
                require(row["Value"] == "" and row["Status"] == "missing",
                        f"Nil/negligible source became zero or observed: {stem}/{indicator}")
            else:
                require(int(row["Value"]) == value and row["Status"] == "observed",
                        f"Export/source value mismatch: {stem}/{indicator}")
        for suffix in ("diagnostic.html", "diagnostic.md", "planning.html", "evidence.csv"):
            body = (out / f"{stem}-{suffix}").read_text(encoding="utf-8")
            require("undefined" not in body and "NaN" not in body, f"Invalid {stem}/{suffix}")
            require(data["territories"][next(i for i,t in enumerate(data["territories"]) if t["id"] == tid)]["name"] in body,
                    f"Selected territory missing from {stem}/{suffix}")
        checks.append({"case": stem, "territory_id": tid, "source_cells_checked": 6,
                       "outputs": 5})
    national_csv = out / "SGP-diagnostic.csv"
    with national_csv.open(encoding="utf-8-sig", newline="") as handle:
        compared = [r for r in csv.DictReader(handle)
                    if r["Record scope"] == "within_area" and
                       r["Indicator ID"] == "SGP_SINGSTAT_RESIDENT_TOTAL"]
    require(len(compared) == 55, f"National area comparison has {len(compared)} rows")
    require(len({r["Territory ID"] for r in compared}) == 55, "Duplicate area comparison")
    require("Master Plan 2025" in (out / "SGP-planning.html").read_text(encoding="utf-8"),
            "National MP2025 plan missing")
    require("Master Plan 2025 Written Statement" not in
            (out / "Tampines-planning.html").read_text(encoding="utf-8"),
            "National plan inherited as an area-specific document")
    report = {"dataset_sha256": hashlib.sha256(dashboard).hexdigest(),
        "checked_cases": checks, "national_comparison_rows": len(compared),
        "comparable_rows": export_state["comparison"],
        "observed_new": source_observed, "missing_nil_or_negligible": source_missing,
        "independent_acceptance": False}
    (project / "evidence/SGP_OUTPUT_VERIFICATION.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"checked_cases": len(checks), "national_comparison_rows": len(compared),
                      "dataset_sha256": report["dataset_sha256"]}))


if __name__ == "__main__":
    main()
