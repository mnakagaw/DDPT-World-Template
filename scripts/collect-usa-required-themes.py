#!/usr/bin/env python3
"""Collect official USA urban/rural and nutrition outcome supplements."""
from __future__ import annotations
import argparse,csv,hashlib,json
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook

URBAN_PAGE='https://www.census.gov/programs-surveys/geography/guidance/geo-areas/urban-rural.html'
COUNTY_XLSX='https://www2.census.gov/geo/docs/reference/ua/2020_UA_COUNTY.xlsx'
STATE_XLSX='https://www2.census.gov/geo/docs/reference/ua/State_Urban_Rural_Pop_2020_2010.xlsx'
CDC_PAGE='https://data.cdc.gov/d/fu4u-a9bh'
CDC_QUERY="https://data.cdc.gov/resource/fu4u-a9bh.csv?$limit=5000&$where=measureid%3D%27OBESITY%27%20AND%20datavaluetypeid%3D%27AgeAdjPrv%27"

def sha(path):
 h=hashlib.sha256()
 with path.open('rb') as f:
  for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
 return h.hexdigest()

def load_areas(path):
 data=json.loads(path.read_text(encoding='utf-8'))
 areas={r['id']:r for r in data['territories'] if r['id']=='USA' or r.get('country_id')=='USA'}
 expected={'USA'}|{f"USA:ACS2023:STATE:{i:02d}" for i in []}
 if len(areas)!=3196:raise RuntimeError(f'Expected 3196 USA geographies, got {len(areas)}')
 return areas

def obs(area,iid,period,value,source,definition,method,locator,numerator=None,denominator=None):
 return {'territory_id':area,'indicator_id':iid,'period':period,'value':value,'status':'observed' if value is not None else 'missing','source_id':source,'definition':definition,'definition_id':iid,'unit':'%','population':definition,'measurement_method':method,'source_locator':locator,'numerator':numerator,'denominator':denominator}

def parse_urban(county_path,state_path,areas):
 county_ids={r['official_code']:r['id'] for r in areas.values() if r.get('level')=='county'}
 state_ids={r['official_code']:r['id'] for r in areas.values() if r.get('level')=='state'}
 values={};totals={};urban={}
 wb=load_workbook(county_path,read_only=True,data_only=True)
 for sheet_name in ['2020_UA_COUNTY','CT_2022']:
  ws=wb[sheet_name];blank=0
  for row in ws.iter_rows(min_row=2,values_only=True):
   if row[0] in (None,''):
    blank+=1
    if blank>=100:break
    continue
   blank=0;code=f'{str(row[0]).zfill(2)}{str(row[1]).zfill(3)}'
   if code not in county_ids:continue
   area=county_ids[code];total=float(row[4]);upop=float(row[11]);pct=float(row[12])*100
   if area in values:raise RuntimeError(f'Duplicate county urban row {code}')
   # The source marks a few 0.00/100.00 values with an asterisk and rounds
   # them to 0.01/99.99; retain the published percentage within that tolerance.
   if abs(pct-upop/total*100)>0.011:raise RuntimeError(f'County urban calculation mismatch {code}')
   values[area]=pct;totals[area]=total;urban[area]=upop
 if set(values)!=set(county_ids.values()):
  missing=sorted(set(county_ids.values())-set(values));raise RuntimeError(f'County urban coverage {len(values)}/{len(county_ids)} missing={missing[:20]}')
 ws=load_workbook(state_path,read_only=True,data_only=True)['States']
 for row in ws.iter_rows(min_row=2,values_only=True):
  if row[0] in (None,''):continue
  code=str(row[0]).zfill(2)
  if code not in state_ids:continue
  area=state_ids[code];total=float(row[3]);upop=float(row[4]);pct=float(row[5])
  if abs(pct-upop/total*100)>1e-7:raise RuntimeError(f'State urban calculation mismatch {code}')
  values[area]=pct;totals[area]=total;urban[area]=upop
 if not set(state_ids.values()).issubset(values):raise RuntimeError('State urban coverage is incomplete')
 national_total=sum(totals[x] for x in state_ids.values());national_urban=sum(urban[x] for x in state_ids.values())
 values['USA']=national_urban/national_total*100;totals['USA']=national_total;urban['USA']=national_urban
 if len(values)!=3196:raise RuntimeError(f'Urban observations {len(values)}/3196')
 observations=[obs(area,'USA_C2020_URBAN_POP_PCT','2020',round(values[area],4),'usa-census-2020-urban-rural','2020 Census population classified as urban as a share of total population.','source_reported' if area!='USA' else 'derived_from_source_counts','Official 2020 county/state urban-rural table',urban[area],totals[area]) for area in sorted(values)]
 return observations,{'country_state_county_observations':len(observations),'state_and_dc':len(state_ids),'counties_and_equivalents':len(county_ids),'national_total_population':int(national_total),'national_urban_population':int(national_urban),'all_rows_reconciled':True}

def parse_cdc(path,areas):
 county_ids={r['official_code']:r['id'] for r in areas.values() if r.get('level')=='county'};rows={};national=None;excluded=[]
 with path.open(encoding='utf-8-sig',newline='') as f:
  for row in csv.DictReader(f):
   if row.get('stateabbr')=='US' and row.get('locationid')=='59':
    national=row;continue
   code=row['locationid'].zfill(5)
   if code not in county_ids:
    excluded.append(code);continue
   if code in rows:raise RuntimeError(f'Duplicate CDC county {code}')
   if row['measureid']!='OBESITY' or row['datavaluetypeid']!='AgeAdjPrv':raise RuntimeError('Unexpected CDC measure row')
   rows[code]=row
 if set(rows)!=set(county_ids) or national is None:raise RuntimeError(f'CDC coverage incomplete: counties {len(rows)}/{len(county_ids)}, national={national is not None}, missing={sorted(set(county_ids)-set(rows))[:20]} excluded={sorted(set(excluded))}')
 observations=[]
 for code,row in sorted(rows.items()):
  observations.append(obs(county_ids[code],'USA_CDC_ADULT_OBESITY_AGE_ADJ_PCT_2022','2022',float(row['data_value']),'usa-cdc-places-2024-county-obesity','Model-based age-adjusted prevalence of obesity among adults.','model_based_small_area_estimate',f"PLACES 2024 county release, FIPS {code}; 95% CI {row['low_confidence_limit']}-{row['high_confidence_limit']}"))
 observations.append(obs('USA','USA_CDC_ADULT_OBESITY_AGE_ADJ_PCT_2022','2022',float(national['data_value']),'usa-cdc-places-2024-county-obesity','Model-based age-adjusted prevalence of obesity among adults.','source_reported',f"PLACES 2024 national row; 95% CI {national['low_confidence_limit']}-{national['high_confidence_limit']}"))
 return observations,{'matched_counties':len(rows),'national_row':True,'excluded_non_scope_fips':sorted(set(excluded)),'all_candidate_counties_matched':True}

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--dataset',required=True);ap.add_argument('--county-urban',required=True);ap.add_argument('--state-urban',required=True);ap.add_argument('--cdc',required=True);ap.add_argument('--out',required=True);a=ap.parse_args()
 paths={k:Path(v).resolve() for k,v in vars(a).items() if k!='out'};out=Path(a.out).resolve();out.mkdir(parents=True,exist_ok=True);areas=load_areas(paths['dataset'])
 urban_obs,urban_audit=parse_urban(paths['county_urban'],paths['state_urban'],areas);cdc_obs,cdc_audit=parse_cdc(paths['cdc'],areas)
 sources=[
  {'id':'usa-census-2020-urban-rural','name':'2020 Census county and state urban-rural population tables','publisher':'United States Census Bureau','url':URBAN_PAGE,'status':'ready','retrieved_at':'2026-09-18','reference_period':'2020','geographic_level':'United States, state/DC and county/county equivalent','source_locator':'2020_UA_COUNTY and State Urban Rural Population tables','license':'United States Government public data','license_url':'https://www.census.gov/about/policies/open-gov/open-data.html','note':'County and state tables are joined by exact Census FIPS. Connecticut 2022 planning-region county equivalents are taken from the dedicated CT_2022 sheet.'},
  {'id':'usa-cdc-places-2024-county-obesity','name':'PLACES: County Data 2024 release — adult obesity','publisher':'Centers for Disease Control and Prevention','url':CDC_PAGE,'status':'ready','retrieved_at':'2026-09-18','reference_period':'2022','geographic_level':'County and county equivalent','source_locator':'measureid=OBESITY; datavaluetypeid=AgeAdjPrv','license':'United States Government public data; see CDC dataset metadata','note':'Official model-based small-area estimate. It is a nutrition-related health outcome, not a census count and not evidence of diet quality or food security.'}
 ]
 indicators=[
  {'id':'USA_C2020_URBAN_POP_PCT','name':'Urban population, 2020 Census','theme':'Settlement','unit':'%','definition':'2020 Census population classified as urban as a share of total population.','definition_id':'USA-C2020-URBAN','population':'2020 Census total population','measurement_method':'source_reported','aggregation':'none','period_policy':'fixed_source_period','series_family':'census','display_role':'primary','source_id':sources[0]['id'],'source_locator':'County/state urban-rural population tables'},
  {'id':'USA_CDC_ADULT_OBESITY_AGE_ADJ_PCT_2022','name':'Adult obesity, age-adjusted prevalence','theme':'Nutrition','unit':'%','definition':'Model-based age-adjusted prevalence of obesity among adults.','definition_id':'USA-CDC-PLACES-2024-OBESITY','population':'Adults in the county; model inputs include 2022 BRFSS and Census/ACS population covariates.','measurement_method':'model_based_small_area_estimate','aggregation':'none','period_policy':'fixed_source_period','series_family':'survey','display_role':'primary','source_id':sources[1]['id'],'source_locator':'PLACES 2024 release, OBESITY, AgeAdjPrv'}
 ]
 receipts=[{'label':k,'path':str(v),'bytes':v.stat().st_size,'sha256':sha(v)} for k,v in paths.items() if k!='dataset']
 audit={'schema_version':'1.0','generated_at':datetime.now(timezone.utc).isoformat(),'country_area_id':'USA','status':'urban_and_nutrition_integrated','themes':{'urban_rural':urban_audit,'nutrition':cdc_audit},'counts':{'urban_observations':len(urban_obs),'nutrition_observations':len(cdc_obs),'total_observations':len(urban_obs)+len(cdc_obs)},'receipts':receipts,'evidence_policy':'2020 Census urban/rural values and CDC modeled obesity estimates remain distinct by period, method and geography. No state or national obesity value is inferred from counties.'}
 bundle={'schema_version':'1.0','country_area_id':'USA','sources':sources,'indicators':indicators,'observations':urban_obs+cdc_obs,'audit':audit}
 (out/'usa-required-themes-bundle.json').write_text(json.dumps(bundle,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');(out/'usa-required-themes-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(json.dumps({'counts':audit['counts'],'urban':urban_audit,'nutrition':cdc_audit},indent=2))

if __name__=='__main__':main()
