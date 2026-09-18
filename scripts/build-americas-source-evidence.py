#!/usr/bin/env python3
"""Build auditable source-preflight and source-disposition evidence for AreaData Americas.

This does not claim that 57 country editions are complete. It records where research can
start, what has actually been acquired and inspected, and what the regional gateway adopts.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from pypdf import PdfReader


COMMON = {
    "unsd_nso": "https://unstats.un.org/home/nso_sites/",
    "unsd_census": "https://unstats.un.org/unsd/demographic-social/census/censusdates",
    "eclac_planning": "https://observatorioplanificacion.cepal.org/en",
    "caricom": "https://statistics.caricom.org/",
    "eclac_data": "https://www.cepal.org/en/data-and-statistics",
}

ALIASES = {
    "Bolivia (Plurinational State of)": "Bolivia",
    "Venezuela (Bolivarian Republic of)": "Venezuela",
    "United States of America": "United States of America",
    "Saint Kitts and Nevis": "St. Kitts and Nevis",
    "Saint Lucia": "St. Lucia",
    "Saint Vincent and the Grenadines": "St. Vincent and the Grenadines",
    "Saint Pierre and Miquelon": "St. Pierre and Miquelon",
    "Saint Barthélemy": "Saint-Barthélemy",
    "Saint Martin (French Part)": "Saint-Martin",
    "Sint Maarten (Dutch part)": "Sint Maarten",
    "Greenland": "Greenland (3)",
    "Bonaire, Sint Eustatius and Saba": "Bonaire, Sint Eustatius and Saba",
    "Curaçao": "Curaçao",
}

MANUAL_NSO = {
    "ATG": ("Antigua and Barbuda National Bureau of Statistics", "https://statistics.gov.ag/"),
    "DMA": ("Dominica Central Statistics Office", "https://stats.gov.dm/"),
    "LCA": ("Saint Lucia Central Statistical Office", "https://stats.gov.lc/"),
    "USA": ("United States Census Bureau", "https://www.census.gov/"),
    "BES": ("Statistics Netherlands, Caribbean Netherlands", "https://www.cbs.nl/en-gb/our-services/caribbean-netherlands"),
    "VGB": ("Virgin Islands Central Statistics Office", "https://bvi.gov.vg/statistics"),
    "GLP": ("INSEE Antilles-Guyane", "https://www.insee.fr/en/accueil"),
    "MTQ": ("INSEE Antilles-Guyane", "https://www.insee.fr/en/accueil"),
    "BLM": ("INSEE", "https://www.insee.fr/en/accueil"),
    "MAF": ("INSEE", "https://www.insee.fr/en/accueil"),
    "GUF": ("INSEE Antilles-Guyane", "https://www.insee.fr/en/accueil"),
    "SPM": ("INSEE", "https://www.insee.fr/en/accueil"),
    "PRI": ("U.S. Census Bureau Island Areas / Puerto Rico", "https://www.census.gov/programs-surveys/decennial-census/about/rdo/island-areas.html"),
    "VIR": ("U.S. Census Bureau Island Areas", "https://www.census.gov/programs-surveys/decennial-census/about/rdo/island-areas.html"),
    "MSR": ("Montserrat Statistics Department", "https://statistics.gov.ms/"),
    "GRD": ("Grenada Central Statistical Office", "https://stats.gov.gd/"),
    "KNA": ("Saint Kitts and Nevis Department of Statistics", "https://www.stats.gov.kn/"),
    "VCT": ("Saint Vincent and the Grenadines Statistical Office", "https://stats.gov.vc/"),
    "SXM": ("Department of Statistics Sint Maarten", "https://stats.sintmaartengov.org/"),
    "BVT": ("Statistics Norway", "https://www.ssb.no/en/"),
    "SGS": ("Government of South Georgia and the South Sandwich Islands", "https://www.gov.gs/"),
    "FLK": ("Falkland Islands Government Statistics", "https://www.falklands.gov.fk/policy/statistics"),
    "GRL": ("Statistics Greenland", "https://stat.gl/"),
}

MANUAL_CENSUS = {
    "ATG": ["https://statistics.gov.ag/subjects/population-and-demography/"],
    "DMA": ["https://stats.gov.dm/census/", "https://stats.gov.dm/subjects/demographic-statistics/"],
    "LCA": ["https://stats.gov.lc/wp-content/uploads/2024/11/St-Lucia-Census-2022.pdf"],
    "USA": ["https://www.census.gov/programs-surveys/decennial-census/decade/2020/2020-census-results.html", "https://www.census.gov/data/developers/data-sets/decennial-census.2020.html"],
}

GOV_START = {
    "CAN": "https://www.canada.ca/en.html", "USA": "https://www.usa.gov/", "BMU": "https://www.gov.bm/",
    "GRL": "https://naalakkersuisut.gl/", "SPM": "https://www.saint-pierre-et-miquelon.gouv.fr/",
    "AIA": "https://www.gov.ai/", "ATG": "https://ab.gov.ag/", "ABW": "https://www.gobierno.aw/",
    "BHS": "https://www.bahamas.gov.bs/", "BRB": "https://www.gov.bb/", "BES": "https://www.rijksdienstcn.com/",
    "VGB": "https://bvi.gov.vg/", "CYM": "https://www.gov.ky/", "CUW": "https://gobiernu.cw/",
    "DMA": "https://dominica.gov.dm/", "GRD": "https://www.gov.gd/", "GLP": "https://www.guadeloupe.gouv.fr/",
    "JAM": "https://www.gov.jm/", "MTQ": "https://www.martinique.gouv.fr/", "MSR": "https://www.gov.ms/",
    "PRI": "https://www.pr.gov/", "BLM": "https://www.saint-barth-saint-martin.gouv.fr/", "KNA": "https://www.gov.kn/",
    "LCA": "https://www.govt.lc/", "MAF": "https://www.saint-barth-saint-martin.gouv.fr/", "VCT": "https://www.gov.vc/",
    "SXM": "https://www.sintmaartengov.org/", "TTO": "https://www.ttconnect.gov.tt/", "TCA": "https://gov.tc/",
    "VIR": "https://www.vi.gov/", "GUF": "https://www.guyane.gouv.fr/", "FLK": "https://www.falklands.gov.fk/",
    "SGS": "https://www.gov.gs/", "BVT": "https://www.regjeringen.no/en/",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def clean_markup(value: str) -> str:
    value = re.sub(r"<[^>]+>", " ", value)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def first_link(value: str) -> str | None:
    match = re.search(r'href\s*=\s*["\']\s*([^"\']+)', value, re.I)
    return html.unescape(match.group(1).strip()) if match else None


def parse_nso(path: Path) -> dict[str, dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    result = {}
    pattern = re.compile(r'<li class="leaf"><strong>(.*?)</strong>(.*?)(?=<li class="leaf"><strong>|</ul>)', re.I | re.S)
    for match in pattern.finditer(text):
        name = clean_markup(match.group(1))
        body = match.group(2)
        links = re.findall(r'<a[^>]+href\s*=\s*["\']\s*([^"\']+)["\'][^>]*>(.*?)</a>', body, re.I | re.S)
        if links:
            url, title = links[0]
            result[name] = {"name": clean_markup(title), "url": html.unescape(url).strip()}
    return result


class CensusDateParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.div_depth = 0
        self.section_depth = None
        self.row_depth = None
        self.cell_depth = None
        self.current_cell = None
        self.current_row = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "div":
            self.div_depth += 1
        depth = self.div_depth
        if tag == "div" and attrs.get("id") in {"North", "South"}:
            self.section_depth = depth
        if self.section_depth and tag == "div" and "row" in attrs.get("class", "").split() and self.row_depth is None:
            self.row_depth = depth
            self.current_row = []
        elif self.row_depth and tag == "div" and "col-md-2" in attrs.get("class", "").split() and self.cell_depth is None:
            self.cell_depth = depth
            self.current_cell = {"text": [], "urls": []}
        if self.current_cell is not None and tag == "a" and attrs.get("href"):
            self.current_cell["urls"].append(attrs["href"].strip())

    def handle_data(self, data):
        if self.current_cell is not None:
            self.current_cell["text"].append(data)

    def handle_endtag(self, tag):
        depth = self.div_depth
        if tag == "div" and self.cell_depth == depth:
            self.current_cell["label"] = re.sub(r"\s+", " ", " ".join(self.current_cell.pop("text"))).strip()
            self.current_row.append(self.current_cell)
            self.current_cell = None
            self.cell_depth = None
        if tag == "div" and self.row_depth == depth:
            if self.current_row:
                self.rows.append(self.current_row)
            self.current_row = None
            self.row_depth = None
        if tag == "div" and self.section_depth == depth:
            self.section_depth = None
        if tag == "div":
            self.div_depth -= 1


def parse_census_dates(path: Path) -> dict[str, list[dict]]:
    parser = CensusDateParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    result = {}
    rounds = ["1990", "2000", "2010", "2020", "2030"]
    for cells in parser.rows:
        if len(cells) < 6:
            continue
        country = cells[0]["label"]
        if "Countries or areas" in country:
            country = country.split("Countries or areas", 1)[1].strip()
        if not country:
            continue
        entries = []
        for round_name, cell in zip(rounds, cells[1:6]):
            label = cell["label"]
            date_only = re.sub(r'\b\d{4}\s+round\s*\(\d{4}-\d{4}\)', '', label, flags=re.I)
            years = re.findall(r'(?<!\d)(?:19|20)\d{2}(?!\d)', date_only)
            if not years:
                continue
            entries.append({
                "round": round_name,
                "date_text": label,
                "year": int(years[-1]),
                "url": cell["urls"][0] if cell["urls"] else None,
                "source": COMMON["unsd_census"],
            })
        result[country] = entries
    return result


def state(status: str, *, urls=None, note=None, evidence=None) -> dict:
    order = ["identified", "accessed", "acquired", "inspected", "geography_matched", "adopted"]
    flags = {name: False for name in order}
    if status in order:
        for name in order[: order.index(status) + 1]:
            flags[name] = True
    flags.update({"unavailable": status == "unavailable", "restricted": status == "restricted", "failed_with_evidence": status == "failed_with_evidence"})
    result = {"status": status, **flags, "urls": [u for u in (urls or []) if u]}
    if note:
        result["note"] = note
    if evidence:
        result["evidence"] = evidence
    return result


def source_map(research: dict) -> dict:
    return {item["id"]: item for item in research["sources"]}


def latin_country_map(research: dict) -> dict:
    return {item["iso"]: item for item in research["countries"]}


def central_map(history: dict) -> dict:
    return {item["country_id"]: item for item in history["countries"]}


def cell_preview(ws, max_rows=12, max_cols=30):
    header_row = None
    fields = []
    numeric = 0
    nonempty = 0
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols), values_only=True), 1):
        values = [v for v in row]
        non = [v for v in values if v not in (None, "")]
        if header_row is None and non:
            header_row = row_no
            fields = [str(v)[:160] for v in non]
    for row in ws.iter_rows(values_only=True):
        for value in row:
            if value not in (None, ""):
                nonempty += 1
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric += 1
    return header_row, fields, nonempty, numeric


def inspect_files(project: Path, wpp_path: Path) -> list[dict]:
    inventory = []
    raw = project / "raw"
    paths = sorted(p for p in raw.rglob("*") if p.is_file() and "source-discovery" not in p.parts)
    for p in paths:
        rel = p.relative_to(project).as_posix()
        ext = p.suffix.lower()
        base = {
            "path": rel, "bytes": p.stat().st_size, "sha256": sha256(p), "format": ext.lstrip(".") or "none",
            "acquisition_status": "acquired", "inspection_status": "structurally_inspected",
            "publication": "evidence_only_not_in_public_site",
        }
        try:
            if ext == ".xlsx":
                wb = load_workbook(p, read_only=True, data_only=True)
                sheets = []
                for ws in wb.worksheets:
                    header_row, fields, nonempty, numeric = cell_preview(ws)
                    sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "first_nonempty_row": header_row, "field_preview": fields, "nonempty_cells": nonempty, "numeric_cells": numeric})
                base["sheets"] = sheets
                base["semantic_inventory_status"] = "all_sheets_structurally_inventoried; adopted fields documented in SOURCE_DISPOSITION.csv"
            elif ext == ".pdf":
                reader = PdfReader(str(p))
                base["pages"] = len(reader.pages)
                base["semantic_inventory_status"] = "all pages accounted for; full table semantics not inferred automatically"
            elif ext == ".json":
                data = json.loads(p.read_text(encoding="utf-8"))
                base["top_level_type"] = type(data).__name__
                if isinstance(data, dict):
                    base["top_level_keys"] = list(data.keys())[:200]
                elif isinstance(data, list):
                    base["record_count"] = len(data)
                base["semantic_inventory_status"] = "top-level structure inspected; adopted fields documented in SOURCE_DISPOSITION.csv"
            else:
                base["semantic_inventory_status"] = "stored source payload; no semantic field adoption"
        except Exception as exc:
            base["inspection_status"] = "inspection_failed_with_evidence"
            base["inspection_error"] = f"{type(exc).__name__}: {exc}"
        inventory.append(base)

    # Keep the evidence package self-contained. The command may receive the
    # acquisition-cache path, while the same workbook is copied under raw/.
    # Match by content hash, enrich that project-relative row, and never record
    # an external working path in the deliverable inventory.
    p = wpp_path.resolve()
    wpp_hash = sha256(p)
    matching = next((item for item in inventory if item["sha256"] == wpp_hash), None)
    try:
        wpp_relative = p.relative_to(project.resolve()).as_posix()
    except ValueError:
        wpp_relative = None
    if matching is None and wpp_relative is None:
        raise RuntimeError("The adopted WPP workbook must be copied inside the project raw directory before evidence is built")

    wb = load_workbook(p, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        header_values = [str(v)[:200] if v is not None else "" for v in next(ws.iter_rows(min_row=17, max_row=17, values_only=True))]
        sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "declared_header_row": 17, "fields": header_values})
    if matching is not None:
        matching["inspection_status"] = "inspected"
        matching["sheets"] = sheets
        matching["semantic_inventory_status"] = "all workbook sheets and row-17 fields inventoried; total population field adoption is in UN_WPP_AMERICAS_ADOPTION_AUDIT.json"
    else:
        inventory.append({
            "path": wpp_relative, "bytes": p.stat().st_size, "sha256": wpp_hash, "format": "xlsx",
            "acquisition_status": "acquired", "inspection_status": "inspected", "sheets": sheets,
            "semantic_inventory_status": "all workbook sheets and row-17 fields inventoried; total population field adoption is in UN_WPP_AMERICAS_ADOPTION_AUDIT.json",
            "publication": "evidence_only_not_in_public_site",
        })
    return inventory


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True, type=Path)
    ap.add_argument("--unsd-nso-html", required=True, type=Path)
    ap.add_argument("--unsd-census-html", required=True, type=Path)
    ap.add_argument("--latin-research", required=True, type=Path)
    ap.add_argument("--central-history", required=True, type=Path)
    ap.add_argument("--wpp", required=True, type=Path)
    args = ap.parse_args()

    project = args.project.resolve()
    evidence = project / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    dataset = json.loads((project / "data/dashboard.json").read_text(encoding="utf-8"))
    research = json.loads(args.latin_research.read_text(encoding="utf-8"))
    history = json.loads(args.central_history.read_text(encoding="utf-8"))
    smap, lmap, cmap = source_map(research), latin_country_map(research), central_map(history)
    nso, dates = parse_nso(args.unsd_nso_html), parse_census_dates(args.unsd_census_html)
    country_rows = [t for t in dataset["territories"] if t.get("level") == "country"]
    boundary_ids = {f["properties"]["territory_id"] for f in dataset["boundaries"]["features"]}
    source_by_id = {s["id"]: s for s in dataset["sources"]}
    observations = dataset["observations"]

    discovery_dir = project / "raw/source-discovery"
    discovery_dir.mkdir(parents=True, exist_ok=True)
    receipts = []
    for src, name, url in ((args.unsd_nso_html, "unsd-nso-directory.html", COMMON["unsd_nso"]), (args.unsd_census_html, "unsd-census-dates.html", COMMON["unsd_census"])):
        dst = discovery_dir / name
        if src.resolve() != dst.resolve():
            shutil.copy2(src, dst)
        receipts.append({"url": url, "retrieved_at": datetime.fromtimestamp(src.stat().st_mtime, timezone.utc).isoformat(), "path": dst.relative_to(project).as_posix(), "bytes": dst.stat().st_size, "sha256": sha256(dst), "terms_note": "Official discovery-page snapshot retained as audit evidence; not included in public site."})
    (discovery_dir / "receipt.json").write_text(json.dumps({"schema_version": "0.1", "files": receipts}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    preflight = []
    planning_rows = []
    geo_rows = []
    for territory in country_rows:
        iso = territory["id"]
        name = territory["name"]
        lookup = ALIASES.get(name, name)
        nso_record = nso.get(lookup)
        if iso in MANUAL_NSO:
            nso_record = {"name": MANUAL_NSO[iso][0], "url": MANUAL_NSO[iso][1]}
        census_entries = dates.get(lookup, [])
        latin = lmap.get(iso)
        central = cmap.get(iso)
        census_urls = []
        if latin:
            census_urls = [smap[s]["url"] for s in latin.get("census_sources", []) if s in smap]
        if central:
            census_urls = [central["official_census_url"], central["adopted_source_url"]] + census_urls
        census_urls = list(dict.fromkeys(MANUAL_CENSUS.get(iso, []) + [u for u in census_urls if u]))
        current_year = datetime.now(timezone.utc).year
        future_rounds = sorted((row for row in census_entries if row.get("year", 0) > current_year), key=lambda x: x["year"], reverse=True)
        completed_rounds = sorted((row for row in census_entries if row.get("year", 0) <= current_year), key=lambda x: x["year"], reverse=True)
        if central:
            census_status = "adopted"
            recent = central["recent_rounds"][:3]
            latest_year = central["adopted_data_year"]
            scheduled = f" A later scheduled round ({future_rounds[0]['year']}) is identified, but its results are not acquired or usable." if future_rounds else ""
            census_note = f"Adopted census data year: {latest_year}. The result was acquired, inspected, geographically matched and integrated.{scheduled}"
        elif latin:
            census_status = "accessed"
            recent = sorted(census_entries, key=lambda x: x["year"], reverse=True)[:3]
            if not recent and latin.get("census_year"):
                recent = [{"year": latin["census_year"], "status": "research_catalog_round", "url": census_urls[0] if census_urls else None}]
            latest_year = latin.get("usable_year") or latin.get("census_year")
            scheduled = f" A later scheduled round ({future_rounds[0]['year']}) is identified, but its results are not acquired or usable." if future_rounds else ""
            census_note = f"Latest reviewed census/result year: {latest_year if latest_year else 'not yet confirmed'}. Result location accessed; data not integrated.{scheduled}"
        elif census_entries:
            census_status = "accessed"
            recent = sorted(census_entries, key=lambda x: x["year"], reverse=True)[:3]
            latest_year = completed_rounds[0]["year"] if completed_rounds else None
            if future_rounds:
                census_note = f"Latest scheduled/identified round: {future_rounds[0]['year']}. Results are not acquired or usable. Latest non-future round identified: {latest_year if latest_year else 'not yet confirmed'}."
            else:
                census_note = f"Latest identified non-future census year: {latest_year if latest_year else 'not yet confirmed'}. Result tables have not been acquired or inspected."
        else:
            census_status = "identified"
            recent = []
            latest_year = None
            census_note = "No census year has yet been confirmed. Start location identified; result tables not acquired or inspected."

        planning_urls = []
        laws = []
        if latin:
            for law in latin.get("laws", []):
                source = smap.get(law.get("source"), {})
                row = {"name": law.get("name"), "year": law.get("year"), "summary": law.get("summary"), "source_id": law.get("source"), "url": source.get("url")}
                laws.append(row)
                if row["url"]:
                    planning_urls.append(row["url"])
            for sid in latin.get("planning_profile", {}).get("relation_sources", []):
                if sid in smap:
                    planning_urls.append(smap[sid]["url"])
        gov_start = GOV_START.get(iso)
        if not gov_start and nso_record:
            parsed = urlparse(nso_record["url"])
            gov_start = f"{parsed.scheme or 'https'}://{parsed.netloc}/" if parsed.netloc else None
        planning_urls = list(dict.fromkeys(planning_urls or [gov_start, COMMON["eclac_planning"]]))
        planning_status = "accessed" if laws else "identified"
        planning_note = latin.get("planning_profile", {}).get("gaps") if latin else "Country/area-specific planning law, guidance and adopted plan documents have not yet been accessed; listed URLs are research starting points, not evidence of absence."

        admin_note = latin.get("published_geography") if latin else "Official administrative codes and ADM1/ADM2 statistical correspondence have not yet been inspected for the country adapter."
        machine_note = latin.get("formats") if latin else "Machine-readable census tables have not yet been inspected for the country adapter."
        entry = {
            "country_area_id": iso, "m49": territory.get("official_code"), "name": name, "iso2": territory.get("iso2"),
            "scope_role": "regional_gateway_member; not a completed country edition",
            "official_statistics_office": state("accessed" if nso_record else "identified", urls=[nso_record["url"]] if nso_record else [COMMON["unsd_nso"]], note=nso_record["name"] if nso_record else "UNSD NSO directory is the discovery start; country-specific office record requires follow-up."),
            "latest_census": state(census_status, urls=census_urls or [COMMON["unsd_census"]], note=census_note),
            "recent_census_rounds": recent,
            "census_results": state("adopted" if central else ("accessed" if latin else "identified"), urls=census_urls or ([nso_record["url"]] if nso_record else [COMMON["unsd_census"]]), note="Integrated into the AreaData census series." if central else "Result location reviewed in prior Latin America research; data not integrated." if latin else "Start location identified; result tables not yet acquired or inspected."),
            "table_catalog": state("inspected" if central else ("accessed" if latin else "identified"), urls=census_urls or [nso_record["url"] if nso_record else COMMON["unsd_census"]], note="Acquired files are itemized in SOURCE_TABLE_INVENTORY.json." if central else "Catalog/location only; no claim of complete table inventory."),
            "machine_readable_data": state("adopted" if central else ("accessed" if latin else "identified"), urls=census_urls or [nso_record["url"] if nso_record else COMMON["unsd_census"]], note=machine_note),
            "administrative_codes": state("geography_matched" if central else ("accessed" if latin else "identified"), urls=census_urls or [nso_record["url"] if nso_record else COMMON["unsd_nso"]], note=admin_note),
            "adm1_adm2_boundaries": state("geography_matched" if central else "identified", urls=[source_by_id.get("natural-earth", {}).get("url")], note="Country reference boundary joined exactly." if iso in boundary_ids else "Reference country polygon not joined; official ADM1/ADM2 boundaries remain a country-adapter task."),
            "planning_law": state(planning_status, urls=planning_urls, note=planning_note),
            "planning_guidance": state("accessed" if latin else "identified", urls=planning_urls, note="Prior Latin America research describes the planning chain but does not complete all subnational guidance." if latin else planning_note),
            "plans_budgets_implementation_evaluation": state("identified", urls=planning_urls, note="Start URLs only. No plan, budget, implementation report or evaluation is adopted in this regional gateway."),
            "country_adapter_status": "integrated_census_and_local_hierarchy" if central else "preflight_only",
        }
        preflight.append(entry)
        planning_rows.append({
            "country_area_id": iso, "name": name, "research_status": planning_status,
            "law_count_reviewed": len(laws), "law_names": " | ".join(str(x["name"]) for x in laws),
            "law_urls": " | ".join(str(x["url"]) for x in laws if x.get("url")),
            "planning_start_urls": " | ".join(u for u in planning_urls if u),
            "guidance_status": "accessed_prior_research" if latin else "not_accessed_start_only",
            "plans_budgets_implementation_evaluation_status": "not_acquired",
            "adoption_status": "not_adopted_in_regional_gateway", "note": planning_note,
        })
        geo_rows.append({
            "country_area_id": iso, "name": name, "m49": territory.get("official_code"),
            "reference_country_boundary": "exact_join" if iso in boundary_ids else "not_joined",
            "reference_boundary_source": "Natural Earth Admin-0 map units",
            "official_adm1_adm2_status": "matched_for_integrated_census" if central else "not_yet_matched",
            "local_hierarchy_integrated": "yes" if central else "no",
            "statistical_to_administrative_crosswalk": "verified_for_adopted_population_series" if central else "not_yet_verified",
            "note": "Country reference geometry is display-only and is not used as proof of legal or statistical subnational boundaries.",
        })

    preflight_doc = {
        "schema_version": "0.2", "scope_id": "M49:019", "title": "AreaData Americas exploration gateway source preflight",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "scope_statement": "This inventory covers 57 UN M49 country/area members as research starts. It does not claim 57 completed country dashboards.",
        "state_contract": "identified, accessed, acquired, inspected, geography_matched and adopted are separate states. unavailable, restricted and failed_with_evidence require evidence. Unresearched is never described as nonexistent.",
        "summary": {"country_area_records": len(preflight), "integrated_country_adapters": sum(x["country_adapter_status"].startswith("integrated") for x in preflight), "prior_latin_america_research_records": sum(x["country_area_id"] in lmap for x in preflight)},
        "common_start_sources": COMMON, "countries": preflight,
    }
    (evidence / "SOURCE_PREFLIGHT.json").write_text(json.dumps(preflight_doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [
        "# AreaData Americas exploration gateway — source preflight", "",
        "This is a research and evidence register for all 57 UN M49 Americas country/area members. It is not a declaration that 57 country editions are complete.", "",
        "Statuses are deliberately separate: identified → accessed → acquired → inspected → geography matched → adopted. Unresearched does not mean unavailable or nonexistent.", "",
        f"- Country/area records: {len(preflight)}", f"- Census/local adapters integrated: {sum(x['country_adapter_status'].startswith('integrated') for x in preflight)}", f"- Prior Latin America research profiles: {sum(x['country_area_id'] in lmap for x in preflight)}", "",
        "| ID | Country/area | Census year | Census evidence | Local adapter | Planning evidence |", "|---|---|---:|---|---|---|",
    ]
    for x in preflight:
        year = next((str(row.get("year")) for row in x["recent_census_rounds"] if row.get("status") == "results_adopted"), None)
        if not year:
            non_future = [row.get("year") for row in x["recent_census_rounds"] if isinstance(row.get("year"), int) and row.get("year") <= datetime.now(timezone.utc).year]
            year = str(max(non_future)) if non_future else "not yet confirmed"
        lines.append(f"| {x['country_area_id']} | {x['name']} | {year} | {x['census_results']['status']} | {x['country_adapter_status']} | {x['planning_law']['status']} |")
    lines += ["", "See SOURCE_TABLE_INVENTORY.json, SOURCE_DISPOSITION.csv, GEOGRAPHY_CROSSWALK.csv, PLANNING_LEGAL_INVENTORY.csv and raw/source-discovery/receipt.json for traceability.", ""]
    (evidence / "SOURCE_PREFLIGHT.md").write_text("\n".join(lines), encoding="utf-8")

    inventory = inspect_files(project, args.wpp)
    inventory_doc = {
        "schema_version": "0.2", "generated_at": datetime.now(timezone.utc).isoformat(),
        "scope_statement": "Every acquired raw file in the project is accounted for structurally. Structural PDF inventory is not represented as a complete semantic table inventory.",
        "file_count": len(inventory), "format_counts": dict(Counter(x["format"] for x in inventory)), "files": inventory,
    }
    (evidence / "SOURCE_TABLE_INVENTORY.json").write_text(json.dumps(inventory_doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    dispositions = []
    obs_by_source = defaultdict(list)
    for obs in observations:
        obs_by_source[obs.get("source_id")].append(obs)
    for source in dataset["sources"]:
        obs = obs_by_source.get(source["id"], [])
        dispositions.append({
            "record_type": "dataset_source", "country_area_id": "MULTI" if len({o["territory_id"].split(":")[0] for o in obs}) > 1 else (obs[0]["territory_id"].split(":")[0] if obs else ""),
            "source_id": source["id"], "source_title": source.get("title", ""), "source_url": source.get("url", ""),
            "source_table_or_field": source.get("indicator_id", "multiple/see observations"),
            "indicator_id": " | ".join(sorted({o["indicator_id"] for o in obs})), "periods": " | ".join(sorted({str(o["period"]) for o in obs})),
            "geography": " | ".join(sorted({o["territory_id"] for o in obs})[:10]) + (" | …" if len({o["territory_id"] for o in obs}) > 10 else ""),
            "observation_count": len(obs), "disposition": "adopted" if obs else "not_adopted", "reason": "Used by dashboard observations." if obs else "No observations adopted from this source record.",
            "terms_status": source.get("terms_review_status") or source.get("license") or "not_stated", "raw_redistribution": source.get("raw_redistribution_status") or "not_publicly_deployed",
        })
    wpp_audit_path = evidence / "UN_WPP_AMERICAS_ADOPTION_AUDIT.json"
    if wpp_audit_path.exists():
        wpp_audit = json.loads(wpp_audit_path.read_text(encoding="utf-8"))
        rows = wpp_audit.get("countries") or wpp_audit.get("country_areas") or wpp_audit.get("records") or []
        wpp_source_id = next((source["id"] for source in dataset["sources"] if source["id"] == "un-wpp2024-demographic-indicators-rev1"), "un-wpp2024-demographic-indicators-rev1")
        wpp_source = source_by_id.get(wpp_source_id, {})
        for item in rows:
            country_area_id = item.get("country_id") or item.get("country_area_id") or item.get("id") or item.get("iso3")
            dispositions.append({
                "record_type": "wpp_country_area_adoption", "country_area_id": country_area_id,
                "source_id": wpp_source_id, "source_title": wpp_source.get("name", "UN World Population Prospects 2024, demographic indicators, Rev.1"), "source_url": wpp_source.get("url", ""),
                "source_table_or_field": item.get("value_column") or item.get("source_column") or item.get("field") or "Total Population, as of 1 July (thousands)",
                "indicator_id": "UN_WPP_POP_TOTAL", "periods": "2023 | 2024 | 2025 | 2026", "geography": country_area_id or "",
                "observation_count": item.get("observation_count", 4 if item.get("status") == "adopted" else 0), "disposition": item.get("status", "documented_in_wpp_audit"),
                "reason": item.get("reason", "See UN_WPP_AMERICAS_ADOPTION_AUDIT.json"), "terms_status": "UN publication; attribution retained", "raw_redistribution": "not_in_public_site",
            })
    fields = ["record_type", "country_area_id", "source_id", "source_title", "source_url", "source_table_or_field", "indicator_id", "periods", "geography", "observation_count", "disposition", "reason", "terms_status", "raw_redistribution"]
    write_csv(evidence / "SOURCE_DISPOSITION.csv", fields, dispositions)
    write_csv(evidence / "GEOGRAPHY_CROSSWALK.csv", list(geo_rows[0].keys()), geo_rows)
    write_csv(evidence / "PLANNING_LEGAL_INVENTORY.csv", list(planning_rows[0].keys()), planning_rows)

    summary = {
        "preflight_records": len(preflight), "inventory_files": len(inventory), "source_disposition_rows": len(dispositions),
        "geography_crosswalk_rows": len(geo_rows), "planning_inventory_rows": len(planning_rows),
        "discovery_receipts": len(receipts), "generated_at": datetime.now(timezone.utc).isoformat(),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
