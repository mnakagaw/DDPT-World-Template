"""Read-only comparison of AreaData's UBOS-located observations to the official XLSX.

This checks source cells and recorded numerator/denominator inputs. It does not
independently adjudicate indicator meaning, geography joins, PDF tables or laws.
"""

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / "generated" / "uganda-areadata-20260924"
DATASET = PROJECT / "data" / "dashboard.json"
WORKBOOK = PROJECT / "raw" / "NPHC-2024-Subcounty-Profiles-Excel-Tables.xlsx"
CELL = re.compile(r"^[A-Z]+[1-9][0-9]*$")


def main():
    dataset = json.loads(DATASET.read_text(encoding="utf-8"))
    checks = defaultdict(list)
    observation_count = 0
    without_workbook_locator = 0
    for observation in dataset["observations"]:
        if observation.get("status") != "observed":
            continue
        locator = observation.get("source_locator") or {}
        if not locator.get("sheet") or not CELL.fullmatch(locator.get("cell", "")):
            without_workbook_locator += 1
            continue
        observation_count += 1
        label = f"{observation['territory_id']}/{observation['indicator_id']}"

        def add(sheet, cell, expected, kind):
            if sheet and cell and expected is not None:
                checks[(sheet, *coordinate_to_tuple(cell))].append((float(expected), label, kind))

        numerator_cell = locator.get("numeratorCell")
        denominator_cell = locator.get("denominatorCell")
        numerator = observation.get("numerator")
        denominator = observation.get("denominator")
        if numerator_cell:
            add(locator.get("numeratorSheet", locator["sheet"]), numerator_cell, numerator, "numerator")
        if denominator_cell:
            add(locator.get("denominatorSheet", locator["sheet"]), denominator_cell, denominator, "denominator")
        if denominator is not None and not numerator_cell:
            add(locator["sheet"], locator["cell"], numerator, "numerator at source cell")
        elif denominator is None or numerator_cell:
            add(locator["sheet"], locator["cell"], observation["value"], "published value")

    targets = defaultdict(lambda: defaultdict(dict))
    for (sheet, row, col), expected in checks.items():
        targets[sheet][row][col] = expected
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    mismatches = []
    matched = 0
    for sheet_name, rows in targets.items():
        if sheet_name not in workbook:
            mismatches.append(f"missing worksheet {sheet_name}")
            continue
        sheet = workbook[sheet_name]
        max_row = max(rows)
        for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=max_row, values_only=True), 1):
            if row_number not in rows:
                continue
            for col, expected_items in rows[row_number].items():
                actual = values[col - 1] if col <= len(values) else None
                for expected, label, kind in expected_items:
                    if isinstance(actual, (int, float)) and abs(actual - expected) <= 0.00002:
                        matched += 1
                    elif len(mismatches) < 30:
                        mismatches.append(f"{sheet_name}!{row_number}:{col} {label} {kind}: expected {expected}, got {actual!r}")
        del sheet
    workbook.close()
    report = {
        "workbook_located_observations": observation_count,
        "other_observed_without_workbook_cell": without_workbook_locator,
        "source_value_checks": sum(len(items) for items in checks.values()),
        "matched_checks": matched,
        "mismatch_examples": mismatches,
        "verdict": "PASS_SOURCE_CELLS_ONLY" if matched == sum(len(items) for items in checks.values()) and not mismatches else "FAIL",
    }
    print(json.dumps(report, indent=2))
    return 0 if report["verdict"] == "PASS_SOURCE_CELLS_ONLY" else 1


if __name__ == "__main__":
    sys.exit(main())
