"""Extract exact M49 GDP and GDP-per-person rows from UNSD AMA workbooks.

The workbooks include published world, region, and country/area values. No
country aggregation or per-capita averaging is performed here.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


FILES = (
    ("current", "UN_AMA_GDP_CURRENT_USD", "GDP, current US dollars", "Economy", "US$", "https://unstats.un.org/unsd/amaapi/api/file/4", "Gross Domestic Product (GDP)"),
    ("per_capita", "UN_AMA_GDP_PER_CAPITA_USD", "GDP per capita, current US dollars", "Economy", "US$ per person", "https://unstats.un.org/unsd/amaapi/api/file/11", None),
    ("constant", "UN_AMA_GDP_CONSTANT_2020_USD", "GDP, constant 2020 US dollars", "Economy", "2020 US$", "https://unstats.un.org/unsd/amaapi/api/file/8", "Gross Domestic Product (GDP)"),
)


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    for kind, *_ in FILES:
        parser.add_argument(f"--{kind.replace('_', '-')}", type=Path, required=True)
    parser.add_argument("--registry", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    registry = json.loads(args.registry.read_text(encoding="utf-8"))
    code_to_id = {int(area["official_code"]): area["id"] for area in registry["territories"] if area.get("official_code") and area["id"] != "CUSTOM:CAM-CAR"}
    code_to_id[1] = "WLD"
    observations, indicators, files, counts = [], [], [], {}
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    for kind, indicator_id, name, theme, unit, url, item_name in FILES:
        path: Path = getattr(args, kind)
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        title = str(next(rows)[0] or "")
        next(rows)
        header = next(rows)
        if "all countries and regions" not in title or not any(year == 2024 for year in header):
            raise ValueError(f"Unexpected UNSD AMA workbook: {path}")
        years = [(index, str(value)) for index, value in enumerate(header) if isinstance(value, int) and 2015 <= value <= 2024]
        if not years or years[-1][1] != "2024":
            raise ValueError(f"Expected 2015–2024 in {path}")
        seen = set()
        for row in rows:
            if item_name and row[2] != item_name:
                continue
            try:
                area_id = code_to_id[int(row[0])]
            except (ValueError, TypeError, KeyError):
                continue
            for index, year in years:
                raw = row[index] if index < len(row) else None
                if not isinstance(raw, (int, float)) or isinstance(raw, bool):
                    continue
                value = float(raw)
                if value < 0 or value != value or value == float("inf"):
                    continue
                key = area_id, indicator_id, year
                if key in seen:
                    raise ValueError(f"Duplicate UNSD AMA observation {key}")
                seen.add(key)
                observations.append({"territory_id": area_id, "indicator_id": indicator_id, "period": year, "value": value, "status": "observed", "source_id": "unsd-ama-2024"})
        counts[indicator_id] = len(seen)
        if not seen or ("WLD", indicator_id, "2024") not in seen:
            raise ValueError(f"No published World 2024 value: {indicator_id}")
        indicators.append({"id": indicator_id, "name": name, "theme": theme, "unit": unit,
            "definition": f"{title}; selected item {item_name or 'per capita GDP'}; published exact geographic row in UNSD National Accounts Main Aggregates. Current-price and constant-price concepts are separate. Country per-capita values are never averaged into a region.",
            "definition_id": f"unsd-ama-{kind}-2024", "population": "World, M49 region or country/area represented by the published row", "measurement_method": "UNSD National Accounts Main Aggregates, published aggregate", "aggregation": "official_only", "series_family": "international_reference", "display_role": "supplementary", "period_policy": "same_period", "display_decimals": 0 if kind != "per_capita" else 2, "visualization": "trend", "source_id": "unsd-ama-2024"})
        files.append({"url": url, "sha256": digest(path), "bytes": path.stat().st_size, "source_title": title})
        workbook.close()
    source = {"id": "unsd-ama-2024", "name": "UNSD National Accounts Main Aggregates — GDP through 2024", "publisher": "United Nations Statistics Division", "url": "https://unstats.un.org/unsd/snaama/downloads", "status": "ready", "retrieved_at": now, "reference_period": "2015–2024", "geographic_level": "world_multi_scope_series", "license": "UNSD reuse terms; original country data and footnotes remain with source workbooks", "raw_redistribution_status": "not_in_public_site", "source_files": files, "note": "Only exact M49-matched published rows are adopted; IMF inflation and GDP are separate source families. The custom AreaData region remains missing unless a separately approved same-period complete cover is configured."}
    result = {"schema_version": "1.0", "generated_at": now, "source": source, "indicators": indicators, "observations": sorted(observations, key=lambda row: (row["territory_id"], row["indicator_id"], row["period"])), "summary": {"observations": len(observations), "by_indicator": counts, "geographic_ids": len({row["territory_id"] for row in observations}), "periods": [str(year) for year in range(2015, 2025)]}}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
