#!/usr/bin/env python3
"""Create a field-level review queue for every acquired census workbook column.

The script is intentionally conservative: extracted fields remain `unreviewed` until a
human or country adapter supplies a terminal disposition and a reason.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


THEMES = {
    "age_sex": r"age|edad|sex|sexo|sexe",
    "households_housing": r"house|hogar|viviend|dwelling|tenure|room|cuarto",
    "drinking_water": r"water|agua|acueduct",
    "sanitation": r"sanit|toilet|servicio sanitario|alcantar",
    "electricity": r"electric|alumbrado|lighting",
    "education_literacy": r"educ|school|escolar|alfabet|literacy",
    "employment": r"employ|ocupaci|labor|econ[oó]mic.*activ|trabaj",
    "disability": r"disab|dificultad|limitaci",
    "migration": r"migr|birthplace|lugar de nacimiento|residencia anterior",
    "urban_rural": r"urban|rural|área|area",
    "ethnicity": r"ethnic|etnia|pueblo|ind[ií]gen|afro|raza|language|idioma",
    "population_total": r"population|poblaci[oó]n|habitantes|total",
    "health": r"health|salud|seguro|mortal|fecund|nacim",
    "nutrition": r"nutri|aliment|food",
    "poverty": r"poverty|pobreza|necesidades b[aá]sicas|mpi",
}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def country_id(path: Path) -> str:
    parts = [part.upper() for part in path.parts]
    known = {
        "AIA", "ATG", "ABW", "BHS", "BRB", "BES", "VGB", "CYM", "CUB", "CUW", "DMA", "DOM", "GRD", "GLP", "HTI", "JAM", "MTQ", "MSR", "PRI", "BLM", "KNA", "LCA", "MAF", "VCT", "SXM", "TTO", "TCA", "VIR",
        "BLZ", "CRI", "SLV", "GTM", "HND", "MEX", "NIC", "PAN",
        "ARG", "BOL", "BVT", "BRA", "CHL", "COL", "ECU", "FLK", "GUF", "GUY", "PRY", "PER", "SGS", "SUR", "URY", "VEN",
        "BMU", "CAN", "GRL", "SPM", "USA",
    }
    return next((part for part in reversed(parts) if part in known), "MULTI")


def infer_theme(text: str) -> str:
    for theme, pattern in THEMES.items():
        if re.search(pattern, text, re.I):
            return theme
    return "unclassified"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True, type=Path)
    args = parser.parse_args()
    project = args.project.resolve()
    files = sorted((project / "raw").rglob("*.xlsx"))
    records = []
    workbooks = []
    duplicate_files = []
    seen_hashes = {}
    for book_path in files:
        country = country_id(book_path)
        source_hash = digest(book_path)
        if source_hash in seen_hashes:
            duplicate_files.append({"path": book_path.relative_to(project).as_posix(), "same_content_as": seen_hashes[source_hash], "sha256": source_hash})
            continue
        seen_hashes[source_hash] = book_path.relative_to(project).as_posix()
        source_id = f"{country}_XLSX_{source_hash[:12].upper()}"
        workbook_record = {
            "country_area_id": country,
            "source_id": source_id,
            "path": book_path.relative_to(project).as_posix(),
            "sha256": source_hash,
            "tables": [],
        }
        try:
            workbook = load_workbook(book_path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                preview_rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15), values_only=True))
                numeric_by_column = {}
                for row in sheet.iter_rows(values_only=True):
                    for index, value in enumerate(row, 1):
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            numeric_by_column[index] = numeric_by_column.get(index, 0) + 1
                title_text = " | ".join(
                    str(value).strip()
                    for row in preview_rows[:5]
                    for value in row
                    if isinstance(value, str) and value.strip()
                )[:2000]
                table_id = sheet.title
                fields = []
                for column, count in sorted(numeric_by_column.items()):
                    header_parts = []
                    for row in preview_rows:
                        if column <= len(row):
                            value = row[column - 1]
                            if isinstance(value, str) and value.strip():
                                header_parts.append(value.strip())
                    header = " | ".join(dict.fromkeys(header_parts))[:1000]
                    semantic_text = f"{book_path.name} {sheet.title} {title_text} {header}"
                    field_id = get_column_letter(column)
                    theme = infer_theme(semantic_text)
                    fields.append({"field_id": field_id, "numeric_cell_count": count, "header": header, "inferred_theme": theme})
                    records.append(
                        {
                            "country_area_id": country,
                            "source_id": source_id,
                            "source_path": book_path.relative_to(project).as_posix(),
                            "table_id": table_id,
                            "table_title": title_text,
                            "field_id": field_id,
                            "field_label": header,
                            "numeric_cell_count": count,
                            "theme": theme,
                            "disposition": "unreviewed",
                            "reason": "Extracted into the semantic review queue; definition, denominator, geography and intended dashboard use have not yet been approved.",
                            "coverage_complete": False,
                        }
                    )
                workbook_record["tables"].append({"table_id": table_id, "title": title_text, "numeric_fields": fields})
            workbook.close()
            workbook_record["inspection_status"] = "field_queue_built"
        except Exception as exc:
            workbook_record["inspection_status"] = "failed_with_evidence"
            workbook_record["error"] = f"{type(exc).__name__}: {exc}"
        workbooks.append(workbook_record)
    output = project / "evidence/COUNTRY_SEMANTIC_INVENTORY.json"
    output.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "scope": "All acquired XLSX census evidence currently stored in the project raw directory.",
                "completion_warning": "unreviewed is deliberately non-terminal and cannot satisfy the country completion gate.",
                "workbook_count": len(workbooks),
                "record_count": len(records),
                "duplicate_file_count": len(duplicate_files),
                "duplicate_files": duplicate_files,
                "workbooks": workbooks,
                "records": records,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"workbooks": len(workbooks), "records": len(records), "failed_workbooks": sum(row.get("inspection_status") == "failed_with_evidence" for row in workbooks), "output": str(output)}, indent=2))


if __name__ == "__main__":
    main()
