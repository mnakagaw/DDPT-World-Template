"""Audit every numeric cell in ORGI's 2011 state/district PCA workbook.

This is the 2011 statistical geography, not the later LGD or 2027 census.
"""

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


SOURCE = "https://censusindia.gov.in/nada/index.php/catalog/6191/download/9268/DDW_PCA0000_2011_Indiastatedist.xlsx"
EXPECTED_HASH = "7a8f70d46c43b5dd30eefe920f251752f69f6af8589594c21ac06499074646f5"
SELECTED_TOTAL_FIELDS = (
    "No_HH", "TOT_P", "TOT_M", "TOT_F", "P_06", "M_06", "F_06",
    "P_SC", "M_SC", "F_SC", "P_ST", "M_ST", "F_ST",
    "P_LIT", "M_LIT", "F_LIT", "TOT_WORK_P", "TOT_WORK_M", "TOT_WORK_F",
    "MAINWORK_P", "MARGWORK_P", "NON_WORK_P",
)


def require(test, message):
    if not test:
        raise ValueError(message)


def load(project):
    file = project / "raw/DDW_PCA0000_2011_Indiastatedist.xlsx"
    checksum = hashlib.sha256(file.read_bytes()).hexdigest()
    require(checksum == EXPECTED_HASH, "Official PCA workbook hash changed")
    book = openpyxl.load_workbook(file, read_only=True, data_only=True)
    require([(sheet.title, sheet.max_row, sheet.max_column) for sheet in book] ==
            [("Sheet1", 2029, 94), ("Sheet2", 1, 1), ("Sheet3", 1, 1)],
            "PCA sheets or dimensions changed")
    lines = book["Sheet1"].iter_rows(values_only=True)
    fields = next(lines)
    require(len(fields) == 94 and fields[:9] ==
            ("State", "District", "Subdistt", "Town/Village", "Ward", "EB",
             "Level", "Name", "TRU"), "PCA headers changed")
    require(len(set(fields)) == len(fields) and set(SELECTED_TOTAL_FIELDS) <= set(fields[9:]),
            "PCA numeric columns changed")
    entries = {}
    counts = Counter()
    for excel_row, cells in enumerate(lines, start=2):
        require(len(cells) == 94, f"Wrong PCA width: {excel_row}")
        state, district, subdistrict, village, ward, eb, level, name, tru = cells[:9]
        require(isinstance(state, str) and isinstance(district, str) and
                subdistrict == "00000" and village == "000000" and
                ward == "0000" and eb == "000000", f"Unexpected PCA geography: {excel_row}")
        require(level in {"India", "STATE", "DISTRICT"} and tru in {"Total", "Rural", "Urban"},
                f"Unexpected PCA level/TRU: {excel_row}")
        require((level == "India" and state == "00" and district == "000") or
                (level == "STATE" and state != "00" and district == "000") or
                (level == "DISTRICT" and state != "00" and district != "000"),
                f"Unexpected PCA code pattern: {excel_row}")
        numbers = dict(zip(fields[9:], cells[9:]))
        require(all(type(value) is int and value >= 0 for value in numbers.values()),
                f"Missing/noninteger/negative PCA cell: {excel_row}")
        key = (level, state, district, tru)
        require(key not in entries, f"Duplicate PCA area/TRU: {key}")
        entries[key] = {"name": name, "row": excel_row, "values": numbers}
        counts[(level, tru)] += 1
    require(len(entries) == 2028 and counts == Counter({
        (level, tru): expected for level, expected in
        (("India", 1), ("STATE", 35), ("DISTRICT", 640))
        for tru in ("Total", "Rural", "Urban")}), "PCA row ledger changed")
    return fields, entries, checksum


def audit(fields, entries):
    numeric = fields[9:]
    grouped = defaultdict(dict)
    for (level, state, district, tru), row in entries.items():
        grouped[(level, state, district)][tru] = row
    checks = Counter()
    for area, by_tru in grouped.items():
        require(set(by_tru) == {"Total", "Rural", "Urban"}, f"Incomplete TRU: {area}")
        require(len({row["name"] for row in by_tru.values()}) == 1,
                f"TRU labels differ: {area}")
        for field in numeric:
            require(by_tru["Total"]["values"][field] ==
                    by_tru["Rural"]["values"][field] + by_tru["Urban"]["values"][field],
                    f"Rural/urban total differs: {area} {field}")
            checks["rural_urban_field_sums"] += 1
    for field in numeric:
        for state in {key[1] for key in grouped if key[0] == "STATE"}:
            districts = [key for key in grouped if key[0] == "DISTRICT" and key[1] == state]
            require(districts, f"No district in state {state}")
            require(sum(grouped[key]["Total"]["values"][field] for key in districts) ==
                    grouped[("STATE", state, "000")]["Total"]["values"][field],
                    f"District sum differs: {state} {field}")
            checks["district_to_state_field_sums"] += 1
        require(sum(grouped[key]["Total"]["values"][field] for key in grouped if key[0] == "STATE") ==
                grouped[("India", "00", "000")]["Total"]["values"][field],
                f"State sum differs: {field}")
        checks["state_to_india_field_sums"] += 1
    sex_groups = []
    for field in numeric:
        if field.endswith("_P") and field[:-2] + "_M" in numeric and field[:-2] + "_F" in numeric:
            sex_groups.append((field, field[:-2] + "_M", field[:-2] + "_F"))
    sex_groups += [("P_06", "M_06", "F_06"), ("P_SC", "M_SC", "F_SC"),
                   ("P_ST", "M_ST", "F_ST"), ("P_LIT", "M_LIT", "F_LIT"),
                   ("P_ILL", "M_ILL", "F_ILL")]
    sex_groups = list(dict.fromkeys(sex_groups))
    for area, by_tru in grouped.items():
        for tru, row in by_tru.items():
            for people, male, female in sex_groups:
                require(row["values"][people] == row["values"][male] + row["values"][female],
                        f"Sex sum differs: {area} {tru} {people}")
                checks["sex_field_sums"] += 1
    return grouped, dict(checks), sex_groups


def main(project):
    fields, entries, checksum = load(project)
    grouped, checks, sex_groups = audit(fields, entries)
    evidence = project / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    with (evidence / "INDICATOR_INVENTORY.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_id", "source_hash", "sheet", "column", "row_class", "unit",
                         "universe", "period", "geography_type", "role", "indicator_id",
                         "decision", "reason", "locator"])
        for tru in ("Total", "Rural", "Urban"):
            for field in fields[9:]:
                adopted = ((tru == "Total" and field in SELECTED_TOTAL_FIELDS) or
                           (tru in {"Rural", "Urban"} and field == "TOT_P"))
                writer.writerow(["ind-orgi-pca2011-state-district", checksum, "Sheet1", field, tru,
                                 "source-reported count", "2011 census population or households as labeled",
                                 "2011", "India / 2011 state-UT / 2011 district", "direct_count" if adopted else "unassessed_field",
                                 "IND_CENSUS2011_" + field + ("_" + tru.upper() if tru != "Total" else "") if adopted else "",
                                 "adopted" if adopted else "priority_unassessed",
                                 "Alternative classifications retained in raw workbook; no current geography inference.",
                                 f"Sheet1 col {fields.index(field)+1}, Level/State/District/TRU row key"])
    report = {"checked_at": datetime.now(timezone.utc).isoformat(), "source_url": SOURCE,
              "source_sha256": checksum, "sheet_dimensions": [["Sheet1", 2029, 94], ["Sheet2", 1, 1], ["Sheet3", 1, 1]],
              "numeric_columns": len(fields[9:]), "source_rows": len(entries),
              "numeric_cells": len(entries) * len(fields[9:]), "areas": len(grouped),
              "state_ut_2011": 35, "district_2011": 640, "checks": checks,
              "sex_triplets": len(sex_groups), "adopted_field_class_combinations": len(SELECTED_TOTAL_FIELDS)+2,
              "unassessed_field_class_combinations": len(fields[9:])*3-len(SELECTED_TOTAL_FIELDS)-2,
              "india_population_2011": grouped[("India", "00", "000")]["Total"]["values"]["TOT_P"],
              "limitations": ["2011 state and district boundaries do not imply 2026 geography",
                              "No sub-district/village field rows in this workbook", "Other official PCA files remain unassessed"]}
    (evidence / "IND_ORGI_PCA2011_AUDIT.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True, type=Path)
    main(parser.parse_args().project)
