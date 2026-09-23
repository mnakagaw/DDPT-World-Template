"""Extract the Americas country/area population series from UN WPP 2024."""
from __future__ import annotations
import argparse, hashlib, json
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

SOURCE_ID="un-wpp2024-demographic-indicators-rev1"
REGIONAL_SOURCE_ID="un-wpp2024-south-america-region-rev1"
INDICATOR_ID="UN_WPP_POP_TOTAL"
SOURCE_URL="https://population.un.org/wpp/assets/Excel%20Files/1_Indicator%20(Standard)/EXCEL_FILES/1_General/WPP2024_GEN_F01_DEMOGRAPHIC_INDICATORS_COMPACT.xlsx"

def file_sha256(path):
    digest=hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda:handle.read(1024*1024),b""): digest.update(chunk)
    return digest.hexdigest()

def scope_countries(dataset,root_id="M49:019"):
    by_id={area["id"]:area for area in dataset["territories"]}; output=[]
    for area in dataset["territories"]:
        current=area; seen={area["id"]}
        while current:
            if current["id"]==root_id:
                if area.get("type")=="country": output.append(area["id"])
                break
            parent=current.get("parent_id")
            if not parent or parent in seen: break
            seen.add(parent); current=by_id.get(parent)
    return sorted(output)

def extract_sheet(workbook,sheet_name,years,stage,members):
    sheet=workbook[sheet_name]; header=[cell.value for cell in sheet[17]]
    columns={name:index for index,name in enumerate(header) if name is not None}
    required=["ISO3 Alpha-code","Year","Total Population, as of 1 July (thousands)"]
    if any(name not in columns for name in required): raise ValueError(f"Unexpected WPP header in {sheet_name}")
    output=[]
    for values in sheet.iter_rows(min_row=18,max_col=13,values_only=True):
        iso3,year=values[columns["ISO3 Alpha-code"]],values[columns["Year"]]
        if iso3 not in members or year not in years: continue
        thousands=Decimal(str(values[columns["Total Population, as of 1 July (thousands)"]]))
        output.append({"territory_id":iso3,"indicator_id":INDICATOR_ID,"period":str(year),"value":int(thousands*1000),"status":"observed","source_id":SOURCE_ID,"series_stage":stage,"variant":"Estimates" if stage=="estimate" else "Medium","reference_time":"1 July","source_value_thousands":float(thousands),"footnote":"UN WPP 2024 Rev.1 estimate as of 1 July." if stage=="estimate" else "UN WPP 2024 Rev.1 medium-variant projection as of 1 July."})
    return output

def extract_south_america(workbook,sheet_name,years,stage):
    sheet=workbook[sheet_name]; header=[cell.value for cell in sheet[17]]
    columns={name:index for index,name in enumerate(header) if name is not None}
    output=[]
    for values in sheet.iter_rows(min_row=18,max_col=13,values_only=True):
        year=values[columns["Year"]]
        if values[columns["Location code"]]!=931 or year not in years: continue
        if values[columns["Region, subregion, country or area *"]]!="South America" or values[columns["SDMX code**"]]!=5 or values[columns["Type"]]!="Subregion":
            raise ValueError("Unexpected UN WPP South America regional identity")
        thousands=Decimal(str(values[columns["Total Population, as of 1 July (thousands)"]]))
        output.append({"territory_id":"M49:005","indicator_id":INDICATOR_ID,"period":str(year),"value":int(thousands*1000),"status":"observed","source_id":REGIONAL_SOURCE_ID,"series_stage":stage,"variant":"Estimates" if stage=="estimate" else "Medium","reference_time":"1 July","source_value_thousands":float(thousands),"source_location_code":931,"source_sdmx_code":"005","footnote":"UN WPP 2024 Rev.1 source-reported South America subregion; not a sum of the available country rows."})
    if len(output)!=len(years): raise ValueError(f"Missing South America WPP rows in {sheet_name}: {len(output)} of {len(years)}")
    return output

def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input",required=True,type=Path); parser.add_argument("--registry",required=True,type=Path); parser.add_argument("--out",required=True,type=Path); parser.add_argument("--retrieved-at"); args=parser.parse_args()
    retrieved=args.retrieved_at or datetime.now(timezone.utc).isoformat().replace("+00:00","Z")
    members=scope_countries(json.loads(args.registry.read_text(encoding="utf-8")))
    if len(members)!=57: raise ValueError(f"Americas registry must contain 57 country/area entries; found {len(members)}")
    workbook=load_workbook(args.input,read_only=True,data_only=True); member_set=set(members)
    rows=extract_sheet(workbook,"Estimates",{2023},"estimate",member_set)+extract_sheet(workbook,"Medium variant",{2024,2025,2026},"medium_projection",member_set)
    regional_rows=extract_south_america(workbook,"Estimates",{2023},"estimate")+extract_south_america(workbook,"Medium variant",{2024,2025,2026},"medium_projection")
    found=sorted({row["territory_id"] for row in rows}); missing=sorted(member_set-set(found)); keys={(row["territory_id"],row["period"]) for row in rows}; expected={(iso3,str(year)) for iso3 in found for year in (2023,2024,2025,2026)}
    if keys!=expected: raise ValueError(f"Incomplete WPP extraction: expected {len(expected)}, found {len(keys)}")
    audit=[{"country_id":country_id,"status":"adopted" if country_id in found else "not_available_in_source","sheets":["Estimates (2023)","Medium variant (2024-2026)"],"join_column":"ISO3 Alpha-code","value_column":"Total Population, as of 1 July (thousands)","transform":"multiply published thousands by 1000 and store integer people" if country_id in found else None,"reason":"Exact ISO3 row and all four requested periods found." if country_id in found else "No exact ISO3 source row was found in either adopted sheet; the registry entry remains missing and is not treated as zero."} for country_id in members]
    output={"schema_version":"1.0","scope_id":"M49:019","generated_at":retrieved,"replaces_indicator_id":"SP.POP.TOTL","indicator":{"id":INDICATOR_ID,"name":"UN population estimate / medium projection","theme":"Population","unit":"people","definition":"Total population as of 1 July in UN World Population Prospects 2024 Rev.1; 2023 is an estimate and 2024 onward uses the medium projection variant.","definition_id":"un-wpp2024-rev1-midyear-total-population","population":"Total population under the UN WPP 2024 demographic estimation framework","measurement_method":"UN demographic estimate through 2023; medium-variant projection from 2024","aggregation":"sum","series_family":"international_reference","display_role":"context","period_policy":"same_period","series_stage_by_period":{"2023":"estimate","2024":"medium_projection","2025":"medium_projection","2026":"medium_projection"},"display_decimals":0,"source_id":SOURCE_ID},"source":{"id":SOURCE_ID,"name":"World Population Prospects 2024, demographic indicators, Rev.1","publisher":"United Nations, Department of Economic and Social Affairs, Population Division","url":SOURCE_URL,"catalog_url":"https://population.un.org/wpp/","status":"ready","retrieved_at":retrieved,"reference_period":"2023 estimate; 2024-2026 medium-variant projections","sha256":file_sha256(args.input),"bytes":args.input.stat().st_size,"license":"CC BY 3.0 IGO","license_url":"https://creativecommons.org/licenses/by/3.0/igo/","raw_redistribution_status":"not_in_public_site","normalized_observation_publication":"permitted_with_attribution","geographic_level":"world_country_series","note":"Direct extraction from the official compact XLSX. Values are total population as of 1 July and are stored in people after converting the published thousands. Projection values are not census counts."},"observations":sorted(rows,key=lambda row:(row["period"],row["territory_id"])),"summary":{"registry_country_ids":members,"country_ids":found,"missing_country_ids":missing,"periods":["2023","2024","2025","2026"],"estimate_periods":["2023"],"medium_projection_periods":["2024","2025","2026"],"country_period_observations":len(rows),"default_display_period":"2026","coverage_note":f"Official workbook rows were found for {len(found)} of {len(members)} Americas country/area registry entries. Missing registry entries remain missing and are never treated as zero.","country_source_audit":audit}}
    output["regional_source"]={**output["source"],"id":REGIONAL_SOURCE_ID,"name":"World Population Prospects 2024, South America subregion, Rev.1","geographic_level":"world_region_series","territory_id":"M49:005","source_location_code":931,"source_sdmx_code":"005","note":"Direct UN WPP 2024 South America subregion observations. The WPP workbook has no country rows for BVT and SGS; this published subregion value is adopted independently of the incomplete country cover."}
    output["regional_observations"]=regional_rows
    output["summary"]["regional_source_audit"]={"territory_id":"M49:005","source_location_code":931,"source_sdmx_code":"005","row_count":len(regional_rows),"periods":[row["period"] for row in regional_rows],"method":"Exact source-reported regional rows, distinct from country-value calculations."}
    args.out.parent.mkdir(parents=True,exist_ok=True); args.out.write_text(json.dumps(output,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(f"Wrote {args.out}: {len(rows)} observations; {len(found)}/{len(members)} countries/areas; missing {', '.join(missing) or 'none'}")

if __name__=="__main__": main()
