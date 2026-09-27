"""Audit the NESDC 2024p GPP workbook before adopting provincial values."""

import argparse
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from thailand_gpp2024_fields import SECTOR_FIELDS, SECTOR_IDS, SECTOR_LABELS


ROOT = Path(__file__).resolve().parents[1]
SHEETS = ("WK", "GRP", "Regions to GDP", "NE", "NO", "SO", "EA", "WE",
          "CE", "BKK&VIC", "CLUSTERS", "PER CAPITA")
GROUP_COUNTS = {"NORTHEASTERN": 20, "NORTHERN": 17, "SOUTHERN": 14,
                "EASTERN": 8, "WESTERN": 6, "CENTRAL": 6,
                "BANGKOK AND VICINITIES": 6}
METRICS = ("gpp_million_baht", "population_thousand_persons", "gpp_per_capita_baht")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def as_values(row, locator):
    vals = row[2:5]
    require(len(vals) == 3 and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                                   and v > 0 for v in vals), f"Invalid C:E at {locator}")
    require(abs(vals[2] - vals[0] / vals[1] * 1000) < 0.001,
            f"Published GPP per capita does not reconcile at {locator}")
    return dict(zip(METRICS, vals))


def close(actual, expected, locator):
    require(abs(actual - expected) <= max(0.0001, abs(expected) * 1e-10),
            f"GPP/population subtotal mismatch: {locator}: {actual} vs {expected}")


def main(project):
    manifest = json.loads((ROOT / "config/thailand-gpp2024-source-manifest.json").read_text(
        encoding="utf-8"))
    for source in manifest["sources"]:
        p = project / source["raw_path"]
        require(p.is_file() and p.stat().st_size == source["bytes"] and
                digest(p) == source["sha256"], f"Pinned source mismatch: {source['id']}")
    source = next(s for s in manifest["sources"] if s["id"] == "tha-nesdc-gpp2024-workbook")
    workbook = load_workbook(project / source["raw_path"], read_only=True, data_only=True)
    require(tuple(workbook.sheetnames) == SHEETS, "Workbook sheet collection changed")
    inventory = []
    workbook_province_blocks = []
    for sheet in workbook:
        number_columns = defaultdict(list)
        numeric_field_labels = set()
        blocks = []
        table_title = None
        header_examples = {}
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            label = row[0] if row else None
            if isinstance(label, str) and label.startswith("GROSS ") and label == label.upper():
                table_title = label
            if isinstance(label, str):
                if re.match(r"[0-9]{4} - [A-Z]", label):
                    blocks.append({"row": row_number, "label": label, "table_title": table_title})
                    workbook_province_blocks.append((sheet.title, row_number, label[:4], table_title))
                elif (re.match(r"[1-7] - [A-Z]", label) or "SUBREGION" in label or
                      label == "WHOLE KINGDOM") and label == label.upper():
                    blocks.append({"row": row_number, "label": label, "table_title": table_title})
            row_has_numeric_measure = False
            for col_index, value in enumerate(row, 1):
                if row_number <= 6 and isinstance(value, str) and value.strip():
                    header_examples.setdefault(col_index, value.strip()[:100])
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    number_columns[col_index].append(row_number)
                    if col_index > 1:
                        row_has_numeric_measure = True
            if row_has_numeric_measure and isinstance(label, str) and label.strip():
                numeric_field_labels.add(label.strip())
        inventory.append({"sheet": sheet.title, "rows": sheet.max_row,
                          "columns": sheet.max_column,
                          "numeric_cells": sum(map(len, number_columns.values())),
                          "numeric_columns": [{"column": get_column_letter(col),
                                               "numeric_cells": len(locations),
                                               "first_numeric_row": locations[0],
                                               "last_numeric_row": locations[-1],
                                               "top_header": header_examples.get(col)}
                                              for col, locations in sorted(number_columns.items())],
                          "table_blocks": blocks,
                          "numeric_row_labels": sorted(numeric_field_labels),
                          "adoption": ("C:E 2024p country and 77 provinces"
                                       if sheet.title == "PER CAPITA" else
                                       "AE 2024p current-price sector country and 77 province rows; "
                                       "other years and CVM priority_unassessed"
                                       if sheet.title in ("WK", "GRP", "NE", "NO", "SO", "EA", "WE", "CE", "BKK&VIC")
                                       else "priority_unassessed")})
    sheet = workbook["PER CAPITA"]
    rows = list(sheet.values)
    require(len(rows) == 97 and "2024p" in rows[1][0], "PER CAPITA edition/layout changed")
    require("GRP 2024p" in rows[2][2] and "POPULATION 2024p" in rows[2][3]
            and "PER CAPITA 2024p" in rows[2][4], "2024p measure headers changed")
    country = {"name": "WHOLE KINGDOM", "row": 11,
               "values": as_values(rows[10], "PER CAPITA!C11:E11")}
    require(rows[10][1] == "WHOLE KINGDOM", "Country row moved")
    regions = []
    for number, row in enumerate(rows[3:10], 4):
        require(row[1] in GROUP_COUNTS, f"Unexpected economic region at row {number}")
        regions.append({"name": row[1], "row": number,
                        "values": as_values(row, f"PER CAPITA!C{number}:E{number}")})
    require(len({x["name"] for x in regions}) == 7, "Seven regions expected")
    provinces = []
    by_region = defaultdict(list)
    current_region = None
    for number, row in enumerate(rows[13:], 14):
        if row[0] == "":
            require(row[1] in GROUP_COUNTS, f"Unexpected region separator at row {number}")
            current_region = row[1]
            continue
        require(isinstance(row[0], int) and current_region,
                f"Unexpected province row structure at row {number}")
        match = re.fullmatch(r"([0-9]{4}) ([A-Z][A-Z .-]+)", row[1] or "")
        require(match, f"Missing NESDC workbook code/name at row {number}")
        require(int(match[1][:2]) == list(GROUP_COUNTS).index(current_region) + 1,
                f"Economic region prefix differs at row {number}")
        entry = {"nesdc_code": match[1], "name": match[2], "region": current_region,
                 "row": number, "values": as_values(row, f"PER CAPITA!C{number}:E{number}")}
        provinces.append(entry)
        by_region[current_region].append(entry)
        require(row[0] == len(by_region[current_region]),
                f"Province ordinal changed at row {number}")
    require(len(provinces) == 77 and len({p["nesdc_code"] for p in provinces}) == 77,
            "Expected 77 uniquely coded NESDC province rows")
    require({k: len(v) for k, v in by_region.items()} == GROUP_COUNTS,
            "Province roster differs from the 77-row regional count")
    require(len(workbook_province_blocks) == 154 and
            {code for _, _, code, _ in workbook_province_blocks} ==
            {p["nesdc_code"] for p in provinces},
            "Detailed current-price/CVM province table blocks differ from PER CAPITA roster")
    for metric in METRICS[:2]:
        close(sum(p["values"][metric] for p in provinces), country["values"][metric], metric)
        for region in regions:
            close(sum(p["values"][metric] for p in by_region[region["name"]]),
                  region["values"][metric], region["name"] + " " + metric)
    # The current-price table of each country, region and province has the
    # same 23 published sector/subtotal rows, with 2024p in AE. Keep the full
    # row audit, then select 21 distinct indicators. Earlier years and the
    # non-additive chain-volume tables remain separately unassessed.
    sector_blocks = {}
    for sheet_name in ("WK", "GRP", "NE", "NO", "SO", "EA", "WE", "CE", "BKK&VIC"):
        detail = list(workbook[sheet_name].values)
        for row_number, row in enumerate(detail, 1):
            raw_label = row[0] if row else None
            if sheet_name == "WK":
                wanted = raw_label == "WHOLE KINGDOM"
                key = "THA"
            elif sheet_name == "GRP":
                wanted = bool(isinstance(raw_label, str) and re.fullmatch(
                    r"[1-7] - [A-Z][A-Z ]+", raw_label))
                key = raw_label.split(" - ", 1)[1] if wanted else None
            else:
                wanted = bool(isinstance(raw_label, str) and re.fullmatch(
                    r"[0-9]{4} - [A-Z][A-Z .-]+", raw_label))
                key = raw_label[:4] if wanted else None
            if not wanted or "CURRENT MARKET PRICES" not in str(detail[row_number - 2][0]):
                continue
            require(key not in sector_blocks, f"Duplicate current-price block: {key}")
            require(detail[row_number + 1][30] == "2024p",
                    f"2024p AE year header changed at {sheet_name}!AE{row_number+2}")
            values = {}
            for offset, expected_label in enumerate(SECTOR_LABELS, 3):
                cell_row = row_number + offset
                actual_label = " ".join(str(detail[cell_row - 1][0]).split())
                require(actual_label == expected_label,
                        f"Sector label changed at {sheet_name}!A{cell_row}: {actual_label}")
                value = detail[cell_row - 1][30]
                require(isinstance(value, (int, float)) and not isinstance(value, bool)
                        and value >= 0, f"Invalid current-price sector at {sheet_name}!AE{cell_row}")
                values[expected_label] = {"value": value,
                                          "locator": f"{sheet_name}!AE{cell_row}"}
            total_row = row_number + 26
            total_label = " ".join(str(detail[total_row - 1][0]).split())
            require(total_label in ("Gross domestic product (GDP)",
                                    "Gross regional product (GRP)",
                                    "Gross provincial product (GPP)"),
                    f"Current-price total row moved at {sheet_name}!A{total_row}")
            total = detail[total_row - 1][30]
            close(values["Agriculture"]["value"],
                  values["Agriculture, forestry and fishing"]["value"],
                  f"{sheet_name} {key} agriculture alias")
            close(values["Non-Agriculture"]["value"],
                  values["Industrial"]["value"] + values["Services"]["value"],
                  f"{sheet_name} {key} non-agriculture")
            close(total, sum(values[label]["value"] for label in
                             ("Agriculture", "Industrial", "Services")),
                  f"{sheet_name} {key} sector total")
            for group, parent in (("industry_detail", "Industrial"),
                                  ("services_detail", "Services")):
                close(values[parent]["value"],
                      sum(values[label]["value"] for _, label, g in SECTOR_FIELDS if g == group),
                      f"{sheet_name} {key} {group}")
            sector_blocks[key] = {"sheet": sheet_name, "block_row": row_number,
                                  "total": total, "fields": values,
                                  "detailed_population": detail[row_number + 27][30]}
    require(len(sector_blocks) == 85 and "THA" in sector_blocks and
            all(region["name"] in sector_blocks for region in regions) and
            all(p["nesdc_code"][:4] in sector_blocks for p in provinces),
            "Expected one current-price 2024p table for country, seven regions and 77 provinces")
    close(sector_blocks["THA"]["total"], country["values"][METRICS[0]],
          "Country current-price GDP vs PER CAPITA GPP")
    for region in regions:
        close(sector_blocks[region["name"]]["total"], region["values"][METRICS[0]],
              f"{region['name']} GRP vs PER CAPITA")
    for province in provinces:
        close(sector_blocks[province["nesdc_code"][:4]]["total"],
              province["values"][METRICS[0]], f"{province['name']} GPP vs PER CAPITA")
    for label in SECTOR_LABELS:
        national = sector_blocks["THA"]["fields"][label]["value"]
        close(sum(sector_blocks[p["nesdc_code"][:4]]["fields"][label]["value"]
                  for p in provinces), national, f"Province current-price sum: {label}")
        for region in regions:
            close(sum(sector_blocks[p["nesdc_code"][:4]]["fields"][label]["value"]
                      for p in by_region[region["name"]]),
                  sector_blocks[region["name"]]["fields"][label]["value"],
                  f"Economic-region current-price sum: {region['name']} {label}")
    population_conflicts = []
    for province in provinces:
        block = sector_blocks[province["nesdc_code"]]
        detail_row = block["block_row"] + 28
        detail = block["detailed_population"]
        printed = province["values"]["population_thousand_persons"]
        require(isinstance(detail, (int, float)),
                f"Missing detailed-table population at {block['sheet']}!AE{detail_row}")
        if abs(detail - printed) > 0.0001:
            population_conflicts.append({"nesdc_code": province["nesdc_code"],
                "name": province["name"], "detailed_table_locator": f"{block['sheet']}!AE{detail_row}",
                "detailed_table_value": detail, "adopted_per_capita_locator":
                    f"PER CAPITA!D{province['row']}", "adopted_per_capita_value": printed,
                "difference_thousand_persons": detail - printed,
                "disposition": "detailed_table_population_rejected; PER CAPITA denominator adopted"})
    require(len(population_conflicts) == 1 and
            population_conflicts[0]["nesdc_code"] == "0212" and
            1.5319 < population_conflicts[0]["difference_thousand_persons"] < 1.5320,
            f"Detailed-table population conflict set changed: {population_conflicts}")
    audit = {"schema_version": "1.0", "source_id": source["id"],
             "source_url": source["url"], "source_sha256": source["sha256"],
             "edition": "2024p provisional", "sheet_inventory": inventory,
             "adopted_sheets": ["PER CAPITA", "WK", "GRP", "NE", "NO", "SO", "EA", "WE", "CE", "BKK&VIC"],
             "adopted_numeric_columns": {"PER CAPITA": "C:E", "current_price_2024p": "AE"},
             "sector_field_disposition": [{"label": label,
                 "status": "adopted" if label in SECTOR_IDS else "audited_duplicate_or_subtotal_not_adopted",
                 "indicator_id": SECTOR_IDS.get(label)} for label in SECTOR_LABELS],
             "numeric_cells_audited": 85 * (3 + len(SECTOR_LABELS)),
             "adopted_direct_observations": 78 * (3 + len(SECTOR_FIELDS)),
             "additional_detailed_population_cells_inspected": 77,
             "source_population_conflicts": population_conflicts,
             "country": country, "regions_audit_only": regions, "provinces": provinces,
             "sector_2024p_current_price": sector_blocks,
             "notes": [
                 "The 77 province rows and seven economic-region rows each reconcile to the printed whole-kingdom GPP and population, within source floating precision.",
                 "Published per-capita cells reconcile to GPP million Baht / population thousand persons times 1000; per-capita values are not summed or averaged across provinces.",
                 "Population is the NESDC 2024p denominator for GPP per capita, not BORA registered population or the NSO 2025 resident census.",
                 "The detailed NO current-price table prints Kam Phaeng Phet population 780.260023 thousand at AE691, whereas PER CAPITA D47 gives 778.728070019944 thousand and its published per-capita value reconciles to the latter. Only the PER CAPITA denominator is adopted; the detailed-table cell is preserved as a source conflict.",
                 "The four-character NESDC workbook codes embed economic-region ordering; they are not confirmed DOPA legal administrative codes or polygon join keys.",
                 "The 2024p AE current-price sector/subtotal cells in the whole-kingdom, seven-region and 77-province tables are audited. Twenty-one distinct sector indicators for country and provinces are adopted; the duplicate agriculture alias and overlapping Non-Agriculture subtotal remain audit-only.",
                 "The 1995-2023 historic series, chain-volume tables, Regions to GDP and CLUSTERS sheets are inventoried but remain priority_unassessed; they are not claimed absent or adopted.",
                 "The source explicitly warns that chain-volume series are not additive. No chain-volume regional or country sum is adopted."
             ]}
    out = project / "evidence/THA_GPP2024_AUDIT.json"
    out.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"country_gpp_million_baht": country["values"][METRICS[0]],
                      "country_population_thousand": country["values"][METRICS[1]],
                      "provinces": len(provinces), "economic_regions_audit_only": 7,
                      "numeric_cells_audited": audit["numeric_cells_audited"],
                      "current_price_sector_fields": len(SECTOR_FIELDS),
                      "direct_observations": audit["adopted_direct_observations"]}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
