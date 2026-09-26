"""Inventory official Oman sources and cross-check eCensus labels with MOI Arabic names."""

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


ROOT = "raw/oman-ecensus-moi/"
MOI_HASH = "63225e0a09ad41a41d5bd3b98b506e9117defdd67b8a7bbc0770ed686c63419b"
EN_FIELDS = ["RESIDENTIAL_STATUS", "LOCATION_GOVERNORATE", "LOCATION_WILAYAT", "RUN_DATE", "AMOUNT"]
AR_FIELDS = ["الجنسية - عماني أو وافد", "الموقع حسب المحافظة", "الموقع حسب الولاية", "التاريخ المرجعي", "المجموع"]


def original(project, name):
    relative = ROOT + name
    body = (project / relative).read_bytes()
    receipt = json.loads((project / (relative + ".receipt.json")).read_text(encoding="utf-8"))
    if receipt["status"] != "acquired" or receipt["sha256"] != hashlib.sha256(body).hexdigest():
        raise ValueError(f"Source receipt/hash mismatch: {name}")
    return body, receipt


def arabic_key(value):
    value = re.sub(r"^(?:محافظة|ولاية)\s+", "", value.strip())
    value = value.replace("ـ", "")
    value = re.sub(r"[\u064b-\u065f]", "", value)
    value = value.translate(str.maketrans("أإآىةؤئ", "ااايهوي"))
    return "".join(char for char in value if char.isalnum())


def series_signature(row):
    return tuple(cell[0] for cell in row)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    moi_body, moi_receipt = original(project, "moi-governorate-wilayat-2025.xlsx")
    if hashlib.sha256(moi_body).hexdigest() != MOI_HASH:
        raise ValueError("MOI workbook version changed")
    book = openpyxl.load_workbook(project / (ROOT + "moi-governorate-wilayat-2025.xlsx"), read_only=True, data_only=True)
    if book.sheetnames != ["المحافظات", "الولايات", "البيانات الوصفية", "المتغيرات"]:
        raise ValueError("MOI workbook sheet structure changed")
    gov_rows = list(book["المحافظات"].values)
    wil_rows = list(book["الولايات"].values)
    if gov_rows[0][:3] != ("RegionId", "RegionName", "RegionNameEN") or \
            wil_rows[0][:4] != ("WilayatId", "WilayatName", "WilayatNameEN", "RegionId"):
        raise ValueError("MOI header changed")
    governors = {row[0]: row for row in gov_rows[1:]}
    wilayats = {row[0]: row for row in wil_rows[1:]}
    if len(governors) != 11 or len(wilayats) != 63 or \
            len({(row[3], arabic_key(row[1])) for row in wilayats.values()}) != 63:
        raise ValueError("MOI hierarchy changed or Arabic names duplicate")
    column_inventory = []
    for sheet in book:
        rows = list(sheet.values)
        for index in range(sheet.max_column):
            cells = [row[index] if index < len(row) else None for row in rows[1:]]
            adopted = (sheet.title == "المحافظات" and index in (0, 1, 2)) or \
                      (sheet.title == "الولايات" and index in (0, 1, 2, 3))
            column_inventory.append({"sheet": sheet.title, "column": index + 1,
                "header": rows[0][index] if index < len(rows[0]) else None,
                "nonempty_cells": sum(value is not None and value != "" for value in cells),
                "numeric_cells": sum(isinstance(value, (int, float)) and not isinstance(value, bool) for value in cells),
                "decision": "adopted_hierarchy_only" if adopted else "not_adopted_or_unassessed",
                "reason": "MOI 2025 names and IDs only; not a 2020 census geography/boundary certificate" if adopted else
                          "Metadata, sort/seat count or other field not adopted as a statistical indicator"})
    book.close()
    en_meta = json.loads(original(project, "ecensus-population-metadata.json")[0])
    ar_meta = json.loads(original(project, "ecensus-population-arabic-metadata.json")[0])
    if en_meta["table_name"] != "v_public_ds_population_en" or \
            ar_meta["table_name"] != "v_public_ds_population_ar" or \
            not set(EN_FIELDS).issubset({item["field"] for item in en_meta["columns"]} | {item["field"] for item in en_meta["measures"]}) or \
            not set(AR_FIELDS).issubset({item["field"] for item in ar_meta["columns"]} | {item["field"] for item in ar_meta["measures"]}):
        raise ValueError("eCensus bilingual metadata changed")
    datasets = {kind: json.loads(original(project, f"ecensus-population-{kind}-pivot.json")[0])
                for kind in ("national", "governorate", "wilayat", "arabic-wilayat")}
    dates = [item[0] for item in datasets["national"]["headers"]["column_headers"]]
    if dates != ["2020-12-12", "2021-12-31", "2023-01-01", "2024-01-01", "2025-01-01", "2026-01-01"]:
        raise ValueError("eCensus reference dates changed")
    for name, data in datasets.items():
        if [item[0] for item in data["headers"]["column_headers"]] != dates or \
                len(data["headers"]["row_headers"]) != len(data["data"]):
            raise ValueError(f"eCensus pivot structure changed: {name}")
    en = datasets["wilayat"]
    ar = datasets["arabic-wilayat"]
    if len(en["data"]) != 130 or len(ar["data"]) != 130 or \
            Counter(map(series_signature, en["data"])) != Counter(map(series_signature, ar["data"])):
        raise ValueError("Bilingual wilayat values/coverage differ")
    arabic_by_signature = defaultdict(list)
    for label, values in zip(ar["headers"]["row_headers"], ar["data"]):
        arabic_by_signature[series_signature(values)].append(label)
    gov_by_arabic = {arabic_key(row[1]): key for key, row in governors.items()}
    wil_by_arabic = {(row[3], arabic_key(row[1])): key for key, row in wilayats.items()}
    rows = []
    unpaired = []
    for labels, values in zip(en["headers"]["row_headers"], en["data"]):
        status, governor, wilayat = labels
        if governor == "Unknown" or wilayat.upper() == "UNKNOWN":
            rows.append({"english_status": status, "english_governorate": governor,
                         "english_wilayat": wilayat, "disposition": "source_exception_unassigned"})
            continue
        matches = arabic_by_signature[series_signature(values)]
        arabic_status = "عماني" if status == "Omani" else "وافد"
        matches = [item for item in matches if item[0] == arabic_status]
        if len(matches) != 1:
            unpaired.append((labels, "bilingual series signature", matches))
            continue
        _, gov_ar, wil_ar = matches[0]
        gov_id = gov_by_arabic.get(arabic_key(gov_ar))
        wil_id = wil_by_arabic.get((gov_id, arabic_key(wil_ar)))
        if gov_id is None or wil_id is None:
            unpaired.append((labels, (gov_ar, wil_ar), (gov_id, wil_id)))
            continue
        rows.append({"english_status": status, "arabic_status": arabic_status,
                     "english_governorate": governor, "arabic_governorate": gov_ar, "moi_region_id": gov_id,
                     "english_wilayat": wilayat, "arabic_wilayat": wil_ar, "moi_wilayat_id": wil_id,
                     "census_2020_value": values[0][0], "recent_2025_value": values[4][0],
                     "disposition": "2020_observed" if values[0][0] is not None else "2020_not_reported"})
    if unpaired:
        raise ValueError(f"Unmatched Arabic source/MOI names: {unpaired}")
    if len(rows) != 130 or sum(row.get("disposition") == "2020_observed" for row in rows) != 122 or \
            sum(row.get("disposition") == "2020_not_reported" for row in rows) != 4:
        raise ValueError("2020 caza/wilayat availability changed")
    evidence = project / "evidence"
    with (evidence / "OMN_MOI_COLUMN_INVENTORY.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(column_inventory[0]))
        writer.writeheader()
        writer.writerows(column_inventory)
    with (evidence / "OMN_ECENSUS_MOI_WILAYAT_CROSSWALK.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    report = {"status": "structural_inventory_partial_semantic_adoption", "moi_sha256": moi_receipt["sha256"],
              "moi_sheets": book.sheetnames, "moi_columns": len(column_inventory),
              "moi_governorates": len(governors), "moi_wilayats": len(wilayats),
              "ecensus_reference_dates": dates, "ecensus_en_metadata_fields": en_meta["columns"] + en_meta["measures"],
              "ecensus_ar_metadata_fields": ar_meta["columns"] + ar_meta["measures"],
              "ecensus_english_rows": len(en["data"]), "ecensus_arabic_rows": len(ar["data"]),
              "2020_observed_nationality_wilayat_rows": 122,
              "2020_not_reported_nationality_wilayat_rows": 4,
              "source_exception_rows": 4,
              "other_fields_status": "unassessed_not_zero",
              "name_join_status": "bilingual_series_signature_plus_arabic_name_match; 2020_vs_2025_legal_boundary_equivalence_unverified"}
    (evidence / "OMN_ECENSUS_MOI_STRUCTURE.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"moi_governorates": len(governors), "moi_wilayats": len(wilayats),
                      "moi_columns": len(column_inventory), "2020_observed_rows": 122,
                      "crosswalk_rows": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
