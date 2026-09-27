"""Build an unpublished Uzbekistan candidate from pinned 2026 official sources.

The SIAT permanent-population estimates (1 January) and the preliminary census
(15 January) remain separate indicators. The 2017 reference shapes are removed.
"""

import argparse
from collections import Counter, defaultdict
import csv
from decimal import Decimal
import hashlib
import html
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/uzbekistan-2026-source-manifest.json"
SIAT_KEYS = ("total", "female", "male", "urban", "rural")
SIAT_CODES = {"total": "246", "female": "244", "male": "245", "urban": "248", "rural": "247"}
CODE_SYSTEM = "NSC SIAT COATO/MHOBT administrative-territorial classifier"
GEOGRAPHY = "NSC SIAT 2026 reporting roster; dated official polygon edition pending"
PREFIX = "UZB_2026_"
CHECKED_AT = "2026-09-27"
INTEGER = re.compile(r"^\d{1,3}(?: \d{3})*$")
PERCENT = re.compile(r"^\d{1,3},\d$")


def require(value, message):
    if not value:
        raise ValueError(message)


def pinned(project):
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    require(manifest["country_id"] == "UZB" and len(manifest["sources"]) == 9,
            "Source manifest changed")
    sources = {}
    for source in manifest["sources"]:
        path = project / source["raw_path"]
        require(path.is_file(), f"Missing source: {path}")
        body = path.read_bytes()
        require(len(body) == source["bytes"] and
                hashlib.sha256(body).hexdigest() == source["sha256"],
                f"Source bytes differ from manifest: {path}")
        if path.suffix == ".pdf":
            require(body.startswith(b"%PDF"), f"Invalid PDF: {path}")
        if path.suffix == ".xlsx":
            require(body.startswith(b"PK"), f"Invalid XLSX: {path}")
        sources[source["id"]] = (source, path)
    return manifest, sources


def text_pages(path):
    result = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
                            capture_output=True, check=True)
    return result.stdout.decode("utf-8", errors="replace").split("\f")


def read_siat(sources):
    rows = {}
    inventory = []
    roster = None
    all_numeric = 0
    for key in SIAT_KEYS:
        source_id = f"uzb-siat-permanent-{key}"
        with sources[source_id][1].open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            header = reader.fieldnames
            raw = list(reader)
        require(header is not None and header[:5] == ["Code", "Klassifikator", "Klassifikator_ru",
                                                       "Klassifikator_en", "Klassifikator_uzc"],
                f"Unexpected SIAT header: {key}")
        years = [str(year) for year in range(2014 if key == "female" else 2010, 2027)]
        require(header[5:] == years, f"Unexpected SIAT years: {key}")
        require(len(raw) == 221 and len({r["Code"] for r in raw}) == 221,
                f"SIAT roster count/duplicate changed: {key}")
        current_roster = [(r["Code"], r["Klassifikator_en"], r["Klassifikator"]) for r in raw]
        if roster is None:
            roster = current_roster
        else:
            require(current_roster == roster, f"SIAT roster/name order differs: {key}")
        numeric = 0
        blank = 0
        zeros_before_2026 = 0
        historical_precision_artifacts = []
        for row in raw:
            for year in years:
                value = row[year].strip()
                if not value:
                    blank += 1
                    require(year != "2026", f"Missing 2026 SIAT value: {key}/{row['Code']}")
                    continue
                number = Decimal(value)
                require(number.is_finite() and number >= 0,
                        f"Unexpected SIAT numeric value: {key}/{row['Code']}/{year}")
                if number.as_tuple().exponent != -1:
                    require(year != "2026", f"Unexpected 2026 SIAT precision: {key}/{row['Code']}")
                    historical_precision_artifacts.append({"code": row["Code"], "year": year,
                                                           "raw_value": value})
                numeric += 1
                if year != "2026" and number == 0:
                    zeros_before_2026 += 1
        rows[key] = {r["Code"]: r for r in raw}
        all_numeric += numeric
        inventory.append({"source_id": source_id, "indicator_code": sources[source_id][0]["indicator_code"],
                          "row_count": len(raw), "year_columns": years,
                          "numeric_cells": numeric, "blank_cells": blank,
                          "pre_2026_zero_cells": zeros_before_2026,
                          "historical_precision_artifacts": historical_precision_artifacts,
                          "adopted_columns": ["2026"],
                          "earlier_years": "priority_unassessed for administrative changes and time comparison"})
    assert roster is not None
    codes = [code for code, _, _ in roster]
    require(len([c for c in codes if len(c) == 4]) == 15 and
            len([c for c in codes if len(c) == 7]) == 206 and codes[0] == "1700",
            "SIAT administrative-code length/count changed")
    require(all(len(c) in (4, 7) and c.isdigit() for c in codes), "Unexpected SIAT code")
    require(all(c[:4] in rows["total"] for c in codes if len(c) == 7),
            "SIAT child without parent")
    checks = Counter()
    for code in codes:
        values = {key: Decimal(rows[key][code]["2026"]) for key in SIAT_KEYS}
        require(values["total"] == values["female"] + values["male"],
                f"SIAT sex subtotal mismatch: {code}")
        require(values["total"] == values["urban"] + values["rural"],
                f"SIAT urban/rural subtotal mismatch: {code}")
        checks["sex_equalities"] += 1
        checks["urban_rural_equalities"] += 1
    parents = [c for c in codes if len(c) == 4]
    for key in SIAT_KEYS:
        for parent in parents:
            children = [c for c in codes if len(c) == 7 and c[:4] == parent]
            if children:
                require(Decimal(rows[key][parent]["2026"]) ==
                        sum(Decimal(rows[key][c]["2026"]) for c in children),
                        f"SIAT child subtotal mismatch: {key}/{parent}")
                checks["child_parent_equalities"] += 1
        require(Decimal(rows[key]["1700"]["2026"]) ==
                sum(Decimal(rows[key][p]["2026"]) for p in parents if p != "1700"),
                f"SIAT national subtotal mismatch: {key}")
        checks["national_equalities"] += 1
    require(Decimal(rows["total"]["1700"]["2026"]) == Decimal("38236.7"),
            "SIAT 2026 national estimate changed")
    return rows, roster, inventory, dict(checks), all_numeric


def parse_census_page(page, expected_names, title):
    require(title in page, f"Census table title changed: {title}")
    lines = page.splitlines()
    found = []
    for index, line in enumerate(lines):
        cells = re.split(r"\s{2,}", line.strip())
        if len(cells) not in (5, 6) or not all(INTEGER.fullmatch(v) for v in cells[-5:-2]) or \
                not all(PERCENT.fullmatch(v) for v in cells[-2:]):
            continue
        name = cells[0] if len(cells) == 6 else (lines[index - 1] + " " + lines[index + 1]).strip()
        normalized = " ".join(name.casefold().split())
        expected = " ".join(expected_names[len(found)].casefold().split()) if len(found) < len(expected_names) else ""
        require(normalized == expected, f"Census row name/order differs: {name!r} vs {expected!r}")
        found.append({"name": expected_names[len(found)],
                      "counts": [int(x.replace(" ", "")) for x in cells[-5:-2]],
                      "percentages": [Decimal(x.replace(",", ".")) for x in cells[-2:]]})
    require(len(found) == 15, f"Census table has {len(found)} rows; expected 15")
    for row in found:
        total, first, second = row["counts"]
        require(total == first + second, f"Census component sum: {row['name']}")
        require(abs(row["percentages"][0] - Decimal(first * 100) / Decimal(total)) <= Decimal("0.05") and
                abs(row["percentages"][1] - Decimal(second * 100) / Decimal(total)) <= Decimal("0.05"),
                f"Census printed percentages: {row['name']}")
    for column in range(3):
        require(found[0]["counts"][column] == sum(r["counts"][column] for r in found[1:]),
                f"Census 14-region subtotal, column {column}")
    return found


def read_census(sources, roster):
    pages = text_pages(sources["uzb-nsc-census2026-preliminary-report"][1])
    require(len(pages) - 1 == 90 and "NOTES TO THE TABLES" in pages[23],
            "Census preliminary PDF page edition changed")
    require(re.search(r"15 January\s+2026", pages[23]) and "preliminary nature" in pages[23],
            "Census reference date/release status changed")
    names = [name for code, name, _ in roster if len(code) == 4]
    sex = parse_census_page(pages[25], names, "Distribution of the population by sex, by region")
    settlement = parse_census_page(pages[26], names,
                                   "Distribution of the population across urban and rural")
    require([r["counts"][0] for r in sex] == [r["counts"][0] for r in settlement],
            "Census p.4/p.5 population differs")
    require(sex[0]["counts"] == [39047321, 19766166, 19281155] and
            settlement[0]["counts"] == [39047321, 21276729, 17770592],
            "Census national benchmark changed")
    return sex, settlement


def read_press(sources, roster, siat):
    pages = text_pages(sources["uzb-nsc-demography-2026-01-27"][1])
    require(len(pages) - 1 == 7 and "January 1, 2026" in pages[1],
            "NSC demographic release edition changed")
    table = pages[1]
    names = [name for code, name, _ in roster if len(code) == 4]
    found = []
    for line in table.splitlines():
        cells = re.split(r"\s{2,}", line.strip())
        if len(cells) != 4 or not re.fullmatch(r"[\d ]+,\d", cells[1]) or \
                not re.fullmatch(r"[\d ]+,\d", cells[2]):
            continue
        found.append(cells)
    require(len(found) == 15, f"NSC press release regional table row count: {len(found)}")
    for (code, name, _), cells in zip((r for r in roster if len(r[0]) == 4), found):
        label = cells[0].casefold()
        expected = name.casefold()
        if expected.endswith(" region"):
            expected = expected[:-7]
        require(label == expected, f"Press release name/order changed: {label} vs {expected}")
        value = Decimal(cells[2].replace(" ", "").replace(",", "."))
        require(value == Decimal(siat["total"][code]["2026"]),
                f"Press release/SIAT 2026 value mismatch: {code}")
    return len(found)


def audit_agriculture_workbook(sources):
    path = sources["uzb-nsc-census2026-preliminary-agriculture-tables"][1]
    workbook = load_workbook(path, read_only=True, data_only=True)
    require(workbook.sheetnames == ["Contents", *(f"1.{i}" for i in range(1, 12))],
            "Preliminary census workbook sheet roster changed")
    inventory = []
    for sheet in workbook:
        numeric = 0
        zero = 0
        dashes = 0
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                    numeric += 1
                    zero += cell == 0
                elif isinstance(cell, str) and cell.strip() == "-":
                    dashes += 1
        inventory.append({"sheet": sheet.title, "title": str(sheet.cell(1, 1).value or "").strip(),
                          "rows": sheet.max_row, "columns": sheet.max_column,
                          "numeric_cells": numeric, "zero_cells": zero, "dash_cells": dashes,
                          "assessment": "priority_unassessed; agriculture only; no local population observation adopted"})
    workbook.close()
    return inventory


def audit_budget_html(sources):
    body = sources["uzb-mof-local-budget-2026-table"][1].read_text(encoding="utf-8")
    table = body[body.find("<table"):body.find("</table>") + len("</table>")]
    require(table.startswith("<table") and table.endswith("</table>"), "MOF budget table missing")
    rows = []
    for tr in re.findall(r"<tr.*?</tr>", table, re.S | re.I):
        cells = [html.unescape(re.sub(r"<[^>]+>", "", value)).strip()
                 for value in re.findall(r"<(?:td|th)[^>]*>(.*?)</(?:td|th)>", tr, re.S | re.I)]
        rows.append(cells)
    require(len(rows) == 17 and len(rows[2]) == 7 and len(rows[3:]) == 14 and
            all(len(row) == 8 for row in rows[3:]), "MOF budget row/column structure changed")
    require(rows[2][0] == "Total" and [row[0] for row in rows[3:]] ==
            [str(i) for i in range(1, 15)], "MOF budget roster changed")
    for row in rows[2:]:
        cells = row[1:] if row[0] == "Total" else row[2:]
        require(len(cells) == 6 and all(Decimal(x.replace(",", "")).is_finite()
                                        for x in cells), "MOF budget numeric column changed")
    national = [Decimal(value.replace(",", "")) for value in rows[2][1:]]
    children = [[Decimal(value.replace(",", "")) for value in row[2:]]
                for row in rows[3:]]
    first_subtotals = [sum(row[index] for row in children) - national[index]
                       for index in range(3)]
    return {"rows": 15, "numeric_columns": 6, "numeric_cells": 90,
            "assessment": "priority_unassessed: pinned HTML has two Revenue/Transfer/Expenditure triplets without a clear period for each; do not adopt as district, regional or national execution values",
            "first_three_subtotal_difference_billion_uzs": [str(value) for value in first_subtotals],
            "second_three_columns": "period and scope unverified"}


def pdf_table_inventory():
    entries = [
        ("population-sex-region", 4), ("population-urban-rural-region", 5),
        ("population-urban-rural-sex-region", 6), ("residence-status", 7),
        ("permanent-registration-sex", 8), ("temporary-absence-reason", 12),
        ("permanent-registration-urban-rural", 20), ("age-sex-region", 24),
        ("main-age-groups", 39), ("birthplace-region-1", 40),
        ("birthplace-region-2", 41), ("present-population-country-of-birth", 42),
        ("ethnicity-region", 43), ("ethnicity-sex-region", 44),
        ("mother-tongue-region", 52), ("dwelling-type-region", 53),
        *((f"agriculture-{i}", page) for i, page in enumerate(
            (55, 57, 57, 58, 59, 60, 61, 63, 65, 66, 67), 1))
    ]
    return [{"table": name, "printed_page": page,
             "assessment": "selected_columns_audited" if page in (4, 5) else "priority_unassessed",
             "adopted_columns": ["total", "male", "female"] if page == 4 else
                                ["urban", "rural"] if page == 5 else []}
            for name, page in entries]


def territory_id(code):
    return "UZB" if code == "1700" else f"UZB:COATO:{code}"


def build(project, manifest, sources, siat, roster, siat_inventory, siat_checks,
          all_siat_numeric, census_sex, census_settlement, press_rows,
          workbook_inventory, budget_inventory):
    path = project / "data/dashboard.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    require(data["country"]["id"] == "UZB" and data["schema_version"] == "0.2",
            "Wrong country or dataset schema")
    data["generated_at"] = manifest["retrieved_at"]
    data["country"]["geography_note"] = (
        "SIAT's 2026 COATO reporting roster supplies 14 first-level units and 206 districts/cities. "
        "Its official code column is retained, but a dated polygon edition is not acquired; the initial "
        "2017 geoBoundaries layer is removed. The census PDF omits codes; its 14 first-level rows are "
        "linked provisionally by unique name and administrative type. Preliminary 15 January census, "
        "1 January SIAT resident estimates and WDI midyear estimates have distinct definitions and dates.")
    territories = []
    type_counts = Counter()
    for code, name, _ in roster:
        if code == "1700":
            level, kind, parent = "national", "country", None
        elif len(code) == 4:
            level, parent = "adm1", "UZB"
            kind = "republic" if code == "1735" else "city" if code == "1726" else "region"
        else:
            level, parent = "adm2", territory_id(code[:4])
            kind = "city" if name.lower().endswith("city") else "district" if name.lower().endswith("district") else ""
            require(kind, f"Unknown SIAT unit type: {code} {name}")
            type_counts[kind] += 1
        territories.append({"id": territory_id(code), "name": name, "level": level,
                            "type": kind, "parent_id": parent, "official_code": code,
                            "code_system": CODE_SYSTEM, "boundary_version": GEOGRAPHY,
                            "source_id": "uzb-siat-permanent-total",
                            "reconciliation_status": "2026 SIAT source-code reporting row; compatible official polygon unverified"})
    require(type_counts == {"city": 31, "district": 175}, "ADM2 type counts changed")
    data["territories"] = territories
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    source_meta = {
        "uzb-nsc-demography-2026-01-27": ("NSC demographic situation, January–December 2025 release", "National Statistics Committee", "as at 2026-01-01", "nation and 14 first-level units"),
        "uzb-nsc-census2026-preliminary-report": ("Preliminary results of the 2026 Population and Agriculture Census", "National Statistics Committee", "census moment 2026-01-15; preliminary", "nation and 14 first-level units in adopted tables"),
        "uzb-nsc-census2026-preliminary-agriculture-tables": ("Preliminary census agriculture tables 1.1–1.11", "National Statistics Committee", "2025 and 2026-01-15 by table", "national agriculture; local rows absent"),
        "uzb-mof-local-budget-2026-table": ("2026 local-budget revenue and expenditure table", "Ministry of Economy and Finance", "2026; column quarter labels unverified", "14 first-level units; district/city coverage unverified"),
    }
    sources_by_id = {s["id"]: s for s in data["sources"] if s["id"] != "geoboundaries-adm1"}
    for item in manifest["sources"]:
        sid = item["id"]
        if sid.startswith("uzb-siat-"):
            key = sid.removeprefix("uzb-siat-permanent-")
            title = "Permanent population - " + key
            period = "as at 2026-01-01; earlier annual columns retained only in raw CSV"
            geographic = "nation, 14 first-level units, 206 districts/cities"
            publisher = "National Statistics Committee, SIAT"
            license_text = "SIAT catalogue marks CC BY 4.0; attribute the NSC and preserve indicator metadata"
        else:
            title, publisher, period, geographic = source_meta[sid]
            license_text = "Official public source; reuse and redistribution terms require source-specific review"
        sources_by_id[sid] = {"id": sid, "name": title, "url": item["url"],
            "catalogue_url": item.get("catalogue_url", item["url"]), "publisher": publisher,
            "reference_period": period, "geographic_level": geographic,
            "status": "ready", "retrieved_at": manifest["retrieved_at"],
            "raw_path": item["raw_path"], "sha256": item["sha256"], "license": license_text,
            "note": "Original bytes and SHA-256 pinned; source acquisition does not mean every table or year is adopted."}
    for sid, name, url, publisher, geography, note in [
        ("uzb-lex-pf21-2030", "Presidential Decree PF-21 and Uzbekistan–2030 Strategy",
         "https://test.lex.uz/docs/8050769", "Lex.uz official legal information", "national",
         "Legal source location; the strategy is a national document, not proof of approval of each regional strategy."),
        ("uzb-asdr-regional-strategies-2026", "Task to prepare 2030 strategies for 14 regions",
         "https://asdr.gov.uz/en/tasks-set-for-the-development-of-regional-strategies/",
         "Agency for Strategic Development and Reforms", "14 first-level units",
         "Agency reports preparation underway; no individual regional strategy body or adoption notice acquired."),
        ("uzb-nsc-census2026-status", "2026 census programme and preliminary-results status",
         "https://gov.uz/en/pages/2026", "Government of Uzbekistan", "national",
         "Census ran 15 January–28 February 2026; preliminary publication is not final census adoption."),
    ]:
        sources_by_id[sid] = {"id": sid, "name": name, "url": url,
            "publisher": publisher, "reference_period": "checked 2026-09-27",
            "geographic_level": geography, "status": "not_collected",
            "retrieved_at": manifest["retrieved_at"],
            "license": "Link location only; terms unverified", "note": note}
    data["sources"] = list(sources_by_id.values())
    data["indicators"] = [i for i in data["indicators"] if not i["id"].startswith(PREFIX)]
    data["observations"] = [o for o in data["observations"] if not o["indicator_id"].startswith(PREFIX)]
    indicator_specs = [
        ("SIAT_TOTAL", "2026 permanent population estimate (1 Jan)", "Population", "thousand people", "uzb-siat-permanent-total", "SIAT registered permanent population estimate, not 2026 census enumeration"),
        ("SIAT_MALE", "2026 permanent male population estimate (1 Jan)", "Population", "thousand people", "uzb-siat-permanent-male", "SIAT male permanent population estimate"),
        ("SIAT_FEMALE", "2026 permanent female population estimate (1 Jan)", "Population", "thousand people", "uzb-siat-permanent-female", "SIAT female permanent population estimate"),
        ("SIAT_URBAN", "2026 urban permanent population estimate (1 Jan)", "Population", "thousand people", "uzb-siat-permanent-urban", "SIAT residents of officially urban settlements"),
        ("SIAT_RURAL", "2026 rural permanent population estimate (1 Jan)", "Population", "thousand people", "uzb-siat-permanent-rural", "SIAT residents of rural settlements; source zero remains zero"),
        ("CENSUS_TOTAL", "2026 preliminary census population (15 Jan)", "Census", "people", "uzb-nsc-census2026-preliminary-report", "Preliminary census persons at the 15 January 2026 census moment; not SIAT estimate"),
        ("CENSUS_MALE", "2026 preliminary census males (15 Jan)", "Census", "people", "uzb-nsc-census2026-preliminary-report", "Preliminary census male persons"),
        ("CENSUS_FEMALE", "2026 preliminary census females (15 Jan)", "Census", "people", "uzb-nsc-census2026-preliminary-report", "Preliminary census female persons"),
        ("CENSUS_URBAN", "2026 preliminary census urban population (15 Jan)", "Census", "people", "uzb-nsc-census2026-preliminary-report", "Preliminary census persons classified as urban"),
        ("CENSUS_RURAL", "2026 preliminary census rural population (15 Jan)", "Census", "people", "uzb-nsc-census2026-preliminary-report", "Preliminary census persons classified as rural"),
    ]
    for key, name, theme, unit, source_id, definition in indicator_specs:
        data["indicators"].append({"id": PREFIX + key, "name": name, "theme": theme,
            "unit": unit, "definition": definition,
            "population": "2026 SIAT registered permanent residents" if key.startswith("SIAT") else
                          "2026 preliminary census persons",
            "source_id": source_id, "aggregation": "none",
            "measurement_method": "NSC SIAT annual administrative permanent-population estimate" if key.startswith("SIAT") else
                                  "NSC preliminary census PDF table cell"})
    observations = []
    for code, _, _ in roster:
        for key, name in (("TOTAL", "total"), ("MALE", "male"), ("FEMALE", "female"),
                          ("URBAN", "urban"), ("RURAL", "rural")):
            sid = f"uzb-siat-permanent-{name}"
            observations.append({"territory_id": territory_id(code), "indicator_id": PREFIX + "SIAT_" + key,
                "period": "2026", "value": float(Decimal(siat[name][code]["2026"])),
                "status": "observed", "source_id": sid, "provenance": "source_reported",
                "population_scope": "SIAT permanent residents as at 2026-01-01",
                "measurement_method": "NSC SIAT annual administrative permanent-population estimate",
                "boundary_version": GEOGRAPHY,
                "source_locator": f"SIAT {sources[sid][0]['indicator_code']}, COATO {code}, column 2026",
                "footnote": "Original source unit: thousand people; do not equate with the 15 January preliminary census or WDI midyear estimates."})
    first_codes = [code for code, _, _ in roster if len(code) == 4]
    for index, code in enumerate(first_codes):
        sex = census_sex[index]["counts"]
        settlement = census_settlement[index]["counts"]
        for key, value, page, column in [
            ("TOTAL", sex[0], 4, "Total population"),
            ("MALE", sex[1], 4, "males"),
            ("FEMALE", sex[2], 4, "females"),
            ("URBAN", settlement[1], 5, "urban areas"),
            ("RURAL", settlement[2], 5, "rural areas"),
        ]:
            observations.append({"territory_id": territory_id(code),
                "indicator_id": PREFIX + "CENSUS_" + key, "period": "2026", "value": value,
                "status": "observed", "source_id": "uzb-nsc-census2026-preliminary-report",
                "provenance": "source_reported",
                "population_scope": "Preliminary 2026 census population, reference moment 2026-01-15",
                "measurement_method": "NSC preliminary census PDF table cell",
                "boundary_version": GEOGRAPHY,
                "source_locator": f"Preliminary report printed p.{page}, {first_codes[index]} name row, {column}",
                "footnote": "Preliminary census row; PDF has no code. Joined to SIAT first-level ID by unique name and administrative type only; exact polygon edition pending."})
    require(len(observations) == 1105 + 75, "Adopted observation count changed")
    data["observations"].extend(observations)
    data["documents"] = [d for d in data["documents"] if not d["id"].startswith("uzb-")]
    data["documents"].extend([
        {"id":"uzb-census2026-preliminary-report","territory_id":"UZB","category":"reference",
         "kind":"preliminary_census_report","title":"2026 preliminary Population and Agriculture Census report",
         "url":sources_by_id["uzb-nsc-census2026-preliminary-report"]["catalogue_url"],
         "source_id":"uzb-nsc-census2026-preliminary-report","period":"2026-01-15",
         "availability":"body_acquired","official_status":"unverified",
         "note":"Printed pp.4–5 selected columns checked; other preliminary population/agriculture tables remain priority_unassessed. No district census rows in the adopted tables."},
        {"id":"uzb-census2026-preliminary-agriculture-tables","territory_id":"UZB",
         "category":"reference","kind":"preliminary_agriculture_tables",
         "title":"2026 census preliminary agriculture Excel tables 1.1–1.11",
         "url":sources_by_id["uzb-nsc-census2026-preliminary-agriculture-tables"]["catalogue_url"],
         "source_id":"uzb-nsc-census2026-preliminary-agriculture-tables","period":"2025 and 2026-01-15 by sheet",
         "availability":"body_acquired","official_status":"unverified",
         "note":"Acquired workbook contains eleven national agriculture sheets, not a district population workbook. Numeric fields remain unassessed for adoption."},
        {"id":"uzb-uzbekistan2030-decree","territory_id":"UZB","category":"plan",
         "kind":"national_strategy_location","title":"Uzbekistan–2030 national strategy, PF-21",
         "url":"https://test.lex.uz/docs/8050769","source_id":"uzb-lex-pf21-2030",
         "period":"2026–2030","availability":"link_verified","official_status":"unverified",
         "note":"National legal source location. This is not an adopted 2030 regional plan for any one of the 14 regions."},
    ])
    for code in first_codes[1:]:
        tid = territory_id(code)
        name = siat["total"][code]["Klassifikator_en"]
        data["documents"].extend([
            {"id":f"uzb-region-strategy-task-{code}","territory_id":tid,"category":"plan",
             "kind":"regional_strategy_preparation_location",
             "title":f"{name}: 2030 strategy preparation task",
             "url":"https://asdr.gov.uz/en/tasks-set-for-the-development-of-regional-strategies/",
             "source_id":"uzb-asdr-regional-strategies-2026","period":"target 2030",
             "availability":"link_verified","official_status":"unverified",
             "note":"Agency describes a task to prepare strategies for all 14 first-level units; this is no proof that this region's draft, final plan, approval or implementation has been acquired."},
            {"id":f"uzb-local-budget-location-{code}","territory_id":tid,
             "category":"budget","kind":"regional_local_budget_table_location",
             "title":f"{name}: MOF 2026 local-budget table location",
             "url":"https://gov.uz/en/imv/sections/view/190028",
             "source_id":"uzb-mof-local-budget-2026-table","period":"2026; quarter unresolved",
             "availability":"body_acquired","official_status":"unverified",
             "note":"Table lists this first-level unit, but its pinned HTML has two numeric triplets with ambiguous quarter labels. No execution value is adopted; district-level budget, allocations and evaluations still require separate source audit."},
        ])
    analysis = data["analysis"]
    analysis["default_indicator_id"] = PREFIX + "SIAT_TOTAL"
    analysis["population_context"] = {
        "primary_indicator_id": PREFIX + "SIAT_TOTAL",
        "reference_indicator_id": PREFIX + "CENSUS_TOTAL",
    }
    analysis["latest_values_only"] = True
    analysis["terminal_territory_ids"] = [territory_id(c) for c in siat["total"] if len(c) == 7]
    analysis["comparisons"] = [{"parent_id":"UZB",
        "member_ids":[territory_id(c) for c in first_codes[1:]],
        "label":"14 SIAT 2026 first-level units",
        "membership_note":"All 14 units have direct SIAT 2026 cells; census p.4–5 uses the same unique reporting names without printed codes. No polygon join.",
        "source_ids":["uzb-siat-permanent-total","uzb-nsc-census2026-preliminary-report"]}]
    analysis["comparisons"].extend({"parent_id":territory_id(parent),
        "member_ids":[territory_id(c) for c in siat["total"] if len(c) == 7 and c[:4] == parent],
        "label":siat["total"][parent]["Klassifikator_en"] + " SIAT 2026 districts/cities",
        "membership_note":"Complete direct SIAT 2026 child roster by COATO prefix; preliminary census has no district rows in adopted tables.",
        "source_ids":["uzb-siat-permanent-total"]} for parent in first_codes[1:])
    data["planning"] = {"title":"Regional strategy and budget source evidence",
        "purpose":"Review available statistics, the regional strategy preparation task and budget-source locations for the selected unit without assuming approval or execution.",
        "system":{"label":"Uzbekistan–2030 and 14 regional strategy preparation tasks",
            "scope":"PF-21 is a national strategy source location; ASDR reports a task to develop strategies for Karakalpakstan, Tashkent city and 12 regions. District/city plan authority and applicable forms remain unverified.",
            "cycle":"Regional strategies are targeted to 2030; individual plan editions, approval and review cycle have not been acquired.",
            "source_ids":["uzb-lex-pf21-2030","uzb-asdr-regional-strategies-2026"]},
        "sections":[{"id":"plan","label":"Strategy and planning documents"},
                    {"id":"budget","label":"Budget and allocations"},
                    {"id":"implementation","label":"Budget execution and implementation"},
                    {"id":"evaluation","label":"Official evaluation"},
                    {"id":"reference","label":"Population and census sources"}]}
    data["collection"]["status"] = "partial"
    adapter = "nsc-siat-census2026-preliminary-partial-2026-09-27"
    data["collection"]["adapters"] = [a for a in data["collection"]["adapters"] if a != adapter] + [adapter]
    data["collection"]["notes"] = [note for note in data["collection"]["notes"] if not note.startswith((
        "Nine official source bodies are byte/hash-pinned", "SIAT 1 January permanent population",
        "The 2017 geoBoundaries layer is excluded"))] + [
        "Nine official source bodies are byte/hash-pinned. Five SIAT CSV 2026 columns and census preliminary PDF printed pp.4–5 selected columns are adopted; all other years/tables remain unassessed.",
        "SIAT 1 January permanent population, census 15 January preliminary population and WDI midyear estimates are separate series. Source zeros remain zero. No population-count difference is labelled growth.",
        "The 2017 geoBoundaries layer is excluded because it has no verified COATO 2026 polygon correspondence. MOF local-budget HTML column periods are ambiguous; no fiscal value is adopted."]
    data["gaps"] = [g for g in data["gaps"] if g.get("category") not in (
        "boundary_reconciliation", "subnational_statistics", "official_geography",
        "census_table_coverage", "planning_documents")]
    data["gaps"].extend([
        {"category":"official_geography","status":"partial","detail":"SIAT COATO codes and hierarchy are acquired, but a dated 2026 official polygon release and census-code crosswalk are not. Initial 2017 reference polygons are excluded.","next_action":"Acquire official dated COATO code list and compatible ADM1/ADM2 polygons; reconcile 2026 census region rows and local plan units."},
        {"category":"census_table_coverage","status":"partial","detail":"Preliminary census PDF p.4–5 selected five population columns are adopted for nation +14 first-level units; other 25 indexed report tables remain priority_unassessed. Excel sheets 1.1–1.11 are agriculture, not district population.","next_action":"Audit every remaining report table and numeric field, especially age, residence, housing and regional characteristics; inspect final 2027 census edition separately."},
        {"category":"planning_documents","status":"partial","detail":"PF-21 national strategy and ASDR regional-strategy task locations found; 14 regional plan bodies, approval, district authority and applicable forms not acquired. MOF local-budget HTML quarter columns unresolved.","next_action":"Acquire each regional strategy/approval and clearly dated regional/district budget, actual expenditure and official evaluation; audit subject and period before adoption."},
    ])
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit_path = project / "evidence/UZBEKISTAN_2026_SOURCE_AUDIT.json"
    audit = {"source_manifest":"config/uzbekistan-2026-source-manifest.json",
        "checked_at":CHECKED_AT,"siat_source_inventory":siat_inventory,
        "siat_all_year_numeric_cells":all_siat_numeric,
        "siat_2026_adopted_cells":1105,"siat_checks":siat_checks,
        "siat_2026_code_counts":{"national":1,"first_level":14,"district_or_city":206,
                                  "cities":31,"districts":175},
        "census_pdf_physical_pages":90,"census_pdf_printed_pages":67,
        "census_table_inventory":pdf_table_inventory(),
        "census_selected_numeric_cells_checked":150,
        "census_direct_observations":75,"census_reference_moment":"2026-01-15",
        "census_national_total":39047321,
        "siat_estimate_reference_moment":"2026-01-01",
        "siat_national_total_thousand":38236.7,
        "press_release_first_level_rows_cross_checked":press_rows,
        "preliminary_agriculture_workbook_inventory":workbook_inventory,
        "mof_budget_table_inventory":budget_inventory,
        "adopted_observations":len(observations),
        "warnings":["The preliminary census report remains preliminary; final results are separate.",
                    "2017 geoBoundaries polygons and ambiguous MOF budget columns are not adopted.",
                    "All SIAT older years and 25 other census report tables remain priority_unassessed."]}
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return audit_path, len(observations)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve()
    manifest, sources = pinned(project)
    siat, roster, siat_inventory, siat_checks, all_siat_numeric = read_siat(sources)
    census_sex, census_settlement = read_census(sources, roster)
    press_rows = read_press(sources, roster, siat)
    workbook_inventory = audit_agriculture_workbook(sources)
    budget_inventory = audit_budget_html(sources)
    audit_path, count = build(project, manifest, sources, siat, roster, siat_inventory,
                              siat_checks, all_siat_numeric, census_sex, census_settlement,
                              press_rows, workbook_inventory, budget_inventory)
    print(json.dumps({"country":"UZB","territories":221,"adopted_observations":count,
                      "audit":str(audit_path)}, ensure_ascii=True))


if __name__ == "__main__":
    main()
