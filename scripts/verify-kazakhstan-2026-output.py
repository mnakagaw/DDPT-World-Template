"""Crosscheck Kazakhstan UI exports against independent pinned workbook cells."""

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/kazakhstan-2026-source-manifest.json"
PREFIX = "KAZ_2026_"
JULY_KEYS = ("TOTAL", "MEN", "WOMEN", "URBAN", "URBAN_MEN", "URBAN_WOMEN",
             "RURAL", "RURAL_MEN", "RURAL_WOMEN")
AUGUST_KEYS = ("AUG_TOTAL", "AUG_GROWTH", "AUG_NATURAL", "AUG_MIGRATION")
CENSUS = {"TOTAL": 19186015, "MEN": 9324840, "WOMEN": 9861175,
          "URBAN": 11741342, "RURAL": 7444673}
# Explicit source-sheet row and 2026 published values, transcribed separately.
CASES = {
    "Kazakhstan": ("KAZ", 6, 7, 20590589, 7400004, 20604819),
    "AbayRegion": ("KAZ:KATO:100000000", 8, 8, 592895, 218584, 591903),
    "AbaiDistrict": ("KAZ:KATO:103200000", 11, None, 13016, 13016, None),
    "Akmola": ("KAZ:KATO:110000000", 22, 9, 793020, 333235, 793929),
    "BurabayDistrict": ("KAZ:KATO:117000000", 42, None, 69977, 22387, None),
    "Ulytau": ("KAZ:KATO:620000000", 229, 23, 218015, 44517, 217670),
    "AlmatyOblast": ("KAZ:KATO:190000000", 59, 11, 1612393, 1300473, 1614234),
    "AlmatyCity": ("KAZ:KATO:750000000", 259, 26, 2373772, 0, 2379793),
    "AstanaCity": ("KAZ:KATO:710000000", 251, 25, 1682721, 0, 1690605),
    "AlmatyDistrictAstana": ("KAZ:KATO:711110000", 252, None, 256353, 0, None),
    "ShymkentCity": ("KAZ:KATO:790000000", 269, 27, 1310681, 0, 1313348),
}


def require(ok, message):
    if not ok:
        raise AssertionError(message)


def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def verify(project):
    data = json.loads((project / "data/dashboard.json").read_text(encoding="utf-8"))
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    sources = {s["id"]: s for s in manifest["sources"]}
    for item in manifest["sources"]:
        body = (project / item["raw_path"]).read_bytes()
        require(len(body) == item["bytes"] and hashlib.sha256(body).hexdigest() == item["sha256"],
                f"Pinned original changed: {item['id']}")
    require(len(data["territories"]) == 249 and data["boundaries"]["features"] == [],
            "2026 reporting roster or excluded polygons changed")
    require(len([row for row in data["observations"] if row["indicator_id"].startswith(PREFIX)]) == 2330,
            "Adopted direct observation count changed")
    require(not any(row["source_id"] in ("kaz-akmola-budget-2026-07-location",
                    "kaz-ulytau-plan-2026-2030", "kaz-burabay-plan-2026-2030")
                    for row in data["observations"]), "Plan/budget targets leaked into numeric statistics")
    july = load_workbook(project / sources["kaz-bns-pop-2026-07-01-sex-locality"]["raw_path"],
                         read_only=True, data_only=True)["2"]
    august = load_workbook(project / sources["kaz-bns-pop-2026-08-01"]["raw_path"],
                           read_only=True, data_only=True)["1."]
    out = project / "evidence/output-verification"
    checked = 0
    for stem, (tid, july_row, august_row, total, rural, august_total) in CASES.items():
        source_values = [int(july.cell(july_row, column).value) for column in range(2, 11)]
        require(source_values[0] == total and source_values[6] == rural,
                f"Hardcoded July original anchor changed: {stem}")
        expected_aug = None
        if august_row is not None:
            expected_aug = {"AUG_TOTAL": int(august.cell(august_row, 6).value),
                            "AUG_GROWTH": int(august.cell(august_row, 3).value),
                            "AUG_NATURAL": int(august.cell(august_row, 4).value),
                            "AUG_MIGRATION": int(august.cell(august_row, 5).value)}
            require(expected_aug["AUG_TOTAL"] == august_total,
                    f"Hardcoded August original anchor changed: {stem}")
        diagnostic = {r["Indicator ID"]: r for r in csv_rows(out / f"{stem}-diagnostic.csv")
                      if r["Record scope"] == "overall"}
        evidence = {r["Indicator ID"]: r for r in csv_rows(out / f"{stem}-evidence.csv")}
        for rows in (diagnostic, evidence):
            require(all(r["Territory ID"] == tid for r in rows.values()),
                    f"Other area leaked into selected export: {stem}")
            for key, expected in zip(JULY_KEYS, source_values):
                row = rows[PREFIX + "JUL_" + key]
                require(row["Status"] == "observed" and row["Value"] == str(expected),
                        f"July {stem}/{key} differs from original workbook: {row['Value']}")
                checked += 1
            for key in AUGUST_KEYS:
                row = rows[PREFIX + key]
                if expected_aug is None:
                    require(row["Status"] == "not_collected" and row["Value"] == "",
                            f"August first-level value leaked into district: {stem}/{key}")
                else:
                    require(row["Status"] == "observed" and row["Value"] == str(expected_aug[key]),
                            f"August {stem}/{key} differs from original workbook")
                checked += 1
            for key in CENSUS:
                row = rows[PREFIX + "CENSUS2021_" + key]
                require(row["Status"] == ("missing" if tid == "KAZ" else "not_collected") and
                        row["Value"] == "",
                        f"2021 census was silently copied into the 2026 period: {stem}/{key}")
                checked += 1
            if tid != "KAZ":
                row = rows["SP.POP.TOTL"]
                require(row["Status"] == "not_collected" and row["Value"] == "",
                        f"WDI national value leaked into local area: {stem}")
        for suffix in ("diagnostic.html", "diagnostic.md", "planning.html"):
            body = (out / f"{stem}-{suffix}").read_text(encoding="utf-8")
            require("undefined" not in body and PREFIX + "JUL_TOTAL" in body,
                    f"Broken {suffix} export: {stem}")
    census_export = {r["Indicator ID"]: r for r in csv_rows(out / "Kazakhstan2021-evidence.csv")}
    for key, expected in CENSUS.items():
        row = census_export[PREFIX + "CENSUS2021_" + key]
        require(row["Status"] == "observed" and row["Value"] == str(expected),
                f"2021 census national printed cell differs: {key}")
        checked += 1
    national = [r for r in csv_rows(out / "Kazakhstan-diagnostic.csv")
                if r["Record scope"] == "within_area" and r["Indicator ID"] == PREFIX + "JUL_TOTAL"]
    require(len(national) == 20 and all(r["Status"] == "observed" for r in national) and
            national[0]["Territory ID"] == CASES["AbayRegion"][0] and
            national[0]["Value"] == "592895" and
            national[-1]["Territory ID"] == CASES["ShymkentCity"][0] and
            national[-1]["Value"] == "1310681",
            "National comparison first/last/count or values differ")
    planning = {stem: (out / f"{stem}-planning.html").read_text(encoding="utf-8")
                for stem in CASES}
    require("Burabay district development plan" in planning["BurabayDistrict"] and
            "8C-39/3" in planning["BurabayDistrict"] and
            "Burabay district development plan" not in planning["Akmola"],
            "Approved district plan not selected correctly")
    require("Ulytau oblast development plan" in planning["Ulytau"] and
            "Ulytau oblast development plan" not in planning["BurabayDistrict"],
            "Ulytau plan not scoped correctly")
    require("Akmola oblast budget execution report" in planning["Akmola"] and
            "Akmola oblast budget execution report" not in planning["AbaiDistrict"],
            "Akmola budget-report location not scoped correctly")
    print(json.dumps({"cases": len(CASES), "original_workbook_export_cells_checked": checked,
                      "comparison_first_last_count": len(national), "pinned_sources": len(manifest["sources"]),
                      "status": "partial_candidate_unpublished"}))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    verify(args.project.resolve())


if __name__ == "__main__":
    main()
