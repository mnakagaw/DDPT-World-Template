"""Inventory the official PBS 2023 census Table 1 source, without geo joins.

The catalogue and workbook hashes are pinned. Every populated numerical cell in
the six acquired workbooks is counted; only reviewed columns/rows are candidates
for dashboard adoption. A census reporting area is not a local government body.
"""

import argparse
import csv
import hashlib
import html
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


CATALOGUE_URL = "https://www.pbs.gov.pk/result-excel/"
BASE = "https://www.pbs.gov.pk/wp-content/uploads/2020/07/"
HASHES = {
    "pbs2023-result-excel-catalog.html": "ab311032ff02e2f74b0f7681b04af91d0f59a12c74a0db27507086abb9c9b035",
    "pbs2023-table-1-national.xlsx": "a15954fb456881e7cdb3cd02692183f949a6b46085a9ca8a0df18bb970b62fc6",
    "pbs2023-table-1-kp-districts.xlsx": "8c869da97389e7dd4b58ff59b9fc249acab4578197125ce66f028c8a72e4c432",
    "pbs2023-table-1-punjab-districts.xlsx": "e082a05c5485908a3fba22392ae222db084458842e26704b4193352ed768a074",
    "pbs2023-table-1-sindh-districts.xlsx": "082e7a735e782ece2380d50f204ee29b1b0222fddbcd772b3d80c073b1757734",
    "pbs2023-table-1-balochistan-districts.xlsx": "318c7182f180b83aa272a9087ed8f62cb3712f85e2f2958c14291ad9a7b797db",
    "pbs2023-table-1-islamabad.xlsx": "99faf74186fd0c7699ba5df008ed927f5cb4533383715785aad0e23d5e1b22b5",
}
FILES = {
    "national": "pbs2023-table-1-national.xlsx",
    "kp": "pbs2023-table-1-kp-districts.xlsx",
    "punjab": "pbs2023-table-1-punjab-districts.xlsx",
    "sindh": "pbs2023-table-1-sindh-districts.xlsx",
    "balochistan": "pbs2023-table-1-balochistan-districts.xlsx",
    "islamabad": "pbs2023-table-1-islamabad.xlsx",
}
SOURCE_NAMES = {key: BASE + ("table_1_national.xlsx" if key == "national" else
                              "table_1_islamabad.xlsx" if key == "islamabad" else
                              f"table_1_{key}_districts.xlsx") for key in FILES}
FIELDS = ("area_km2", "population_2023", "male_2023", "female_2023", "transgender_2023",
          "sex_ratio", "density_per_km2", "urban_proportion_pct", "average_household_size",
          "population_2017", "annual_growth_rate_pct")
ADOPTED = {("total", field) for field in
           ("population_2023", "male_2023", "female_2023", "transgender_2023")}
ADOPTED.update({("rural", "population_2023"), ("urban", "population_2023")})


def require(test, message):
    if not test:
        raise ValueError(message)


def sha256(file):
    return hashlib.sha256(file.read_bytes()).hexdigest()


def source_rows(path, key):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    require(len(book.sheetnames) == 1, f"Workbook sheet count changed: {key}")
    sheet = book.active
    lines = list(sheet.iter_rows(values_only=True))
    require(sheet.max_column == 12 and len(lines) >= 12, f"Table 1 shape changed: {key}")
    require(str(lines[0][0]).startswith("TABLE 1 : AREA, POPULATION BY SEX"),
            f"Table identity changed: {key}")
    require(lines[1][0] == "NAME OF ADMINISTRATIVE UNIT" and lines[2][2] == "ALL SEXES",
            f"Table 1 field headers changed: {key}")
    groups = []
    counts = Counter()
    numeric_cells = Counter()
    notes = []
    active_parent = None
    for excel_row, cells in enumerate(lines[4:], start=5):
        label = cells[0]
        if type(cells[2]) not in {int, float}:
            if (isinstance(label, str) and label.strip() and
                    "table" in label.lower() and "population" in label.lower()):
                notes.append({"row": excel_row, "text": " ".join(label.split())})
            continue
        require(isinstance(label, str) and label.strip(), f"Unlabeled numeric row {key}:{excel_row}")
        cleaned = " ".join(label.split())
        if cleaned.upper() in {"RURAL", "URBAN"}:
            require(active_parent is not None and cleaned.lower() not in active_parent["parts"],
                    f"Orphan or duplicate rural/urban row {key}:{excel_row}")
            row_class = cleaned.lower()
            active_parent["parts"][row_class] = {"row": excel_row, "values": cells[1:12]}
        else:
            if key == "national":
                level = "national_scope" if cleaned.upper() == "PAKISTAN" else "province_or_ict"
            elif cleaned.upper().endswith("DISTRICT"):
                level = "district"
            elif key == "islamabad":
                level = "tehsil"
            elif "TEHSIL" in cleaned.upper():
                level = "tehsil"
            elif "TALUKA" in cleaned.upper():
                level = "taluka"
            elif "SUB-DIVISION" in cleaned.upper():
                level = "sub_division"
            elif len(groups) == 0:
                level = "province"
            else:
                level = "other_reporting_unit"
            active_parent = {"row": excel_row, "name": cleaned, "level": level,
                             "parts": {"total": {"row": excel_row, "values": cells[1:12]}}}
            groups.append(active_parent)
            row_class = "total"
            counts[level] += 1
        values = cells[1:12]
        for field, value in zip(FIELDS, values):
            if type(value) in {int, float}:
                numeric_cells[(row_class, field)] += 1
            elif value is not None and value != "":
                raise ValueError(f"Nonnumeric numeric-column value {key}:{excel_row}:{field}={value!r}")
    require(groups and all(set(g["parts"]) == {"total", "rural", "urban"} for g in groups),
            f"Incomplete reporting row triplets: {key}")
    return groups, counts, numeric_cells, notes


def catalogue_inventory(file, acquired_urls):
    content = file.read_text(encoding="utf-8")
    items = []
    for match in re.finditer(r"<tr[^>]*>(.*?)</tr>", content, re.S | re.I):
        tr = match.group(1)
        parts = re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S | re.I)
        if len(parts) != 2:
            continue
        number = re.search(r"Table-\s*([\d]+(?:\([a-z]\))?)", parts[0], re.I)
        if not number:
            continue
        strong = re.search(r"<strong[^>]*>(.*?)</strong>", parts[1], re.S | re.I)
        title = html.unescape(re.sub(r"<[^>]*>", "", strong.group(1))).strip() if strong else ""
        urls = list(dict.fromkeys(html.unescape(x) for x in
                                  re.findall(r'href="([^"]+\.xlsx)"', parts[1], re.I)))
        require(urls, f"Missing links: Table {number.group(1)}")
        table_id = number.group(1)
        valid_acquired = [url for url in urls if table_id == "1" and url in acquired_urls]
        mismatched_links = [url for url in urls if not re.search(
            rf"/table_{re.escape(table_id.replace('(', '').replace(')', ''))}_", url, re.I)]
        items.append({"table": table_id, "title": title, "xlsx_locations": urls,
                      "acquired_urls": valid_acquired, "mismatched_links": mismatched_links,
                      "status": "partly_acquired" if valid_acquired else "official_location_identified"})
    require(len(items) >= 26 and len({item["table"] for item in items}) == len(items),
            "Census catalogue numbered table list changed")
    return items


def audit(project):
    raw = project / "raw"
    for name, expected in HASHES.items():
        require(sha256(raw / name) == expected, f"Source hash changed: {name}")
    catalogue = catalogue_inventory(raw / "pbs2023-result-excel-catalog.html", set(SOURCE_NAMES.values()))
    inventory = {}
    groups_by_source = {}
    district_totals = defaultdict(list)
    province_totals = {}
    check = Counter()
    exceptions = []
    all_cell_counts = Counter()
    for key, name in FILES.items():
        groups, group_counts, cell_counts, notes = source_rows(raw / name, key)
        groups_by_source[key] = groups
        all_cell_counts.update(cell_counts)
        inventory[key] = {"file": name, "source_url": SOURCE_NAMES[key], "sha256": HASHES[name],
                          "sheet": "Sheet1", "groups": len(groups), "group_levels": dict(group_counts),
                          "numeric_cell_count": sum(cell_counts.values()),
                          "numeric_cells_by_row_class_and_field": {
                              f"{r}.{f}": n for (r, f), n in sorted(cell_counts.items())},
                          "notes": notes}
        current_district = None
        for group in groups:
            total = group["parts"]["total"]["values"]
            rural = group["parts"]["rural"]["values"]
            urban = group["parts"]["urban"]["values"]
            if group["level"] == "district":
                current_district = group
                district_totals[key].append(group)
            elif group["level"] == "province":
                province_totals[key] = group
            elif group["level"] not in {"national_scope", "province_or_ict"}:
                require(current_district is not None, f"Reporting unit before district {key}:{group['row']}")
                group["district_row"] = current_district["row"]
            for idx in (1, 2, 3, 4, 9):
                if all(type(row[idx]) in {int, float} for row in (total, rural, urban)):
                    if total[idx] == rural[idx] + urban[idx]:
                        check["total_rural_urban_count_equal"] += 1
                    else:
                        exceptions.append({"source": key, "row": group["row"],
                                           "field": FIELDS[idx], "issue": "rural_plus_urban_differs"})
            for part_name, part in group["parts"].items():
                values = part["values"]
                if all(type(values[idx]) in {int, float} for idx in (1, 2, 3, 4)):
                    if values[1] == sum(values[idx] for idx in (2, 3, 4)):
                        check["sex_count_equal"] += 1
                    else:
                        exceptions.append({"source": key, "row": part["row"],
                                           "field": "population_2023", "issue": "sex_sum_differs"})
    national = groups_by_source["national"]
    require(len(national) == 6 and national[0]["name"] == "PAKISTAN", "National Table 1 scope changed")
    for part_name in ("total", "rural", "urban"):
        for field_index in (1, 2, 3, 4, 9):
            source_value = national[0]["parts"][part_name]["values"][field_index]
            sum_value = sum(g["parts"][part_name]["values"][field_index] for g in national[1:])
            if source_value == sum_value:
                check["five_regions_to_national_count_equal"] += 1
            else:
                exceptions.append({"source": "national", "row": 5, "field": FIELDS[field_index],
                                   "row_class": part_name, "issue": "five_regions_sum_differs",
                                   "source_value": source_value, "sum_value": sum_value})
    for key in ("kp", "punjab", "sindh", "balochistan", "islamabad"):
        groups = groups_by_source[key]
        source_region = groups[0] if key != "islamabad" else groups[0]
        district_groups = district_totals[key]
        if key == "islamabad":
            require(len(district_groups) == 1, "ICT district changed")
        national_region = next(g for g in national[1:] if
                               g["name"].upper() == ("KHYBER PAKHTUNKHWA" if key == "kp" else key.upper()))
        for part_name in ("total", "rural", "urban"):
            for field_index in (1, 2, 3, 4, 9):
                expected = source_region["parts"][part_name]["values"][field_index]
                actual = sum(g["parts"][part_name]["values"][field_index] for g in district_groups)
                if expected == actual:
                    check["districts_to_region_count_equal"] += 1
                else:
                    exceptions.append({"source": key, "row": source_region["row"], "field": FIELDS[field_index],
                                       "row_class": part_name, "issue": "district_sum_differs",
                                       "source_value": expected, "sum_value": actual})
                if expected == national_region["parts"][part_name]["values"][field_index]:
                    check["region_xlsx_to_national_xlsx_equal"] += 1
                else:
                    exceptions.append({"source": key, "row": source_region["row"], "field": FIELDS[field_index],
                                       "row_class": part_name, "issue": "region_xlsx_disagrees_with_national_xlsx"})
    evidence = project / "evidence"
    evidence.mkdir(exist_ok=True)
    with (evidence / "PAK_INDICATOR_INVENTORY.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_id", "source_hash", "sheet", "column", "row_class", "numeric_cells",
                         "period", "source_scope", "adopted_geography", "decision", "reason", "locator"])
        for key, name in FILES.items():
            for row_class in ("total", "rural", "urban"):
                for idx, field in enumerate(FIELDS, start=2):
                    count = inventory[key]["numeric_cells_by_row_class_and_field"].get(f"{row_class}.{field}", 0)
                    adopted = (row_class, field) in ADOPTED and count > 0
                    writer.writerow([f"pak-pbs-census2023-table1-{key}", HASHES[name], "Sheet1", field,
                                     row_class, count, "2023" if field != "population_2017" else "2017",
                                     "Pakistan Table 1 4 provinces and ICT; AJK/GB not in these XLSX",
                                     "Table 1 national scope, four provinces and 136 districts only" if adopted else "",
                                     "adopted_at_named_upper_levels" if adopted else "priority_unassessed",
                                     "Direct source count for 2023 Table 1 coverage only; lower-unit rows remain unadopted" if adopted else
                                     "Retained in pinned source; definition, geography or denominator not yet accepted",
                                     f"Sheet1 col {idx}; Excel row and named reporting-unit key"])
    with (evidence / "PAK_CATALOGUE_TABLES.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["table", "title", "xlsx_locations", "acquired_locations", "mismatched_links", "status"])
        for item in catalogue:
            writer.writerow([item["table"], item["title"], len(item["xlsx_locations"]),
                             len(item["acquired_urls"]), len(item["mismatched_links"]), item["status"]])
    report = {"checked_at_utc": datetime.now(timezone.utc).isoformat(),
              "catalogue_url": CATALOGUE_URL, "catalogue_sha256": HASHES["pbs2023-result-excel-catalog.html"],
              "numbered_catalogue_tables": len(catalogue),
              "catalogue_xlsx_locations": sum(len(x["xlsx_locations"]) for x in catalogue),
              "catalogue_mismatched_links": [
                  {"table": x["table"], "urls": x["mismatched_links"]} for x in catalogue if x["mismatched_links"]],
              "acquired_table1_xlsx": len(FILES), "table1_workbooks": inventory,
              "all_populated_numeric_cells": sum(all_cell_counts.values()),
              "numeric_cells_by_row_class_and_field": {f"{r}.{f}": n for (r, f), n in sorted(all_cell_counts.items())},
              "reporting_units": sum(len(x) for x in groups_by_source.values()),
              "districts": {key: len(x) for key, x in district_totals.items()},
              "checks": dict(check), "exceptions": exceptions,
              "national_table1_population_2023": national[0]["parts"]["total"]["values"][1],
              "source_notes": national[0] and inventory["national"]["notes"],
              "limitations": ["Table 1 national figure covers the four provinces and ICT in the source workbook; AJK/GB are listed separately by PBS GIS.",
                              "Table 4 onward may use a different detailed-characteristic population base per Table 1 footnote.",
                              "Districts and lower reporting units have no official codes in these XLSX; no current local-government or polygon join is implied.",
                              "Only six Table 1 workbooks acquired. Other numbered catalogue tables are official locations, not inspected source bodies.",
                              "Lower reporting units and derived Table 1 columns remain unadopted pending hierarchy and definition checks."]}
    (evidence / "PAK_PBS2023_TABLE1_AUDIT.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("numbered_catalogue_tables", "catalogue_xlsx_locations",
          "acquired_table1_xlsx", "all_populated_numeric_cells", "reporting_units", "districts", "checks", "exceptions")},
          ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True, type=Path)
    audit(parser.parse_args().project)
