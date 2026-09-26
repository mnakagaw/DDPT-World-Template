#!/usr/bin/env python3
"""Prepare a complete source-column review register; never mark it accepted."""

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def normalized(value):
    return " ".join(str(value).split()) if value is not None else ""


parser = argparse.ArgumentParser()
parser.add_argument("--project", required=True)
parser.add_argument("--manifest", required=True)
args = parser.parse_args()
project = Path(args.project).resolve(strict=True)
manifest_path = Path(args.manifest).resolve(strict=True)
if not manifest_path.is_relative_to(project):
    raise SystemExit("Source manifest must be inside project")
manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
receipt = next(r for r in manifest["receipts"] if r["file"] == "Census_Final_Results.xlsx")
source = manifest_path.parent / receipt["file"]
if receipt["status"] != "acquired" or hashlib.sha256(source.read_bytes()).hexdigest() != receipt["sha256"]:
    raise SystemExit("Official workbook receipt mismatch")
inventory_path = project / "evidence/QAT_CENSUS2020_NUMERIC_COLUMNS.csv"
with inventory_path.open(encoding="utf-8-sig", newline="") as stream:
    inventory = list(csv.DictReader(stream))
if len(inventory) != 1588:
    raise SystemExit("Expected 1,588 source numeric columns")
by_sheet = {}
for column in inventory:
    by_sheet.setdefault(column["sheet"], []).append(column)

workbook = load_workbook(source, read_only=True, data_only=True)
rows = []
for sheet_name, columns in by_sheet.items():
    sheet = workbook[sheet_name]
    cells = list(sheet.iter_rows(values_only=True))
    title = columns[0]["title_en"]
    if re.fullmatch(r"[1-9]\d*", sheet_name):
        if "Municipality" in title:
            grain = "2020 census municipality cross-tab"
            reason = "Municipality axis is present; every category and denominator requires source-cell review before any further indicator adoption."
        elif re.search(r"\bZone\b", title):
            grain = "2020 census zone cross-tab"
            reason = "Zone axis is present; four official GIS zone IDs lack workbook table-2 rows and zone 99 population is null. Reconcile table-specific area coverage before adoption."
        else:
            grain = "national cross-tab (no municipality or zone in table title)"
            reason = "No municipality or zone axis is declared in the table title; keep source categories national and do not copy a national value to a local area."
    elif sheet_name == "المحتويات":
        grain = "contents index"
        reason = "Numeric values are table numbers in a contents index, not census observations."
    else:
        grain = "census questionnaire/form"
        reason = "Numeric values are in a source questionnaire/form; do not treat codes, options or examples as published result observations."

    for entry in columns:
        first = int(entry["first_numeric_row"])
        index = 0
        for letter in entry["column"]:
            index = index * 26 + ord(letter) - ord("A") + 1
        index -= 1
        header = []
        for row_number in range(4, min(first, 12)):
            row = cells[row_number - 1]
            if index < len(row) and isinstance(row[index], str):
                value = normalized(row[index])
                if value and value not in header:
                    header.append(value)
        axis_examples = []
        for row_number in range(first, min(first + 4, len(cells) + 1)):
            row = cells[row_number - 1]
            labels = [normalized(value) for value in row[:index]
                      if isinstance(value, str) and normalized(value)]
            if labels:
                axis_examples.append(f"{row_number}:" + " / ".join(labels[-3:]))
        header_text = " | ".join(header)
        if re.search(r"ratio|average|mean|percent|%", f"{title} {header_text}", re.I):
            source_unit = "mixed_or_derived_measure_requires_row_check"
        elif grain in {"contents index", "census questionnaire/form"}:
            source_unit = "not_a_published_statistical_measure"
        else:
            source_unit = "source_native_count_or_measure_not_yet_normalized"
        rows.append({
            "source_sha256": receipt["sha256"],
            "sheet": sheet_name,
            "source_column": entry["column"],
            "title_en": title,
            "source_column_heading": header_text[:500],
            "source_row_axis_examples": " | ".join(axis_examples)[:500],
            "source_header_preview_original": entry["header_preview"],
            "numeric_cells_typed": entry["typed_numeric_cells"],
            "numeric_cells_text": entry["numeric_text_cells"],
            "source_granularity_triage": grain,
            "population_or_universe_from_title": title if title else "not specified in table title",
            "source_unit_triage": source_unit,
            "period": "2020 Census workbook; table-specific reference period not independently audited",
            "existing_adoption_decision": entry["adoption_decision"],
            "r3_triage_disposition": (
                "non_result_index_or_form" if grain in {"contents index", "census questionnaire/form"}
                else "pending_semantic_row_review" if grain.startswith("2020 census")
                else "national_cross_tab_not_assigned_to_municipality"
            ),
            "reason_and_next_action": reason,
            "semantic_acceptance": "not_assessed_by_structural_triage",
        })
workbook.close()

out = project / "evidence/QAT_R3_COLUMN_TRIAGE.csv"
with out.open("w", encoding="utf-8-sig", newline="") as stream:
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
summary = {
    "source_sha256": receipt["sha256"],
    "numeric_columns": len(rows),
    "source_granularity_triage_counts": dict(Counter(r["source_granularity_triage"] for r in rows)),
    "r3_triage_disposition_counts": dict(Counter(r["r3_triage_disposition"] for r in rows)),
    "semantic_acceptance": "not_assessed_by_structural_triage",
    "interpretation": "This register gives every numeric column a source title, heading, sample row labels and conservative granularity/disposition. It is not the required cell-level semantic, unit, denominator or adoption audit.",
}
(project / "evidence/QAT_R3_COLUMN_TRIAGE_SUMMARY.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False))
