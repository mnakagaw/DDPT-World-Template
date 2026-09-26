"""Adopt the governorate rows in Iraq's official 2024 census summary workbook.

The five numeric source columns are separate published concepts. In particular,
the post-adjustment population is not substituted for the original tabulation.
The geoBoundaries polygons are name-matched reference shapes, not legal 2024
census boundaries or official administrative-code matches.
"""

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


SOURCE_ID = "irq-cosit-census-2024-final-governorates"
SOURCE_URL = "https://nogp.gov.iq/DatasetDetails.aspx?id=17"
SOURCE_FILE = "iraq-census-2024.xlsx"
SOURCE_HASH = "2f2446396744ec8c760b797d46b2b177ec8954aa2ef8542c5c319f9b608176f9"
PREFIX = "IRQ_CENSUS_2024_"
NAME_CROSSWALK = {
    "دهوك": "Dohuk", "نينوى": "Ninawa", "السليمانية": "Al-Sulaimaniyah",
    "كركوك": "Kirkuk", "اربيل": "Erbil", "ديالى": "Diyala",
    "الانبار": "Al-Anbar", "بغداد": "Baghdad", "بابل": "Babil",
    "كربلاء": "Karbala", "واسط": "Wasit", "صلاح الدين": "Salah al-Din",
    "النجف": "An-Najaf", "القادسية": "Al-Qadisiyah", "المثنى": "Al-Muthanna",
    "ذي قار": "Dhi Qar", "ميسان": "Maysan", "البصرة": "Al-Basrah",
}
FIELDS = (
    ("POP_TABULATED", 2, "Population of Iraq — original census tabulation", "people"),
    ("POP_ADJUSTED", 3, "Population after Cabinet Decision 24853/2024 adjustment", "people"),
    ("FOREIGN_POP", 5, "Foreign population in adjusted column", "people"),
    ("IRAQI_POP_ADJUSTED", 7, "Iraqi population after adjustment, excluding foreigners", "people"),
    ("IRAQI_SHARE_ADJUSTED", 9, "Governorate share of adjusted Iraqi population", "%"),
)


def positive_integer(value, cell):
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"Expected nonnegative integer in {cell}: {value!r}")
    return value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True)
    args = parser.parse_args()
    project = Path(args.project)
    raw = project / "raw" / SOURCE_FILE
    if hashlib.sha256(raw.read_bytes()).hexdigest() != SOURCE_HASH:
        raise ValueError("Official workbook SHA-256 mismatch")
    book = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    if book.sheetnames != ["Sheet1"]:
        raise ValueError(f"Unexpected sheets: {book.sheetnames}")
    sheet = book.active
    rows = list(sheet.values)
    if len(rows) != 24 or len(rows[4:22]) != 18 or rows[22][0] != "الإجمالي":
        raise ValueError("Unexpected published governorate table layout")
    if "2024" not in rows[0][0] or rows[3][0] != "المحافظة":
        raise ValueError("Census year or governorate header changed")
    national = rows[22]
    if national[2] != 46118793 or national[3] != 46118793:
        raise ValueError("Published national population changed")
    if national[5] != 340131 or national[7] != 45778662:
        raise ValueError("Published national foreign/Iraqi population changed")
    if abs(national[9] - 1.0) > 1e-9:
        raise ValueError("National share does not equal one")
    if set(NAME_CROSSWALK) != {row[0] for row in rows[4:22]}:
        raise ValueError("Governorate names changed or crosswalk incomplete")

    dataset_path = project / "data" / "dashboard.json"
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    governorates = {item["name"]: item for item in dataset["territories"]
                    if item["parent_id"] == "IRQ"}
    if len(governorates) != 18 or set(governorates) != set(NAME_CROSSWALK.values()):
        raise ValueError("Reference shapes differ from the manually checked name crosswalk")
    if len(dataset["boundaries"]["features"]) != 18:
        raise ValueError("Reference geometry count changed")
    if {feature["properties"]["territory_id"] for feature in dataset["boundaries"]["features"]} != {
            item["id"] for item in governorates.values()}:
        raise ValueError("Reference boundary IDs do not cover the registered governorates")

    for column in (2, 3, 5, 7):
        if sum(positive_integer(row[column], f"{openpyxl.utils.get_column_letter(column+1)}{number}")
               for number, row in enumerate(rows[4:22], 5)) != national[column]:
            raise ValueError(f"Governorate/national sum mismatch in column {column}")
    if abs(sum(row[9] for row in rows[4:22]) - 1.0) > 1e-9:
        raise ValueError("Governorate shares do not total one")
    for row in rows[4:23]:
        if row[3] - row[5] != row[7] or abs(row[9] - row[7] / national[7]) > 1e-9:
            raise ValueError(f"Adjusted population / foreign / share identity failed for {row[0]}")

    observations = []
    crosswalk = []

    def add_observations(territory_id, row, row_number):
        for suffix, column, _, unit in FIELDS:
            number = row[column]
            if unit == "people":
                positive_integer(number, f"{openpyxl.utils.get_column_letter(column+1)}{row_number}")
            elif not isinstance(number, (float, int)) or not 0 <= number <= 1:
                raise ValueError(f"Invalid published share in J{row_number}")
            observations.append({
                "territory_id": territory_id, "indicator_id": PREFIX + suffix,
                "period": "2024", "value": number * 100 if unit == "%" else number,
                "status": "observed", "measurement_method": "source_reported",
                "source_id": SOURCE_ID,
                "source_locator": f"Sheet1!{openpyxl.utils.get_column_letter(column+1)}{row_number}",
            })

    add_observations("IRQ", national, 23)
    for row_number, row in enumerate(rows[4:22], 5):
        territory = governorates[NAME_CROSSWALK[row[0]]]
        territory["source_name_ar"] = row[0]
        territory["reconciliation_status"] = "name_correspondence_only_reference_boundary_unverified"
        add_observations(territory["id"], row, row_number)
        crosswalk.append((row_number, row[0], territory["name"], territory["id"],
                          territory["provider_code"], "name_correspondence_only"))

    definitions = {
        "POP_TABULATED": "Source header: سكان العراق. The original governorate population column; do not conflate with the separate Cabinet-decision-adjusted column.",
        "POP_ADJUSTED": "Source header describes adjustment of residents and displaced persons who are not inhabitants of the Kurdistan Region, returning them to their governorates under Cabinet Decision 24853/2024 paragraph 5.",
        "FOREIGN_POP": "Source header: عدد السكان الأجانب. Foreigners in the adjusted column, not a second population total.",
        "IRAQI_POP_ADJUSTED": "Source header: صافي عدد السكان الإجمالي من العراقيين (عدا الأجانب). Adjusted governorate population minus foreign population.",
        "IRAQI_SHARE_ADJUSTED": "Published share of the adjusted Iraqi population across 18 governorates. Stored as percentage points rather than the workbook's 0–1 fraction.",
    }
    indicators = [{"id": PREFIX + suffix, "name": label, "theme": "Census 2024",
                   "unit": unit, "definition": definitions[suffix],
                   "population": "Persons counted in Iraq 2024 census final governorate summary",
                   "source_id": SOURCE_ID, "aggregation": "sum" if unit == "people" else "none",
                   "measurement_method": "source_reported", "series_family": "census",
                   "period_policy": "latest_available_per_indicator", "display_decimals": 0 if unit == "people" else 2}
                  for suffix, _, label, unit in FIELDS]
    dataset["indicators"] = [item for item in dataset["indicators"] if not item["id"].startswith(PREFIX)] + indicators
    dataset["observations"] = [item for item in dataset["observations"]
                               if not item["indicator_id"].startswith(PREFIX)] + observations
    stamp = datetime.now(timezone.utc).isoformat()
    dataset["generated_at"] = stamp
    dataset["sources"] = [item for item in dataset["sources"] if item["id"] != SOURCE_ID]
    dataset["sources"].append({
        "id": SOURCE_ID, "name": "COSIT final results of the 2024 Population and Housing Census",
        "url": SOURCE_URL, "publisher": "Iraq Commission of Statistics and GIS (COSIT)",
        "reference_period": "2024", "geographic_level": "country, governorate",
        "status": "ready", "retrieved_at": stamp, "sha256": SOURCE_HASH,
        "raw_path": f"raw/{SOURCE_FILE}", "license": "CC BY 4.0 portal envelope; resource terms review pending",
        "note": "Official National Open Government Portal workbook. Five numeric source columns are kept separate. GeoBoundaries 2022 polygons are name-matched navigation references only, without certified 2024 boundary/code correspondence.",
    })
    dataset["country"]["geography_note"] = (
        "COSIT 2024 census summary supplies 18 governorate rows. The original and Cabinet-decision-adjusted population columns remain distinct. Governorate labels have a manual name correspondence to geoBoundaries 2022 reference shapes; no official codes or 2024 legal boundary version are certified. The World Bank national series is separate.")
    dataset["gaps"] = [item for item in dataset["gaps"] if item["category"] != "subnational_statistics"]
    dataset["gaps"].append({
        "category": "subnational_statistics", "status": "partial",
        "detail": "Five numeric columns from the official 2024 census summary adopted for all 18 governorates. District tables in the full report, other census themes and official codes remain unreviewed.",
        "next_action": "Audit district/thematic tables, official code and boundary correspondence, planning law/documents and source terms before acceptance.",
    })
    dataset["collection"]["adapters"] = sorted(set(dataset["collection"]["adapters"] + ["iraq-cosit-census-2024-final-summary"]))
    dataset["analysis"]["latest_values_only"] = True
    dataset_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    receipt = {"url": SOURCE_URL, "api_metadata_url": "https://nogp.gov.iq/api/v1/datasets/17",
               "resource_id": 23, "retrieved_at_utc": stamp, "size_bytes": raw.stat().st_size,
               "sha256": SOURCE_HASH, "redistribution_terms": "portal_cc_by_4_0_resource_review_pending",
               "acquisition": "Official site ASP.NET download postback; direct API resource endpoint returned internal error"}
    (project / "raw" / (SOURCE_FILE + ".receipt.json")).write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    evidence = project / "evidence"
    evidence.mkdir(exist_ok=True)
    inventory = [{"sheet": "Sheet1", "column": openpyxl.utils.get_column_letter(i+1),
                  "source_field": rows[3][i], "adoption": "identifier" if i == 0 else
                  "adopted" if i in (2, 3, 5, 7, 9) else "blank_column",
                  "source_rows": "5–23"} for i in range(11)]
    (evidence / "SOURCE_TABLE_INVENTORY.json").write_text(
        json.dumps({"country": "IRQ", "reference_year": 2024, "fields": inventory,
                    "footnote": rows[23][0], "full_report_status": "acquired_not_inventoried"},
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (evidence / "CODE_CROSSWALK.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("source_row", "official_arabic_name", "reference_english_name",
                         "areadata_territory_id", "geoboundaries_shape_id", "match_status"))
        writer.writerows(crosswalk)
    audit = {"country": "IRQ", "reference_year": 2024,
             "status": "partial_import_not_accepted", "source_url": SOURCE_URL,
             "sha256": SOURCE_HASH, "governorates": 18, "indicators": 5,
             "observations": len(observations),
             "national": {"tabulated": national[2], "adjusted": national[3],
                          "foreign": national[5], "adjusted_iraqi": national[7]},
             "reconciliation": "All four additive governorate columns sum to national; adjusted minus foreign equals adjusted Iraqi in every row; shares use adjusted Iraqi denominator and sum to one.",
             "unresolved": ["Official geography codes and 2024 boundaries", "District/thematic tables in 562-page official report", "Planning law and documents", "Resource-specific terms"]}
    (evidence / "IRQ_CENSUS_IMPORT_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"governorates": 18, "indicators": 5,
                      "observations": len(observations), "national": national[2]}))


if __name__ == "__main__":
    main()
