#!/usr/bin/env python3
"""Build an exact-code Guatemala health, nutrition and poverty theme bundle."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


HEALTH_URL = "https://datosabiertos.mspas.gob.gt/dataset/353ffabb-2a3d-4ac1-9816-001b7c3cf62c/resource/ad2416d9-6617-4968-9ed3-5de6efb32b05/download/mec-2024-departamento-municipio.csv"
NUTRITION_URL = "https://portal.siinsan.gob.gt/wp-content/uploads/VULNERABILIDAD_DC_MUNICIPAL_QUINTO_CENSO_2024_2.xlsx"
NUTRITION_REPORT_URL = "https://portal.sesan.gob.gt/wp-content/uploads/2025/10/20251006-01-QUINTO-CENSO-TALLA-2024_.pdf"
POVERTY_URL = "https://datos.segeplan.gob.gt/dataset/bfddf2e8-db9b-408f-b2b6-2519e207b150/resource/40ccb784-5f6d-47ab-8343-342fa2caabeb/download/datos-abiertos_resultados-mapa2023_mupio.csv"


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_semicolon(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=";")
        rows = list(reader)
    header = [cell.strip() for cell in rows[0]]
    return [dict(zip(header, [cell.strip() for cell in row])) for row in rows[1:] if any(cell.strip() for cell in row)]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True)
    parser.add_argument("--raw", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    project, raw_dir, out = Path(args.project), Path(args.raw), Path(args.out)
    dataset = json.loads((project / "data" / "dashboard.json").read_text(encoding="utf-8"))
    territories = dataset["territories"]
    by_id = {row["id"]: row for row in territories}
    municipalities = [row for row in territories if row.get("country_id") == "GTM" and row.get("level") == "municipality"]
    by_code = {str(row["official_code"]): row for row in municipalities}
    if len(by_code) != 340:
        raise SystemExit(f"Expected 340 Guatemala municipalities, found {len(by_code)}")
    exact = {(norm(by_id[row["parent_id"]]["name"]), norm(row["name"])): row for row in municipalities}

    health_file = raw_dir / "mec-2024-departamento-municipio.csv"
    nutrition_file = raw_dir / "VULNERABILIDAD_DC_MUNICIPAL_QUINTO_CENSO_2024.xlsx"
    nutrition_report = raw_dir / "GT-Quinto-Censo-Talla-2024.pdf"
    poverty_file = raw_dir / "GT-poverty-map-2023.csv"
    for required in (health_file, nutrition_file, nutrition_report, poverty_file):
        if not required.is_file():
            raise SystemExit(f"Missing required source: {required}")

    # MSPAS publishes names, not codes. These are explicit, reviewed source-label aliases;
    # every other name must match the adopted INE hierarchy exactly after accent/case folding.
    aliases = {
        ("alta verapaz", "cahabon"): "1612", ("alta verapaz", "la tinta"): "1616",
        ("alta verapaz", "lanquin"): "1611", ("alta verapaz", "tucuru"): "1606",
        ("baja verapaz", "el chol"): "1506", ("chimaltenango", "pochuta"): "408",
        ("chimaltenango", "yepocapa"): "412", ("chiquimula", "quetzaltepeque"): "2009",
        ("chiquimula", "san juan la ermita"): "2003", ("escuintla", "puerto de san jose"): "509",
        ("huehuetenango", "barillas"): "1326", ("huehuetenango", "san ildelfonso ixtahuacan"): "1309",
        ("jutiapa", "quezada"): "2217", ("quetzaltenango", "colomba"): "917",
        ("quetzaltenango", "olintepeque"): "903", ("quiche", "chichicastenango"): "1406",
        ("quiche", "ixcan"): "1420", ("quiche", "nebaj"): "1413",
        ("quiche", "pachalun"): "1421", ("sacatepequez", "alotenango"): "314",
        ("san marcos", "el rodeo"): "1214", ("suchitepequez", "pueblo nuevo suchitepequez"): "1019",
        ("totonicapan", "san bartolo"): "808",
    }
    health_rows = read_semicolon(health_file)
    health_totals: dict[str, int] = defaultdict(int)
    used_aliases: set[tuple[str, str, str]] = set()
    source_keys: set[tuple[str, str]] = set()
    for row in health_rows:
        dep, mun = norm(row["Departamento"]), norm(row["Municipio"])
        if dep == "el peten":
            dep = "peten"
        key = (dep, mun)
        source_keys.add(key)
        area = exact.get(key)
        if area is None:
            code = aliases.get(key)
            if not code or code not in by_code:
                raise SystemExit(f"Unmatched MSPAS geography: {key}")
            area = by_code[code]
            used_aliases.add((key[0], key[1], code))
        value = int(row["Casos"].replace(",", ""))
        health_totals[area["id"]] += value
    if len(health_totals) != 338:
        raise SystemExit(f"Expected 338 MSPAS municipality geographies, found {len(health_totals)}")
    health_missing = sorted(str(row["official_code"]) for row in municipalities if row["id"] not in health_totals)

    workbook = openpyxl.load_workbook(nutrition_file, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    nutrition: dict[str, tuple[float, int, str]] = {}
    for row in sheet.iter_rows(min_row=8, values_only=True):
        if row[0] is None:
            continue
        code = str(int(row[0]) * 100 + int(row[2]))
        if code not in by_code:
            raise SystemExit(f"Nutrition workbook has unknown municipality code {code}")
        nutrition[by_code[code]["id"]] = (round(float(row[6]), 1), int(row[5]), str(row[7]))
    if len(nutrition) != 340:
        raise SystemExit(f"Expected 340 nutrition municipalities, found {len(nutrition)}")

    poverty_rows = read_semicolon(poverty_file)
    poverty: dict[str, float] = {}
    for row in poverty_rows:
        code = str(int(row["Codigo"]))
        if code not in by_code:
            raise SystemExit(f"Poverty CSV has unknown municipality code {code}")
        poverty[by_code[code]["id"]] = float(row["Incidencia de pobreza general"])
    if len(poverty) != 340:
        raise SystemExit(f"Expected 340 poverty municipalities, found {len(poverty)}")

    source_specs = [
        ("GTM_MSPAS_CHRONIC_MORBIDITY_2024", health_file, HEALTH_URL),
        ("GTM_SESAN_SCHOOL_STUNTING_2024", nutrition_file, NUTRITION_URL),
        ("GTM_SESAN_SCHOOL_STUNTING_REPORT_2024", nutrition_report, NUTRITION_REPORT_URL),
        ("GTM_SEGEPLAN_MUNICIPAL_POVERTY_2023", poverty_file, POVERTY_URL),
    ]
    hashes = {source_id: sha256(filename) for source_id, filename, _ in source_specs}
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    sources = [
        {"id": "GTM_MSPAS_CHRONIC_MORBIDITY_2024", "name": "Morbilidad por enfermedades crónicas 2024, por departamento y municipio", "publisher": "Ministerio de Salud Pública y Asistencia Social de Guatemala", "url": HEALTH_URL, "status": "ready", "retrieved_at": now, "reference_period": "2024", "sha256": hashes["GTM_MSPAS_CHRONIC_MORBIDITY_2024"], "bytes": health_file.stat().st_size, "license": "Creative Commons Attribution", "license_url": "http://www.opendefinition.org/licenses/cc-by", "geographic_level": "municipality", "raw_path": "raw/guatemala-supplements/mec-2024-departamento-municipio.csv", "note": "Counts of reported chronic-disease morbidity records by diagnosis, age group and sex. Municipality totals are sums of source rows, not unique people or prevalence rates. Two municipalities have no source row and remain missing, not zero."},
        {"id": "GTM_SESAN_SCHOOL_STUNTING_2024", "name": "Vulnerabilidad por desnutrición crónica municipal, Quinto Censo Nacional de Talla 2024", "publisher": "SESAN / SIINSAN, Gobierno de Guatemala", "url": NUTRITION_URL, "catalog_url": "https://portal.siinsan.gob.gt/censo_2024/", "status": "ready", "retrieved_at": now, "reference_period": "2024", "sha256": hashes["GTM_SESAN_SCHOOL_STUNTING_2024"], "bytes": nutrition_file.stat().st_size, "license": "Reuse terms not stated on the official download page; the public dashboard links to the source and does not redistribute the workbook.", "geographic_level": "municipality", "raw_path": "raw/guatemala-supplements/VULNERABILIDAD_DC_MUNICIPAL_QUINTO_CENSO_2024.xlsx", "note": "Municipal prevalence of height-for-age stunting among children in first grade of public primary school, ages 6 years 0 months to 9 years 11 months. It is not prevalence among all children."},
        {"id": "GTM_SEGEPLAN_MUNICIPAL_POVERTY_2023", "name": "Mapa de Pobreza por Municipio 2023", "publisher": "SEGEPLAN / Instituto Nacional de Estadística Guatemala", "url": POVERTY_URL, "catalog_url": "https://datos.segeplan.gob.gt/dataset/mapa-pobreza-por-municipio-2023", "status": "ready", "retrieved_at": now, "reference_period": "2023", "sha256": hashes["GTM_SEGEPLAN_MUNICIPAL_POVERTY_2023"], "bytes": poverty_file.stat().st_size, "license": "Creative Commons Non-Commercial", "geographic_level": "municipality", "raw_path": "raw/guatemala-supplements/GT-poverty-map-2023.csv", "note": "Municipal small-area poverty estimates produced from ENCOVI 2023 and Census 2018 inputs. Values retain the source period and method and are not represented as Census enumeration counts."},
    ]
    indicators = [
        {"id": "GTM_MSPAS_CHRONIC_MORBIDITY_CASES_2024", "name": "Reported chronic-disease morbidity cases", "theme": "Health", "unit": "reported cases", "definition": "Sum of MSPAS 2024 chronic-disease morbidity case rows across diagnosis, age group and sex for the municipality.", "definition_id": "gtm_mspas_chronic_morbidity_cases_2024", "population": "Reported morbidity events in the MSPAS dataset", "measurement_method": "sum_of_source_rows", "aggregation": "sum", "source_id": "GTM_MSPAS_CHRONIC_MORBIDITY_2024", "series_family": "administrative", "display_role": "primary", "display_decimals": 0},
        {"id": "GTM_SESAN_SCHOOL_STUNTING_PCT_2024", "name": "First-grade public-school children with height-for-age stunting", "theme": "Nutrition", "unit": "% of measured first-grade public-school children", "definition": "Prevalence of height-for-age stunting among measured children in first grade of public primary school, ages 6 years 0 months to 9 years 11 months.", "definition_id": "gtm_sesan_school_stunting_pct_2024", "population": "Measured first-grade public-school children ages 6y0m to 9y11m", "measurement_method": "source_reported_prevalence", "aggregation": "official_only", "source_id": "GTM_SESAN_SCHOOL_STUNTING_2024", "series_family": "census", "display_role": "primary", "display_decimals": 1},
        {"id": "GTM_SEGEPLAN_GENERAL_POVERTY_PCT_2023", "name": "Population in general poverty", "theme": "Poverty", "unit": "% of population", "definition": "Municipal incidence of general poverty in the official 2023 small-area poverty map.", "definition_id": "gtm_segeplan_general_poverty_pct_2023", "population": "Municipal population represented by the 2023 small-area estimate", "measurement_method": "source_reported_small_area_estimate", "aggregation": "official_only", "source_id": "GTM_SEGEPLAN_MUNICIPAL_POVERTY_2023", "series_family": "survey", "display_role": "primary", "display_decimals": 1},
    ]
    observations = []
    for territory_id, value in sorted(health_totals.items()):
        observations.append({"territory_id": territory_id, "indicator_id": indicators[0]["id"], "period": "2024", "value": value, "status": "observed", "source_id": sources[0]["id"], "footnote": "Reported morbidity case total; not unique persons and not a population rate.", "source_precision": "integer"})
    for territory_id, (value, n, category) in sorted(nutrition.items()):
        observations.append({"territory_id": territory_id, "indicator_id": indicators[1]["id"], "period": "2024", "value": value, "status": "observed", "source_id": sources[1]["id"], "footnote": f"School height census denominator N={n}; official vulnerability category: {category}. This population excludes children outside first grade of the public primary sector.", "source_precision": "one decimal"})
    for territory_id, value in sorted(poverty.items()):
        observations.append({"territory_id": territory_id, "indicator_id": indicators[2]["id"], "period": "2023", "value": value, "status": "observed", "source_id": sources[2]["id"], "footnote": "Official municipal small-area estimate using ENCOVI 2023 and Census 2018 inputs; not a direct Census count.", "source_precision": "one decimal"})

    audit = {
        "schema_version": "1.0", "generated_at": now, "country_area_id": "GTM", "status": "three_supplemental_themes_integrated",
        "counts": {"municipality_registry": 340, "health_source_rows": len(health_rows), "health_source_geographies": len(source_keys), "health_observed_municipalities": len(health_totals), "health_missing_municipalities": health_missing, "nutrition_municipalities": len(nutrition), "poverty_municipalities": len(poverty), "observations": len(observations)},
        "identity": {"nutrition_and_poverty": "Exact published municipality codes joined to the adopted INE Census hierarchy.", "health": "Exact normalized department+municipality names plus the listed reviewed aliases; all 338 published source geographies matched uniquely.", "health_aliases": [{"source_department": d, "source_municipality": m, "official_code": c} for d, m, c in sorted(used_aliases)]},
        "meaning_safeguards": ["Health values are reported cases, not people or rates; absent municipality rows remain missing.", "Nutrition applies only to measured first-grade public-school children ages 6y0m to 9y11m.", "Poverty values are official small-area estimates using ENCOVI 2023 and Census 2018, not direct Census counts."],
        "receipts": [{"source_id": source_id, "url": url, "path": str(filename), "sha256": hashes[source_id], "bytes": filename.stat().st_size} for source_id, filename, url in source_specs],
    }
    bundle = {"schema_version": "0.2", "country_area_id": "GTM", "period": "latest-available", "sources": sources, "indicators": indicators, "observations": observations, "audit": audit}
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out.parent / "GTM_SUPPLEMENT_INTEGRATION_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit["counts"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
