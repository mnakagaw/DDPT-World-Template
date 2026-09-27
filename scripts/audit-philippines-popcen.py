"""Inventory PSA POPCEN originals and reconcile Table B to the 2024 Q2 PSGC.

The resulting audit is private to the generated project. This script does not
publish census values or interpret an LGU's legal planning responsibility.
"""

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


REGIONS = {
    "NCR": "13", "CAR": "14", "R01": "01", "R02": "02", "R03": "03",
    "R04A": "04", "MIMAROPA": "17", "R05": "05", "NIR": "18",
    "R06": "06", "R07": "07", "R08": "08", "R09": "09", "R10": "10",
    "R11": "11", "R12": "12", "Caraga": "16", "BARMM": "19",
}

# Source-row exceptions were inspected against Table B notes, Table A notes,
# Q2 PSGC name, geographic level, and 2020 population. Do not infer from a
# present-day catalogue, a name alone, or a later code edition.
OVERRIDES = {
    "NCR:11": ("1380300000", "Makati/Taguig ten-barangay transfer; Table B recasts 2020"),
    "NCR:23": ("1381500000", "Makati/Taguig ten-barangay transfer; Table B recasts 2020"),
    "MIMAROPA:69": ("1705323000", "Table B uses Dr. Jose P. Rizal; PSGC uses Rizal, with same 2020 count"),
    "R09:73": ("0908304000", "Imelda 2020 count differs from Q2 PSGC by 1,850; 2024 code/name checked"),
    "R09:80": ("0908311000", "Payao 2020 count differs from Q2 PSGC by 1,850; 2024 code/name checked"),
    "BARMM:64": ("1908700000", "Table B Maguindanao del Norte reporting total includes Cotabato City"),
    "BARMM:139": (None, "Special Geographic Area is a POPCEN reporting subtotal, not a PSGC LGU"),
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def normalized(value):
    value = unicodedata.normalize("NFKD", str(value))
    value = "".join(c for c in value if not unicodedata.combining(c)).upper()
    value = re.sub(r"\s*[0-9]+\s*\*?\s*$", "", value).replace("*", "")
    value = re.sub(r"\bCITY OF\b|\bCITY\b", "", value)
    return re.sub(r"[^A-Z0-9]+", "", value)


def int_cell(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and int(value) == value:
        return int(value)
    return None


def main(project):
    manifest_path = project / "evidence/PHL_SOURCE_MANIFEST.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    pinned = json.loads((Path(__file__).resolve().parents[1] /
        "config/philippines-popcen-source-manifest.json").read_text(encoding="utf-8"))
    require(manifest["source_count"] == 30, "Expected 30 PSA originals")
    by_id = {x["id"]: x for x in manifest["sources"]}
    require(len(by_id) == 30, "Duplicate PSA source ID")
    require(len(pinned["sources"]) == 30, "Expected 30 pinned PSA originals")
    for item in pinned["sources"]:
        live = by_id[item["id"]]
        require(all(live.get(key) == value for key, value in item.items()),
                "Acquired PSA original differs from pinned metadata: " + item["id"])
    for item in manifest["sources"]:
        body = (project / item["raw_path"]).read_bytes()
        require(len(body) == item["bytes"] and hashlib.sha256(body).hexdigest() == item["sha256"],
                "Original hash changed: " + item["id"])

    workbook_inventory = []
    for item in manifest["sources"]:
        if not item["raw_path"].endswith(".xlsx"):
            continue
        book = load_workbook(project / item["raw_path"], read_only=True, data_only=True)
        sheets = []
        for sheet in book:
            by_col = Counter()
            for row in sheet.values:
                for col, value in enumerate(row, 1):
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        by_col[col] += 1
            sheets.append({"sheet": sheet.title, "rows": sheet.max_row,
                           "numeric_cells_by_column": {str(k): v for k, v in sorted(by_col.items())},
                           "disposition": "duplicate_or_working_sheet"
                           if item["id"].startswith("phl-popcen-table-c-") and
                           (sheet.title == "Table B" or sheet.title.startswith("Table C_"))
                           else "inventory_only"})
        workbook_inventory.append({"source_id": item["id"], "sheets": sheets})
        book.close()

    q2 = load_workbook(project / by_id["phl-psgc-2024q2-datafile"]["raw_path"],
                       read_only=True, data_only=True)
    roster = {}
    candidate_index = defaultdict(list)
    levels = Counter()
    for row in q2["PSGC"].values:
        if len(row) < 9 or row[3] not in ("Reg", "Prov", "City", "Mun", "Bgy"):
            continue
        code = str(row[0]).zfill(10)
        require(re.fullmatch(r"\d{10}", code) and code not in roster, "Invalid or duplicate PSGC code")
        entry = {"code": code, "name": str(row[1]).strip(), "level": row[3],
                 "city_class": row[5], "population_2020": int_cell(row[8])}
        roster[code] = entry
        levels[row[3]] += 1
        if row[3] != "Bgy":
            candidate_index[(code[:2], normalized(row[1]), entry["population_2020"])].append(entry)
    q2.close()
    require(levels == {"Reg": 18, "Prov": 82, "City": 149, "Mun": 1493, "Bgy": 42004},
            f"Unexpected June PSGC roster: {levels}")

    source = load_workbook(project / by_id["phl-popcen-table-b"]["raw_path"],
                           read_only=True, data_only=True)
    rows = []
    exception_rows = []
    for sheet in source:
        require(sheet.title in REGIONS, f"Unexpected Table B sheet {sheet.title}")
        region_code = REGIONS[sheet.title] + "00000000"
        for row_number, cells in enumerate(sheet.values, 1):
            if len(cells) <= 5 or not isinstance(cells[0], str) or int_cell(cells[5]) is None:
                continue
            numbers = {"2010": int_cell(cells[2]), "2015": int_cell(cells[3]),
                       "2020": int_cell(cells[4]), "2024": int_cell(cells[5])}
            require(all(x is not None and x >= 0 for x in numbers.values()),
                    f"Incomplete Table B population {sheet.title}:{row_number}")
            original_name = cells[0].strip()
            key = f"{sheet.title}:{row_number}"
            if row_number == 7:
                code = region_code
                entry = roster[code]
                level = "Reg"
                method = "region_code_and_table_position"
            elif key in OVERRIDES:
                code, reason = OVERRIDES[key]
                entry = roster[code] if code else None
                level = entry["level"] if entry else "SGA"
                method = "documented_exception"
                exception_rows.append({"source_row": key, "code": code, "reason": reason,
                                       "table_b_2020": numbers["2020"],
                                       "psgc_2020": entry["population_2020"] if entry else None})
            else:
                matches = candidate_index[(REGIONS[sheet.title], normalized(original_name), numbers["2020"])]
                require(len(matches) == 1,
                        f"Unresolved or ambiguous Table B to Q2 PSGC match: {key} {original_name} {matches}")
                entry = matches[0]
                code = entry["code"]
                level = entry["level"]
                method = "region_name_and_2020_population_unique"
            if entry and level == "Reg" and entry["population_2020"] != numbers["2020"]:
                require(sheet.title in {"R06", "R07", "BARMM"},
                        f"Unexpected region historical count mismatch {sheet.title}")
                exception_rows.append({"source_row": key, "code": code,
                    "reason": "PSGC 2020 population not restated to Table B 2024 reporting geography",
                    "table_b_2020": numbers["2020"], "psgc_2020": entry["population_2020"]})
            rows.append({"sheet": sheet.title, "row": row_number, "source_name": original_name,
                         "psgc_name": entry["name"] if entry else None, "code": code,
                         "level": level, "city_class": entry["city_class"] if entry else None,
                         "population": numbers, "q2_psgc_2020": entry["population_2020"] if entry else None,
                         "match_method": method, "source_locator": f"Table B/{sheet.title}!A{row_number}:J{row_number}"})
    source.close()
    require(len(rows) == 1743 and len({x["code"] for x in rows if x["code"]}) == 1742,
            "Unexpected Table B reporting area count or duplicate PSGC code")
    require(Counter(x["level"] for x in rows) ==
            {"Reg": 18, "Prov": 82, "City": 149, "Mun": 1493, "SGA": 1},
            "Table B reporting levels differ from PSGC")
    require(len(exception_rows) == 10, "Unexpected source-code exceptions")

    # Table B defines its 2024 *reporting* hierarchy. A province subtotal
    # excludes HUCs, while the SGA is a separate non-LGU reporting subtotal.
    current_province = None
    current_sheet = None
    for row in rows:
        if row["sheet"] != current_sheet:
            current_sheet = row["sheet"]
            current_province = None
        level, code = row["level"], row["code"]
        region_id = "PHL:PSGC:" + REGIONS[row["sheet"]] + "00000000"
        if level == "Reg":
            row["parent_id"] = "PHL"
        elif level == "Prov":
            row["parent_id"] = region_id
            current_province = code
        elif level == "SGA":
            row["parent_id"] = region_id
            current_province = "SGA"
        elif code.startswith("19999"):
            require(current_province == "SGA", "SGA municipality outside Table B SGA block")
            row["parent_id"] = "PHL:POPCEN:SGA-2024"
        elif row["city_class"] == "HUC" or current_province is None or current_province == "SGA":
            row["parent_id"] = region_id
        elif code[:4] == current_province[:4]:
            row["parent_id"] = "PHL:PSGC:" + current_province
        else:
            row["parent_id"] = region_id

    id_for = lambda row: "PHL:PSGC:" + row["code"] if row["code"] else "PHL:POPCEN:SGA-2024"
    by_parent = defaultdict(list)
    for row in rows:
        by_parent[row["parent_id"]].append(row)
    reconciliation = []
    for parent in rows:
        children = by_parent.get(id_for(parent), [])
        if not children:
            continue
        subtotal = sum(x["population"]["2024"] for x in children)
        reconciliation.append({"parent_id": id_for(parent), "source_name": parent["source_name"],
                               "observed_2024": parent["population"]["2024"],
                               "children_2024": subtotal,
                               "difference": parent["population"]["2024"] - subtotal,
                               "child_count": len(children)})
    mismatches = [x for x in reconciliation if x["difference"]]

    table_a = load_workbook(project / by_id["phl-popcen-table-a"]["raw_path"],
                            read_only=True, data_only=True)
    a = table_a.active
    country_value = int_cell(a["I7"].value)
    note = str(a["B172"].value)
    table_a.close()
    require(country_value == 112729484 and "1,708" in note,
            "National count or diplomatic-persons note changed")
    region_sum = sum(x["population"]["2024"] for x in rows if x["level"] == "Reg")
    require(country_value - region_sum == 1708, "Table A region/national geography mismatch")

    audit = {"schema_version": "1.0", "country": "PHL", "census_reference_date": "2024-07-01",
             "psgc_edition": "2024-06-30", "source_count": 30,
             "workbook_inventory": workbook_inventory,
             "psgc_roster_levels": dict(levels), "table_b_rows": rows,
             "table_b_levels": dict(Counter(x["level"] for x in rows)),
             "source_code_exceptions": exception_rows,
             "table_a_national_2024": country_value, "table_b_region_sum_2024": region_sum,
             "non_territorial_diplomatic_persons": 1708,
             "reporting_parent_reconciliation": reconciliation,
             "reporting_parent_mismatches": mismatches,
             "adoption": {"table_a": "2024 national population only",
                          "table_b": "2024 direct population column F for all reporting rows; exceptions documented",
                          "table_c": "priority_unassessed: full workbook/column inventory only",
                          "older_population_and_pgr": "priority_unassessed for cross-period comparable indicators",
                          "psgc_q4": "historical code edition lead, not applied to July 2024 census"}}
    out = project / "evidence/PHL_POPCEN_AUDIT.json"
    out.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inventory_path = project / "evidence/PHL_INDICATOR_INVENTORY.csv"
    with inventory_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=["source_id", "source_sha256", "sheet",
            "column_number", "original_field", "numeric_cells", "role", "decision", "reason"])
        writer.writeheader()
        for workbook in workbook_inventory:
            source_id = workbook["source_id"]
            for sheet in workbook["sheets"]:
                for col, count in sheet["numeric_cells_by_column"].items():
                    number = int(col)
                    label = f"Numeric column {number}; inspect original sheet header"
                    role, decision, reason = "unclassified", "priority_unassessed", "Original numeric column inventoried; full semantics pending"
                    if source_id == "phl-popcen-table-a":
                        label = {3:"2010 population",5:"2015 population",7:"2020 population",
                            9:"2024 population",11:"2010–2015 PGR",12:"2015–2020 PGR",
                            13:"2015–2024 PGR",14:"2020–2024 PGR"}.get(number,label)
                        if number == 9:
                            role, decision, reason = "direct_count", "adopted_national_row_only", "Table A!I7 national 2024 count"
                    elif source_id == "phl-popcen-table-b":
                        label = {3:"2010 population",4:"2015 population",5:"2020 population",
                            6:"2024 population",7:"2010–2015 PGR",8:"2015–2020 PGR",
                            9:"2015–2024 PGR",10:"2020–2024 PGR"}.get(number,label)
                        if number == 6:
                            role, decision, reason = "direct_count", "adopted_all_reporting_rows", "2024 column F direct count with Q2 PSGC crosswalk"
                        elif number == 5:
                            role, decision, reason = "historical_join_check", "support_only", "Historical count checked for identity, but recasts differ in ten rows"
                    elif source_id.startswith("phl-popcen-table-c-"):
                        if number == 4:
                            label = "2024 population by province, city, municipality or barangay"
                        if sheet["disposition"] == "duplicate_or_working_sheet":
                            role, decision, reason = "duplicate_or_working", "not_adopted", "Table B duplicate or Table C working sheet; not a new independent table"
                        else:
                            role, decision, reason = "direct_count", "priority_unassessed", "Barangay-level code/parent, all source rows and possible subtotals need semantic crosswalk"
                    elif source_id == "phl-psgc-2024q2-datafile" and sheet["sheet"] == "PSGC":
                        label = {1:"10-digit PSGC code",3:"Correspondence code",9:"2020 population"}.get(number,label)
                        if number in (1,3,9):
                            role, decision, reason = "code_or_historical_check", "support_only", "Q2 2024 code identity and historical check, not an adopted 2024 population"
                    writer.writerow({"source_id": source_id, "source_sha256": by_id[source_id]["sha256"],
                        "sheet": sheet["sheet"], "column_number": number,
                        "original_field": label, "numeric_cells": count,
                        "role": role, "decision": decision, "reason": reason})
    print(json.dumps({"table_b_rows": len(rows), "levels": audit["table_b_levels"],
                      "source_code_exceptions": len(exception_rows),
                      "national": country_value, "region_sum": region_sum,
                      "reporting_parent_mismatches": mismatches,
                      "workbooks": len(workbook_inventory)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
