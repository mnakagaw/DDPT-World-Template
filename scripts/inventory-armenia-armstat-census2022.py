"""Inventory every numeric column in ArmStat's nine Census 2022 chapter archives.

The archive originals and this mechanical inventory stay in the ignored country
project. Column presence is not semantic acceptance of any observation.
"""

import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import xlrd


BASE_URL = "https://armstat.am/file/article/section_{}.7z"
SELECTED = {"table 1.1.1.xlsx": {4, 7, 10},
            "table 1.1.2.xlsx": {4, 7, 10}}
VALIDATION_ONLY = {"table 1.2.xls": {2, 5}}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inspect_sheet(name, rows, column_count):
    numeric = Counter()
    nonempty_rows = 0
    title = None
    for row_no, row in enumerate(rows, 1):
        values = list(row)
        if any(value not in (None, "") for value in values):
            nonempty_rows += 1
        if row_no <= 4 and not title:
            title = next((str(value).strip().replace("\n", " ")
                          for value in values if isinstance(value, str)
                          and ("table " in value.lower() or "աղյուսակ" in value.lower())), None)
        for col, value in enumerate(values, 1):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric[col] += 1
    return {"sheet": name, "nonempty_rows": nonempty_rows, "max_columns": column_count,
            "title": title, "numeric_cells_including_headers": sum(numeric.values()),
            "numeric_columns": {str(col): {"cell_count_including_headers": count,
                "disposition": "selected_rows_adopted_other_rows_unassessed"
                if col in SELECTED.get(name, set()) else
                "used_for_validation_not_adopted" if col in VALIDATION_ONLY.get(name, set())
                else "semantically_unassessed"}
                for col, count in sorted(numeric.items())}}


def inspect_book(path):
    sheets = []
    if path.suffix == ".xls":
        workbook = xlrd.open_workbook(str(path))
        for sheet in workbook.sheets():
            entry = inspect_sheet(path.name, (sheet.row_values(row) for row in range(sheet.nrows)),
                                  sheet.ncols)
            entry["worksheet"] = sheet.name
            entry["rows"] = sheet.nrows
            sheets.append(entry)
    else:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                entry = inspect_sheet(path.name, sheet.iter_rows(values_only=True), sheet.max_column)
                entry["worksheet"] = sheet.title
                entry["rows"] = sheet.max_row
                sheets.append(entry)
        finally:
            workbook.close()
    return {"path": str(path).replace("\\", "/"), "bytes": path.stat().st_size,
            "sha256": sha256(path), "sheets": sheets}


def main(project):
    root = project / "raw/armstat"
    chapters = []
    books = []
    for chapter in range(1, 10):
        archive = root / f"section_{chapter}.7z"
        if not archive.is_file():
            raise FileNotFoundError(archive)
        chapter_root = root / f"section_{chapter}"
        files = sorted(path for path in chapter_root.rglob("*") if path.is_file()
                       and not path.name.startswith("~$"))
        narratives = [path for path in files if path.suffix.lower() == ".docx"]
        workbooks = [path for path in files if path.suffix.lower() in (".xls", ".xlsx")]
        if len(narratives) != 1 or not workbooks:
            raise ValueError(f"Unexpected chapter {chapter} contents: {files}")
        chapters.append({"chapter": chapter, "url": BASE_URL.format(chapter),
                         "archive": str(archive.relative_to(project)).replace("\\", "/"),
                         "archive_bytes": archive.stat().st_size,
                         "archive_sha256": sha256(archive),
                         "narrative": str(narratives[0].relative_to(project)).replace("\\", "/"),
                         "narrative_sha256": sha256(narratives[0]),
                         "workbook_count": len(workbooks)})
        for path in workbooks:
            entry = inspect_book(path)
            entry["path"] = str(path.relative_to(project)).replace("\\", "/")
            entry["chapter"] = chapter
            entry["source_url"] = BASE_URL.format(chapter)
            books.append(entry)
    numeric_cells = sum(sheet["numeric_cells_including_headers"]
                        for book in books for sheet in book["sheets"])
    numeric_columns = sum(len(sheet["numeric_columns"])
                          for book in books for sheet in book["sheets"])
    result = {"inventoried_at": datetime.now(timezone.utc).isoformat(),
              "catalogue_url": "https://armstat.am/en/?nid=82&id=2623",
              "scope": "All nine English chapter archives; numeric cells include year and table headers."
                       " Only specified cells from Chapter 1 are adopted separately.",
              "chapters": chapters, "workbooks": books,
              "totals": {"chapters": len(chapters), "workbooks": len(books),
                         "worksheets": sum(len(book["sheets"]) for book in books),
                         "numeric_columns": numeric_columns,
                         "numeric_cells_including_headers": numeric_cells}}
    output = project / "evidence/ARM_ARMSTAT2022_SOURCE_INVENTORY.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({**result["totals"], "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True, type=Path)
    main(parser.parse_args().project)
