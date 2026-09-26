"""Read-only inventory of both NPC 2020 census workbooks.

The second URL contains ``2022`` in its filename, but the worksheet contents
must be checked before interpreting that string as a data reference year.
Originals and the generated inventory stay in the ignored country project.
"""

import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


URLS = {
    "Census_Final_Results.xlsx": "https://www.npc.qa/en/statistics/census2020/results/Documents/Census_Final_Results.xlsx",
    "Qatar_Census_2022_Final_Results.xlsx": "https://www.npc.qa/Style%20Library/NPC/CensusDetailedResults/documents/Qatar_Census_2022_Final_Results.xlsx",
    "Census_Final_Results.pdf": "https://www.npc.qa/en/statistics/census2020/results/Documents/Census_Final_Results.pdf",
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_summary(sheet):
    if not hasattr(sheet, "iter_rows"):
        return {"kind": "chart", "sha256_cell_values": None, "numeric_cells": 0}
    digest = hashlib.sha256()
    columns = Counter()
    title = None
    row_count = 0
    nonempty_rows = 0
    for row_count, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = list(row)
        digest.update(json.dumps(values, ensure_ascii=True, default=str).encode("utf-8"))
        if any(value is not None for value in values):
            nonempty_rows += 1
        if row_count <= 4:
            for value in values:
                if isinstance(value, str) and value.isascii() and len(value) > 18 and " by " in value.lower():
                    title = value.replace("\n", " ").strip()
                    break
        for col, value in enumerate(values, 1):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                columns[col] += 1
    return {"kind": "worksheet", "rows": row_count, "columns": sheet.max_column,
            "nonempty_rows": nonempty_rows, "title": title,
            "sha256_cell_values": digest.hexdigest(), "numeric_cells": sum(columns.values()),
            "numeric_columns": {str(col): n for col, n in sorted(columns.items())}}


def inspect(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {name: sheet_summary(workbook[name]) for name in workbook.sheetnames}
        return {"url": URLS[path.name], "path": str(path), "bytes": path.stat().st_size,
                "sha256": sha256(path), "sheet_count": len(workbook.sheetnames),
                "sheet_order": workbook.sheetnames, "sheets": sheets}
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    root = args.project / "raw/npc-census2020"
    originals = {name: inspect(root / name) for name in list(URLS)[:2]}
    complete, alternate = originals.values()
    common = set(complete["sheets"]) & set(alternate["sheets"])
    changed = sorted(name for name in common if
                     complete["sheets"][name]["sha256_cell_values"] !=
                     alternate["sheets"][name]["sha256_cell_values"])
    only_complete = [name for name in complete["sheet_order"] if name not in alternate["sheets"]]
    only_alternate = [name for name in alternate["sheet_order"] if name not in complete["sheets"]]
    assert len(complete["sheet_order"]) == 209 and len(alternate["sheet_order"]) == 206
    assert len(common) == 206 and len(only_complete) == 3 and not only_alternate and not changed
    assert complete["sheets"]["1"]["numeric_cells"] == 27
    assert complete["sheets"]["2"]["numeric_cells"] == 431
    pdf = root / "Census_Final_Results.pdf"
    assert pdf.read_bytes().startswith(b"%PDF")
    result = {"inventoried_at": datetime.now(timezone.utc).isoformat(),
              "meaning": "Mechanical cell inventory; semantic adoption is separate.",
              "originals": originals,
              "pdf": {"url": URLS[pdf.name], "path": str(pdf), "bytes": pdf.stat().st_size,
                      "sha256": sha256(pdf)},
              "comparison": {"common_sheets": len(common), "changed_cell_value_sheets": changed,
                             "only_complete_sheets": only_complete,
                             "only_alternate_sheets": only_alternate,
                             "common_numeric_cells": sum(complete["sheets"][name]["numeric_cells"]
                                                         for name in common),
                             "reference_year_from_numbered_sheets": 2020}}
    out = args.project / "evidence/QAT_NPC2020_WORKBOOK_INVENTORY.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"workbooks": [v["sheet_count"] for v in originals.values()],
                      "common": len(common), "changed": changed,
                      "numeric_cells": result["comparison"]["common_numeric_cells"],
                      "output": str(out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
