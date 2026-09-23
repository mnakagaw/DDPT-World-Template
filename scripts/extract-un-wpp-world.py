"""Normalize source-reported UN WPP 2024 Rev.1 world, region and country data.

The July demographic workbook and the male/female five-year-age workbooks
are kept as raw evidence outside Git. Exact M49 region codes are joined by
SDMX code; the custom Central America + Caribbean group is the documented sum
of the two complete, disjoint UN subregions for count indicators only.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

YEARS = {2000, 2023, 2024, 2025, 2026}
COUNT_FIELDS = {
    "UN_WPP_POP_TOTAL": ("Total Population, as of 1 July (thousands)", "Total population", "people"),
    "UN_WPP_POP_MALE": ("Male Population, as of 1 July (thousands)", "Male population", "people"),
    "UN_WPP_POP_FEMALE": ("Female Population, as of 1 July (thousands)", "Female population", "people"),
    "UN_WPP_BIRTHS": ("Births (thousands)", "Births", "people"),
    "UN_WPP_DEATHS": ("Total Deaths (thousands)", "Deaths", "people"),
    "UN_WPP_NET_MIGRANTS": ("Net Number of Migrants (thousands)", "Net migration", "people"),
}
RATE_FIELDS = {
    "UN_WPP_MEDIAN_AGE": ("Median Age, as of 1 July (years)", "Median age", "years"),
    "UN_WPP_FERTILITY": ("Total Fertility Rate (live births per woman)", "Total fertility rate", "births per woman"),
    "UN_WPP_LIFE_EXPECTANCY": ("Life Expectancy at Birth, both sexes (years)", "Life expectancy at birth", "years"),
    "UN_WPP_INFANT_MORTALITY": ("Infant Mortality Rate (infant deaths per 1,000 live births)", "Infant mortality", "per 1,000 live births"),
    "UN_WPP_UNDER5_MORTALITY": ("Under-Five Mortality (deaths under age 5 per 1,000 live births)", "Under-five mortality", "per 1,000 live births"),
    "UN_WPP_GROWTH": ("Population Growth Rate (percentage)", "Population growth rate", "%"),
    "UN_WPP_DENSITY": ("Population Density, as of 1 July (persons per square km)", "Population density", "people/km²"),
}
BASE_URL = "https://population.un.org/wpp/assets/Excel%20Files/1_Indicator%20(Standard)/EXCEL_FILES/"
DEMOGRAPHIC_URL = BASE_URL + "1_General/WPP2024_GEN_F01_DEMOGRAPHIC_INDICATORS_COMPACT.xlsx"
AGE_URLS = {
    "male": BASE_URL + "2_Population/WPP2024_POP_F02_2_POPULATION_5-YEAR_AGE_GROUPS_MALE.xlsx",
    "female": BASE_URL + "2_Population/WPP2024_POP_F02_3_POPULATION_5-YEAR_AGE_GROUPS_FEMALE.xlsx",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def value_number(value, multiplier=1):
    try:
        number = Decimal(str(value)) * multiplier
        return int(number) if multiplier == 1000 else float(number)
    except (InvalidOperation, TypeError):
        return None


def territory_id(row, columns, allowed, m49_groups):
    kind = row[columns["Type"]]
    if kind == "World":
        return "WLD"
    if kind == "Country/Area":
        country = row[columns["ISO3 Alpha-code"]]
        return country if country in allowed else None
    if kind in {"Region", "Subregion", "Intermediate region"}:
        code = row[columns["SDMX code**"]]
        try:
            region = f"M49:{int(code):03d}"
        except (TypeError, ValueError):
            return None
        return region if region in m49_groups else None
    return None


def rows_from_sheet(book, sheet_name):
    sheet = book[sheet_name]
    # WPP Rev.1 age files declare A1:L18 despite containing many more cells.
    sheet.reset_dimensions()
    iterator = sheet.iter_rows(values_only=True)
    for _ in range(16):
        next(iterator)
    columns = {name: index for index, name in enumerate(next(iterator)) if name is not None}
    return columns, iterator


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--demographic", required=True, type=Path)
    parser.add_argument("--age-male", required=True, type=Path)
    parser.add_argument("--age-female", required=True, type=Path)
    parser.add_argument("--registry", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--retrieved-at")
    args = parser.parse_args()
    retrieved = args.retrieved_at or datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    registry = json.loads(args.registry.read_text(encoding="utf-8"))
    allowed = {item["id"] for item in registry["territories"] if item["type"] == "country"}
    groups = {item["id"] for item in registry["territories"] if item["type"] == "exploration_scope"}
    # The world navigation combines M49 013 and 029 into one custom node.
    # Extract their exact published UN rows as calculation inputs even though
    # those two component nodes are not displayed in the AreaData registry.
    source_groups = groups | {"M49:013", "M49:029"}
    assert len(allowed) == 248 and "WLD" in groups
    fields = {**COUNT_FIELDS, **RATE_FIELDS}
    observed = []
    raw_details = {}
    book = load_workbook(args.demographic, read_only=True, data_only=True)
    for sheet_name, years, stage in [("Estimates", {2000, 2023}, "estimate"), ("Medium variant", {2024, 2025, 2026}, "medium_projection")]:
        columns, iterator = rows_from_sheet(book, sheet_name)
        for field, _, _ in fields.values():
            if field not in columns:
                raise ValueError(f"Missing WPP field {field!r}")
        for row in iterator:
            if len(row) <= columns["Year"] or row[columns["Year"]] not in years:
                continue
            area = territory_id(row, columns, allowed, source_groups)
            if not area:
                continue
            year = str(row[columns["Year"]])
            raw_details[(area, year)] = {"location_code": row[columns["Location code"]], "sdmx_code": row[columns["SDMX code**"]], "name": row[columns["Region, subregion, country or area *"]], "type": row[columns["Type"]]}
            for indicator_id, (field, _, _) in fields.items():
                number = value_number(row[columns[field]], 1000 if indicator_id in COUNT_FIELDS else 1)
                if number is None:
                    continue
                observed.append({"territory_id": area, "indicator_id": indicator_id, "period": year, "value": number, "status": "observed", "source_id": "un-wpp2024-global-rev1"})
    book.close()
    age_profiles = {}
    age_hashes = {}
    for sex, input_path in [("male", args.age_male), ("female", args.age_female)]:
        age_hashes[sex] = sha256(input_path)
        book = load_workbook(input_path, read_only=True, data_only=True)
        for sheet_name, years in [("Estimates", {2000, 2023}), ("Medium variant", {2024, 2025, 2026})]:
            columns, iterator = rows_from_sheet(book, sheet_name)
            ages = [name for name in columns if name == "100+" or isinstance(name, str) and len(name) <= 5 and "-" in name and name[0].isdigit()]
            if ages != [f"{i}-{i+4}" for i in range(0, 100, 5)] + ["100+"]:
                raise ValueError(f"Unexpected age groups: {ages}")
            for row in iterator:
                if len(row) <= columns["Year"] or row[columns["Year"]] not in years:
                    continue
                area = territory_id(row, columns, allowed, source_groups)
                if not area:
                    continue
                key = f"{area}@{row[columns['Year']]}"
                if key not in age_profiles:
                    age_profiles[key] = {"territory_id": area, "period": str(row[columns["Year"]]), "ages": ages}
                values = [value_number(row[columns[age]], 1000) for age in ages]
                if any(value is None for value in values):
                    raise ValueError(f"Age profile has missing value: {key} {sex}")
                age_profiles[key][sex] = values
        book.close()
    if any("male" not in row or "female" not in row for row in age_profiles.values()):
        raise ValueError("Male/female age profiles do not align")

    # This is an AreaData navigation grouping, not a UN-published region.
    # Summation is permitted for exact, disjoint M49 013 and 029 count fields.
    by_key = {(row["territory_id"], row["indicator_id"], row["period"]): row for row in observed}
    for year in sorted(YEARS):
        for target, component_ids in [("CUSTOM:CAM-CAR", ["M49:013", "M49:029"]), ("M49:019", ["M49:021", "M49:013", "M49:029", "M49:005"])]:
            for indicator_id in COUNT_FIELDS:
                members = [by_key.get((component, indicator_id, str(year))) for component in component_ids]
                if all(members):
                    observed.append({**members[0], "territory_id": target, "value": sum(member["value"] for member in members), "provenance": "calculated", "components": component_ids, "footnote": "AreaData sum of complete, disjoint source-reported UN WPP M49 regions for one year and variant; this is not a source-published UN aggregate."})
            profiles = [age_profiles.get(f"{component}@{year}") for component in component_ids]
            if all(profiles):
                age_profiles[f"{target}@{year}"] = {"territory_id": target, "period": str(year), "ages": profiles[0]["ages"], "male": [sum(profile["male"][index] for profile in profiles) for index in range(len(profiles[0]["ages"]))], "female": [sum(profile["female"][index] for profile in profiles) for index in range(len(profiles[0]["ages"]))], "provenance": "areadata_calculated", "components": component_ids}
    observed = [row for row in observed if row["territory_id"] in allowed or row["territory_id"] in groups]
    age_profiles = {key: profile for key, profile in age_profiles.items() if profile["territory_id"] in allowed or profile["territory_id"] in groups}

    indicators = []
    for indicator_id, (field, name, unit) in fields.items():
        is_count = indicator_id in COUNT_FIELDS
        indicators.append({"id": indicator_id, "name": name, "theme": "Population" if indicator_id.startswith("UN_WPP_POP_") or indicator_id in {"UN_WPP_MEDIAN_AGE", "UN_WPP_GROWTH", "UN_WPP_DENSITY"} else "Vital events and migration", "unit": unit, "definition": field + "; UN World Population Prospects 2024 Rev.1. Estimates through 2023 and medium-variant projections from 2024.", "definition_id": f"un-wpp2024-rev1-{indicator_id.lower()}", "population": "UN WPP covered countries and areas; source-published regional and world estimates", "measurement_method": "UN WPP 2024 demographic estimate or medium-variant projection", "aggregation": "sum" if is_count else "official_only", "series_family": "international_reference", "display_role": "primary" if indicator_id == "UN_WPP_POP_TOTAL" else "supplementary", "period_policy": "same_period", "series_stage_by_period": {str(year): "estimate" if year <= 2023 else "medium_projection" for year in YEARS}, "display_decimals": 0 if is_count else 1, "source_id": "un-wpp2024-global-rev1"})
    source = {"id": "un-wpp2024-global-rev1", "name": "World Population Prospects 2024 Rev.1 — demographic and age-sex indicators", "publisher": "United Nations, Department of Economic and Social Affairs, Population Division", "url": "https://population.un.org/wpp/", "status": "ready", "retrieved_at": retrieved, "reference_period": "2000, 2023 estimates; 2024–2026 medium projections", "geographic_level": "world_multi_scope_series", "license": "CC BY 3.0 IGO", "license_url": "https://creativecommons.org/licenses/by/3.0/igo/", "raw_redistribution_status": "not_in_public_site", "source_files": [{"url": DEMOGRAPHIC_URL, "sha256": sha256(args.demographic), "bytes": args.demographic.stat().st_size}, *[{"url": AGE_URLS[sex], "sha256": age_hashes[sex], "bytes": path.stat().st_size} for sex, path in [("male", args.age_male), ("female", args.age_female)]]], "note": "Exact WPP world, region, subregion and country/area rows. Americas and custom Central America + Caribbean count and age cells are complete non-overlapping AreaData sums of exact UN regional rows."}
    actual_countries = sorted({row["territory_id"] for row in observed if row["territory_id"] in allowed})
    result = {"schema_version": "1.0", "generated_at": retrieved, "source": source, "indicators": indicators, "observations": sorted(observed, key=lambda row: (row["territory_id"], row["indicator_id"], row["period"])), "age_profiles": age_profiles, "summary": {"registry_country_area_count": len(allowed), "wpp_country_area_count": len(actual_countries), "missing_country_area_ids": sorted(allowed - set(actual_countries)), "region_ids": sorted({row["territory_id"] for row in observed if row["territory_id"] in groups}), "periods": sorted(YEARS), "age_profile_count": len(age_profiles), "source_region_identity": {key[0]: detail for key, detail in raw_details.items() if key[1] == "2023" and key[0] in groups}}}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Wrote {args.out}: {len(observed)} observations, {len(age_profiles)} age-sex profiles, {len(actual_countries)}/{len(allowed)} country/area matches")


if __name__ == "__main__":
    main()
