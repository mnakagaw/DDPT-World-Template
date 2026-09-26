"""Import direct, rounded CAS LFHLCS 2018-19 Table 1.1 / HL.5 observations.

This is a household *survey* of residents of residential dwellings, not a
population census. The report's average population refers to mid-2018.
Source labels are not adopted as official administrative codes or polygons.
"""

import argparse
import json
import runpy
from datetime import datetime, timezone
from pathlib import Path

import xlrd
from openpyxl import load_workbook


fetch = runpy.run_path(str(Path(__file__).with_name("fetch-lebanon-cas-sources.py")))
FILES, digest = fetch["FILES"], fetch["digest"]
PDF = "lbn-cas-lfhlcs2018-report-table-1-1"
XLS = "lbn-cas-lfhlcs2018-demography-hl5"
MICS = "lbn-cas-mics2023-equitable-chance-ch11"
PLAN = "lbn-dgu-zoning-procedure"
PLAN_OFFICE = "lbn-mopw-dgu-office"
PLAN_OVERVIEW = "lbn-dgu-national-masterplan-overview"
PROFILE_CATALOG = "lbn-cas-district-profiles-2018-19"

# (row in original HL.5, geographic type, source label, parent governorate,
#  printed population, printed households, printed report page). Every row of
#  report Table 1.1 is represented, including the source's coextensive Beirut
#  and Akkar caza/governorate pairs. This is not an administrative code list.
ROWS = [
    (7, "caza", "Beirut", "Beirut", 341700, 100000, 18),
    (8, "governorate", "Beirut", None, 341700, 100000, 18),
    (9, "caza", "Baabda", "Mount Lebanon", 553800, 151600, 18),
    (10, "caza", "Matn", "Mount Lebanon", 511000, 147700, 18),
    (11, "caza", "Chouf", "Mount Lebanon", 277000, 76300, 18),
    (12, "caza", "Aley", "Mount Lebanon", 300800, 74900, 18),
    (13, "caza", "Keserwan", "Mount Lebanon", 260500, 77000, 18),
    (14, "caza", "Jbeil", "Mount Lebanon", 129500, 35100, 18),
    (15, "governorate", "Mount Lebanon", None, 2032600, 562600, 18),
    (16, "caza", "Tripoli", "North Lebanon", 243800, 55600, 18),
    (17, "caza", "Koura", "North Lebanon", 84600, 22200, 18),
    (18, "caza", "Zgharta", "North Lebanon", 87700, 21800, 18),
    (19, "caza", "Batroun", "North Lebanon", 58900, 17500, 18),
    (20, "caza", "Bcharre", "North Lebanon", 22100, 6300, 18),
    (21, "caza", "Minieh-Danniyeh", "North Lebanon", 140800, 29800, 18),
    (22, "governorate", "North Lebanon", None, 637900, 153200, 18),
    (23, "caza", "Akkar", "Akkar", 324000, 68200, 19),
    (24, "governorate", "Akkar", None, 324000, 68200, 19),
    (25, "caza", "Zahleh", "Bekaa", 177400, 46000, 19),
    (26, "caza", "West Beqaa", "Bekaa", 86400, 20400, 19),
    (27, "caza", "Rachaya", "Bekaa", 33800, 9800, 19),
    (28, "governorate", "Bekaa", None, 297700, 76200, 19),
    (29, "caza", "Baalbek", "Baalbek-Hermel", 214600, 51200, 19),
    (30, "caza", "Hermel", "Baalbek-Hermel", 30500, 7200, 19),
    (31, "governorate", "Baalbek-Hermel", None, 245100, 58400, 19),
    (32, "caza", "Saida", "South Lebanon", 296600, 74800, 19),
    (33, "caza", "Tyr", "South Lebanon", 255700, 63200, 19),
    (34, "caza", "Jezzine", "South Lebanon", 32100, 9800, 19),
    (35, "governorate", "South Lebanon", None, 584400, 147800, 19),
    (36, "caza", "Nabatieh", "Nabatieh", 180200, 45500, 19),
    (37, "caza", "Bint Jbeil", "Nabatieh", 96200, 26000, 19),
    (38, "caza", "Marjaayoun", "Nabatieh", 74000, 20600, 19),
    (39, "caza", "Hasbaya", "Nabatieh", 28700, 8100, 19),
    (40, "governorate", "Nabatieh", None, 379200, 100200, 19),
    (41, "national", "Lebanon", None, 4842500, 1266700, 19),
]
GOVERNORATES = [row[2] for row in ROWS if row[1] == "governorate"]
INDICATORS = [
    ("LBN_CAS18_RESIDENTS", "LFHLCS residents, mid-2018 estimate", "Population", "people, rounded to 100", PDF,
     "CAS survey estimate for people living in residential dwellings; excludes non-residential units including camps and informal settlements. Average over four 2018-19 rounds refers to mid-2018. Not a census count.", "Residents in residential dwellings"),
    ("LBN_CAS18_HOUSEHOLDS", "LFHLCS households, mid-2018 estimate", "Households", "households, rounded to 100", PDF,
     "CAS survey estimate of households in residential dwellings, rounded to 100 in published Table 1.1.", "Households in residential dwellings"),
    ("LBN_CAS18_WOMEN", "LFHLCS women residents, mid-2018 estimate", "Demography", "people, rounded to 100", XLS,
     "HL.5 survey estimate of women residents. Source values are in thousands of people and are converted to people, then rounded to 100 for display; the survey population excludes non-residential units.", "Women residents in residential dwellings"),
    ("LBN_CAS18_MEN", "LFHLCS men residents, mid-2018 estimate", "Demography", "people, rounded to 100", XLS,
     "HL.5 survey estimate of men residents. Source values are in thousands of people and are converted to people, then rounded to 100 for display; the survey population excludes non-residential units.", "Men residents in residential dwellings"),
]


def slug(label):
    return "-".join("".join(c if c.isalnum() else " " for c in label.upper()).split())


def area_id(kind, name, parent=None):
    if kind == "national":
        return "LBN"
    if kind == "governorate":
        return "LBN:CAS18:GOV:" + slug(name)
    return "LBN:CAS18:CAZA:" + slug(parent) + ":" + slug(name)


def rounded_100(thousands):
    return int(round(float(thousands) * 10)) * 100


def main(project):
    raw = project / "raw" / "official"
    evidence = project / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    for name, spec in FILES.items():
        if digest(raw / name) != spec["sha256"]:
            raise ValueError("Changed CAS original: " + name)
    book = xlrd.open_workbook(raw / "LFHLCS_2018_2019_Demography.xls", on_demand=True)
    if book.nsheets != 26:
        raise ValueError("Changed CAS demography sheet count")
    sheet = book.sheet_by_name("HL5")
    if (sheet.nrows, sheet.ncols) != (42, 7):
        raise ValueError("Changed CAS HL.5 shape")
    if "governorates,caza and sex" not in str(sheet.cell_value(0, 0)):
        raise ValueError("Changed CAS HL.5 title")
    for row, kind, name, parent, pop, households, page in ROWS:
        source_name = sheet.cell_value(row - 1, 1 if kind == "caza" else 0).strip()
        if kind == "governorate":
            # The XLS total row carries the parent name in the previous group
            # header and says only Total in its own row.
            source_name = sheet.cell_value(row - 1, 1).strip()
            if source_name != "Total":
                raise ValueError(f"Missing governorate total at row {row}")
        elif source_name != name:
            raise ValueError(f"CAS source caza label changed at row {row}: {source_name}")
        if rounded_100(sheet.cell_value(row - 1, 4)) != pop:
            raise ValueError(f"Published PDF population differs from HL.5 source row {row}")
        if households < 0 or households % 100:
            raise ValueError(f"Invalid printed household anchor row {row}")
    if len(ROWS) != 35 or len(GOVERNORATES) != 8 or sum(row[1] == "caza" for row in ROWS) != 26:
        raise ValueError("CAS source row universe changed")
    # Published counts are independently rounded to hundreds; the source
    # explicitly says its rows may not add exactly to the nationwide line.
    if sum(row[4] for row in ROWS if row[1] == "governorate") != 4842600:
        raise ValueError("Rounded governorate population subtotal changed")
    if sum(row[5] for row in ROWS if row[1] == "governorate") != 1266600:
        raise ValueError("Rounded governorate household subtotal changed")

    path = project / "data" / "dashboard.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    if data["country"]["id"] != "LBN":
        raise ValueError("Expected Lebanon project")
    own_sources = {PDF, XLS, MICS, PLAN, PLAN_OFFICE, PLAN_OVERVIEW, PROFILE_CATALOG}
    own_indicators = {x[0] for x in INDICATORS}
    data["territories"] = [x for x in data["territories"] if x["id"] == "LBN"]
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    data["observations"] = [x for x in data["observations"] if x["indicator_id"] not in own_indicators]
    data["indicators"] = [x for x in data["indicators"] if x["id"] not in own_indicators]
    data["sources"] = [x for x in data["sources"] if x["id"] not in own_sources]
    gov_ids = [area_id("governorate", name) for name in GOVERNORATES]
    caza_ids = []
    for _, kind, name, parent, _, _, _ in ROWS:
        if kind == "national":
            continue
        aid = area_id(kind, name, parent)
        if kind == "caza":
            caza_ids.append(aid)
        coextensive = kind == "caza" and name == parent and name in ("Beirut", "Akkar")
        data["territories"].append({
            "id": aid, "name": name, "level": "cas18_" + kind,
            "type": "CAS LFHLCS " + kind, "parent_id": "LBN" if kind == "governorate" else area_id("governorate", parent),
            "official_code": None, "code_system": "CAS LFHLCS 2018-19 source row; no official code adopted",
            "boundary_version": None, "source_id": PDF,
            "reconciliation_status": "Coextensive source caza and governorate kept as separate types; no additive parent-child use" if coextensive else "CAS survey row label only; official code and period-matched polygon unverified",
        })
    for indicator_id, name, theme, unit, source, definition, population in INDICATORS:
        data["indicators"].append({"id": indicator_id, "name": name, "theme": theme, "unit": unit,
            "definition": definition, "population": population, "source_id": source, "aggregation": "none",
            "measurement_method": "source_reported", "display_decimals": 0})
    for row, kind, name, parent, pop, households, page in ROWS:
        aid = area_id(kind, name, parent)
        values = [pop, households, rounded_100(sheet.cell_value(row - 1, 2)), rounded_100(sheet.cell_value(row - 1, 3))]
        for idx, (indicator_id, _, _, _, source, _, _) in enumerate(INDICATORS):
            locator = (f"LFHLCS main report Table 1.1, printed page {page} (PDF page {page+2}), {kind} {name}, "
                       f"{'Number of individuals' if idx == 0 else 'Number of households'}") if idx < 2 else \
                      f"LFHLCS Demography.xls sheet HL5, Excel row {row}, column {'C' if idx == 2 else 'D'} ({'Women' if idx == 2 else 'Men'}); source thousands rounded to 100 people"
            data["observations"].append({"territory_id": aid, "indicator_id": indicator_id,
                "period": "2018", "value": values[idx], "status": "observed", "measurement_method": "source_reported",
                "source_id": source, "source_locator": locator})
    comparisons = [{"parent_id": "LBN", "member_ids": gov_ids,
                    "label": "Eight CAS LFHLCS 2018-19 governorate reporting rows",
                    "membership_note": "Direct published survey estimates for eight 2018-19 governorates; counts rounded to 100, mid-2018 reference. The 2017 nine-shape geoBoundaries reference is not joined. Do not mix with a newer administrative edition or WDI midyear series.",
                    "source_ids": [PDF, XLS]}]
    for gov in GOVERNORATES:
        children = [area_id("caza", name, parent) for _, kind, name, parent, _, _, _ in ROWS if kind == "caza" and parent == gov]
        if len(children) > 1:
            comparisons.append({"parent_id": area_id("governorate", gov), "member_ids": children,
                "label": f"CAS 2018-19 caza rows within {gov}",
                "membership_note": "Direct 2018-19 survey rows. Beirut and Akkar have one coextensive caza each and are terminal at governorate level; no current legal code or polygon is inferred.",
                "source_ids": [PDF, XLS]})
    data["analysis"]["comparisons"] = comparisons
    data["analysis"]["terminal_territory_ids"] = sorted(caza_ids + [area_id("governorate", x) for x in ("Beirut", "Akkar")])
    data["analysis"]["default_indicator_id"] = INDICATORS[0][0]
    data["analysis"]["latest_values_only"] = True
    data["analysis"].pop("population_context", None)
    data["country"]["geography_note"] = "CAS LFHLCS 2018-19 reports eight governorates and 26 caza, with Beirut and Akkar source caza coextensive with their governorates. It covers residents of residential dwellings only and refers to mid-2018. The 2017 geoBoundaries nine-shape reference is not joined to this survey geography. Official codes, current boundaries and legal planning units remain unverified."
    for name, source, title, geographic in [
        ("LFHLCS_2018_2019_Report.pdf", PDF, "CAS LFHLCS 2018-19 main report Table 1.1", "national, eight governorates and 26 source caza"),
        ("LFHLCS_2018_2019_Demography.xls", XLS, "CAS LFHLCS 2018-19 demography workbook HL.5", "national, eight governorates and 26 source caza"),
        ("MICS6_Ch11_EQ_2023.xlsx", MICS, "CAS Sub-National MICS 2023 Chapter 11", "five governorates plus separately sampled displaced/refugee domains; no national result"),
    ]:
        spec = FILES[name]
        data["sources"].append({"id": source, "name": title, "url": spec["url"],
            "publisher": "Lebanon Central Administration of Statistics", "reference_period": "mid-2018 survey population" if source != MICS else "2023 limited-area MICS",
            "geographic_level": geographic, "status": "partial", "license": "terms_review_required",
            "raw_path": "raw/official/" + name, "sha256": spec["sha256"],
            "retrieved_at": datetime.fromtimestamp((raw / name).stat().st_mtime, timezone.utc).isoformat(),
            "note": "Only report Table 1.1 adopted; other report tables not assessed." if source == PDF else
                    "Only HL.5 residents by sex and area adopted; other 25 workbook sheets unassessed." if source == XLS else
                    "Chapter 11 acquired for inventory only. Published MICS covers five governorates and separate displaced/refugee domains; it excludes three governorates and supplies no adopted national estimate."})
    data["sources"] += [
        {"id": PLAN, "name": "DGU zoning-plan submission and decision procedure",
         "url": "https://www.dgu.gov.lb/dgup_30.html", "publisher": "Lebanon Directorate General for Urban Planning",
         "reference_period": "page checked 2026-09-26", "geographic_level": "town / municipality process; actual selected-area plan unverified",
         "status": "not_collected", "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required", "note": "Official procedural lead only; no plan body, final decision or local fiscal record adopted."},
        {"id": PLAN_OFFICE, "name": "Ministry of Public Works and Transport DGU responsibilities",
         "url": "https://mopw.gov.lb/%D8%A7%D9%84%D9%88%D8%AD%D8%AF%D8%A7%D8%AA-%D8%A7%D9%84%D8%A7%D8%AF%D8%A7%D8%B1%D9%8A%D8%A9-%281%29/%D8%A7%D9%84%D9%85%D8%AF%D9%8A%D8%B1%D9%8A%D8%A9-%D8%A7%D9%84%D8%B9%D8%A7%D9%85%D8%A9-%D9%84%D9%84%D8%AA%D9%86%D8%B8%D9%8A%D9%85-%D8%A7%D9%84%D9%85%D8%AF%D9%86%D9%8A",
         "publisher": "Lebanon Ministry of Public Works and Transport", "reference_period": "page checked 2026-09-26",
         "geographic_level": "national office description and local planning support", "status": "not_collected",
         "retrieved_at": datetime.now(timezone.utc).isoformat(), "license": "terms_review_required",
         "note": "Official office-responsibility lead; no legal plan unit or actual plan is adopted from it."},
        {"id": PLAN_OVERVIEW, "name": "DGU national physical master-plan overview",
         "url": "https://dgu.gov.lb/MasterPlan.html", "publisher": "Lebanon Directorate General for Urban Planning",
         "reference_period": "page checked 2026-09-26", "geographic_level": "national overview; local application unverified",
         "status": "not_collected", "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required",
         "note": "Overview mentions the 2009/2366 decree. Plan body, current legal force, local scope, finance and results remain unassessed."},
        {"id": PROFILE_CATALOG, "name": "CAS district statistical profiles 2018-19 catalogue",
         "url": "https://www.cas.gov.lb/Publications/district-statistical-profiles-2018-2019-%D8%A5%D8%AD%D8%B5%D8%A7%D8%A1%D8%A7%D8%AA-%D8%A7%D9%84%D8%A3%D9%82%D8%B6%D9%8A%D8%A9-2018-2019/",
         "publisher": "Lebanon Central Administration of Statistics", "reference_period": "2018-19 catalogue, checked 2026-09-26",
         "geographic_level": "district profile documents; contents and full coverage unassessed",
         "status": "not_collected", "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required",
         "note": "Profile catalogue located only; profile bodies, variables, dates and district coverage have not been audited."},
    ]
    data["gaps"] = [g for g in data["gaps"] if g["category"] not in ("subnational_statistics", "boundary_reconciliation", "planning_documents")]
    data["gaps"] += [
        {"category": "subnational_statistics", "status": "partial",
         "detail": "Four direct mid-2018 LFHLCS survey indicators in national, eight governorate and 26 source caza rows. This is not a census and excludes non-residential dwellings, camps and informal settlements. Only PDF Table 1.1 and XLS HL.5 are adopted. 2023 MICS Chapter 11 and other CAS tables remain unassessed; MICS has no full-country coverage.",
         "next_action": "Inventory all PDF tables and all numeric fields in all acquired/located CAS workbooks; inspect complementary district profiles and newer survey domains, methods and denominators."},
        {"category": "boundary_reconciliation", "status": "partial",
         "detail": "CAS 2018-19 reports eight governorates and 26 caza, with two coextensive pairs; 2017 geoBoundaries has nine ADM1 shapes. No official code or period-compatible local polygon joined.",
         "next_action": "Acquire official administrative code and boundary editions plus historical correspondence for the survey reference geography."},
        {"category": "planning_documents", "status": "not_collected",
         "detail": "DGU office and zoning procedure locations identified. Actual selected-area plans, current law/authority, municipal budgets, implementation and evaluation not acquired.",
         "next_action": "Retrieve current legal text, plan body/decision, district/municipal scope, fiscal and implementation evidence for contrasting areas."},
    ]
    data["collection"]["adapters"] = sorted(set(data["collection"]["adapters"] + ["lebanon-cas-lfhlcs2018-table1-1-hl5-partial-v1"]))
    data["generated_at"] = datetime.now(timezone.utc).isoformat()
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    mics_book = load_workbook(raw / "MICS6_Ch11_EQ_2023.xlsx", read_only=True, data_only=True)
    audit = {"sources": {name: {"sha256": spec["sha256"], "bytes": (raw/name).stat().st_size} for name,spec in FILES.items()},
             "pdf_adopted_table": "Table 1.1 printed pages 18-19 (PDF pages 20-21), four fields; population and household counts adopted, percentages inventoried but not adopted",
             "demography_sheets": [{"sheet": t.name, "rows": t.nrows, "columns": t.ncols,
                 "decision": "HL5 selected rows and columns only" if t.name == "HL5" else "unassessed_not_adopted"} for t in book.sheets()],
             "hl5_fields": [{"column": c, "label": label, "decision": decision} for c,label,decision in [
                 ("A","governorate label","identity_only"),("B","caza/Total label","identity_only"),
                 ("C","women, thousand people","adopted_rounded_100"),("D","men, thousand people","adopted_rounded_100"),
                 ("E","women and men, thousand people","crosschecked_pdf_population"),
                 ("F","Arabic caza label","identity_only"),("G","Arabic governorate label","identity_only")]],
             "mics_ch11_sheets": [{"sheet": t.title, "rows": t.max_row, "columns": t.max_column,
                 "decision": "unassessed_not_adopted"} for t in mics_book],
             "source_rows": [{"xls_row": row, "kind": kind, "name": name, "parent": parent,
                 "pdf_page": page, "population": pop, "households": households} for row,kind,name,parent,pop,households,page in ROWS],
             "unmatched_geometry": "2017 provider ADM1 has nine shapes; CAS survey has eight governorate rows; no local shapes joined"}
    audit["published_rounding_differences"] = {"population_governorate_sum_minus_national": 100,
                                                 "households_governorate_sum_minus_national": -100,
                                                 "treatment": "retain national and governorates as direct reported values; never replace either with subtotal"}
    (evidence / "LBN_CAS_SOURCE_FIELD_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    book.release_resources()
    print(json.dumps({"governorates":8,"caza_source_rows":26,"indicators":4,"direct_observations":len(ROWS)*4,
                      "demography_sheets":26,"mics_ch11_sheets":len(audit["mics_ch11_sheets"])},indent=2))


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--project",required=True)
    main(Path(parser.parse_args().project))
