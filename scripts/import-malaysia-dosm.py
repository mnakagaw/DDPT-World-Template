"""Audit pinned DOSM originals and build an unpublished Malaysian partial candidate.

The source CSVs retain every row in raw/. Only explicitly selected reporting
dimensions and state HIES estimates enter the dashboard. Source names are IDs,
not an official UPI-code or geometry crosswalk.
"""

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/malaysia-dosm-source-manifest.json"
ROSTER = "DOSM population CSV 2020-2025 reporting roster; official boundary version unverified"
SINGLETONS = {"Perlis", "W.P. Kuala Lumpur", "W.P. Labuan", "W.P. Putrajaya"}
YEARS = {"2020", "2024"}
SEXES = {"both": "Both sexes", "male": "Male", "female": "Female"}
HIES = {
    "income_mean": ("Mean monthly household income", "MYR/month", 0, "Income"),
    "income_median": ("Median monthly household income", "MYR/month", 0, "Income"),
    "expenditure_mean": ("Mean monthly household expenditure", "MYR/month", 0, "Income"),
    "gini": ("Household income Gini coefficient", "ratio", 3, "Inequality"),
    "poverty": ("Households below the absolute poverty line", "percent", 1, "Poverty"),
}
AMENITIES = {
    "piped_water": "Households with piped water at home",
    "sanitation": "Households covered by sanitary latrines",
    "electricity": "Households with electricity at home",
}
SCHEMAS = {
    "population_district.csv": ("state", "district", "date", "sex", "age", "ethnicity", "population"),
    "population_state.csv": ("state", "date", "sex", "age", "ethnicity", "population"),
    "population_malaysia.csv": ("date", "sex", "age", "ethnicity", "population"),
    "hies_state.csv": ("date", "state", *HIES),
    "hies_district.csv": ("date", "state", "district", *HIES),
    "hh_access_amenities.csv": ("date", "state", "district", *AMENITIES),
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def slug(value):
    return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def state_id(name):
    return "MYS:DOSM:STATE:" + slug(name)


def district_id(state, district):
    return "MYS:DOSM:DISTRICT:" + slug(state) + ":" + slug(district)


def decimal(raw, locator):
    try:
        value = Decimal(raw)
    except (InvalidOperation, TypeError):
        raise ValueError(f"Non-numeric source value at {locator}: {raw!r}") from None
    require(value.is_finite(), f"Non-finite value at {locator}")
    return value


def read_csv(path, fields, number_fields):
    inventory = {"filename": path.name, "rows": 0, "columns": list(fields),
                 "numeric_columns": {field: {"numeric_cells": 0, "blank_cells": 0,
                                              "min": None, "max": None}
                                     for field in number_fields},
                 "rows_by_year": Counter()}
    selected = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        require(tuple(reader.fieldnames or ()) == fields,
                f"CSV schema changed: {path.name}: {reader.fieldnames}")
        for rownum, row in enumerate(reader, 2):
            inventory["rows"] += 1
            year = row["date"][:4]
            inventory["rows_by_year"][year] += 1
            for field in number_fields:
                raw = row[field]
                item = inventory["numeric_columns"][field]
                if raw == "":
                    item["blank_cells"] += 1
                    continue
                value = decimal(raw, f"{path.name}:{rownum}:{field}")
                item["numeric_cells"] += 1
                item["min"] = value if item["min"] is None else min(item["min"], value)
                item["max"] = value if item["max"] is None else max(item["max"], value)
            if path.name.startswith("population_"):
                keep = (year in YEARS and row["sex"] in SEXES and
                        row["age"] == "overall" and row["ethnicity"] == "overall")
            elif path.name == "hies_state.csv":
                keep = year in {"2022", "2024"}
            elif path.name == "hh_access_amenities.csv":
                keep = year in {"2022", "2024"} and (row["district"] == "All Districts"
                    or row["state"] in SINGLETONS and row["district"] == row["state"])
            else:
                keep = year in {"2022", "2024"}
            if keep:
                selected.append({**row, "_row": rownum})
    inventory["rows_by_year"] = dict(sorted(inventory["rows_by_year"].items()))
    for item in inventory["numeric_columns"].values():
        item["min"] = str(item["min"]) if item["min"] is not None else None
        item["max"] = str(item["max"]) if item["max"] is not None else None
    return inventory, selected


def audit(project, manifest):
    for source in manifest["sources"]:
        path = project / source["raw_path"]
        require(path.is_file() and path.stat().st_size == source["bytes"] and
                sha(path) == source["sha256"], f"Pinned original changed: {source['id']}")
    inventories, rows = {}, {}
    for name, fields in SCHEMAS.items():
        number_fields = (["population"] if name.startswith("population_") else
                         list(HIES) if name.startswith("hies_") else list(AMENITIES))
        inventories[name], rows[name] = read_csv(project / "raw" / name,
                                                  fields, number_fields)
    require(inventories["population_district.csv"]["rows"] == 383040 and
            inventories["population_state.csv"]["rows"] == 270063 and
            inventories["population_malaysia.csv"]["rows"] == 17814 and
            inventories["hies_state.csv"]["rows"] == 32 and
            inventories["hies_district.csv"]["rows"] == 322 and
            inventories["hh_access_amenities.csv"]["rows"] == 688,
            "Pinned DOSM row counts changed")
    states = {r["state"] for r in rows["population_state.csv"] if r["date"] == "2024-01-01"}
    require(len(states) == 16 and SINGLETONS <= states, "Expected 16 DOSM states/territories")
    district_rows = rows["population_district.csv"]
    district_sets = {year: {(r["state"], r["district"]) for r in district_rows
                             if r["date"] == f"{year}-01-01" and r["sex"] == "both"}
                     for year in YEARS}
    require(len(district_sets["2020"]) == 160 and
            len(district_sets["2024"]) == 160 and
            {(state, state) for state in SINGLETONS} <= district_sets["2020"],
            "DOSM district reporting row count changed")
    require(len(district_sets["2020"] - district_sets["2024"]) == 4 and
            len(district_sets["2024"] - district_sets["2020"]) == 4,
            "Unexpected 2020/2024 district label drift")
    actual_districts = district_sets["2024"] - {(state, state) for state in SINGLETONS}
    require(len(actual_districts) == 156, "Expected 156 subdivided districts")
    for scope, expected in (("population_malaysia.csv", 6),
                            ("population_state.csv", 96),
                            ("population_district.csv", 960)):
        require(len(rows[scope]) == expected, f"Selected population dimensions changed: {scope}")
    require(len(rows["hies_state.csv"]) == 32 and
            len(rows["hh_access_amenities.csv"]) == 32,
            "State HIES or amenities coverage changed")
    amenity_selected_blanks = [{"state": r["state"], "date": r["date"],
                                "field": field, "line": r["_row"]}
                               for r in rows["hh_access_amenities.csv"]
                               for field in AMENITIES if r[field] == ""]
    require(len(amenity_selected_blanks) == 24 and
            all(x["date"] == "2024-01-01" and
                x["field"] in {"piped_water", "sanitation"}
                for x in amenity_selected_blanks),
            "2024 missing amenity field pattern changed")
    for year in ("2020", "2024"):
        selected = [r for r in district_rows if r["date"] == f"{year}-01-01"
                    and r["sex"] == "both"]
        state_selected = [r for r in rows["population_state.csv"]
                          if r["date"] == f"{year}-01-01" and r["sex"] == "both"]
        national = next(r for r in rows["population_malaysia.csv"]
                        if r["date"] == f"{year}-01-01" and r["sex"] == "both")
        # The published files are rounded independently to 0.1 thousand.
        district_sum = sum(decimal(r["population"], "district") for r in selected)
        state_sum = sum(decimal(r["population"], "state") for r in state_selected)
        national_value = decimal(national["population"], "national")
        require(abs(district_sum - national_value) <= Decimal("1.0") and
                abs(state_sum - national_value) <= Decimal("0.5"),
                f"Rounded population totals diverged in {year}")
    hies_district = {(r["state"], r["district"]) for r in rows["hies_district.csv"]
                     if r["date"] == "2024-01-01"}
    amenities_district_all = set()
    with (project / "raw/hh_access_amenities.csv").open(encoding="utf-8-sig", newline="") as handle:
        for r in csv.DictReader(handle):
            if r["date"] == "2024-01-01" and r["district"] != "All Districts":
                amenities_district_all.add((r["state"], r["district"]))
    # UPI is a land-administration register, not a proven DOSM census code version.
    workbook = load_workbook(project / "raw/mygeoportal-upi-codes.xlsx", read_only=True,
                             data_only=True)
    require(len(workbook.sheetnames) == 1, "UPI workbook shape changed")
    sheet = workbook.active
    require(sheet.max_row == 3323 and sheet.max_column == 10, "UPI workbook size changed")
    upi_state_codes = set()
    upi_district_keys = set()
    for index, values in enumerate(sheet.values):
        if index == 0:
            require(values[:6] == ("KOD", "NEGERI", "KOD", "DAERAH/JAJAHAN/BAHAGIAN",
                                    "KOD", "MUKIM/BANDAR/PEKAN/LAND DISTRICT"),
                    "UPI workbook headers changed")
            continue
        if values[0] and values[1]:
            upi_state_codes.add((str(values[0]), str(values[1])))
        if values[0] and values[2] and values[3]:
            upi_district_keys.add((str(values[0]), str(values[2]), str(values[3])))
    census_pdf = PdfReader(project / "raw/mycensus2020-district-key-findings.pdf")
    gazette_pdf = PdfReader(project / "raw/dbkl-local-plan-adoption-2025.pdf")
    plan_pdf = PdfReader(project / "raw/dbkl-local-plan2040-volume2-en.pdf")
    budget_pdf = PdfReader(project / "raw/dbkl-budget-speech-2025.pdf")
    annual_pdf = PdfReader(project / "raw/dbkl-annual-report-2023.pdf")
    require(len(census_pdf.pages) == 216 and len(gazette_pdf.pages) == 5 and
            len(plan_pdf.pages) == 576 and len(budget_pdf.pages) == 17 and
            len(annual_pdf.pages) == 467, "PDF pagination changed")
    gazette_text = gazette_pdf.pages[3].extract_text()
    require("5 May 2025" in gazette_text and "11 June 2025" in gazette_text and
            "Federal Territory of Kuala Lumpur" in gazette_text,
            "Gazette adoption evidence changed")
    plan_cover = plan_pdf.pages[2].extract_text().replace("\n", " ")
    require("KUALA LUMPUR LOCAL PLAN 2040" in plan_cover and
            "VOLUME 2" in plan_cover,
            "DBKL plan body identity changed")
    budget_page = budget_pdf.pages[15].extract_text()
    require("Anggaran hasil 2025" in budget_page and "RM2.45 bilion" in budget_page and
            "Anggaran Perbelanjaan 2025" in budget_page and
            "RM2.77 bilion" in budget_page and "RM325.9 juta" in budget_page,
            "DBKL 2025 budget paragraph 54 changed")
    budget_workbook = load_workbook(project / "raw/dbkl-budget-summary.xlsx",
                                    read_only=True, data_only=True)
    require(len(budget_workbook.sheetnames) == 1 and
            budget_workbook.active.max_row == 5 and
            budget_workbook.active.max_column == 11,
            "DBKL historical budget workbook changed")
    budget_rows = list(budget_workbook.active.values)
    require(budget_rows[0][1] == "2014\n(RM)" and
            budget_rows[0][-1] == "2023\n(RM)" and
            all(isinstance(value, (int, float)) for row in budget_rows[1:]
                for value in row[1:]), "DBKL budget matrix changed")
    audit_result = {
        "source_hashes": {s["id"]: s["sha256"] for s in manifest["sources"]},
        "csv_inventory": inventories,
        "field_dispositions": {
            "population_malaysia.csv": {"population": "adopted only both/male/female × overall age × overall ethnicity × 2020/2024; remaining dimensions and years priority_unassessed"},
            "population_state.csv": {"population": "adopted same selected dimensions for all 16 state rows; remaining dimensions and years priority_unassessed"},
            "population_district.csv": {"population": "adopted 2024 dimensions for 156 subdivided district keys and 2020 dimensions for 152 exact keys shared with 2024; four 2020 renamed keys held rather than merged. Four singleton self-district rows duplicate direct state rows; remaining dimensions and years priority_unassessed"},
            "hies_state.csv": {field: "adopted 2022 and 2024 direct state survey values; no district or national imputation" for field in HIES},
            "hies_district.csv": {field: "priority_unassessed: 2024 has 162 source rows, drops four singleton areas and adds/newly names Sabah-Sarawak districts; official dated code/boundary crosswalk needed" for field in HIES},
            "hh_access_amenities.csv": {field: "adopted 2022 and 2024 state totals only; 2016/2019 and district values priority_unassessed pending geography and survey-method checks" for field in AMENITIES},
        },
        "selected_amenity_missing_cells": amenity_selected_blanks,
        "geography": {
            "state_names": sorted(states),
            "population_reporting_districts": 160,
            "subdivided_districts": 156,
            "singleton_state_as_district_rows": sorted(SINGLETONS),
            "population_2020_only_unjoined_keys": sorted([list(pair) for pair in district_sets["2020"] - district_sets["2024"]]),
            "population_2024_only_keys": sorted([list(pair) for pair in district_sets["2024"] - district_sets["2020"]]),
            "hies_2024_rows": len(hies_district),
            "hies_2024_exact_2024_population_district_matches": len(hies_district & actual_districts),
            "hies_2024_unjoined_keys": sorted([list(pair) for pair in hies_district - district_sets["2024"]]),
            "amenities_2024_exact_population_district_matches": len(amenities_district_all & actual_districts),
            "amenities_2024_unjoined_keys": sorted([list(pair) for pair in amenities_district_all - district_sets["2024"]]),
            "upi_state_code_rows": len(upi_state_codes),
            "upi_land_district_code_rows": len(upi_district_keys),
            "official_code_crosswalk_status": "not_adopted: 2023 land administration UPI register is not itself a dated DOSM 2020/2024 statistical geography crosswalk",
            "provider_polygon_status": "2017 geoBoundaries ADM1 removed; no verified corresponding official 2020/2024 boundary",
        },
        "pdf_inventory": {
            "mycensus2020-district-key-findings.pdf": {"pages": 216, "decision": "acquired_not_adopted: main tabulations and Kinta exemplar need full table/field semantic inventory"},
            "dbkl-local-plan-adoption-2025.pdf": {"pages": 5, "decision": "adopted only English Gazette page 4 adoption date, effective date and authority"},
            "dbkl-local-plan2040-volume2-en.pdf": {"pages": 576, "decision": "body acquired; cover/title verified; numerical tables, plans, outcomes and maps priority_unassessed"},
            "dbkl-budget-speech-2025.pdf": {"pages": 17, "decision": "paragraph 54 printed page 16 revenue/expenditure estimates adopted as findings; other paragraphs and numbers priority_unassessed"},
            "dbkl-annual-report-2023.pdf": {"pages": 467, "decision": "body acquired; financial statement section is image based and actual spending/evaluation not extracted or adopted"},
        },
        "dbkl_budget_workbook_inventory": {"sheet": budget_workbook.active.title,
            "rows": 5, "columns": 11, "year_columns": list(range(2014, 2024)),
            "numeric_cells": 40,
            "row_labels": [row[0] for row in budget_rows[1:]],
            "decision": "historical 2014-2023 budget estimates inventoried, not adopted as actual revenue/expenditure"},
        "adopted_indicators": 14,
        "adopted_population_observations": 1026,
        "adopted_hies_state_observations": 160,
        "adopted_amenity_state_observation_slots": 96,
        "adopted_amenity_state_observed": 72,
        "adopted_amenity_state_missing": 24,
        "adopted_observation_slots": 1282,
        "adopted_observations": 1258,
    }
    path = project / "evidence/MYS_DOSM_AUDIT.json"
    path.write_text(json.dumps(audit_result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return audit_result, rows


def source_record(item, retrieved_at):
    name = item["id"].replace("mys-", "").replace("-", " ").title()
    publisher = ("Department of Statistics Malaysia" if "dosm" in item["id"] else
                 "National Geospatial Centre Malaysia" if "mygeoportal" in item["id"] else
                 "Kuala Lumpur City Hall / Federal Gazette")
    level = ("country, state and administrative district" if "population" in item["id"] else
             "state and district" if "hies" in item["id"] or "amenities" in item["id"] else
             "Federal Territory of Kuala Lumpur" if "dbkl" in item["id"] else
             "source-specific geographic reporting")
    result = {"id": item["id"], "name": name, "url": item["url"],
              "publisher": publisher, "reference_period": "source-specific, see locator and period",
              "geographic_level": level, "status": "ready", "retrieved_at": retrieved_at,
              "sha256": item["sha256"], "raw_path": item["raw_path"],
              "license": "CC BY 4.0" if "dosm" in item["id"] and item["id"].endswith(("district", "state", "malaysia", "amenities")) and "census2020" not in item["id"] else "Official public source; reuse terms to be confirmed",
              "note": "Original acquired and hash-pinned; adopted dimensions and unresolved fields are in MYS_DOSM_AUDIT.json."}
    if "catalogue_url" in item:
        result["catalogue_url"] = item["catalogue_url"]
    return result


def build(project, manifest, audit_result, rows):
    path = project / "data/dashboard.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    require(data["country"]["id"] == "MYS", "Expected Malaysia candidate")
    require(audit_result["adopted_observation_slots"] == 1282 and
            audit_result["adopted_observations"] == 1258,
            "Audit observation contract changed")
    data["country"]["geography_note"] = (
        "DOSM population CSVs report 16 state/federal-territory rows and 160 district rows; "
        "Perlis and three Federal Territories are whole-area district duplicates. The other "
        "156 districts are source reporting units, not verified local planning authorities. "
        "Four district names differ between 2020 and 2024; their 2020 rows are held. Names, UPI 2023 codes and 2017 geoBoundaries polygons have no validated dated "
        "crosswalk. Population 2020 is census-adjusted, 2024 is an intercensal estimate; "
        "HIES household survey statistics and WDI are separate evidence.")
    data["territories"] = [t for t in data["territories"] if t["id"] == "MYS"]
    data["territories"][0].update({"source_id": "mys-dosm-population-malaysia",
                                    "boundary_version": ROSTER})
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    state_names = audit_result["geography"]["state_names"]
    districts = sorted({(r["state"], r["district"])
                        for r in rows["population_district.csv"]
                        if r["date"] == "2024-01-01" and r["sex"] == "both" and
                        r["state"] not in SINGLETONS})
    for name in state_names:
        data["territories"].append({"id": state_id(name), "name": name,
            "level": "adm1", "type": "Federal Territory" if name.startswith("W.P.") else "State",
            "parent_id": "MYS", "official_code": None,
            "code_system": "DOSM state reporting name; UPI crosswalk pending",
            "boundary_version": ROSTER, "source_id": "mys-dosm-population-state",
            "reconciliation_status": "DOSM source row verified; current official code and polygon not joined"})
    for state, district in districts:
        data["territories"].append({"id": district_id(state, district),
            "name": district, "level": "adm2", "type": "Administrative District",
            "parent_id": state_id(state), "official_code": None,
            "code_system": "DOSM state plus district reporting name; UPI crosswalk pending",
            "boundary_version": ROSTER, "source_id": "mys-dosm-population-district",
            "reconciliation_status": "Exact DOSM population CSV name pair; local planning authority and polygon unverified"})
    sources = {s["id"]: s for s in data["sources"]}
    for item in manifest["sources"]:
        sources[item["id"]] = source_record(item, manifest["retrieved_at"])
    sources["mys-planmalaysia-act172-location"] = {"id": "mys-planmalaysia-act172-location",
        "name": "PLANMalaysia Local Plan Act 172 sections 12-16 overview",
        "url": "https://www.planmalaysia.gov.my/main/article/rancangan-tempatan",
        "publisher": "PLANMalaysia", "reference_period": "web page checked 2026-09-27",
        "geographic_level": "Peninsular Malaysia planning regime; applicability differs elsewhere",
        "status": "partial", "retrieved_at": manifest["retrieved_at"],
        "license": "Official public page; reuse terms not checked",
        "note": "Official summary located; direct 2021 Act PDF URL returned 404 and full current legal text has not been verified."}
    sources["mys-planmalaysia-manual-location"] = {"id": "mys-planmalaysia-manual-location",
        "name": "PLANMalaysia planning manuals catalogue",
        "url": "https://www.planmalaysia.gov.my/main/documents?type=manual-perancangan",
        "publisher": "PLANMalaysia", "reference_period": "catalogue checked 2026-09-27",
        "geographic_level": "Peninsular local planning guidance",
        "status": "partial", "retrieved_at": manifest["retrieved_at"],
        "license": "Official public catalogue; item terms not checked",
        "note": "Lists Manual Rancangan Tempatan Edisi 2022 (Versi 2); older direct PDF returned 404. Manual body not acquired."}
    data["sources"] = list(sources.values())
    pop_names = {"both": "population, both sexes", "male": "male population",
                 "female": "female population"}
    indicators = [i for i in data["indicators"] if not i["id"].startswith("MYS_DOSM_")]
    observations = [o for o in data["observations"]
                    if not o["indicator_id"].startswith("MYS_DOSM_")]
    base_observation_count = len(observations)
    for year, label, method in (("2020", "2020 census-adjusted", "DOSM census-adjusted population table, 0.1 thousand rounding"),
                                ("2024", "2024 intercensal estimated", "DOSM cohort-component intercensal estimate, 0.1 thousand rounding")):
        for sex in SEXES:
            indicators.append({"id": f"MYS_DOSM_POP_{year}_{sex.upper()}",
                "name": f"{label} {pop_names[sex]}", "theme": "Population",
                "unit": "people", "display_decimals": 0,
                "definition": f"DOSM {year} population table: age overall, ethnicity overall, sex {sex}. Source is printed in thousands to one decimal; AreaData multiplies by 1,000. "
                              + ("The 2020 value is adjusted census-based, not an unrounded enumerated count." if year == "2020" else
                                 "The 2024 value is an intercensal estimate, not a new census count."),
                "population": "DOSM population table resident population, all ages and ethnicities",
                "source_id": "mys-dosm-population-malaysia", "aggregation": "none",
                "measurement_method": method})
    for field, (name, unit, digits, theme) in HIES.items():
        indicators.append({"id": "MYS_DOSM_HIES_" + field.upper(),
            "name": name, "theme": theme, "unit": unit,
            "display_decimals": digits,
            "definition": "DOSM HIES state survey estimate for the reported year. " +
                ("Poverty is the proportion of surveyed households below the official Poverty Line Income; it is not population poverty." if field == "poverty" else
                 "State survey statistic, not an administrative count. Relative standard errors have not been separately audited."),
            "population": "Households in DOSM Household Income and Expenditure Survey",
            "source_id": "mys-dosm-hies-state", "aggregation": "none",
            "measurement_method": "DOSM HIES state household survey estimate"})
    for field, name in AMENITIES.items():
        indicators.append({"id": "MYS_DOSM_AMENITY_" + field.upper(),
            "name": name, "theme": "Basic amenities", "unit": "percent",
            "display_decimals": 1,
            "definition": f"DOSM HIES state total, proportion of surveyed households with {field.replace('_', ' ')}. Do not average district rates into a state rate.",
            "population": "Households in DOSM Household Income and Expenditure Survey",
            "source_id": "mys-dosm-amenities", "aggregation": "none",
            "measurement_method": "DOSM HIES state household survey estimate"})
    source_for_scope = {"population_malaysia.csv": "mys-dosm-population-malaysia",
                        "population_state.csv": "mys-dosm-population-state",
                        "population_district.csv": "mys-dosm-population-district"}
    adopted_district_keys = set(districts)
    for filename in source_for_scope:
        for row in rows[filename]:
            if filename == "population_malaysia.csv":
                tid = "MYS"
            elif filename == "population_state.csv":
                tid = state_id(row["state"])
            elif row["state"] in SINGLETONS:
                continue  # Exact duplicate reporting area; direct state source used.
            elif (row["state"], row["district"]) not in adopted_district_keys:
                continue  # Four 2020-only labels need a dated official crosswalk.
            else:
                tid = district_id(row["state"], row["district"])
            year = row["date"][:4]
            value = decimal(row["population"], f"{filename}:{row['_row']}") * 1000
            require(value == value.to_integral_value(), "Population unit conversion lost precision")
            observations.append({"territory_id": tid,
                "indicator_id": f"MYS_DOSM_POP_{year}_{row['sex'].upper()}",
                "period": year, "value": int(value), "status": "observed",
                "source_id": source_for_scope[filename], "provenance": "calculated",
                "calculation": "published 0.1 thousand persons × 1,000; no distribution or imputation",
                "population_scope": "all ages and ethnicities, " + SEXES[row["sex"]],
                "measurement_method": ("DOSM census-adjusted population table, 0.1 thousand rounding"
                    if year == "2020" else "DOSM cohort-component intercensal estimate, 0.1 thousand rounding"),
                "boundary_version": ROSTER,
                "source_locator": f"{filename}:line {row['_row']} (state={row.get('state','Malaysia')}, district={row.get('district','n/a')}, sex={row['sex']}, age=overall, ethnicity=overall, date={row['date']})"})
    for row in rows["hies_state.csv"]:
        for field in HIES:
            observations.append({"territory_id": state_id(row["state"]),
                "indicator_id": "MYS_DOSM_HIES_" + field.upper(),
                "period": row["date"][:4], "value": float(decimal(row[field], "HIES")),
                "status": "observed", "source_id": "mys-dosm-hies-state",
                "provenance": "source_reported", "population_scope": "HIES households",
                "measurement_method": "DOSM HIES state household survey estimate",
                "boundary_version": ROSTER,
                "source_locator": f"hies_state.csv:line {row['_row']}:{field}, state={row['state']}, date={row['date']}"})
    for row in rows["hh_access_amenities.csv"]:
        for field in AMENITIES:
            raw_value = row[field]
            observations.append({"territory_id": state_id(row["state"]),
                "indicator_id": "MYS_DOSM_AMENITY_" + field.upper(),
                "period": row["date"][:4],
                "value": float(decimal(raw_value, "amenities")) if raw_value else None,
                "status": "observed" if raw_value else "missing",
                "source_id": "mys-dosm-amenities",
                "provenance": "source_reported", "population_scope": "HIES households",
                "measurement_method": "DOSM HIES state household survey estimate",
                "boundary_version": ROSTER,
                "source_locator": f"hh_access_amenities.csv:line {row['_row']}:{field}, state={row['state']}, district={row['district']}, date={row['date']}"})
    require(len(observations) - base_observation_count == 1282,
            "Unexpected count of imported DOSM observations")
    data["indicators"] = indicators
    data["observations"] = observations
    kl_id = state_id("W.P. Kuala Lumpur")
    kl_match = {"territory_id": kl_id, "country_id": "MYS",
        "type": "Federal Territory",
        "code_system": "DOSM state reporting name; UPI crosswalk pending",
        "official_code": None, "boundary_version": ROSTER,
        "method": "DBKL Gazette identifies the Federal Territory of Kuala Lumpur; DOSM separately lists W.P. Kuala Lumpur as a single whole-area state/district row. No smaller PBT-to-district join is inferred.",
        "source_id": "mys-dbkl-kl-local-plan-gazette",
        "locator": "P.U. (B) 206/2025 English page 4",
        "checked_at": "2026-09-27"}
    data["documents"] = [d for d in data["documents"] if not d["id"].startswith("mys-")]
    data["documents"].extend([
        {"id": "mys-dbkl-kl-local-plan-2040", "territory_id": kl_id,
         "category": "plan", "kind": "local_plan", "title": "Kuala Lumpur Local Plan 2040, Volume 2",
         "url": sources["mys-dbkl-kl-local-plan-2040-vol2"]["url"],
         "source_id": "mys-dbkl-kl-local-plan-2040-vol2", "period": "2040",
         "availability": "body_acquired", "official_status": "adopted_effective_2025",
         "official_evidence": {"source_id": "mys-dbkl-kl-local-plan-gazette",
            "locator": "P.U. (B) 206/2025 English page 4; adopted 5 May 2025, effective 11 June 2025",
            "checked_at": "2026-09-27", "authority": "Federal Government Gazette / DBKL"},
         "territory_match": kl_match,
         "note": "Gazette confirms adoption and effect for the Federal Territory. The acquired 576-page Volume 2 body has only title/initial pages checked; numerical tables and implementation outcomes remain unassessed. This is not a district plan for other Malaysian areas."},
        {"id": "mys-dbkl-kl-gazette-document", "territory_id": kl_id,
         "category": "reference", "kind": "legal_notice", "title": "P.U. (B) 206/2025 Kuala Lumpur Local Plan adoption notice",
         "url": sources["mys-dbkl-kl-local-plan-gazette"]["url"],
         "source_id": "mys-dbkl-kl-local-plan-gazette", "period": "2025",
         "availability": "content_verified", "official_status": "unverified",
         "territory_match": kl_match,
         "content": {"summary": "The Commissioner with ministerial approval adopted the Kuala Lumpur Local Plan 2040 with modifications on 5 May 2025; it took effect on 11 June 2025.",
             "evidence": {"source_id": "mys-dbkl-kl-local-plan-gazette",
                 "locator": "P.U. (B) 206/2025 English page 4",
                 "checked_at": "2026-09-27"}}},
        {"id": "mys-dbkl-budget-speech-document", "territory_id": kl_id,
         "category": "budget", "kind": "budget_speech",
         "title": "DBKL Budget 2025 speech and published estimates",
         "url": sources["mys-dbkl-budget-speech-2025"]["url"],
         "source_id": "mys-dbkl-budget-speech-2025", "period": "2025",
         "target_period": {"label": "2025", "kind": "calendar_year",
            "start": "2025-01-01", "end": "2025-12-31"},
         "availability": "content_verified", "official_status": "unverified",
         "territory_match": kl_match,
         "content": {"summary": "DBKL's 2025 budget speech paragraph 54 reports estimated revenue RM2.45 billion, estimated expenditure RM2.77 billion and a reported deficit RM325.9 million. The two rounded billion figures do not exactly reproduce the reported deficit; it is not recalculated here. These are budget estimates, not actual spending.",
             "evidence": {"source_id": "mys-dbkl-budget-speech-2025",
                 "locator": "PDF page 16, paragraph 54", "checked_at": "2026-09-27"}},
         "findings": [
            {"kind": "revenue", "label": "Estimated 2025 revenue",
             "value": 2450000000, "value_status": "observed", "unit": "MYR",
             "definition": "DBKL speech reports a rounded budget estimate of RM2.45 billion; this is not actual revenue.",
             "scope": "DBKL budget 2025, Federal Territory of Kuala Lumpur",
             "period": {"label": "2025", "kind": "calendar_year", "start": "2025-01-01", "end": "2025-12-31"},
             "evidence": {"source_id": "mys-dbkl-budget-speech-2025", "locator": "PDF page 16, paragraph 54, Anggaran hasil 2025", "checked_at": "2026-09-27"}},
            {"kind": "budget", "label": "Estimated 2025 expenditure",
             "value": 2770000000, "value_status": "observed", "unit": "MYR",
             "definition": "DBKL speech reports a rounded budget estimate of RM2.77 billion; this is not actual spending or execution.",
             "scope": "DBKL budget 2025, Federal Territory of Kuala Lumpur",
             "period": {"label": "2025", "kind": "calendar_year", "start": "2025-01-01", "end": "2025-12-31"},
             "evidence": {"source_id": "mys-dbkl-budget-speech-2025", "locator": "PDF page 16, paragraph 54, Anggaran Perbelanjaan 2025", "checked_at": "2026-09-27"}},
         ]},
        {"id": "mys-dbkl-historical-budget-summary-document", "territory_id": kl_id,
         "category": "budget", "kind": "historical_budget_estimates",
         "title": "DBKL historical budget estimates 2014-2023",
         "url": sources["mys-dbkl-budget-summary"]["url"],
         "source_id": "mys-dbkl-budget-summary", "period": "2014-2023",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "Four numeric rows across ten years inventoried; source is estimates, not actual expenditure, and is not an execution finding."},
        {"id": "mys-dbkl-annual-report-2023-document", "territory_id": kl_id,
         "category": "implementation", "kind": "annual_report",
         "title": "DBKL Annual Report 2023 (financial statements in scanned section)",
         "url": sources["mys-dbkl-annual-report-2023"]["url"],
         "source_id": "mys-dbkl-annual-report-2023", "period": "2023",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "467-page body acquired. The financial-statement pages are image based; actual expenditure, completion and evaluation figures have not been verified or adopted."},
        {"id": "mys-dosm-census-report-reference", "territory_id": "MYS",
         "category": "reference", "kind": "census_reference",
         "title": "MyCensus 2020 administrative district key findings (acquired sample/report)",
         "url": sources["mys-dosm-census2020-district-report"]["url"],
         "source_id": "mys-dosm-census2020-district-report", "period": "2020",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "216-page report body acquired. Its unrounded counts, 19 numbered tables and local examples have not been adopted in the dashboard; OpenDOSM adjusted CSV series is separately identified."},
    ])
    data["planning"] = {"title": "Plans and local evidence",
        "purpose": "Use the selected DOSM reporting area for a statistical baseline, then obtain the actual responsible planning authority's plan, budget, spending and evaluation evidence.",
        "system": {"label": "Malaysian planning frameworks differ by jurisdiction",
            "scope": "PLANMalaysia describes Act 172 structure/local plans for Peninsular states. Kuala Lumpur uses the separate Federal Territory (Planning) Act 1982 (Act 267); its 2040 local plan is confirmed by Gazette P.U. (B) 206/2025. Sabah, Sarawak and other Federal Territories require their own legal and authority checks. A DOSM district is not automatically a local planning authority (PBT).",
            "cycle": "Only Kuala Lumpur Local Plan 2040's Gazette effective date is confirmed here; no universal planning cycle is assigned.",
            "source_ids": ["mys-planmalaysia-act172-location", "mys-dbkl-kl-local-plan-gazette"]},
        "sections": [{"id": key, "label": label} for key, label in (
            ("plan", "Local and structure plans"), ("budget", "Budgets and estimates"),
            ("implementation", "Spending and implementation"),
            ("evaluation", "Official evaluation"),
            ("reference", "Census and legal references"))]}
    comparisons = [{"parent_id": "MYS", "member_ids": [state_id(n) for n in state_names],
        "label": "DOSM sixteen state/federal-territory reporting rows",
        "membership_note": "All 16 DOSM state population rows are presented once. States and Federal Territories differ legally. Rounded subnational population values do not exactly sum to the separately published national row; no rate or HIES value is aggregated.",
        "source_ids": ["mys-dosm-population-state"]}]
    for state in state_names:
        members = [district_id(s, d) for s, d in districts if s == state]
        if members:
            comparisons.append({"parent_id": state_id(state), "member_ids": members,
                "label": f"DOSM administrative districts of {state}",
                "membership_note": "Exact DOSM population state+district reporting keys from the 2020/2024 CSV. UPI official code, polygon, local-plan authority and changes since 2024 remain unverified. Do not sum independently rounded district values into the published state population.",
                "source_ids": ["mys-dosm-population-district"]})
    data["analysis"]["comparisons"] = comparisons
    data["analysis"]["terminal_territory_ids"] = [district_id(s, d) for s, d in districts]
    data["analysis"]["default_indicator_id"] = "MYS_DOSM_POP_2024_BOTH"
    data["analysis"]["population_context"] = {"primary_indicator_id": "MYS_DOSM_POP_2024_BOTH",
        "reference_indicator_id": "SP.POP.TOTL"}
    data["analysis"]["latest_values_only"] = True
    data["collection"]["status"] = "partial"
    data["collection"]["adapters"] = list(dict.fromkeys([
        *data["collection"].get("adapters", []), "dosm-open-data-population-hies-state-2026-09-27"]))
    data["collection"]["notes"] = [
        "Six DOSM CSV originals, one census PDF, one UPI code workbook, three DBKL PDFs, one Gazette PDF and one DBKL budget workbook are hash-pinned. Fourteen new indicators have 1,282 source slots: 1,258 observed and 24 explicitly missing amenity cells, only for declared population dimensions and state HIES fields.",
        "DOSM 2020 population is adjusted census-based; 2024 is intercensal estimated. Original unit thousand people is exactly multiplied by 1,000, with independent rounding differences retained. WDI is a separate nationwide series.",
        "Four 2020 district labels differ from the 2024 population roster; their older rows are held. HIES district 2024 has 162 rows and inconsistent names/coverage against the 160-row population roster; none is cross-joined. Amenities district rows are likewise held. All unadopted age/ethnicity/history combinations remain in originals and audit inventory.",
        "UPI 2023 land-code workbook is acquired but no official dated statistical code/2017 provider polygon crosswalk is established. All reference polygons are withheld.",
        "Kuala Lumpur 2040 Local Plan Gazette confirms adoption/effective dates only for the Federal Territory. One 576-page plan volume is acquired but full content is unassessed. DBKL 2025 budget estimates are verified from speech page 16; 2023 annual financial statement images and actual spending/evaluation remain unassessed.",
        *[x for x in data["collection"].get("notes", []) if not x.startswith((
            "Initial national-data site inputs only", "Six DOSM CSV originals",
            "DOSM 2020 population is adjusted", "Four 2020 district labels",
            "UPI 2023 land-code workbook", "Kuala Lumpur 2040 Local Plan Gazette"))],
    ]
    data["gaps"] = [g for g in data["gaps"] if g["category"] not in
                    {"boundary_reconciliation", "subnational_statistics", "planning_documents"}]
    data["gaps"].extend([
        {"category": "boundary_reconciliation", "status": "partial",
         "detail": "Sixteen state and 156 subdivided district statistical rows have no assigned official UPI code or matching polygon. DOSM HIES 2024 district roster and names differ; 2017 geoBoundaries shapes are withheld.",
         "next_action": "Acquire dated DOSM-to-UPI code and official boundary crosswalk, including Sabah/Sarawak district changes and PBT jurisdiction."},
        {"category": "subnational_statistics", "status": "partial",
         "detail": "Population 2020/2024 selected sex totals and HIES 2022/2024 state five fields plus amenity three fields are adopted. Age/ethnicity, other years, district survey rows and PDF tables remain priority_unassessed.",
         "next_action": "Audit full census district/PBT/mukim tables and HIES district survey definitions and geography; retain exact source denominators and error information."},
        {"category": "planning_documents", "status": "partial",
         "detail": "Kuala Lumpur 2040 local-plan Gazette adoption and DBKL 2025 budget speech estimates are verified. Other states/PBT plans, actual spending and official evaluations are uncollected or unextracted.",
         "next_action": "Identify each legal planning jurisdiction under the applicable state/Federal Territory/Sabah/Sarawak law and obtain actual current plan, budget and assessment bodies."},
    ])
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(data['territories'])} territories, 14 indicators and 1,282 observations; {path}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    audit_result, rows = audit(args.project, manifest)
    build(args.project, manifest, audit_result, rows)


if __name__ == "__main__":
    main()
