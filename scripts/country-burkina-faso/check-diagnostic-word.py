"""Cross-check the BFA diagnostic Word against the selected dataset records."""
from argparse import ArgumentParser
from hashlib import sha256
from pathlib import Path
import json

from docx import Document

ap = ArgumentParser()
ap.add_argument("--project", required=True)
ap.add_argument("--territory", default="BFA")
ap.add_argument("--docx", required=True)
ap.add_argument("--render-dir", required=True)
args = ap.parse_args()
project = Path(args.project)
dataset_path = project / "data/dashboard.json"
dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
area = next(row for row in dataset["territories"] if row["id"] == args.territory)
word = Path(args.docx)
doc = Document(word)
assert doc.paragraphs[0].text == "Territorial Development Diagnostic"
assert area["name"] in doc.paragraphs[1].text
assert len(doc.inline_shapes) >= 2, (area["id"], len(doc.inline_shapes))
table = doc.tables[0]
assert [cell.text for cell in table.rows[0].cells] == ["Indicator", "Value", "Year", "Source table"]
indicators = {item["id"]: item for item in dataset["indicators"]}
observations = {(row["territory_id"], row["indicator_id"], row["period"]): row
                for row in dataset["observations"] if row["status"] == "observed"}
matched = []
for row in table.rows[1:]:
    name, value, year, source_table = [cell.text for cell in row.cells]
    matching_indicators = [item for item in indicators.values() if item["name"] == name]
    assert len(matching_indicators) == 1, name
    indicator = matching_indicators[0]
    observation = observations[(area["id"], indicator["id"], year)]
    assert value == f"{observation['value']:,.1f} {indicator['unit']}", (name, value, observation)
    assert source_table == (indicator.get("upstream_table") or next(source["name"][:35]
                           for source in dataset["sources"] if source["id"] == observation["source_id"]))
    matched.append({"indicator_id": indicator["id"], "value": observation["value"], "year": year,
                    "source_id": observation["source_id"]})
assert len(matched) == 7, (area["id"], len(matched))
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
population = observations[(area["id"], "BFA_RGPH2019_POP_TOTAL", "2019")]["value"]
assert f"{population:,}" in text
assert "No signed area-specific plan" in text
assert "not an adopted local development plan" in text
pages = sorted(Path(args.render_dir).glob("page-*.png"))
assert len(pages) >= 3 and all(page.stat().st_size > 50_000 and
                               page.stat().st_mtime_ns >= word.stat().st_mtime_ns for page in pages), \
    "Missing, empty or stale page render"
report = {
    "area": area["name"], "territory_id": area["id"],
    "dataset_sha256": sha256(dataset_path.read_bytes()).hexdigest(),
    "docx_sha256": sha256(word.read_bytes()).hexdigest(),
    "matched_indicator_rows": matched,
    "population_2019": population,
    "inline_figures": len(doc.inline_shapes), "rendered_pages": len(pages),
    "page_images": [page.name for page in pages],
    "review_limit": "Automated value/year/source checks and PNG page count; human visual inspection of all rendered pages is separately recorded. No browser download, print dialogue, or independent audit is certified.",
}
suffix = "" if area["level"] == "national" else "-" + sha256(area["id"].encode()).hexdigest()[:10]
out = project / f"evidence/WORD_QA{suffix}.json"
out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"matched_rows": len(matched), "pages": len(pages), "docx_sha256": report["docx_sha256"]}))
