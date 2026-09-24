"""Verify every static Word report is bound to the current BFA dataset and area."""
from argparse import ArgumentParser
from hashlib import sha256
from pathlib import Path
import json

from docx import Document


ap = ArgumentParser()
ap.add_argument("--project", required=True)
args = ap.parse_args()
project = Path(args.project)
data_path = project / "data/dashboard.json"
data = json.loads(data_path.read_text(encoding="utf-8"))
manifest = json.loads((project / "evidence/WORD_REPORT_MANIFEST.json").read_text(encoding="utf-8"))
assert manifest["dataset_sha256"] == sha256(data_path.read_bytes()).hexdigest()
assert len(manifest["files"]) == len(data["territories"]) == 410
population = {row["territory_id"]: row["value"] for row in data["observations"]
              if row["indicator_id"] == "BFA_RGPH2019_POP_TOTAL" and row["period"] == "2019"
              and row["status"] == "observed"}
checked = []
for index, (area, entry) in enumerate(zip(data["territories"], manifest["files"])):
    assert entry["index"] == index and entry["territory_id"] == area["id"]
    assert entry["path"] == f"exports/diagnostic/area-{index:04d}.docx"
    target = project / "site" / entry["path"]
    assert sha256(target.read_bytes()).hexdigest() == entry["sha256"]
    document = Document(target)
    assert document.paragraphs[0].text == "Territorial Development Diagnostic"
    assert area["name"] in document.paragraphs[1].text
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert manifest["dataset_sha256"] in body
    assert f"{population[area['id']]:,}" in body
    assert "not an adopted local development plan" in body
    assert document.tables and document.inline_shapes
    checked.append(entry["sha256"])
print(json.dumps({"checked_reports": len(checked), "unique_docx_hashes": len(set(checked)),
                  "dataset_sha256": manifest["dataset_sha256"]}))
