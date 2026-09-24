"""Build a source-grounded AreaData diagnostic DOCX for one BFA territory.

Usage: python build-diagnostic-word.py --project generated/burkina-faso-20260924
       --territory BFA --out ".../Territorial Development Diagnostic.docx"
The output is a review artifact, not an approved development plan.
"""
from argparse import ArgumentParser
from collections import defaultdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import json

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ap = ArgumentParser()
ap.add_argument("--project", required=True)
ap.add_argument("--territory", default="BFA")
ap.add_argument("--out", required=True)
args = ap.parse_args()
project = Path(args.project)
data = json.loads((project / "data/dashboard.json").read_text(encoding="utf-8"))
area = next(item for item in data["territories"] if item["id"] == args.territory)
indicators = {item["id"]: item for item in data["indicators"]}
sources = {item["id"]: item for item in data["sources"]}
territories = {item["id"]: item for item in data["territories"]}
observations = defaultdict(list)
for row in data["observations"]:
    observations[(row["territory_id"], row["indicator_id"])].append(row)


def latest(indicator_id, area_id=None):
    rows = [row for row in observations[(area_id or area["id"], indicator_id)]
            if row["status"] == "observed" and isinstance(row.get("value"), (int, float))]
    return max(rows, key=lambda row: row["period"]) if rows else None


def number(value, decimals=0):
    return f"{value:,.{decimals}f}" if value is not None else "Not available"


def font(size=23, bold=False):
    font_path = Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf")
    try:
        return ImageFont.truetype(str(font_path), size)
    except OSError:
        return ImageFont.load_default()


def bar_chart(values, title, unit):
    width, margin, row_h = 1160, 45, 48
    height = 92 + row_h * len(values)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((margin, 15), title, fill="#182f42", font=font(28, True))
    vmax = max(value for _, value in values) if values else 1
    plot_x, plot_w = 400, 500
    for i, (label, value) in enumerate(values):
        y = 62 + i * row_h
        draw.text((margin, y), label[:25], fill="#273f4e", font=font(21))
        draw.rectangle((plot_x, y, plot_x + round(value / vmax * plot_w), y + 25), fill="#347667")
        draw.text((plot_x + round(value / vmax * plot_w) + 12, y),
                  f"{number(value, 1 if unit == '%' else 0)} {unit}", fill="#172c3b", font=font(19))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def pyramid_chart(profile):
    ages = profile["ages"]
    males, females = profile["male"], profile["female"]
    width, height = 1160, 730
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((45, 12), "Population by age and sex 2019", fill="#182f42", font=font(27, True))
    draw.text((200, 54), "Male", fill="#255b72", font=font(22, True))
    draw.text((855, 54), "Female", fill="#b66862", font=font(22, True))
    center, span, top, row_h = 580, 440, 94, 32
    maximum = max(males + females)
    for i in range(len(ages) - 1, -1, -1):
        y = top + (len(ages) - 1 - i) * row_h
        mw = round(males[i] / maximum * span)
        fw = round(females[i] / maximum * span)
        draw.rectangle((center - 28 - mw, y, center - 28, y + 24), fill="#286986")
        draw.rectangle((center + 28, y, center + 28 + fw, y + 24), fill="#c9857b")
        draw.text((center - 22, y), ages[i], fill="#172c3b", font=font(17))
    draw.text((45, 686), "Counts reported in INSD RGPH 2019 Table I.20; each bar uses the same scale.",
              fill="#4c5e68", font=font(17))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Title"].font.size = Pt(20)
doc.styles["Heading 1"].font.size = Pt(13)
doc.styles["Heading 2"].font.size = Pt(11)
title_ppr = doc.styles["Title"].element.pPr
if title_ppr is not None:
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

doc.add_paragraph("Territorial Development Diagnostic", style="Title")
doc.add_paragraph(f"{area['name']}  |  AreaData review build  |  Latest available values with source year")
doc.add_paragraph("This document assembles available census and related evidence for territorial planning review. It is not an adopted local development plan or a statement of current legal boundaries. The selected statistical geography is the 2019 RGPH geography; Burkina Faso changed its administrative regions and provinces in 2025.")

population = latest("BFA_RGPH2019_POP_TOTAL")
assert population and population["period"] == "2019"
doc.add_heading("Population and geographic scope", 1)
doc.add_paragraph(f"INSD reported {number(population['value'])} residents for {area['name']} in the 2019 census. Population estimates were used in some security-affected areas according to the source. This value is not a 2025 or current population estimate.")
doc.add_paragraph(f"Area identity: {area['type']} · {area['id']} · boundary edition {area.get('boundary_version') or 'unverified'}. The 2017 geoBoundaries shapes are reference maps, not certified present-day boundaries. Commune polygons are not joined where an official correspondence has not been verified.")

profile = data.get("analysis", {}).get("census_population_pyramids", {}).get(f"{area['id']}@2019")
if profile:
    doc.add_picture(pyramid_chart(profile), width=Inches(6.8))
    male, female = sum(profile["male"]), sum(profile["female"])
    assert male + female == population["value"]
    doc.add_paragraph(f"The 2019 source counts {number(male)} males and {number(female)} females. These figures sum to the selected area's reported census population. Source: INSD RGPH 2019, Table I.20.")
else:
    age_bands = [(label, latest(indicator_id)) for label, indicator_id in (
        ("Age 0-4", "BFA_RGPH2019_AGE_0_4"),
        ("Age 5-14", "BFA_RGPH2019_AGE_5_14"),
        ("Age 15-24", "BFA_RGPH2019_AGE_15_24"),
        ("Age 25-64", "BFA_RGPH2019_AGE_25_64"),
        ("Age 65+", "BFA_RGPH2019_AGE_65_PLUS"),
    )]
    if all(row and row["period"] == "2019" for _, row in age_bands):
        assert sum(row["value"] for _, row in age_bands) == population["value"]
        doc.add_picture(bar_chart([(label, row["value"]) for label, row in age_bands],
                                  "Population by broad age group, 2019", "people"), width=Inches(6.8))
        doc.add_paragraph("The five age groups sum to the selected area's 2019 census population. These are source-reported counts, not a projection. Source: INSD RGPH 2019, communal disparities age-group table.")

if area["level"] == "national" or profile:
    doc.add_page_break()
national_selected = [
    "BFA_RGPH2019_SCHOOL_ATTENDING_6_16_REGIONAL",
    "BFA_RGPH2019_NET_POSTPRIMARY_ENROLMENT_REGIONAL",
    "BFA_RGPH2019_NET_SECONDARY_ENROLMENT_REGIONAL",
    "BFA_RGPH2019_ILO_UNEMPLOYMENT_RATE_REGIONAL",
    "BFA_RGPH2019_WASTEWATER_STREET_NATURE_REGIONAL",
    "BFA_RGPH2019_DISABILITY_PREVALENCE_5PLUS_REGIONAL",
    "BFA_INSD_POVERTY_INCIDENCE_MODEL",
]
local_selected = [
    "BFA_INSD_LITERACY_15_64",
    "BFA_INSD_OUT_OF_SCHOOL_6_11",
    "BFA_INSD_POST_PRIMARY_NET_ATTENDANCE",
    "BFA_INSD_YOUTH_EMPLOYMENT_15_24",
    "BFA_INSD_HOUSEHOLD_WATER_IMPROVED",
    "BFA_INSD_HOUSEHOLD_ELECTRIC_LIGHT",
    "BFA_INSD_POVERTY_INCIDENCE_MODEL",
]
selected = national_selected if area["level"] == "national" else local_selected
available = [(indicators[id], latest(id)) for id in selected if id in indicators and latest(id)]
doc.add_heading("Selected social and economic evidence", 1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, value in zip(table.rows[0].cells, ("Indicator", "Value", "Year", "Source table")):
    cell.text = value
for indicator, row in available:
    source = sources[row["source_id"]]
    cells = table.add_row().cells
    cells[0].text = indicator["name"]
    cells[1].text = f"{number(row['value'], 1)} {indicator['unit']}"
    cells[2].text = row["period"]
    cells[3].text = indicator.get("upstream_table") or source["name"][:35]
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
for cell in table.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

doc.add_heading("Regional comparison", 1)
if area["level"] == "national":
    children = [child for child in data["territories"] if child["parent_id"] == area["id"] and child["level"] == "adm1"]
    values = [(child["name"], latest("BFA_RGPH2019_POP_TOTAL", child["id"])["value"])
              for child in children if latest("BFA_RGPH2019_POP_TOTAL", child["id"])]
    values.sort(key=lambda pair: pair[1], reverse=True)
    assert len(values) == 13
    doc.add_picture(bar_chart(values[:8], "Largest 2019 census regions by population", "people"), width=Inches(6.8))
    doc.add_paragraph("The chart shows eight of the 13 historical census regions. It is not a complete 2025 region ranking. The full 2019 regional set remains available in the AreaData dataset and CSV export.")
else:
    children = [child for child in data["territories"] if child["parent_id"] == area["id"]]
    compare_children = len(children) >= 2
    peers = children if compare_children else [peer for peer in data["territories"] if peer["parent_id"] == area["parent_id"]]
    values = [(peer["name"], latest("BFA_RGPH2019_POP_TOTAL", peer["id"])["value"])
              for peer in peers if latest("BFA_RGPH2019_POP_TOTAL", peer["id"])]
    values.sort(key=lambda pair: pair[1], reverse=True)
    if values:
        shown = values[:8]
        selected_pair = next((pair for pair in values if pair[0] == area["name"]), None)
        if selected_pair and selected_pair not in shown:
            shown = shown[:7] + [selected_pair]
        doc.add_picture(bar_chart(shown, "2019 census population of comparable areas", "people"), width=Inches(6.8))
        peer_note = " The selected area is included." if not compare_children else ""
        doc.add_paragraph(f"The figure displays {len(shown)} of {len(values)} registered {'child' if compare_children else 'peer'} areas with a 2019 census count.{peer_note} It does not turn incomplete child coverage into a parent total; consult the full AreaData comparison CSV for every registered area and missing value.")
    else:
        doc.add_paragraph("No same-level 2019 census population comparison is available for this area. Missing members are not treated as zero.")

if area["level"] == "national":
    doc.add_page_break()
doc.add_heading("Evidence based interpretation", 1)
if profile:
    young = sum(profile["male"][:3]) + sum(profile["female"][:3])
    share = young / population["value"] * 100
    doc.add_paragraph(f"Children aged 0 to 14 account for {number(young)} of {number(population['value'])} residents, or {share:.1f}% by an AreaData calculation from the 2019 age-sex table. This describes age structure in the census period; it does not by itself establish school demand, migration or a policy priority.")
elif all(latest(indicator_id) for indicator_id in ("BFA_RGPH2019_AGE_0_4", "BFA_RGPH2019_AGE_5_14")):
    young = latest("BFA_RGPH2019_AGE_0_4")["value"] + latest("BFA_RGPH2019_AGE_5_14")["value"]
    share = young / population["value"] * 100
    doc.add_paragraph(f"Residents aged 0 to 14 account for {number(young)} of {number(population['value'])} residents, or {share:.1f}% by an AreaData calculation from the two 2019 census age groups. This is a historical age composition, not a current service-demand forecast.")
education = latest("BFA_RGPH2019_NET_POSTPRIMARY_ENROLMENT_REGIONAL")
if education:
    doc.add_paragraph(f"The published 2019 net post-primary enrolment rate for {area['name']} is {education['value']:.1f}%. Its source population is ages 12 to 15. The rate should be reviewed beside local school records before setting a target; it is not the same measure as primary attendance or gross enrolment.")
else:
    education = latest("BFA_INSD_POST_PRIMARY_NET_ATTENDANCE")
    if education:
        doc.add_paragraph(f"The 2019 post-primary net attendance measure for {area['name']} is {education['value']:.1f}%. It comes from the communal-disparities source and is kept separate from the regional enrolment measure; this document does not equate their denominators or definitions.")
poverty = latest("BFA_INSD_POVERTY_INCIDENCE_MODEL")
if poverty:
    doc.add_paragraph(f"The monetary poverty figure of {poverty['value']:.1f}% refers to a 2018 small-area model that uses the 2019 census frame. It is displayed as a modelled estimate, not as a directly enumerated 2019 census count or a current poverty rate.")
doc.add_paragraph("Recorded differences can inform questions for local review. Causes, priorities, budget decisions and community agreement require separate evidence and responsible-authority confirmation; they are not inferred from a map colour or rank.")

doc.add_heading("Planning source and unresolved materials", 1)
doc.add_paragraph("The selected articles of the 2025 territorial code assign communal and regional development-plan responsibilities to communes and regions. This report uses historical statistical areas and does not certify which present-day authority corresponds to each selected area. The 2024 MEF guides are methodological references. No signed area-specific plan, adopted budget, implementation report or evaluation was verified for this selected area.")
doc.add_paragraph("Before using this diagnostic in a plan, obtain the current consolidated law and official territorial correspondence, check the latest plan and budget with the responsible authority, and confirm each statistical definition, period and boundary edition.")

doc.add_heading("Indicator definitions used in this document", 1)
for indicator, row in available:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{indicator['name']} ({row['period']}). ").bold = True
    paragraph.add_run(indicator.get("definition") or "Definition not recorded in this dataset.")

if area["level"] == "national":
    doc.add_page_break()
doc.add_heading("Sources and reproducibility", 1)
planning_ids = {"bfa-mef-pcd-guide-2024", "bfa-mef-prd-guide-2024", "bfa-collectivities-code-2025"}
for source_id in sorted({row["source_id"] for _, row in available} | {population["source_id"]} | planning_ids):
    source = sources[source_id]
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(f"{source['name']} — {source['url']} — retrieved {source.get('retrieved_at', 'not recorded')}").font.size = Pt(8.5)
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_after = Pt(3)
paragraph.add_run(f"AreaData dataset edition {data['generated_at']}. Dataset SHA-256 {sha256((project / 'data/dashboard.json').read_bytes()).hexdigest()}. Source PDF hashes and selected numeric-column locators are preserved in the project evidence register. Values in this document are copied from the selected area observations; derived values are labelled as AreaData calculations.").font.size = Pt(8.5)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.text = "AreaData  |  Source-grounded review material  |  Not an approved development plan"
out = Path(args.out)
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(json.dumps({"file": str(out), "sha256": sha256(out.read_bytes()).hexdigest(),
                  "territory_id": area["id"], "area": area["name"],
                  "observed_values_in_table": len(available), "pyramid": bool(profile)}, ensure_ascii=False))
