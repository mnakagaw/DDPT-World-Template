"""Read-only structural inventory of acquired CBS census workbooks."""

import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


# Explicit header/data rows keep the first published data row in the inventory.
LAYOUT = {
    "Big Localities": (3, 4, 5),
    "Statistical area": (2, 3, 4),
    "Little Localities": (4, 5, 6),
    "District Sub-District Natural  ": (2, 3, 4),
    "Metropolitan": (3, 4, 5),
    "Type of Locality": (3, 4, 5),
    "Municipal Status": (3, 4, 5),
    "Cluster of local authorities": (3, 4, 5),
    "Comparison between censuses": (3, 4, 5),
    "selected data 2022": (2, 3, 4),
}
ADOPTED_DISTRICT_COLUMNS = {
    8: "pop_approx", 9: "DependencyRatio", 10: "Foreign_pcnt",
    200: "Acadm1Cert_pcnt", 219: "walk5_pcnt", 235: "WrkY_pcnt",
    385: "employeesAnnual_medWage", 407: "hh_total_approx",
    438: "HousingDens2_pcnt", 440: "rent_pcnt",
    450: "Vehicle1up_pcnt", 454: "Computer_avg",
}
OTHER_LAYOUT = {
    "households-selected-characteristics.xlsx": (6, 0, 7),
    "population-group-religion-age-sex.xlsx": (6, 7, 8),
    "population-households-locality.xlsx": (6, 8, 9),
    "population-selected-characteristics.xlsx": (6, 0, 7),
}


def short(value):
    if value is None:
        return None
    value = str(value).replace("\n", " ").strip()
    return value[:180]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    dataset = json.loads((project / "data/dashboard.json").read_text(encoding="utf-8"))
    if dataset["country"]["id"] != "ISR":
        raise ValueError("Expected Israel candidate")
    raw = project / "raw/israel-cbs-census-2022"
    inventory = []
    columns = []
    for path in sorted(raw.glob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook:
            sample = []
            for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 18), values_only=True), 1):
                cells = {str(index): short(value) for index, value in enumerate(row[:min(sheet.max_column, 24)], 1)
                         if value is not None and str(value).strip()}
                if cells:
                    sample.append({"row": number, "cells": cells})
            inventory.append({"file": path.name, "sheet": sheet.title, "rows": sheet.max_row,
                              "columns": sheet.max_column, "sample": sample})
            first_rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True))
            layout = LAYOUT.get(sheet.title, OTHER_LAYOUT.get(path.name))
            if layout is None:
                raise ValueError(f"Unreviewed workbook layout: {path.name}/{sheet.title}")
            label_row, code_row, data_start = layout
            nonempty = [0] * sheet.max_column
            numeric = [0] * sheet.max_column
            examples = [None] * sheet.max_column
            for row in sheet.iter_rows(min_row=data_start, values_only=True):
                for index, value in enumerate(row):
                    if value is None or str(value).strip() == "":
                        continue
                    nonempty[index] += 1
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric[index] += 1
                    if examples[index] is None:
                        examples[index] = short(value)
            for index in range(sheet.max_column):
                code = short(first_rows[code_row - 1][index]) if code_row and len(first_rows) >= code_row else ""
                adopted_code = (ADOPTED_DISTRICT_COLUMNS.get(index + 1)
                                if path.name == "broad-geographical-units.xlsx"
                                and sheet.title == "District Sub-District Natural  " else None)
                if adopted_code and code != adopted_code:
                    raise ValueError(f"Adopted source column changed: {index + 1}/{code}")
                columns.append((path.name, sheet.title, index + 1,
                                short(first_rows[label_row - 1][index]) if len(first_rows) >= label_row else "",
                                code, nonempty[index], numeric[index], examples[index] or "",
                                "adopted_country_six_districts" if adopted_code else "unassessed",
                                "Only the Nationwide and CBS district codes 1-6 rows are adopted; code 7 remains source-only."
                                if adopted_code else "Semantic definition, geography and adoption still to review."))
        workbook.close()
    output = project / "evidence/ISR_CBS_WORKBOOK_STRUCTURE.json"
    output.write_text(json.dumps(inventory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (project / "evidence/ISR_CBS_COLUMN_INVENTORY.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("file", "sheet", "column_index", "display_label", "field_code_or_subheader", "nonempty_data_cells",
                         "numeric_data_cells", "first_data_value", "decision", "reason"))
        writer.writerows(columns)
    print(json.dumps([{"file": item["file"], "sheet": item["sheet"],
                       "rows": item["rows"], "columns": item["columns"]} for item in inventory],
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
