"""Inventory every numeric column in two acquired Kuwait census tables."""

import argparse
import csv
import hashlib
import json
import re
import subprocess
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


RAW = "raw/kuwait-csb-2021/"
EXPECTED = {
    "table1-governorate-nationality-gender.pdf": "368bbce7a09929686fb9777ad36e6ab3a8f99dd2d48987bd02a61611130a96b3",
    "table1-governorate-nationality-gender.xlsx": "1ef400667d1cd91b68aba1c446fb63861a65cd45d269b4090707b27893f89bd7",
    "table51-usual-residence.pdf": "23292eaaaa4e70bef182b11bfb082e6eb589e45a00fa0b4b1382628c4dd8a18a",
    "table51-usual-residence.xlsx": "1b3f865e3ec52558a1a516176fc53b73436b27659c2579bf9c1a1fc4468daa51",
}
FIELDS = ["kuwaiti_male", "kuwaiti_female", "kuwaiti_total", "nonkuwaiti_male",
          "nonkuwaiti_female", "nonkuwaiti_total", "all_male", "all_female", "all_total"]
NAMES = ["The Capital Governorate", "Hawalli Governorate", "Al-Ahmadi Governorate",
         "Al-Jahra Governorate", "Al-Farwaniya Governorate",
         "Mubarak Al-Kabeer Governorate", "Not Stated", "Total"]


def pdf_text(path):
    return subprocess.check_output(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"]).decode("utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    evidence = project / "evidence"
    receipts = {}
    for name, digest in EXPECTED.items():
        path = project / RAW / name
        receipt = json.loads((project / RAW / (name + ".receipt.json")).read_text(encoding="utf-8"))
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest or receipt["sha256"] != digest:
            raise ValueError(f"Kuwait original/receipt mismatch: {name}")
        receipts[name] = receipt

    sheet1 = openpyxl.load_workbook(project / RAW / "table1-governorate-nationality-gender.xlsx",
                                    read_only=True, data_only=True).active
    if (sheet1.max_row, sheet1.max_column) != (28, 12):
        raise ValueError("Table 1 workbook dimensions changed")
    table1 = []
    for row_index in range(17, 25):
        row = [sheet1.cell(row_index, col).value for col in range(1, 13)]
        name = str(row[11]).replace("\xa0", " ").strip()
        if name != NAMES[row_index - 17] or any(not isinstance(v, int) for v in row[2:11]):
            raise ValueError(f"Table 1 source row differs: {row_index}: {name}")
        record = {"source_excel_row": row_index, "source_name_en": name,
                  "source_name_ar": row[1], **dict(zip(FIELDS, row[2:11]))}
        for prefix in ("kuwaiti", "nonkuwaiti", "all"):
            if record[prefix + "_male"] + record[prefix + "_female"] != record[prefix + "_total"]:
                raise ValueError(f"Sex subtotal differs: {name}/{prefix}")
        for suffix in ("male", "female", "total"):
            if record["kuwaiti_" + suffix] + record["nonkuwaiti_" + suffix] != record["all_" + suffix]:
                raise ValueError(f"Nationality subtotal differs: {name}/{suffix}")
        table1.append(record)
    if table1[-1]["all_total"] != 4385717 or table1[-2]["all_total"] != 4578:
        raise ValueError("2021 census Table 1 national/unassigned controls differ")
    for field in FIELDS:
        if sum(row[field] for row in table1[:-1]) != table1[-1][field]:
            raise ValueError(f"Six governorates plus Not Stated differ: {field}")
    pdf1 = pdf_text(project / RAW / "table1-governorate-nationality-gender.pdf")
    pdf_numbers = [[int(n) for n in re.findall(r"\b\d+\b", line)]
                   for line in pdf1.splitlines() if len(re.findall(r"\b\d+\b", line)) == 9]
    if len(pdf_numbers) != 8 or any(pdf_numbers[i] != [table1[i][field] for field in reversed(FIELDS)]
                                    for i in range(8)):
        raise ValueError("Table 1 PDF and Excel numeric columns differ")

    sheet51 = openpyxl.load_workbook(project / RAW / "table51-usual-residence.xlsx",
                                     read_only=True, data_only=True).active
    if (sheet51.max_row, sheet51.max_column) != (173, 11):
        raise ValueError("Table 51 workbook dimensions changed")
    table51 = []
    for row_index in range(12, 171):
        name_ar, value, name_en = (sheet51.cell(row_index, col).value for col in (1, 2, 3))
        if not isinstance(value, int) or not isinstance(name_en, str):
            raise ValueError(f"Table 51 source row differs: {row_index}")
        table51.append({"source_excel_row": row_index, "source_name_en": name_en.strip(),
                        "source_name_ar": name_ar, "population": value})
    if len(table51) != 159 or len({row["source_name_en"] for row in table51[:-2]}) != 157 or \
            table51[-2]["source_name_en"] != "NOT STATED" or table51[-1]["source_name_en"] != "TOTAL":
        raise ValueError("Table 51 geography inventory differs")
    if sum(row["population"] for row in table51[:-1]) != 4385717 or \
            table51[-2]["population"] != 4578 or table51[-1]["population"] != 4385717:
        raise ValueError("Table 51 national/unassigned controls differ")
    pdf51 = pdf_text(project / RAW / "table51-usual-residence.pdf")
    pdf_rows = [(match.group(1).strip(), int(match.group(2))) for line in pdf51.splitlines()
                if (match := re.match(r"^\s*([A-Z][A-Z0-9 -]+?)\s{2,}(\d+)\s+", line))]
    if pdf_rows != [(row["source_name_en"], row["population"]) for row in table51]:
        raise ValueError("Table 51 PDF and Excel area rows differ")
    expected_governorate_totals = [row["all_total"] for row in table1[:6]]
    running = 0
    matching_prefix_ends = []
    target = 0
    for row in table51[:-2]:
        running += row["population"]
        if target < 6 and running == sum(expected_governorate_totals[:target + 1]):
            matching_prefix_ends.append(row["source_excel_row"])
            target += 1
    if matching_prefix_ends != [44, 61, 102, 135, 155, 168]:
        raise ValueError("Table 51 contiguous arithmetic partition changed")

    with (evidence / "KWT_CSB_2021_TABLE1_ROWS.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["source_excel_row", "source_name_en", "source_name_ar", *FIELDS])
        writer.writeheader()
        writer.writerows(table1)
    with (evidence / "KWT_CSB_2021_TABLE51_ROWS.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["source_excel_row", "source_name_en", "source_name_ar", "population"])
        writer.writeheader()
        writer.writerows(table51)
    with (evidence / "INDICATOR_INVENTORY.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["source_id", "source_hash", "table_or_sheet", "column_or_variable", "original_label",
                         "unit", "universe", "period", "geography_type", "role", "indicator_id", "decision", "reason", "locator"])
        for field in FIELDS:
            writer.writerow(["kwt-csb-registration-census-2021-table1", EXPECTED["table1-governorate-nationality-gender.xlsx"],
                             "Table 1 / Sheet1", field, field.replace("_", " "), "people",
                             "PACI registration census population; non-Kuwaiti exclusions noted on official table webpage",
                             "2021 census project; exact register extraction date unverified", "nation and governorate source rows",
                             "direct printed count", "KWT_CSB_2021_" + field.upper(), "adopted",
                             "PDF and XLSX agree for all eight rows; Not Stated is retained outside territorial hierarchy",
                             "Excel rows 17-24, columns C-K; PDF Table 1 page 1"])
        writer.writerow(["kwt-csb-registration-census-2021-table51", EXPECTED["table51-usual-residence.xlsx"],
                         "Table 51 / Sheet1", "population", "Population Number", "people",
                         "habitual residence area reporting rows", "2021 census project", "157 area labels plus Not Stated and total",
                         "direct count", "", "unadopted", "Official area codes and explicit parent-governorate keys not verified; arithmetic row blocks alone are insufficient for a geographic join",
                         "Excel rows 12-170, column B; PDF Table 51 pages 1-3"])
    structure = {"checked_at": datetime.now(timezone.utc).isoformat(), "raw_sha256": EXPECTED,
                 "table1_excel_pdf_rows_checked": 8, "table1_numeric_columns_checked": len(FIELDS),
                 "table1_national_total": table1[-1]["all_total"], "table1_not_stated": table1[-2]["all_total"],
                 "table1_six_governorates_total": sum(row["all_total"] for row in table1[:6]),
                 "table51_excel_pdf_rows_checked": len(table51), "table51_named_area_rows": 157,
                 "table51_contiguous_prefix_end_excel_rows": matching_prefix_ends,
                 "table51_geographic_adoption": "withheld_pending_official_area_code_parent_and_boundary_crosswalk",
                 "note": "The row order and exact subtotal matches do not themselves establish 157 official parent-child identities. The homepage count 4,385,291 differs from Table 1 final 4,385,717 and is not used."}
    (evidence / "KWT_CSB_2021_STRUCTURE.json").write_text(
        json.dumps(structure, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"table1_rows": 8, "adoptable_governorates": 6,
                      "numeric_columns": len(FIELDS), "table51_named_areas_unadopted": 157,
                      "not_stated": 4578, "national": 4385717}))


if __name__ == "__main__":
    main()
