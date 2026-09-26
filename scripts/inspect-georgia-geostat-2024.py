"""Inventory acquired Geostat 2024 census sheets and check candidate table arithmetic."""

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl


RAW_DIR = "raw/georgia-geostat-2024"
FIELDS_POP = ("total", "male", "female", "urban_total", "urban_male", "urban_female", "rural_total", "rural_male", "rural_female")
FIELDS_IDP = ("idp_total", "idp_origin_abkhazia", "idp_origin_tskhinvali", "idp_male", "idp_male_origin_abkhazia", "idp_male_origin_tskhinvali", "idp_female", "idp_female_origin_abkhazia", "idp_female_origin_tskhinvali")
ALIASES_IN_TABLE_01 = {
    "Adjara A.R.": "Autonomous Republic of Adjara",
    "Imereti": "ImereTi",
    "Sighnagi Municipality": "Sighnaghi Municipality",
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def numeric(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if value == "-":
        return 0  # Arithmetic only: the source footnote says magnitude nil.
    raise ValueError(f"Unexpected numeric source cell: {value!r}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    raw = project / RAW_DIR
    receipts = json.loads((project / "evidence/GEO_GEOSTAT_2024_ACQUISITION.json").read_text(encoding="utf-8"))["receipts"]
    if len(receipts) != 7:
        raise ValueError("Expected six census XLSX files and one PDF")
    inventories = []
    for receipt in receipts:
        file = project / receipt["raw_path"]
        if not file.is_relative_to(raw) or digest(file) != receipt["sha256"]:
            raise ValueError(f"Missing or changed source original: {file}")
        if file.suffix != ".xlsx":
            continue
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        if len(book.worksheets) != 1:
            raise ValueError(f"Unexpected sheets: {file.name}")
        sheet = book.active
        numeric_cells = Counter()
        markers = Counter()
        row_count_with_number = 0
        for row in sheet.iter_rows(values_only=True):
            seen = False
            for j, value in enumerate(row[1:], 2):
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric_cells[j] += 1
                    seen = True
                elif isinstance(value, str) and value in {"-", "...", "x", "X"}:
                    markers[value] += 1
            row_count_with_number += int(seen)
        header = list(sheet.iter_rows(min_row=1, max_row=min(8, sheet.max_row), values_only=True))
        inventories.append({
            "file": file.name, "source_url": receipt["source_url"], "sha256": receipt["sha256"],
            "sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column,
            "rows_with_numeric_content": row_count_with_number,
            "numeric_columns": [{"excel_column": openpyxl.utils.get_column_letter(col), "column_index": col,
                                 "numeric_cells": count,
                                 "header_cells_rows_1_to_8": [{"row": row_index, "label": str(row[col - 1])}
                                                              for row_index, row in enumerate(header, 1)
                                                              if len(row) >= col and row[col - 1] is not None]}
                                for col, count in sorted(numeric_cells.items())],
            "special_markers": dict(markers),
            "adoption_decision": "candidate_checked_columns_only" if file.name.startswith(("02-", "06-")) else "not_adopted_this_pass",
        })
    pop_file = raw / "02-self-governed-units-urban-rural-sex.xlsx"
    pop_sheet = openpyxl.load_workbook(pop_file, read_only=False, data_only=True).active
    if (pop_sheet.max_row, pop_sheet.max_column) != (95, 10) or pop_sheet.cell(7, 2).value != 3929581:
        raise ValueError("2024 population table changed")
    pop_rows = []
    current_region = None
    for n in range(7, 92):
        name = pop_sheet.cell(n, 1).value
        if not isinstance(name, str) or not name:
            raise ValueError(f"Missing population geography at row {n}")
        values = [pop_sheet.cell(n, c).value for c in range(2, 11)]
        if any(not isinstance(v, (int, float)) and v != "-" for v in values):
            raise ValueError(f"Unexpected population value at row {n}")
        if n == 7:
            area_type, parent = "country", ""
        elif pop_sheet.cell(n, 1).font.bold:
            area_type, parent, current_region = "region_or_self_governing_city", "Georgia", name
        elif current_region == "C. Tbilisi":
            area_type, parent = "tbilisi_district", current_region
        else:
            area_type, parent = "self_governing_unit", current_region
        if numeric(values[0]) != numeric(values[1]) + numeric(values[2]):
            raise ValueError(f"Sex columns differ at population row {n}")
        if numeric(values[0]) != numeric(values[3]) + numeric(values[6]):
            raise ValueError(f"Urban/rural columns differ at population row {n}")
        for a, b, c in ((3, 4, 5), (6, 7, 8)):
            if numeric(values[a]) != numeric(values[b]) + numeric(values[c]):
                raise ValueError(f"Settlement sex columns differ at population row {n}")
        pop_rows.append({"source_excel_row": n, "source_name_en": name, "area_type": area_type,
                         "parent_source_name_en": parent, **dict(zip(FIELDS_POP, values))})
    regions = [r for r in pop_rows if r["area_type"] == "region_or_self_governing_city"]
    municipalities = [r for r in pop_rows if r["area_type"] == "self_governing_unit"]
    districts = [r for r in pop_rows if r["area_type"] == "tbilisi_district"]
    if (len(pop_rows), len(regions), len(municipalities), len(districts)) != (85, 11, 63, 10):
        raise ValueError("Unexpected geography inventory")
    for field in FIELDS_POP:
        if sum(numeric(r[field]) for r in regions) != numeric(pop_rows[0][field]):
            raise ValueError(f"National region cover fails: {field}")
        for region in regions:
            children = [r for r in pop_rows if r["parent_source_name_en"] == region["source_name_en"]]
            if not children or sum(numeric(r[field]) for r in children) != numeric(region[field]):
                raise ValueError(f"Region child cover fails: {region['source_name_en']} {field}")

    detailed = openpyxl.load_workbook(raw / "01-administrative-territorial-units-sex.xlsx", read_only=True, data_only=True).active
    detailed_counts = {(str(r[0]).strip(), r[1], r[2], r[3]) for i, r in enumerate(detailed.iter_rows(values_only=True), 1)
                       if i >= 6 and isinstance(r[0], str)}
    for r in pop_rows:
        key = (ALIASES_IN_TABLE_01.get(r["source_name_en"], r["source_name_en"]), r["total"], r["male"], r["female"])
        if key not in detailed_counts:
            raise ValueError(f"Population broad row absent from detailed Table 01: {key}")

    idp_sheet = openpyxl.load_workbook(raw / "06-idp-residence-origin-sex.xlsx", read_only=True, data_only=True).active
    idp_rows = []
    joined = {(r["source_name_en"], r["total"]): r for r in pop_rows if r["area_type"] != "tbilisi_district"}
    if idp_sheet.max_row != 85:
        raise ValueError("IDP sheet changed")
    for n in range(8, 83):
        name = idp_sheet.cell(n, 1).value
        values = [idp_sheet.cell(n, c).value for c in range(2, 11)]
        if not isinstance(name, str) or any(not isinstance(v, (int, float)) for v in values):
            raise ValueError(f"Unexpected IDP row {n}")
        match = next((r for (label, _), r in joined.items() if label == name), None)
        if match is None or values[0] > match["total"]:
            raise ValueError(f"IDP geography missing or exceeds census population: {name}")
        if values[0] != values[1] + values[2] or values[0] != values[3] + values[6] or \
                values[3] != values[4] + values[5] or values[6] != values[7] + values[8]:
            raise ValueError(f"IDP origin/sex arithmetic fails at row {n}")
        idp_rows.append({"source_excel_row": n, "source_name_en": name,
                         "area_type": match["area_type"], "parent_source_name_en": match["parent_source_name_en"],
                         **dict(zip(FIELDS_IDP, values))})
    if len(idp_rows) != 75 or len({r["source_name_en"] for r in idp_rows}) != 75:
        raise ValueError("IDP area inventory changed")
    for field in FIELDS_IDP:
        if sum(r[field] for r in idp_rows if r["area_type"] == "region_or_self_governing_city") != idp_rows[0][field]:
            raise ValueError(f"IDP national region cover fails: {field}")
        for region in regions:
            if region["source_name_en"] == "C. Tbilisi":
                continue  # No district-level IDP rows were published in Table 06.
            children = [r for r in idp_rows if r["parent_source_name_en"] == region["source_name_en"]]
            parent = next(r for r in idp_rows if r["source_name_en"] == region["source_name_en"])
            if sum(r[field] for r in children) != parent[field]:
                raise ValueError(f"IDP municipality cover fails: {region['source_name_en']} {field}")

    evidence = project / "evidence"
    for name, rows, fields in (("GEO_GEOSTAT_2024_POPULATION_ROWS.csv", pop_rows, FIELDS_POP),
                               ("GEO_GEOSTAT_2024_IDP_ROWS.csv", idp_rows, FIELDS_IDP)):
        with (evidence / name).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["source_excel_row", "source_name_en", "area_type", "parent_source_name_en", *fields])
            writer.writeheader()
            writer.writerows(rows)
    structure = {
        "status": "source_structure_checked_adoption_pending", "raw_sha256": {r["raw_path"].split("/")[-1]: r["sha256"] for r in receipts},
        "sheet_inventory": inventories, "population_rows": 85, "population_regions": 11,
        "population_municipalities": 63, "tbilisi_internal_districts": 10,
        "population_numeric_fields": 9, "population_nil_cells": 27,
        "population_national": 3929581, "population_national_urban": 2455444,
        "population_national_rural": 1474137,
        "all_population_rows_reproduced_in_detailed_sheet": True,
        "table01_spelling_variants": ALIASES_IN_TABLE_01,
        "idp_rows": 75, "idp_numeric_fields": 9, "idp_national": 210628,
        "idp_tbilisi_district_rows": 0,
        "source_note": "Both tables exclude occupied territories. A dash in Table 02 means magnitude nil; no zero observation is created by the adapter for those cells. IDPs are a population subset, never added to total census population.",
    }
    (evidence / "GEO_GEOSTAT_2024_STRUCTURE.json").write_text(json.dumps(structure, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: structure[k] for k in ("population_rows", "population_regions", "population_municipalities", "tbilisi_internal_districts", "population_nil_cells", "idp_rows", "idp_national")}))


if __name__ == "__main__":
    main()
