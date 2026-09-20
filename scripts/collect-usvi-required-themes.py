#!/usr/bin/env python3
"""Extract an auditable U.S. Virgin Islands 2020 Census island bundle."""
from __future__ import annotations
import argparse,csv,hashlib,json,sys,zipfile
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook

FIELDS={
 'VIR_C2020_POP_TOTAL':('DPVI0010001','Population','people','Total resident population.'),
 'VIR_C2020_FEMALE_PCT':(('DPVI0010039','DPVI0010001'),'Population','%','Female population as a share of total population.'),
 'VIR_C2020_HOUSING_UNITS':('DPVI0930001','Households and housing','housing units','Total housing units.'),
 'VIR_C2020_PUBLIC_WATER_PCT':(('DPVI1040002','DPVI1040001'),'Water','% of housing units','Housing units whose source of water is a public system.'),
 'VIR_C2020_PUBLIC_SEWER_PCT':(('DPVI1050002','DPVI1050001'),'Sanitation','% of housing units','Housing units using public sewer for sewage disposal.'),
 'VIR_C2020_HIGH_SCHOOL_OR_HIGHER_PCT':('DPVI0190001','Education','% of population age 25+ in households','Census-published percent high school graduate or higher.'),
 'VIR_C2020_LABOR_FORCE_PCT':(('DPVI0490002','DPVI0490001'),'Employment','% of population age 16+ in households','Population age 16+ in households in the labor force.'),
 'VIR_C2020_DISABILITY_PCT':(('DPVI0370002','DPVI0370001'),'Disability','% of civilian population in households','Civilian population in households with a disability.'),
 'VIR_C2020_BORN_ELSEWHERE_PCT':(('DPVI0410005','DPVI0410001'),'Migration','% of population in households','Population in households born elsewhere, using the Census table category.'),
 'VIR_C2020_HISPANIC_LATINO_PCT':(('DPVI0090002','DPVI0090001'),'Race and ethnicity','%','Hispanic or Latino population of any race.'),
 'VIR_C2020_UNINSURED_PCT':(('DPVI0600005','DPVI0600001'),'Health','% of civilian population in households','Civilian population in households with no health insurance coverage.'),
 'VIR_C2020_SNAP_HOUSEHOLDS_PCT':(('DPVI0780002','DPVI0780001'),'Food assistance','% of households','Households receiving Food Stamp/SNAP benefits in 2019; a program-participation measure, not a nutrition outcome.'),
 'VIR_C2020_POVERTY_PCT':(('DPVI0890002','DPVI0890001'),'Poverty','% of families','Families with income in 2019 below the poverty level.'),
}

def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()

def number(v):
 if v in (None,'','.'):return None
 try:return float(v)
 except ValueError:return None

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--raw',required=True);ap.add_argument('--boundary-zip',required=True);ap.add_argument('--out',required=True);a=ap.parse_args()
 raw=Path(a.raw).resolve();out=Path(a.out).resolve();out.mkdir(parents=True,exist_ok=True);matrix=raw/'2020-iac-usvi-table-matrix.xlsx';geo_layout=raw/'2020-iac-geographic-header-record-usvi.xlsx';ex=raw/'extracted'
 ws=load_workbook(matrix,read_only=True,data_only=True)['Table Matrix'];segments={i:[] for i in range(1,5)};labels={}
 for rowno,row in enumerate(ws.iter_rows(values_only=True),1):
  ref=row[2] if len(row)>2 else None;seg=row[3] if len(row)>3 else None
  if isinstance(ref,str) and ref.startswith('DPVI') and isinstance(seg,(int,float)):
   # The official workbook has one evident segment typo: VI86 belongs to file 03.
   if ref=='DPVI0860001':seg=3
   segments[int(seg)].append(ref);labels[ref]=str(row[1] or ref)
 expected={1:247,2:242,3:247,4:29}
 if {k:len(v) for k,v in segments.items()}!=expected:raise RuntimeError(f'Unexpected matrix segment counts: {{k:len(v) for k,v in segments.items()}}')
 ws=load_workbook(geo_layout,read_only=True,data_only=True).active;geo_fields=[]
 for row in ws.iter_rows(values_only=True):
  ref=row[1] if len(row)>1 else None
  if isinstance(ref,str) and ref.strip() and ref.strip()!='DATA DICTIONARY REFERENCE':geo_fields.append(ref.strip())
 geos={}
 with (ex/'vigeo2020.dp').open(encoding='utf-8',newline='') as f:
  for values in csv.reader(f,delimiter='|'):
   row=dict(zip(geo_fields,values))
   if row['SUMLEV'] in {'040','050'}:geos[row['LOGRECNO']]=row
 if len(geos)!=4:raise RuntimeError(f'Expected territory plus three island rows, got {len(geos)}')
 data={log:{} for log in geos}
 for seg,refs in segments.items():
  with (ex/f'vi0000{seg}2020.dp').open(encoding='utf-8',newline='') as f:
   for values in csv.reader(f,delimiter='|'):
    log=values[4]
    if log in data:data[log].update(dict(zip(refs,values[5:])))
 ids={log:('VIR' if row['SUMLEV']=='040' else f"VIR:C2020:ISLAND:{row['COUNTY']}") for log,row in geos.items()}
 territories=[{'id':ids[log],'country_id':'VIR','name':row['BASENAME'],'level':'island','type':'island','parent_id':'VIR','official_code':f"78{row['COUNTY']}",'code_system':'Census ANSI/FIPS','boundary_version':'2020 Census TIGER/Line county-equivalent boundary','valid_from':'2020-04-01'} for log,row in geos.items() if row['SUMLEV']=='050']
 source={'id':'vir-census-2020-demographic-profile','name':'2020 Island Areas Census U.S. Virgin Islands Demographic Profile Summary File','publisher':'United States Census Bureau','url':'https://www.census.gov/data/tables/2020/dec/2020-us-virgin-islands.html','status':'ready','retrieved_at':datetime.now(timezone.utc).date().isoformat(),'reference_period':'2020 Census; income and SNAP fields refer to 2019','geographic_level':'U.S. Virgin Islands and islands','license':'United States Government public data','license_url':'https://www.census.gov/about/policies/open-gov/open-data.html','note':'Official fixed summary-file ZIP, matrix and geographic layout are archived. Social and economic table universes exclude the group-quarters population because of documented COVID-19 collection limits.'}
 indicators=[];observations=[]
 for iid,(spec,theme,unit,definition) in FIELDS.items():
  refs=(spec,) if isinstance(spec,str) else spec;locator=' / '.join(refs) if len(refs)>1 else refs[0]
  indicators.append({'id':iid,'name':labels[refs[0]],'theme':theme,'unit':unit,'definition':definition,'definition_id':refs[0],'population':definition,'measurement_method':'source_reported' if len(refs)==1 else 'derived_from_source_counts','aggregation':'sum' if unit in {'people','housing units'} else 'none','period_policy':'fixed_source_period','series_family':'census','display_role':'primary','source_id':source['id'],'source_locator':locator})
  for log,row in geos.items():
   if len(refs)==1:value=number(data[log].get(refs[0]));num=den=None
   else:
    num=number(data[log].get(refs[0]));den=number(data[log].get(refs[1]));value=None if num is None or den in (None,0) else round(num/den*100,4)
   observations.append({'territory_id':ids[log],'indicator_id':iid,'period':'2020','value':value,'status':'observed' if value is not None else 'missing','source_id':source['id'],'definition':definition,'definition_id':refs[0],'unit':unit,'population':definition,'measurement_method':'source_reported' if len(refs)==1 else 'derived_from_source_counts','source_locator':locator,'numerator':num,'denominator':den})
 # Settlement context from same-vintage Census population and land area; not an urban/rural classification.
 density_id='VIR_C2020_POP_DENSITY';indicators.append({'id':density_id,'name':'Population density (settlement context)','theme':'Settlement','unit':'people per km²','definition':'2020 Census population divided by Census land area. This is a settlement-density context measure, not an urban/rural classification.','definition_id':'VIR-C2020-POP-AREALAND','population':'2020 Census resident population and geographic header land area','measurement_method':'derived_from_source_counts','aggregation':'none','period_policy':'fixed_source_period','series_family':'census','display_role':'context','source_id':source['id'],'source_locator':'POP100 / AREALAND * 1,000,000'})
 for log,row in geos.items():
  pop=number(row['POP100']);land=number(row['AREALAND']);value=None if pop is None or not land else round(pop/(land/1e6),4);observations.append({'territory_id':ids[log],'indicator_id':density_id,'period':'2020','value':value,'status':'observed' if value is not None else 'missing','source_id':source['id'],'definition':indicators[-1]['definition'],'definition_id':'VIR-C2020-POP-AREALAND','unit':'people per km²','population':indicators[-1]['population'],'measurement_method':'derived_from_source_counts','source_locator':'POP100 / AREALAND * 1,000,000','numerator':pop,'denominator':land})
 # Exact island boundary join from the retained Census county-equivalent cartographic file.
 sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.work'/'python-packages'));import shapefile
 bz=Path(a.boundary_zip).resolve();ext=out/'boundary-extracted';ext.mkdir(exist_ok=True)
 with zipfile.ZipFile(bz) as z:z.extractall(ext)
 reader=shapefile.Reader(str(next(ext.glob('*.shp'))));features=[];bycode={r['official_code']:r['id'] for r in territories}
 for rec,shape in zip(reader.records(),reader.shapes()):
  code=str(rec.as_dict()['GEOID'])
  if code in bycode:features.append({'type':'Feature','properties':{'territory_id':bycode[code],'source_id':'vir-census-boundaries-2020','official_code':code,'code_system':'Census ANSI/FIPS','geometry_edition':'2020 Census TIGER/Line county-equivalent boundary','join_method':'Exact Census island GEOID/FIPS join','reference_only':True},'geometry':shape.__geo_interface__})
 if len(features)!=3:raise RuntimeError(f'Expected 3 island boundaries, got {len(features)}')
 bsource={'id':'vir-census-boundaries-2020','name':'2020 Census TIGER/Line county-equivalent boundary file','publisher':'United States Census Bureau','url':'https://www2.census.gov/geo/tiger/TIGER2020/COUNTY/','status':'ready','retrieved_at':source['retrieved_at'],'reference_period':'2020','geographic_level':'U.S. Virgin Islands islands','license':'United States Government public data; reference geometry','note':'Three island/county-equivalent features joined by exact FIPS; not a legal-boundary certification.'}
 members=sorted(r['id'] for r in territories);receipts=[{'path':str(p),'bytes':Path(p).stat().st_size,'sha256':sha(p)} for p in [raw/'vi2020.dp.zip',matrix,geo_layout,raw/'variables.json',bz]]
 audit={'schema_version':'1.0','generated_at':datetime.now(timezone.utc).isoformat(),'country_area_id':'VIR','status':'complete_country_depth_bundle','counts':{'territories':3,'indicators':len(indicators),'observations':len(observations),'boundaries':3},'matrix_segment_counts':expected,'source_universe_limit':'Documented 2020 USVI social/economic characteristics exclude group-quarters population; each indicator preserves its table universe.','receipts':receipts}
 bundle={'schema_version':'1.0','country_area_id':'VIR','period':'2020','replace_country_branch':True,'territories':territories,'terminal_territory_ids':members,'comparisons':[{'parent_id':'VIR','member_ids':members,'label':'Three islands','membership_note':'2020 Census island/county-equivalent geography; exact FIPS joins.','source_ids':[source['id'],bsource['id']]}],'indicators':indicators,'observations':observations,'sources':[source,bsource],'boundaries':{'type':'FeatureCollection','features':features},'audit':audit}
 (out/'usvi-country-depth-bundle.json').write_text(json.dumps(bundle,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');(out/'USVI_COLLECTION_AUDIT.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(audit['counts'],indent=2))

if __name__=='__main__':main()
