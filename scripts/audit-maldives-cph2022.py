"""Audit selected Maldives 2022 census source fields before country import.

Only P4, P5, H2, H7, EC3 and ED16 selected fields are adopted. The other
acquired workbooks remain mechanically inventoried, not semantically accepted.
"""

import argparse
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/maldives-cph2022-source-manifest.json"
ATOLLS = "HA HDH SH N R B LH K AA ADH V M F DH TH L GA GDH GN S".split()
POP = ("population", "female", "male", "maldivian_population", "maldivian_female",
       "maldivian_male", "foreign_population", "foreign_female", "foreign_male")
WATER = ("households", "safe_source_reported", "unsafe_source_reported",
         "rain_treated", "rain_untreated", "desalinated", "bottled")
EMPLOYMENT = ("population_15_plus", "labour_force", "employed", "unemployed", "outside_labour_force")
LITERACY = ("maldivian_population_10_plus", "maldivian_literate_mother_tongue")


def require(ok, message):
    if not ok:
        raise ValueError(message)


def digest(body):
    return hashlib.sha256(body).hexdigest()


def counted(code, index, row, columns, names):
    result = {}
    for column, name in zip(columns, names, strict=True):
        value = row[column]
        if value is None and (code, index, column) == ("P5", 17, 6):
            result[name] = None  # Source blank: inferred arithmetic zero is not an observed value.
            continue
        require(isinstance(value, int) and not isinstance(value, bool) and value >= 0,
                f"Expected count {code}!{chr(65 + column)}{index}: {value!r}")
        result[name] = value
    return result


def pop_row(code, index, row, key=None):
    fields = counted(code, index, row, range(2, 11), POP)
    require(fields["population"] == fields["female"] + fields["male"], f"Sex sum {code} row {index}")
    if fields["maldivian_female"] is not None:
        require(fields["maldivian_population"] == fields["maldivian_female"] + fields["maldivian_male"],
                f"Maldivian sex sum {code} row {index}")
    require(fields["foreign_population"] == fields["foreign_female"] + fields["foreign_male"],
            f"Foreign sex sum {code} row {index}")
    for suffix in ("population", "female", "male"):
        if fields[f"maldivian_{suffix}"] is not None:
            require(fields[suffix] == fields[f"maldivian_{suffix}"] + fields[f"foreign_{suffix}"],
                    f"Nationality sum {code} row {index} {suffix}")
    return {"row": index, "code": key, "name": str(row[1]).strip(), "fields": fields}


def household_row(index, row, key=None):
    return {"row": index, "code": key, "name": str(row[2]).strip(),
            "fields": counted("H2", index, row, (3,), ("households",))}


def water_row(index, row, key=None):
    fields = counted("H7", index, row, range(3, 10), WATER)
    require(fields["households"] == fields["safe_source_reported"] + fields["unsafe_source_reported"],
            f"Water total {index}")
    require(fields["safe_source_reported"] == fields["rain_treated"] +
            fields["desalinated"] + fields["bottled"], f"Source safe-category sum {index}")
    require(fields["unsafe_source_reported"] == fields["rain_untreated"],
            f"Source unsafe-category sum {index}")
    return {"row": index, "code": key, "name": str(row[2]).strip(), "fields": fields}


def employment_row(index, row, key=None, name=None):
    fields = counted("EC3", index, row, range(2, 7), EMPLOYMENT)
    require(fields["population_15_plus"] == fields["labour_force"] + fields["outside_labour_force"],
            f"Employment population sum {index}")
    require(fields["labour_force"] == fields["employed"] + fields["unemployed"],
            f"Employment labour force sum {index}")
    return {"row": index, "code": key, "name": name or str(row[1]).strip(), "fields": fields}


def literacy_row(index, row, key=None, name=None):
    fields = counted("ED16", index, row, (2, 8), LITERACY)
    require(row[2] == row[3] + row[4] and row[8] == row[9] + row[10],
            f"Maldivian literacy sex sum {index}")
    require(fields["maldivian_literate_mother_tongue"] <= fields["maldivian_population_10_plus"],
            f"Literate exceeds population {index}")
    return {"row": index, "code": key, "name": name or str(row[1]).strip(), "fields": fields}


def each_sum(rows, names, expected, label, permit_source_blank=False):
    for field in names:
        if any(item["fields"][field] is None for item in rows):
            require(permit_source_blank and label == "P5 Maale parts" and
                    field == "maldivian_female", f"Unexpected source blank: {label} {field}")
            continue
        require(sum(item["fields"][field] for item in rows) == expected["fields"][field],
                f"{label} {field}: {sum(item['fields'][field] for item in rows)} != {expected['fields'][field]}")


def main(project):
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    inventory = json.loads((project / "evidence/MDV_CPH2022_WORKBOOK_INVENTORY.json").read_text(encoding="utf-8"))
    require(inventory["manifest_sha256"] == digest(MANIFEST.read_bytes()), "Rerun full workbook inventory")
    source = {entry["code"].upper(): entry for entry in manifest["source_files"]}

    def load(code):
        item = source[code]
        path = project / item["raw_path"]
        body = path.read_bytes()
        require(len(body) == item["bytes"] and digest(body) == item["sha256"], f"Original changed {code}")
        workbook = load_workbook(path, read_only=True, data_only=True)
        require(len(workbook.worksheets) == 1, f"Expected one sheet in {code}")
        return list(enumerate(workbook.active.iter_rows(values_only=True), 1))

    p4, p5, h2, h7, ec3, ed16 = (load(code) for code in ("P4", "P5", "H2", "H7", "EC3", "ED16"))
    atolls = [pop_row("P4", i, row, str(row[0]).upper()) for i, row in p4
               if str(row[0]).upper() in ATOLLS and isinstance(row[2], int)]
    require(len(atolls) == 20 and {x["code"] for x in atolls} == set(ATOLLS), "Missing 2022 atoll row")
    atoll_by_code = {item["code"]: item for item in atolls}

    p5_country = pop_row("P5", 6, p5[5][1])
    p5_maale = pop_row("P5", 8, p5[7][1])
    p5_maale_parts = [pop_row("P5", i, row) for i, row in p5 if 9 <= i <= 17]
    p5_admin = pop_row("P5", 19, p5[18][1])
    p5_nonadmin = pop_row("P5", 207, p5[206][1])
    require([x["name"] for x in (p5_country, p5_maale, p5_admin, p5_nonadmin)] ==
            ["Republic", "Maale", "Admin islands", "Non-admin islands"], "P5 scope labels changed")
    islands = [pop_row("P5", i, row, str(row[0]).upper()) for i, row in p5
               if str(row[0]).upper() in ATOLLS and isinstance(row[2], int)]
    require(len(islands) == 186 and len(p5_maale_parts) == 9, "P5 island/city-part count changed")
    island_keys = {(x["code"], x["name"]): x for x in islands}
    require(len(island_keys) == len(islands), "Duplicate island key")
    for code, parent in atoll_by_code.items():
        each_sum([x for x in islands if x["code"] == code], POP, parent, f"P4/P5 {code}")
    each_sum(atolls, POP, p5_admin, "P4 atolls/P5 admin")
    each_sum(p5_maale_parts, POP, p5_maale, "P5 Maale parts", permit_source_blank=True)
    each_sum([p5_maale, p5_admin, p5_nonadmin], POP, p5_country, "P5 country partition")

    h2_country = household_row(7, h2[6][1]); h2_maale = household_row(9, h2[8][1])
    h2_maale_parts = [household_row(i, row) for i, row in h2 if 10 <= i <= 18]
    h2_admin = household_row(20, h2[19][1]); h2_nonadmin = household_row(208, h2[207][1])
    h2_islands = [household_row(i, row, str(row[1]).upper()) for i, row in h2
                  if str(row[1]).upper() in ATOLLS and isinstance(row[3], int)]
    require(len(h2_islands) == 186 and len(h2_maale_parts) == 9,
            "H2 island/city-part count changed")
    require({(x["code"], x["name"]) for x in h2_islands} == set(island_keys),
            "P5/H2 island identity differs")
    require([x["name"] for x in (h2_country, h2_maale, h2_admin, h2_nonadmin)] ==
            ["Republic", "Maale", "Admin islands", "Non-admin islands"], "H2 scope labels changed")
    each_sum(h2_islands, ("households",), h2_admin, "H2 admin islands")
    each_sum(h2_maale_parts, ("households",), h2_maale, "H2 Maale parts")
    each_sum([h2_maale, h2_admin, h2_nonadmin], ("households",), h2_country, "H2 country partition")
    require([x["name"] for x in h2_maale_parts] == [x["name"] for x in p5_maale_parts],
            "P5/H2 Maale subarea identity differs")

    w_country = water_row(7, h7[6][1]); w_maale = water_row(8, h7[7][1])
    w_all_atolls = water_row(9, h7[8][1]); w_admin = water_row(10, h7[9][1])
    w_atolls = [water_row(i, row, str(row[1]).upper()) for i, row in h7
                 if str(row[1]).upper() in ATOLLS and isinstance(row[3], int)]
    w_nonadmin = water_row(32, h7[31][1])
    require(len(w_atolls) == 20 and {x["code"] for x in w_atolls} == set(ATOLLS),
            "H7 atoll count changed")
    each_sum(w_atolls, WATER, w_admin, "H7 admin-atoll rows")
    each_sum([w_admin, w_nonadmin], WATER, w_all_atolls, "H7 all atolls")
    each_sum([w_maale, w_all_atolls], WATER, w_country, "H7 country")
    require(w_country["fields"]["households"] == h2_country["fields"]["households"] and
            w_maale["fields"]["households"] == h2_maale["fields"]["households"] and
            w_nonadmin["fields"]["households"] == h2_nonadmin["fields"]["households"],
            "H2/H7 household universes differ")
    h2_by_code = defaultdict(list)
    for item in h2_islands:
        h2_by_code[item["code"]].append(item)
    for item in w_atolls:
        require(sum(x["fields"]["households"] for x in h2_by_code[item["code"]]) ==
                item["fields"]["households"], f"H2/H7 atoll denominator differs: {item['code']}")

    def second_source(code, rows, constructor, fields, country_i, maale_i, nonadmin_i):
        country = constructor(country_i, rows[country_i - 1][1])
        maale = constructor(maale_i, rows[maale_i - 1][1])
        nonadmin = constructor(nonadmin_i, rows[nonadmin_i - 1][1])
        selected, held = [], []
        seen = set()
        for index, row in rows:
            ab = str(row[0]).upper()
            if ab not in ATOLLS or not isinstance(row[2], int):
                continue
            name = str(row[1]).strip()
            if code == "EC3":
                name = re.sub(rf"^{re.escape(str(row[0]))}\s+", "", name, flags=re.I)
            item = constructor(index, row, ab, name)
            key = (ab, name)
            require(key not in seen, f"Duplicate {code} row {key}")
            seen.add(key)
            if key in island_keys:
                selected.append(item)
            else:
                held.append({**item, "decision": "unmatched_2022_island_name_not_adopted"})
        require(len(selected) == 184 and len(held) == 2 and len(seen) == 186,
                f"Unexpected {code} match coverage {len(selected)} {len(held)}")
        require({(x["code"], x["name"]) for x in selected}.issubset(island_keys),
                f"Unmatched selected {code} row")
        return {"country": country, "maale": maale, "nonadmin": nonadmin,
                "islands": selected, "held_rows": held, "selected_fields": fields}

    emp = second_source("EC3", ec3, employment_row, EMPLOYMENT, 8, 9, 199)
    lit = second_source("ED16", ed16, literacy_row, LITERACY, 7, 8, 198)
    require({(x["code"], x["name"]) for x in emp["held_rows"]} ==
            {("L", "Mundhoo"), ("GDH", "Madeveli")}, "Employment name conflicts changed")
    require({(x["code"], x["name"]) for x in lit["held_rows"]} ==
            {("L", "Mundhoo"), ("GDH", "Madeveli")}, "Literacy name conflicts changed")

    audit = {"status": "selected_numeric_fields_audited_other_52_workbooks_or_fields_priority_unassessed",
             "manifest_sha256": digest(MANIFEST.read_bytes()),
             "selected_source_codes": ["P4", "P5", "H2", "H7", "EC3", "ED16"],
             "population_fields": POP, "water_fields": WATER,
             "population": {"country": p5_country, "maale": p5_maale,
                            "maale_parts": p5_maale_parts, "admin_subtotal": p5_admin,
                            "atolls": atolls, "islands": islands, "nonadmin": p5_nonadmin},
             "households": {"country": h2_country, "maale": h2_maale,
                            "maale_parts": h2_maale_parts, "admin_subtotal": h2_admin,
                            "islands": h2_islands, "nonadmin": h2_nonadmin},
             "water": {"country": w_country, "maale": w_maale,
                       "all_atolls_subtotal": w_all_atolls, "admin_subtotal": w_admin,
                       "atolls": w_atolls, "nonadmin": w_nonadmin},
             "employment": emp, "literacy": lit,
             "cautions": ["P4 HDh includes an overlapping Kulhudhuffushi City subrow; only coded atoll total used",
                          "P5 Harbours row G17 Maldivian female is blank; implied arithmetic zero is not adopted as observed",
                          "2022 census atoll reporting scope is not a 2026 atoll council jurisdiction",
                          "Two EC3 and two ED16 source island names differ from P5/H2; values held pending crosswalk",
                          "H7 source safe/unsafe categories are not WHO/JMP service classifications",
                          "Unassessed original columns and 52 other acquired workbook bodies remain priority_unassessed"]}
    (project / "evidence/MDV_CPH2022_SELECTED_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"atolls": len(atolls), "administrative_islands": len(islands),
                      "maale_parts": len(p5_maale_parts), "population": p5_country["fields"]["population"],
                      "households": h2_country["fields"]["households"],
                      "atoll_household_sum": w_admin["fields"]["households"],
                      "employment_islands_matched": len(emp["islands"]),
                      "literacy_islands_matched": len(lit["islands"]),
                      "unmatched_names_each": len(emp["held_rows"])}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
