#!/usr/bin/env python3
"""Build a source-faithful bundle for Panama's four remaining themes.

The four official products have different publication geographies.  This
collector keeps those geographies explicit instead of filling lower areas from
national or health-region values:

* Census 2023 urban-place table: country through corregimiento;
* Census 2023 recent interprovincial migration: country and province/comarca;
* MINSA under-five malnutrition: country total only;
* MIDES multidimensional-poverty incidence: all 699 corregimientos.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

INEC_CENSUS_PAGE = "https://www.inec.gob.pa/publicaciones/Default3.aspx?ID_CATEGORIA=19&ID_PUBLICACION=1231"
INEC_URBAN_XLSX = "https://www.inec.gob.pa/archivos/P053342420240202111608Cuadro%202.xlsx"
INEC_MIGRATION_PAGE = "https://www.inec.gob.pa/publicaciones/Default3.aspx?ID_CATEGORIA=19&ID_PUBLICACION=1233&ID_SUBCATEGO="
MINSA_PAGE = "https://www.minsa.gob.pa/node/34746"
MIDES_PDF = "https://www.mides.gob.pa/wp-content/uploads/2024/05/Informe-IPM2.pdf"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).upper()
    text = re.sub(r"[^A-Z0-9]+", " ", text).strip()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+P$", "", text)
    text = re.sub(r"\s+CABECERA$", "", text)
    aliases = {
        "NANCE DEL RISCO": "NANCE DE RISCO",
        "VALLE DEL RISCO": "VALLE DE RISCO",
        "JADABERI": "JADEBERI",
        "FEUILLET": "FEULLIET",
        "CERRO DE PLATA": "CERRO PLATA",
        "CALIDONIA O LA EXPOSICION": "LA EXPOSICION O CALIDONIA",
        "COMARCA KUNA YALA": "KUNA YALA",
        "COMARCA EMBERA": "EMBERA",
        "COMARCA NGABE BUGLE": "NGABE BUGLE",
        "PANAMA OESTE 1": "PANAMA OESTE",
    }
    return aliases.get(text, text)


def number(value: object) -> float | None:
    if value in (None, "", "-"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_panama_areas(dataset_path: Path):
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    areas = {row["id"]: row for row in dataset["territories"] if row["id"] == "PAN" or row.get("country_id") == "PAN"}
    if len(areas) != 795:
        raise RuntimeError(f"Expected PAN plus 794 subnational areas; found {len(areas)}")
    populations = {
        row["territory_id"]: float(row["value"])
        for row in dataset["observations"]
        if row["indicator_id"] == "PAN_C2023_POP_TOTAL" and row.get("value") is not None
    }
    if set(populations) != set(areas):
        raise RuntimeError(f"Population denominator coverage mismatch: {len(populations)}/{len(areas)}")
    children = defaultdict(list)
    for area in areas.values():
        if area["id"] != "PAN":
            children[area["parent_id"]].append(area["id"])
    return dataset, areas, populations, children


def area_keys(areas: dict[str, dict]):
    keys = {"country": {"TOTAL": "PAN"}, "province": {}, "district": {}, "corregimiento": {}}
    for area in areas.values():
        if area["id"] == "PAN":
            continue
        if area["level"] == "province":
            keys["province"][norm(area["name"])] = area["id"]
        elif area["level"] == "district":
            parent = areas[area["parent_id"]]
            keys["district"][f"{norm(parent['name'])}|{norm(area['name'])}"] = area["id"]
        elif area["level"] == "corregimiento":
            district = areas[area["parent_id"]]
            province = areas[district["parent_id"]]
            keys["corregimiento"][f"{norm(province['name'])}|{norm(district['name'])}|{norm(area['name'])}"] = area["id"]
    return keys


def parse_urban(path: Path, areas: dict, populations: dict, children: dict):
    ws = load_workbook(path, read_only=True, data_only=True).active
    keys = area_keys(areas)
    urban = {area_id: 0.0 for area_id in areas}
    seen = set()
    province = district = None
    unmatched = []
    for row_no, values in enumerate(ws.iter_rows(min_row=9, values_only=True), 9):
        values = list(values)
        total = number(values[17] if len(values) > 17 else None)
        if total is None:
            continue
        a, b, c = values[:3]
        area_id = None
        if norm(a) == "TOTAL" and not norm(b) and not norm(c):
            area_id = "PAN"
        elif norm(a) and not norm(b) and not norm(c):
            province, district = norm(a), None
            area_id = keys["province"].get(province)
        elif norm(b) and not norm(a) and not norm(c):
            district = norm(b)
            area_id = keys["district"].get(f"{province}|{district}")
        elif norm(c) and not norm(a) and not norm(b):
            area_id = keys["corregimiento"].get(f"{province}|{district}|{norm(c)}")
        else:
            continue
        if not area_id:
            unmatched.append({"row": row_no, "a": a, "b": b, "c": c, "value": total})
            continue
        if area_id in seen:
            raise RuntimeError(f"Duplicate urban aggregate row for {area_id} at workbook row {row_no}")
        urban[area_id] = total
        seen.add(area_id)
    if unmatched:
        raise RuntimeError(f"Unmatched urban aggregate rows: {json.dumps(unmatched[:20], ensure_ascii=False)}")
    # The official table is exhaustive for urban places.  Territories without
    # an aggregate row contain zero urban places; every published parent total
    # must still reconcile to its published children.
    for parent_id, member_ids in children.items():
        child_sum = sum(urban[child] for child in member_ids)
        if abs(urban[parent_id] - child_sum) > 0.01:
            raise RuntimeError(f"Urban hierarchy does not reconcile for {parent_id}: {urban[parent_id]} != {child_sum}")
    if urban["PAN"] != 2675550:
        raise RuntimeError(f"Unexpected national urban population: {urban['PAN']}")
    observations = []
    for area_id in sorted(areas):
        pct = round(urban[area_id] / populations[area_id] * 100, 4) if populations[area_id] else None
        observations.append(observation(area_id, "PAN_C2023_URBAN_POP_PCT", "2023", pct, "pan-inec-census-2023-urban-places", "%", "Population living in urban places as a share of Census 2023 population.", "derived_from_complete_official_enumeration", "Cuadro 2, urban population total / Census 2023 total population"))
    return observations, {"published_aggregate_rows": len(seen), "zero_urban_territories": len(areas) - len(seen), "national_urban_population": urban["PAN"], "all_hierarchy_totals_reconciled": True}


def parse_migration(path: Path, areas: dict):
    ws = load_workbook(path, read_only=True, data_only=True).active
    province_by_name = {norm(row["name"]): row["id"] for row in areas.values() if row.get("level") == "province"}
    expected = {"TOTAL": "PAN", **province_by_name}
    found = {}
    rows = {}
    for row_no, values in enumerate(ws.iter_rows(min_row=7, values_only=True), 7):
        values = list(values)
        label_raw = values[1] if len(values) > 1 else None
        label = norm(label_raw)
        if label not in expected or (isinstance(label_raw, str) and label_raw != label_raw.lstrip()):
            continue
        total, undeclared, foreign = number(values[2]), number(values[16]), number(values[17])
        if total is None:
            continue
        value = total - (undeclared or 0) - (foreign or 0)
        area_id = expected[label]
        if area_id in found:
            raise RuntimeError(f"Duplicate migration summary for {area_id}")
        found[area_id] = value
        rows[area_id] = row_no
    if set(found) != set(expected.values()):
        missing = sorted(set(expected.values()) - set(found))
        raise RuntimeError(f"Migration geography coverage mismatch; missing {missing}")
    if abs(found["PAN"] - sum(value for area_id, value in found.items() if area_id != "PAN")) > 0.01:
        raise RuntimeError("Migration province/comarca values do not reconcile to the country total")
    observations = [observation(area_id, "PAN_C2023_RECENT_INTERPROVINCIAL_MIGRANTS_2018_2023", "2018-2023", value, "pan-inec-census-2023-migration-volume-3", "people", "People whose usual province/comarca in 2023 differed from their province/comarca of residence five years earlier; foreign and undeclared previous residence excluded.", "derived_from_source_counts", f"Cuadro 2 row {rows[area_id]}: total minus no declarada and extranjero") for area_id, value in sorted(found.items())]
    return observations, {"geographies": len(found), "country_value": found["PAN"], "province_comarca_sum_reconciled": True}


def parse_nutrition(path: Path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    row = [cell.value for cell in ws[5]]
    numerator, denominator, value = number(row[1]), number(row[2]), number(row[3])
    if numerator != 4604 or denominator != 368659 or value is None:
        raise RuntimeError(f"Unexpected MINSA total row: {row}")
    calculated = numerator / denominator * 100
    if abs(value - calculated) > 1e-9:
        raise RuntimeError("MINSA nutrition prevalence does not reconcile to numerator and denominator")
    obs = observation("PAN", "PAN_MINSA_UNDER5_MALNUTRITION_PCT_2022", "2022", round(value, 4), "pan-minsa-health-statistics-2022-under5-malnutrition", "%", "Prevalence of malnutrition among the estimated population under age five.", "source_reported", "Cuadro 47, total row")
    return [obs], {"national_numerator": int(numerator), "national_denominator": int(denominator), "national_prevalence_pct": value, "health_regions_not_crosswalked_to_administrative_areas": True}


POVERTY_ALIASES = {
    "BOCAS DEL DRAGO": "BOCA DEL DRAGO",
    "NANCE DE RISCO": "NANCE DEL RISCO",
    "VALLE DE RISCO": "VALLE DEL RISCO",
    "CERRO PLATA": "CERRO DE PLATA",
    "JADEBERI": "JADABERI",
    "FEULLIET": "FEUILLET",
    "EL PIRO NO 2 MUAKWATA KUBU": "EL PIRO NO 2",
    "SANTA CATALINA O CALOVEBORA BLEDESHIA": "SANTA CATALINA O CALOVEVORA",
}


def poverty_name(value: str) -> str:
    normalized = norm(value)
    if normalized == "KUNA YALA":
        return "COMARCA KUNA YALA"
    return POVERTY_ALIASES.get(normalized, normalized)


def parse_poverty(text_path: Path, areas: dict, populations: dict):
    raw_lines = text_path.read_text(encoding="utf-8", errors="replace").splitlines()
    lines = [norm(line) for line in raw_lines]
    values = {}
    evidence = {}
    unmatched = []
    duplicates = {}
    for area in sorted((row for row in areas.values() if row.get("level") == "corregimiento"), key=lambda row: row["id"]):
        district = areas[area["parent_id"]]
        district_name = poverty_name(district["name"])
        corr_name = poverty_name(area["name"])
        # The report omits the long alternative district name and uses only the
        # official short form for its cabecera.
        if district_name == "SANTA CATALINA O CALOVEVORA":
            if corr_name.startswith("SANTA CATALINA "):
                corr_name = "SANTA CATALINA"
        prefix = f"{district_name} {corr_name}"
        candidates = []
        for index, line in enumerate(lines):
            if not re.match(rf"^{re.escape(prefix)}(?:\s+CABECERA)?(?:\s|$)", line):
                continue
            remainder = re.sub(rf"^{re.escape(prefix)}(?:\s+CABECERA)?\s*", "", line)
            nums = [float(token) for token in re.findall(r"-?\d+(?:\.\d+)?", remainder)]
            if len(nums) >= 2:
                candidates.append((index + 1, nums[0], nums[1]))
        # A single PDF row is split after the population value.
        if not candidates and area["id"] == "PAN:C2023:CORR:070301":
            for index, line in enumerate(lines[:-1]):
                if line == "LOS SANTOS 9676" and lines[index + 1].startswith("LA VILLA DE LOS SANTOS CABECERA"):
                    nums = [float(token) for token in re.findall(r"-?\d+(?:\.\d+)?", lines[index + 1].replace("LA VILLA DE LOS SANTOS CABECERA", "", 1))]
                    if nums:
                        candidates.append((index + 2, 9676.0, nums[0]))
        if not candidates:
            unmatched.append({"territory_id": area["id"], "district": district["name"], "corregimiento": area["name"], "prefix": prefix})
            continue
        # Repeated PDF headers/layout can duplicate a row.  Select the record
        # whose report population is closest to the full Census total and
        # require all plausible copies to agree on incidence.
        candidates.sort(key=lambda item: abs(item[1] - populations[area["id"]]))
        selected = candidates[0]
        close = [item for item in candidates if item[1] == selected[1]]
        if len({item[2] for item in close}) != 1:
            raise RuntimeError(f"Conflicting poverty incidence copies for {area['id']}: {close}")
        values[area["id"]] = selected[2]
        evidence[area["id"]] = {"text_line": selected[0], "report_population": selected[1], "census_population": populations[area["id"]], "incidence_pct": selected[2], "candidate_count": len(candidates)}
        if len(candidates) > 1:
            duplicates[area["id"]] = candidates
    if unmatched or len(values) != 699:
        raise RuntimeError(f"Poverty rows matched {len(values)}/699; unmatched={json.dumps(unmatched, ensure_ascii=False)}")
    if any(value < 0 or value > 100 for value in values.values()):
        raise RuntimeError("Poverty incidence outside 0-100 range")
    observations = [observation(area_id, "PAN_MIDES_MPI_INCIDENCE_PCT_2023", "2023", value, "pan-mides-mpi-corregimiento-2024", "%", "Incidence (H) of multidimensional poverty among the population with complete deprivation information, using Census 2023 data.", "source_reported", f"IPM por corregimiento 2024, normalized text line {evidence[area_id]['text_line']}") for area_id, value in sorted(values.items())]
    return observations, {"matched_corregimientos": len(values), "duplicate_layout_rows": len(duplicates), "min_incidence_pct": min(values.values()), "max_incidence_pct": max(values.values()), "crosswalk": evidence}


def observation(area_id, indicator_id, period, value, source_id, unit, definition, method, locator):
    return {"territory_id": area_id, "indicator_id": indicator_id, "period": period, "value": value, "status": "observed" if value is not None else "missing", "source_id": source_id, "definition": definition, "definition_id": indicator_id, "unit": unit, "population": definition, "measurement_method": method, "source_locator": locator}


def indicator(indicator_id, name, theme, unit, definition, source_id, locator, method="source_reported"):
    family = "administrative" if source_id.startswith("pan-minsa-") else "census"
    return {"id": indicator_id, "name": name, "theme": theme, "unit": unit, "definition": definition, "definition_id": indicator_id, "population": definition, "measurement_method": method, "aggregation": "none", "period_policy": "fixed_source_period", "series_family": family, "display_role": "primary", "source_id": source_id, "source_locator": locator}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", required=True)
    parser.add_argument("--urban", required=True)
    parser.add_argument("--migration", required=True)
    parser.add_argument("--nutrition", required=True)
    parser.add_argument("--poverty-pdf", required=True)
    parser.add_argument("--poverty-text", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    paths = {key: Path(value).resolve() for key, value in vars(args).items() if key not in {"out"}}
    out = Path(args.out).resolve(); out.mkdir(parents=True, exist_ok=True)
    _, areas, populations, children = load_panama_areas(paths["dataset"])
    urban_obs, urban_audit = parse_urban(paths["urban"], areas, populations, children)
    migration_obs, migration_audit = parse_migration(paths["migration"], areas)
    nutrition_obs, nutrition_audit = parse_nutrition(paths["nutrition"])
    poverty_obs, poverty_audit = parse_poverty(paths["poverty_text"], areas, populations)
    sources = [
        {"id": "pan-inec-census-2023-urban-places", "name": "Censo 2023, Cuadro 2: población de lugares poblados urbanos", "publisher": "Instituto Nacional de Estadística y Censo (INEC), Contraloría General de la República de Panamá", "url": INEC_URBAN_XLSX, "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2023", "geographic_level": "República, provincia/comarca, distrito y corregimiento", "source_locator": "Cuadro2_Barr, urban population total", "license": "Reuse terms not stated in the source publication; AreaData publishes source-attributed factual observations only.", "terms_review_status": "terms_not_stated_in_source", "raw_redistribution_status": "withheld_pending_terms_review", "note": "Official exhaustive table of urban places. Published aggregates reconcile through the complete Census 2023 hierarchy; an adopted area absent from the table therefore has zero urban-place population."},
        {"id": "pan-inec-census-2023-migration-volume-3", "name": "Censo 2023, Volumen III, Cuadro 2: migrantes interprovinciales 2018-23", "publisher": "Instituto Nacional de Estadística y Censo (INEC), Contraloría General de la República de Panamá", "url": INEC_MIGRATION_PAGE, "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2018-2023", "geographic_level": "República and provincia/comarca of usual residence", "source_locator": "Cuadro 2, summary rows", "license": "Reuse terms not stated in the source publication; AreaData publishes source-attributed factual observations only.", "terms_review_status": "terms_not_stated_in_source", "raw_redistribution_status": "withheld_pending_terms_review", "note": "Country and all 13 province/comarca summary rows retained. Foreign and undeclared previous residence are excluded from the displayed interprovincial count."},
        {"id": "pan-minsa-health-statistics-2022-under5-malnutrition", "name": "Estadísticas de Salud 2022, Cuadro 47: desnutrición en menores de 5 años", "publisher": "Ministerio de Salud de Panamá", "url": MINSA_PAGE, "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2022", "geographic_level": "República and health regions", "source_locator": "C47, total row", "license": "Reuse terms not stated in the source publication; AreaData publishes source-attributed factual observations only.", "terms_review_status": "terms_not_stated_in_source", "raw_redistribution_status": "withheld_pending_terms_review", "note": "Only the national total is integrated. Health regions are not silently equated to administrative provinces."},
        {"id": "pan-mides-mpi-corregimiento-2024", "name": "Índice de Pobreza Multidimensional por Corregimiento 2024", "publisher": "Ministerio de Desarrollo Social de Panamá", "url": MIDES_PDF, "status": "ready", "retrieved_at": "2026-09-18", "reference_period": "2023", "geographic_level": "699 corregimientos", "source_locator": "Incidencia H by district and corregimiento", "license": "Reuse terms not stated in the source publication; AreaData publishes source-attributed factual observations only.", "terms_review_status": "terms_not_stated_in_source", "raw_redistribution_status": "withheld_pending_terms_review", "note": "The 2024 publication uses Census 2023 data. All 699 rows were matched to the adopted Census geography; no parent rate is inferred."},
    ]
    indicators = [
        indicator("PAN_C2023_URBAN_POP_PCT", "Urban population", "Settlement", "%", "Population living in urban places as a share of Census 2023 population.", sources[0]["id"], "Cuadro 2 urban population / Census 2023 population", "derived_from_complete_official_enumeration"),
        indicator("PAN_C2023_RECENT_INTERPROVINCIAL_MIGRANTS_2018_2023", "Recent interprovincial migrants", "Migration", "people", "People whose 2023 province/comarca differed from five years earlier, excluding foreign and undeclared previous residence.", sources[1]["id"], "Cuadro 2 total minus no declarada and extranjero", "derived_from_source_counts"),
        indicator("PAN_MINSA_UNDER5_MALNUTRITION_PCT_2022", "Malnutrition among children under five", "Nutrition", "%", "Prevalence of malnutrition among the estimated population under age five.", sources[2]["id"], "Cuadro 47, total row"),
        indicator("PAN_MIDES_MPI_INCIDENCE_PCT_2023", "Multidimensional poverty incidence", "Poverty", "%", "Incidence (H) of multidimensional poverty among the population with complete deprivation information, using Census 2023 data.", sources[3]["id"], "IPM por corregimiento 2024, incidence H"),
    ]
    receipts = [{"label": key, "path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)} for key, path in paths.items() if key != "dataset"]
    bundle = {"schema_version": "1.0", "country_area_id": "PAN", "indicators": indicators, "observations": urban_obs + migration_obs + nutrition_obs + poverty_obs, "sources": sources, "audit": {"schema_version": "1.0", "generated_at": datetime.now(timezone.utc).isoformat(), "country_area_id": "PAN", "status": "four_required_themes_integrated", "themes": {"urban_rural": urban_audit, "migration": migration_audit, "nutrition": nutrition_audit, "poverty": poverty_audit}, "counts": {"urban_observations": len(urban_obs), "migration_observations": len(migration_obs), "nutrition_observations": len(nutrition_obs), "poverty_observations": len(poverty_obs), "total_observations": len(urban_obs + migration_obs + nutrition_obs + poverty_obs)}, "receipts": receipts, "geography_policy": "Each indicator is published only at the exact official source geography. Health regions are not mapped to provinces; poverty is not aggregated to parents."}}
    (out / "panama-required-themes-bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "panama-required-themes-audit.json").write_text(json.dumps(bundle["audit"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"counts": bundle["audit"]["counts"], "urban": urban_audit, "migration": migration_audit, "nutrition": nutrition_audit, "poverty": {key: value for key, value in poverty_audit.items() if key != "crosswalk"}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
