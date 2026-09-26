"""Inventory the five not-yet-adopted sheets in TÜİK's 2025 ADNKS workbook.

Read-only input. This produces a private audit and makes no site/dataset changes.
It deliberately does not aggregate suppressed or omitted small-area records.
"""

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PINNED_SHA256 = "72e36cf8f1eeb2e8c12480e14148b42448c6aad05d19932d94add37f360980c1"
SPECS = {
    "BÜYÜKŞEHİR B. NÜFUSU": {"start": 7, "fields": {3: "total", 4: "male", 5: "female"}, "code": [1]},
    "BELEDİYE NÜFUSU": {"start": 7, "fields": {9: "total", 10: "male", 11: "female"}, "code": [2, 3, 4]},
    "MAHALLE NÜFUSU": {"start": 7, "fields": {11: "total", 12: "male", 13: "female"}, "code": [2, 3, 4, 5]},
    "KÖY NÜFUSU": {"start": 7, "fields": {8: "total", 9: "male", 10: "female"}, "code": [2, 3, 4]},
    "KENT-KIR SINIFLAMASI": {"start": 6, "fields": {}, "code": [2, 3, 4, 5]},
}


def cell(row, column):
    return row[column - 1] if column <= len(row) else None


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    actual_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    if actual_hash != PINNED_SHA256:
        raise ValueError(f"Workbook SHA-256 changed: {actual_hash}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if set(SPECS) - set(workbook.sheetnames):
        raise ValueError("Expected TÜİK sheets are missing")
    inventory = []
    neighborhood_keys = set()
    village_keys = set()
    classification_joins = []
    for title, spec in SPECS.items():
        sheet = workbook[title]
        counts = Counter()
        field_counts = {name: Counter() for name in spec["fields"].values()}
        code_set = set()
        duplicate_codes = 0
        examples = []
        for line_number, row in enumerate(sheet.iter_rows(min_row=spec["start"], values_only=True), spec["start"]):
            if not isinstance(cell(row, 1), (int, float)):
                if any(value is not None for value in row):
                    counts["non_data_tail_rows"] += 1
                continue
            counts["data_rows"] += 1
            code = tuple(cell(row, col) for col in spec["code"])
            if code in code_set:
                duplicate_codes += 1
            code_set.add(code)
            if title == "MAHALLE NÜFUSU":
                neighborhood_keys.add((cell(row, 2), cell(row, 3), cell(row, 5)))
            elif title == "KÖY NÜFUSU":
                village_keys.add((cell(row, 2), cell(row, 3), cell(row, 4)))
            elif title == "KENT-KIR SINIFLAMASI":
                classification_joins.append(((cell(row, 2), cell(row, 3), cell(row, 5)),
                                             (cell(row, 2), cell(row, 3), cell(row, 4))))
            if len(examples) < 2:
                examples.append({"row": line_number, "code": code,
                                 "name_cells": [str(value) for value in row if isinstance(value, str)][:4]})
            if not spec["fields"]:
                classification = cell(row, 13)
                field_counts.setdefault("classification", Counter())[str(classification)] += 1
                continue
            values = {}
            for column, name in spec["fields"].items():
                value = cell(row, column)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    field_counts[name]["numeric"] += 1
                    values[name] = value
                elif value is None:
                    field_counts[name]["blank"] += 1
                else:
                    field_counts[name][str(value)] += 1
            if all(name in values for name in ("total", "male", "female")):
                counts["sex_complete_rows"] += 1
                if values["male"] + values["female"] != values["total"]:
                    counts["sex_sum_conflict_rows"] += 1
            else:
                counts["sex_suppressed_or_missing_rows"] += 1
        expected_header = 5 if title == "KENT-KIR SINIFLAMASI" else 6
        headers = [str(value) if value is not None else "" for value in next(
            sheet.iter_rows(min_row=expected_header, max_row=expected_header, values_only=True))]
        method_row = 3 if title == "KENT-KIR SINIFLAMASI" else 4
        method = str(cell(next(sheet.iter_rows(min_row=method_row, max_row=method_row, values_only=True)), 1))
        inventory.append({"sheet": title, "max_row": sheet.max_row, "max_column": sheet.max_column,
                          "data_start_row": spec["start"], "headers": headers,
                          "numeric_field_columns": {name: column for column, name in spec["fields"].items()},
                          "key_columns": spec["code"], "counts": dict(counts),
                          "distinct_key_tuples": len(code_set), "duplicate_key_tuples": duplicate_codes,
                          "field_state_counts": {name: dict(states) for name, states in field_counts.items()},
                          "first_rows": examples, "method_note": method,
                          "disposition": "not_adopted_pending_geographic_and_planning_review"})
    matched_neighborhoods = set()
    matched_villages = set()
    ambiguous = unmatched = 0
    for neighborhood_key, village_key in classification_joins:
        is_neighborhood = neighborhood_key in neighborhood_keys
        is_village = village_key in village_keys
        if is_neighborhood:
            matched_neighborhoods.add(neighborhood_key)
        if is_village:
            matched_villages.add(village_key)
        ambiguous += is_neighborhood and is_village
        unmatched += not (is_neighborhood or is_village)
    joins = {"classified_rows": len(classification_joins),
             "neighborhood_population_rows": len(neighborhood_keys),
             "village_population_rows": len(village_keys),
             "matched_neighborhoods": len(matched_neighborhoods),
             "matched_villages": len(matched_villages),
             "ambiguous_rows": ambiguous, "unmatched_classification_rows": unmatched,
             "population_rows_missing_classification": len(neighborhood_keys - matched_neighborhoods) + len(village_keys - matched_villages)}
    if (joins["matched_neighborhoods"] != len(neighborhood_keys) or
            joins["matched_villages"] != len(village_keys) or ambiguous or unmatched):
        raise ValueError(f"Cross-sheet code join differs from expected full published-row match: {joins}")
    report = {"source": "TÜİK ADNKS 2025 favorite tables", "reference_date": "2025-12-31",
              "workbook_sha256": actual_hash, "checked_at": datetime.now(timezone.utc).isoformat(),
              "reviewed_sheets": inventory, "classification_code_join": joins,
              "cross_sheet_limits": [
                  "Metropolitan municipality, municipality, district, neighborhood and village have different legal and statistical roles; do not collapse into one hierarchy by name.",
                  "Neighborhoods with population <=10 and organized industrial zones are omitted; sex cells can be C-suppressed.",
                  "Villages without a registered residential address are omitted and small sex cells can be '-' suppressed.",
                  "The settlement classification uses degree of urbanisation grids. It is categorical and cannot be summed as population without an audited code-level join and coverage check.",
                  "The first two province/district sheets are already separately adopted. This audit does not reclassify them or claim the five remaining sheets are integrated.",
              ]}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({item["sheet"]: item["counts"] for item in inventory}, ensure_ascii=False))


if __name__ == "__main__":
    main()
