"""Import a checked subset of Qatar NPC Census 2020 tables into a country candidate.

The official PDF's pp. 43-47 supply the municipality/zone-number crosswalk.
The 2015 geoBoundaries reference shapes are deliberately not joined to 2020
census geographies. Source workbooks and the candidate are kept outside Git.
"""

import argparse
import hashlib
import json
import re
import shutil
import subprocess
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PREFIX = "QAT_NPC2020_"
TABLES = ("1", "2", "3", "7", "31", "131", "133", "137", "139", "143", "149", "152")
MUNICIPALITIES = ("Doha", "Al Rayyan", "Al Wakra", "Umm Slal",
                  "Al Khor and Al Thakhira", "Al Shamal", "Al Daayen", "Al Sheehaniya")
AGE_COLUMN_ORDER = ("Al Sheehaniya", "Al Daayen", "Al Shamal", "Al Khor and Al Thakhira",
                    "Umm Slal", "Al Wakra", "Al Rayyan", "Doha")
MISSING_ZONE_NAMES = {10: "Wadi Al Sail", 11: "Rumaila", 19: "Doha Port",
                      60: "Al Dafna", 62: "Lekhwair"}
AL_SHAMAL_STRATEGY_SHA256 = "4611beee6e67b13001cfc65cbcc0573cd5eb4c373500e2bb95169b9cb0c3e111"


def check(condition, detail):
    if not condition:
        raise ValueError(detail)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def count(ws, row, col):
    value = ws.cell(row, col).value
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    check(isinstance(value, int) and not isinstance(value, bool) and value >= 0,
          f"Expected nonnegative count at {ws.title}!{get_column_letter(col)}{row}: {value!r}")
    return value


def slug(value):
    return re.sub(r"-+", "-", re.sub(r"[^A-Z0-9]+", "-", value.upper())).strip("-")


def municipality_id(name):
    return f"QAT:NPC2020:MUNICIPALITY:{slug(name)}"


def zone_id(code):
    return f"QAT:NPC2020:ZONE:{code:02d}"


def sid(table):
    return f"qat-npc-census2020-table-{table}"


def read_pdf_crosswalk(path):
    command = shutil.which("pdftotext") or "C:/Program Files/Tesseract-OCR/poppler/bin/pdftotext.exe"
    process = subprocess.run([command, "-layout", str(path), "-"], capture_output=True, check=True)
    raw = process.stdout.decode("utf-8")
    first = raw.index("Administrative Division of the State")
    start = raw.index("Administrative Division of the State", first + 1)
    end = raw.index(" Note:", start)
    section = raw[start:end]
    groups = defaultdict(list)
    current = None
    for line in section.splitlines():
        heading = re.search(r"•\s*(.+?)\s+Municipality:", line)
        if heading:
            current = heading.group(1).strip()
            continue
        if current is None:
            continue
        match = re.match(r"^\s*(.*?)\s{4,}(\d{1,2})(?:\s|$)", line)
        # Four long labels wrap onto the preceding line, leaving only the code.
        # Empty page-number lines are excluded explicitly.
        if match and (match.group(1).strip() or int(match.group(2)) in (55, 56, 70, 71)):
            groups[current].append(int(match.group(2)))
    check(set(groups) == set(MUNICIPALITIES), "PDF municipality headings changed")
    all_codes = [code for codes in groups.values() for code in codes]
    check(len(all_codes) == len(set(all_codes)) == 92, "PDF zone-code crosswalk changed")
    check({name: len(groups[name]) for name in MUNICIPALITIES} ==
          {"Doha": 59, "Al Rayyan": 10, "Al Wakra": 7, "Umm Slal": 1,
           "Al Khor and Al Thakhira": 3, "Al Shamal": 3, "Al Daayen": 2,
           "Al Sheehaniya": 7}, "PDF zone group sizes changed")
    return dict(groups)


def read_geographies(wb, groups):
    ws1, ws2 = wb["1"], wb["2"]
    check(ws1.cell(3, 1).value == "Population by Sex and Municipality", "Table 1 title changed")
    check(ws2.cell(3, 1).value == "Population by Sex and Zone", "Table 2 title changed")
    names = ("Total",) + MUNICIPALITIES
    municipality = {}
    for row, name in enumerate(names, 8):
        check(ws1.cell(row, 1).value == name, f"Table 1 geography row {row}")
        female, male, total = [count(ws1, row, col) for col in (2, 3, 4)]
        check(total == female + male, f"Table 1 sex sum {name}")
        municipality[name] = {"row": row, "female": female, "male": male, "total": total}
    for key in ("female", "male", "total"):
        check(municipality["Total"][key] == sum(municipality[n][key] for n in MUNICIPALITIES),
              f"Table 1 national {key}")
    check(municipality["Total"]["total"] == 2846118, "Census 2020 national total changed")
    zones = {}
    for row in range(10, ws2.max_row + 1):
        raw_code = ws2.cell(row, 1).value
        if raw_code is None:
            continue
        code = int(raw_code) if isinstance(raw_code, int) or str(raw_code).strip().isdigit() else None
        if code is None:
            continue
        check(code not in zones and code == int(ws2.cell(row, 7).value), f"Table 2 zone code {row}")
        female, male, total = [count(ws2, row, col) for col in (3, 4, 5)]
        check(total == female + male, f"Table 2 zone sex sum {code}")
        zones[code] = {"row": row, "name": ws2.cell(row, 2).value,
                       "female": female, "male": male, "total": total}
    all_codes = {code for codes in groups.values() for code in codes}
    check(len(zones) == 87 and set(zones).issubset(all_codes), "Table 2 populated-zone coverage changed")
    check(all_codes - set(zones) == set(MISSING_ZONE_NAMES), "Unreported zone set changed")
    check(ws2.cell(9, 5).value == municipality["Total"]["total"], "Table 2 national count")
    for name, codes in groups.items():
        for field in ("female", "male", "total"):
            check(sum(zones[code][field] for code in codes if code in zones) == municipality[name][field],
                  f"Zone-to-municipality {name} {field}")
    return municipality, zones


def read_age(ws, municipality):
    check(ws.cell(3, 1).value == "Population by Municipality & Age Group", "Table 3 title")
    expected = ("Under 1", "1 - 4", "5 - 9", "10 - 14", "15 - 19", "20 - 24", "25 - 29",
                "30 - 34", "35 - 39", "40 - 44", "45 - 49", "50 - 54", "55 - 59", "60 - 64",
                "65 - 69", "70 - 74", "+ 75")
    result = {}
    for col, name in [(10, "Total"), *[(i + 2, n) for i, n in enumerate(AGE_COLUMN_ORDER)]]:
        check(count(ws, 7, col) == municipality[name]["total"], f"Age total {name}")
        values = []
        for row, label in enumerate(expected, 8):
            check(ws.cell(row, 1).value == label, f"Age label row {row}")
            values.append(count(ws, row, col))
        check(sum(values) == municipality[name]["total"], f"Age bands sum {name}")
        result[name] = {"column": col, "under15": sum(values[:4]),
                        "age15_64": sum(values[4:14]), "age65plus": sum(values[14:])}
    return result


def read_municipal_rows(ws, start, name_col=1):
    result = {}
    for offset, name in enumerate(("Total",) + MUNICIPALITIES):
        row = start + offset
        check(ws.cell(row, name_col).value in (name, "Total of Households") if name == "Total"
              else ws.cell(row, name_col).value == name, f"{ws.title} municipality {name} row {row}")
        result[name] = row
    return result


def main(project):
    inventory = json.loads((project / "evidence/QAT_NPC2020_WORKBOOK_INVENTORY.json").read_text(encoding="utf-8"))
    root = project / "raw/npc-census2020"
    main_xlsx, other_xlsx, pdf = (root / filename for filename in
                                  ("Census_Final_Results.xlsx", "Qatar_Census_2022_Final_Results.xlsx",
                                   "Census_Final_Results.pdf"))
    check(digest(main_xlsx) == inventory["originals"][main_xlsx.name]["sha256"], "Main XLSX changed")
    check(digest(other_xlsx) == inventory["originals"][other_xlsx.name]["sha256"], "Alternate XLSX changed")
    check(digest(pdf) == inventory["pdf"]["sha256"], "Administrative PDF changed")
    strategy = root / "Al_Shamal_Dec_2017.pdf"
    check(strategy.is_file() and strategy.read_bytes().startswith(b"%PDF") and
          digest(strategy) == AL_SHAMAL_STRATEGY_SHA256, "Historical Al Shamal strategy PDF missing or changed")
    check(inventory["comparison"]["changed_cell_value_sheets"] == [] and
          inventory["comparison"]["common_sheets"] == 206, "Workbook variants diverge")
    groups = read_pdf_crosswalk(pdf)
    wb = load_workbook(main_xlsx, read_only=True, data_only=True)
    try:
        municipality, zones = read_geographies(wb, groups)
        age = read_age(wb["3"], municipality)
        rows7 = read_municipal_rows(wb["7"], 6)
        rows131 = read_municipal_rows(wb["131"], 7)
        rows133 = read_municipal_rows(wb["133"], 7)
        rows137 = read_municipal_rows(wb["137"], 7)
        rows139 = read_municipal_rows(wb["139"], 8)
        rows143 = read_municipal_rows(wb["143"], 7)
        rows149 = read_municipal_rows(wb["149"], 10)
        source_values = defaultdict(dict)

        def save(name, suffix, value, table, row, col, note=None):
            source_values[name][suffix] = {"value": value, "table": table, "row": row,
                                            "column": col, "note": note}

        for name in ("Total",) + MUNICIPALITIES:
            m = municipality[name]
            for suffix, col, key in (("POP_TOTAL", 4, "total"), ("POP_MALE", 3, "male"),
                                     ("POP_FEMALE", 2, "female")):
                save(name, suffix, m[key], "1", m["row"], col)
            a = age[name]
            for suffix, key, row_span in (("AGE_0_14", "under15", "8-11"),
                                          ("AGE_15_64", "age15_64", "12-21"),
                                          ("AGE_65PLUS", "age65plus", "22-24")):
                save(name, suffix, a[key], "3", row_span, a["column"],
                     "AreaData sum of complete, non-overlapping census age bands; all bands reconcile to Table 1.")
            ws = wb["7"]
            row = rows7[name]
            households, residents = count(ws, row, 4), count(ws, row, 3)
            mean = ws.cell(row, 2).value
            check(isinstance(mean, (int, float)) and abs(mean - residents / households) < 1e-9,
                  f"Table 7 average household size {name}")
            check(residents <= m["total"], f"Household residents exceed population {name}")
            for suffix, value, col in (("HOUSEHOLDS", households, 4),
                                       ("HOUSEHOLD_RESIDENTS", residents, 3),
                                       ("HOUSEHOLD_AVERAGE_SIZE", mean, 2)):
                save(name, suffix, value, "7", row, col,
                     "NPC's published mean is retained; the interface displays two decimals." if col == 2 else None)
            ws = wb["131"]
            row = rows131[name]
            building = [count(ws, row, col) for col in range(2, 6)]
            check(building[3] == sum(building[:3]), f"Table 131 building status {name}")
            for suffix, col in (("BUILDINGS_TOTAL", 5), ("BUILDINGS_UNDER_CONSTRUCTION", 3),
                                ("BUILDINGS_COMPLETE", 4)):
                save(name, suffix, count(ws, row, col), "131", row, col)
            for table, rows, total_col, selected, kind in (
                ("133", rows133, 7, (("OCCUPIED_HOUSING_UNITS", 7),
                                           ("OCCUPIED_APARTMENTS", 5), ("OCCUPIED_VILLAS", 6)), "occupied units"),
                ("137", rows137, 7, (("HOUSEHOLDS_APARTMENT", 5), ("HOUSEHOLDS_VILLA", 6)), "households"),
                ("143", rows143, 7, (("HOUSING_UNITS_TOTAL", 7),
                                           ("HOUSING_APARTMENTS", 5), ("HOUSING_VILLAS", 6)), "housing units")):
                ws = wb[table]
                row = rows[name]
                check(count(ws, row, total_col) == sum(count(ws, row, col) for col in range(2, 7)),
                      f"Table {table} {kind} category subtotal {name}")
                if table == "137":
                    check(count(ws, row, 7) == households, f"Table 137 versus Table 7 households {name}")
                for suffix, col in selected:
                    save(name, suffix, count(ws, row, col), table, row, col)
            ws = wb["139"]
            row = rows139[name]
            sewer_no, sewer_yes, water, electric, residential = [count(ws, row, col) for col in range(2, 7)]
            check(sewer_no + sewer_yes == water == electric == residential,
                  f"Table 139 residential building facilities {name}")
            for suffix, col in (("RESIDENTIAL_BUILDINGS_FACILITY_BASE", 6),
                                ("RESIDENTIAL_BUILDINGS_SEWER_CONNECTED", 3),
                                ("RESIDENTIAL_BUILDINGS_SEWER_NOT_CONNECTED", 2)):
                save(name, suffix, count(ws, row, col), "139", row, col)
            ws = wb["149"]
            row = rows149[name]
            sector = [count(ws, row, col) for col in range(2, 10)]
            check(sector[4] == sum(sector[:4]) and sector[7] == sum(sector[4:7]),
                  f"Table 149 establishment sector subtotals {name}")
            save(name, "OPERATING_ESTABLISHMENTS", sector[7], "149", row, 9)
            save(name, "BUSINESS_ESTABLISHMENTS", sector[4], "149", row, 6)
            ws = wb["152"]
            erow = 8 if name == "Total" else 10 + 2 * MUNICIPALITIES.index(name)
            prow = erow + 1
            check(ws.cell(erow, 1).value == ("Total" if name == "Total" else name) and
                  ws.cell(erow, 2).value == "EST." and ws.cell(prow, 2).value == "P. E.",
                  f"Table 152 paired municipality rows {name}")
            for row in (erow, prow):
                check(count(ws, row, 10) == sum(count(ws, row, col) for col in range(3, 10)),
                      f"Table 152 establishment size subtotals {name} row {row}")
            check(count(ws, erow, 10) == sector[4], f"Table 152 versus 149 business count {name}")
            save(name, "BUSINESS_PERSONS_ENGAGED", count(ws, prow, 10), "152", prow, 10)

        ws = wb["31"]
        check(ws.cell(3, 1).value == "Population (10+) by Municipality, Sex and Educational Attainment",
              "Table 31 title")
        categories = (10, 13, 16, 19, 22, 25)
        for col, name in [(11, "Total"), *[(i + 3, n) for i, n in enumerate(AGE_COLUMN_ORDER)]]:
            total = count(ws, 7, col)
            check(total == sum(count(ws, row, col) for row in categories),
                  f"Table 31 six education classes {name}")
            check(total == municipality[name]["total"] -
                  sum(count(wb["3"], row, age[name]["column"]) for row in range(8, 11)),
                  f"Table 31 age 10+ universe {name}")
            for row in (7, *categories):
                check(count(ws, row, col) == count(ws, row + 1, col) + count(ws, row + 2, col),
                      f"Table 31 sex subtotal {name} row {row}")
            for suffix, row in (("EDU_10PLUS_TOTAL", 7), ("EDU_ILLITERATE", 10),
                                ("EDU_UNIVERSITY_PLUS", 25)):
                save(name, suffix, count(ws, row, col), "31", row, col)

        for code, item in zones.items():
            label = zone_id(code)
            for suffix, col, key in (("POP_TOTAL", 5, "total"), ("POP_MALE", 4, "male"),
                                     ("POP_FEMALE", 3, "female")):
                save(label, suffix, item[key], "2", item["row"], col)
    finally:
        wb.close()

    path = project / "data/dashboard.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    check(data["country"]["id"] == "QAT", "Expected Qatar country candidate")
    data["territories"] = [territory for territory in data["territories"] if territory["id"] == "QAT"]
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    data["indicators"] = [i for i in data["indicators"] if not i["id"].startswith(PREFIX)]
    data["observations"] = [o for o in data["observations"] if not o["indicator_id"].startswith(PREFIX)]
    data["sources"] = [s for s in data["sources"] if not s["id"].startswith("qat-npc-")
                       and not s["id"].startswith("qat-qnmp-")]
    data["country"]["geography_note"] = (
        "NPC/PSA General Census 2020 reports eight municipalities and 92 numbered zones in its PDF administrative "
        "division (pp. 43-47). XLSX Table 2 publishes population for 87 zones; five listed zones have no population "
        "row and remain missing, not zero. All 87 reported zone counts and sex totals exactly sum to their municipality "
        "rows in Table 1. Municipality identities and zone numbers are from the 2020 source, while municipality codes "
        "and compatible official polygons have not been acquired. The 2015 geoBoundaries shapes are unjoined. "
        "The published tables say December 2020; the census methodology describes a moving household reference time. "
        "WDI midyear estimates are a separate series.")
    data["territories"][0]["source_id"] = sid("1")
    data["territories"][0]["reconciliation_status"] = "Direct national row, sum of eight municipalities"
    for name in MUNICIPALITIES:
        mid = municipality_id(name)
        data["territories"].append({"id": mid, "name": name, "level": "municipality",
            "type": "NPC Census 2020 municipality, planning unit candidate",
            "parent_id": "QAT", "official_code": None,
            "code_system": "NPC Census 2020 municipality label; official municipality code not acquired",
            "boundary_version": None, "source_id": sid("1"),
            "reconciliation_status": "Table 1 direct row; complete 2020 zone subtotal checked; official polygon pending"})
        for code in groups[name]:
            data["territories"].append({"id": zone_id(code),
                "name": zones[code]["name"] if code in zones else MISSING_ZONE_NAMES[code],
                "level": "census_zone", "type": "NPC Census 2020 numbered zone for internal diagnosis",
                "parent_id": mid, "official_code": str(code),
                "code_system": "NPC Census 2020 administrative division PDF Zone No.",
                "boundary_version": None, "source_id": "qat-npc-census2020-zone-crosswalk",
                "reconciliation_status": "PDF pp. 43-47 official Zone No. and municipality; 2020 polygon not acquired; population row missing" if code not in zones else
                                         "PDF pp. 43-47 official Zone No. and municipality; XLSX Table 2 row checked; 2020 polygon not acquired"})
    data["analysis"]["comparisons"] = [{"parent_id": "QAT",
        "member_ids": [municipality_id(n) for n in MUNICIPALITIES],
        "label": "Eight Census 2020 municipalities", "source_ids": [sid("1"), "qat-npc-census2020-zone-crosswalk"],
        "membership_note": "All three population/sex fields reconcile to the directly reported national row. No polygon is joined."}]
    for name in MUNICIPALITIES:
        data["analysis"]["comparisons"].append({"parent_id": municipality_id(name),
            "member_ids": [zone_id(code) for code in groups[name]],
            "label": f"NPC Census 2020 zones of {name}",
            "source_ids": [sid("2"), "qat-npc-census2020-zone-crosswalk"],
            "membership_note": "The PDF lists all zones, including five without a population row nationally. Published zone population/sex subtotals reconcile to this municipality's direct Table 1 row; no missing row is filled as zero. Zones are internal diagnostic geography, not separate municipality planning units."})
    data["analysis"]["terminal_territory_ids"] = [zone_id(code) for codes in groups.values() for code in codes]
    data["analysis"]["default_indicator_id"] = PREFIX + "POP_TOTAL"
    data["analysis"]["latest_values_only"] = True
    data["analysis"].pop("population_context", None)

    specs = [
        ("POP_TOTAL", "Census population 2020", "Population", "people", "All persons in the published December 2020 census table", "1", "source_reported"),
        ("POP_MALE", "Census 2020 male population", "Population", "people", "Male persons in the same census geography", "1", "source_reported"),
        ("POP_FEMALE", "Census 2020 female population", "Population", "people", "Female persons in the same census geography", "1", "source_reported"),
        ("AGE_0_14", "Census population aged 0-14", "Population", "people", "All persons in the four source age bands below 15", "3", "areadata_calculated"),
        ("AGE_15_64", "Census population aged 15-64", "Population", "people", "All persons in the ten source age bands 15-64", "3", "areadata_calculated"),
        ("AGE_65PLUS", "Census population aged 65+", "Population", "people", "All persons in the three source age bands 65+", "3", "areadata_calculated"),
        ("HOUSEHOLDS", "Census households", "Households", "households", "Households in NPC Table 7, distinct from all population and all housing units", "7", "source_reported"),
        ("HOUSEHOLD_RESIDENTS", "Persons in census households", "Households", "people", "Persons living in households in Table 7; a subset of total census population", "7", "source_reported"),
        ("HOUSEHOLD_AVERAGE_SIZE", "Average persons per household", "Households", "persons per household", "NPC published household-resident count divided by household count, displayed to two decimals", "7", "source_reported"),
        ("EDU_10PLUS_TOTAL", "Census education-table population aged 10+", "Education", "people", "All persons aged 10+ in NPC Table 31; denominator for education categories", "31", "source_reported"),
        ("EDU_ILLITERATE", "Census illiterate persons aged 10+", "Education", "people", "NPC Table 31 illiterate category among persons aged 10+", "31", "source_reported"),
        ("EDU_UNIVERSITY_PLUS", "Census university and above, aged 10+", "Education", "people", "NPC Table 31 university-and-above category among persons aged 10+", "31", "source_reported"),
        ("BUILDINGS_TOTAL", "Census buildings, all statuses", "Housing", "buildings", "NPC Table 131 buildings: complete, under construction or under demolition", "131", "source_reported"),
        ("BUILDINGS_UNDER_CONSTRUCTION", "Census buildings under construction", "Housing", "buildings", "NPC Table 131 status under construction", "131", "source_reported"),
        ("BUILDINGS_COMPLETE", "Census complete buildings", "Housing", "buildings", "NPC Table 131 status complete; all building types", "131", "source_reported"),
        ("OCCUPIED_HOUSING_UNITS", "Occupied housing units", "Housing", "housing units", "NPC Table 133 occupied housing units, all unit types", "133", "source_reported"),
        ("OCCUPIED_APARTMENTS", "Occupied apartments", "Housing", "housing units", "NPC Table 133 occupied units of apartment type", "133", "source_reported"),
        ("OCCUPIED_VILLAS", "Occupied villas", "Housing", "housing units", "NPC Table 133 occupied units of villa type", "133", "source_reported"),
        ("HOUSEHOLDS_APARTMENT", "Households in apartments", "Households", "households", "NPC Table 137 households by housing-unit type: apartment", "137", "source_reported"),
        ("HOUSEHOLDS_VILLA", "Households in villas", "Households", "households", "NPC Table 137 households by housing-unit type: villa", "137", "source_reported"),
        ("RESIDENTIAL_BUILDINGS_FACILITY_BASE", "Residential buildings in facility table", "Infrastructure", "buildings", "NPC Table 139 completed residential buildings in connection-to-facilities table; its universe is not silently equated with all completed buildings", "139", "source_reported"),
        ("RESIDENTIAL_BUILDINGS_SEWER_CONNECTED", "Residential buildings connected to sewage", "Infrastructure", "buildings", "NPC Table 139 residential buildings connected to sewage in its own denominator", "139", "source_reported"),
        ("RESIDENTIAL_BUILDINGS_SEWER_NOT_CONNECTED", "Residential buildings without sewage connection", "Infrastructure", "buildings", "NPC Table 139 residential buildings not connected to sewage in its own denominator", "139", "source_reported"),
        ("HOUSING_UNITS_TOTAL", "Housing units, all occupancy states", "Housing", "housing units", "NPC Table 143 all housing units, distinct from occupied units", "143", "source_reported"),
        ("HOUSING_APARTMENTS", "Housing units: apartments", "Housing", "housing units", "NPC Table 143 all apartment housing units, regardless of occupancy", "143", "source_reported"),
        ("HOUSING_VILLAS", "Housing units: villas", "Housing", "housing units", "NPC Table 143 all villa housing units, regardless of occupancy", "143", "source_reported"),
        ("OPERATING_ESTABLISHMENTS", "Operating establishments", "Economy", "establishments", "NPC Table 149 all operating establishments including business and government/diplomatic sectors", "149", "source_reported"),
        ("BUSINESS_ESTABLISHMENTS", "Business-sector establishments", "Economy", "establishments", "NPC Table 149 business-sector subtotal; reconciled to Table 152 EST. total", "149", "source_reported"),
        ("BUSINESS_PERSONS_ENGAGED", "Persons engaged in business establishments", "Economy", "people", "NPC Table 152 P.E. count in business establishments; workplace-based, not employed residents", "152", "source_reported"),
    ]
    for suffix, name, theme, unit, definition, table, method in specs:
        indicator = {"id": PREFIX + suffix, "name": name, "theme": theme, "unit": unit,
            "definition": definition + ". Census tables labelled December 2020; household enumeration had a moving reference time.",
            "population": definition, "source_id": sid(table),
            "aggregation": "none", "measurement_method": method,
            "display_decimals": 2 if suffix == "HOUSEHOLD_AVERAGE_SIZE" else 0}
        if suffix == "POP_TOTAL":
            indicator["series_family"] = "census"
        data["indicators"].append(indicator)

    def emit(name, tid):
        for suffix, detail in source_values[name].items():
            table = detail["table"]
            locator = (f"NPC Census 2020 XLSX Table {table}, sheet '{table}', row {detail['row']}, "
                       f"column {get_column_letter(detail['column'])}")
            observation = {"territory_id": tid, "indicator_id": PREFIX + suffix, "period": "2020",
                "value": detail["value"], "status": "observed", "source_id": sid(table),
                "measurement_method": "areadata_calculated" if suffix.startswith("AGE_") else "source_reported",
                "source_locator": locator}
            if suffix.startswith("AGE_"):
                observation["provenance"] = "calculated"
            if detail["note"]:
                observation["footnote"] = detail["note"]
            data["observations"].append(observation)
    emit("Total", "QAT")
    for name in MUNICIPALITIES:
        emit(name, municipality_id(name))
    for code in zones:
        emit(zone_id(code), zone_id(code))

    retrieved_at = datetime.now(timezone.utc).isoformat()
    for table in TABLES:
        data["sources"].append({"id": sid(table),
            "name": f"NPC/PSA General Population, Housing and Establishments Census 2020, Table {table}",
            "url": inventory["originals"][main_xlsx.name]["url"],
            "publisher": "National Planning Council (successor to Planning and Statistics Authority)",
            "reference_period": "Published tables labelled December 2020; household enumeration moving reference time",
            "geographic_level": "national, eight municipalities" + (", 87 reported zones" if table == "2" else ""),
            "status": "partial", "license": "https://www.npc.qa/en/aboutus/pages/TermsOfUse.aspx",
            "raw_path": str(main_xlsx.relative_to(project)).replace("\\", "/"),
            "sha256": inventory["originals"][main_xlsx.name]["sha256"],
            "retrieved_at": retrieved_at,
            "note": f"Original XLSX workbook retained locally; sheet '{table}' has selected cells checked. Other worksheet contents remain unassessed. Attribute NPC when republishing."})
    data["sources"].append({"id": "qat-npc-census2020-zone-crosswalk",
        "name": "NPC/PSA Census 2020 administrative division, PDF pp. 43-47",
        "url": inventory["pdf"]["url"],
        "publisher": "National Planning Council (successor to Planning and Statistics Authority)",
        "reference_period": "Census 2020 administrative division", "geographic_level": "municipality and zone number",
        "status": "partial", "license": "https://www.npc.qa/en/aboutus/pages/TermsOfUse.aspx",
        "raw_path": str(pdf.relative_to(project)).replace("\\", "/"), "sha256": inventory["pdf"]["sha256"],
        "retrieved_at": retrieved_at,
        "note": "Pages 43-47 identify all 92 zone numbers and parent municipalities; only 87 have population rows in XLSX Table 2. No polygon adopted."})
    leads = [
        ("qat-npc-census2020-catalog", "NPC Census 2020 detailed results catalogue", "https://www.npc.qa/en/statistics/census2020/Pages/results/default.aspx", "206 common XLSX sheets inventoried; 12 source tables partially adopted, other 144 numbered tables unassessed."),
        ("qat-npc-census2020-alternate", "NPC alternate Census 2020 workbook URL", inventory["originals"][other_xlsx.name]["url"], "Filename contains 2022, but all 206 common sheet cell values match the fuller workbook and numbered tables state December 2020; not a separate 2022 census."),
        ("qat-npc-census2020-terms", "NPC website terms of use", "https://www.npc.qa/en/aboutus/pages/TermsOfUse.aspx", "Published content requires NPC attribution; raw redistribution and logo conditions to review before hosting."),
        ("qat-qnmp-msdp", "Qatar National Master Plan municipal MSDP directory", "https://www.mm.gov.qa/QatarMasterPlan/English/MSDP-Municipalities.aspx?panel=about", "Official municipality-level spatial-plan location; each plan body, version and current status require individual verification."),
        ("qat-qnmp-zoning", "QNMP municipal zoning rules page", "https://www.mm.gov.qa/QatarMasterPlan/English/msdp-zoning.aspx?panel=about", "Official page distinguishes six currently described zoning deployments from prospective Al Khor/Al Wakra extension; page date and operative rules need audit."),
        ("qat-qnmp-al-rayyan-2017", "Al Rayyan and Al Shahhaniya 2017 municipality spatial-plan strategy", "https://www.mm.gov.qa/QatarMasterPlan/Downloads-qnmp/MunicipalityStrategy/English/Al%20Raayan%20Dec%202017.pdf", "Historical combined-area strategic volume found; do not apply it as current 2020/2026 Al Rayyan-only plan without status and boundary audit."),
        ("qat-qnmp-al-shamal-2017", "Al Shamal Municipality Vision and Development Strategy (2017)", "https://www.mm.gov.qa/QatarMasterPlan/Downloads-qnmp/MunicipalityStrategy/English/Al%20Shamal%20Dec%20%202017.pdf", "Official 36-page historical strategic-context volume acquired. Sections 1.1-1.2 describe the three-part MSDP and intended review; present legal effect, zoning maps and subsequent revisions have not been verified."),
        ("qat-qnmp-qndf", "Qatar National Development Framework", "https://www.mm.gov.qa/QatarMasterPlan/English/qndf.aspx", "National spatial framework location; municipal application, legal force and current amendment to verify."),
        ("qat-npc-gis-atlas", "NPC/PSA digital census atlas", "https://gis.psa.gov.qa/TestQatarAtlas/Index", "Official map portal location only; no 2020 boundary/code dataset downloaded or adopted."),
    ]
    for ident, name, url, note in leads:
        source = {"id": ident, "name": name, "url": url,
            "publisher": "Ministry of Municipality" if ident.startswith("qat-qnmp-") else "National Planning Council",
            "reference_period": "2017 strategy edition" if ident == "qat-qnmp-al-shamal-2017" else "official location checked 2026-09-26",
            "geographic_level": "Al Shamal municipality" if ident == "qat-qnmp-al-shamal-2017" else "Qatar",
            "status": "not_collected", "license": "source_terms_review_required", "retrieved_at": retrieved_at,
            "note": note}
        acquired_file = {"qat-npc-census2020-catalog": "results.html",
                         "qat-npc-census2020-alternate": other_xlsx.name,
                         "qat-npc-census2020-terms": "TermsOfUse.html",
                         "qat-qnmp-al-shamal-2017": "Al_Shamal_Dec_2017.pdf"}.get(ident)
        if acquired_file and (root / acquired_file).exists():
            source.update({"status": "partial", "raw_path": str((root / acquired_file).relative_to(project)).replace("\\", "/"),
                           "sha256": digest(root / acquired_file)})
        data["sources"].append(source)
    data["documents"] = [document for document in data["documents"]
                         if document["id"] != "qat-qnmp-al-shamal-strategy-2017"]
    data["documents"].append({
        "id": "qat-qnmp-al-shamal-strategy-2017", "territory_id": municipality_id("Al Shamal"),
        "category": "reference", "kind": "historical-msdp-strategy-volume",
        "title": "Historical Al Shamal Municipality Vision and Development Strategy (December 2017)",
        "url": "https://www.mm.gov.qa/QatarMasterPlan/Downloads-qnmp/MunicipalityStrategy/English/Al%20Shamal%20Dec%20%202017.pdf",
        "source_id": "qat-qnmp-al-shamal-2017", "period": "2017 edition; strategic horizon to 2032",
        "target_period": {"label": "2017 edition; strategic horizon to 2032", "kind": "multi_year"},
        "availability": "body_acquired", "official_status": "unverified",
        "note": "Only the strategic-context volume was acquired. Current legal effect, later review, zoning regulations/maps and any 2020 zone alignment are unverified; this is not a verified complete current MSDP."})
    data["gaps"] = [g for g in data["gaps"] if g["category"] not in
                    ("subnational_statistics", "boundary_reconciliation", "planning_documents")]
    data["gaps"] += [
        {"category": "subnational_statistics", "status": "partial",
         "detail": "NPC Census 2020 original XLSX and PDF acquired. All 156 numbered tables plus a blank index sheet and supporting sheets/charts mechanically inventoried; 12 tables partially adopted with 29 indicators. Remaining 144 numbered tables require semantic review. Five of 92 official zones have no population row; this is missing, not zero. WDI midyear estimates remain separate.",
         "next_action": "Classify remaining source tables and definitions, seek current municipality/zone series for health, education, labour and public services, and verify publication terms."},
        {"category": "boundary_reconciliation", "status": "partial",
         "detail": "Official PDF Zone No. to municipality crosswalk and 87 populated-zone XLSX row IDs reconcile. Five listed zones have no population row. No municipality code edition or 2020-compatible official polygons; 2015 geoBoundaries reference shapes are unjoined.",
         "next_action": "Acquire NPC/Ministry of Municipality 2020 municipality and zone geometry with license/version and verify code-to-polygon join."},
        {"category": "planning_documents", "status": "partial",
         "detail": "Ministry of Municipality QNMP directory located and the official Al Shamal December 2017 strategic-context volume acquired as a historical reference. No current complete MSDP, revisions, zoning maps, approvals, budget, implementation or evaluation content adopted; no document collected for a selected zone.",
         "next_action": "Audit current QNDF/MSDP legal status and all eight municipal plan bodies, versions, zone maps and actual budget/implementation materials before planning outputs are claimed."},
    ]
    data["collection"]["adapters"] = sorted(set(data["collection"]["adapters"] +
                                                ["qatar-npc-census2020-twelve-table-partial-v1"]))
    data["generated_at"] = datetime.now(timezone.utc).isoformat()
    adopted = [o for o in data["observations"] if o["indicator_id"].startswith(PREFIX)]
    check(len(data["territories"]) == 101 and len(specs) == 29 and len(adopted) == 522,
          f"Import count drift: territories {len(data['territories'])}, indicators {len(specs)}, observations {len(adopted)}")
    check(len({(o["territory_id"], o["indicator_id"], o["period"]) for o in adopted}) == len(adopted),
          "Duplicate adopted observation")
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit = {"selected_tables": list(TABLES), "workbook_sha256": digest(main_xlsx),
             "pdf_sha256": digest(pdf), "crosswalk": groups,
             "missing_population_zone_codes": sorted(MISSING_ZONE_NAMES),
             "municipality_population": {name: municipality[name]["total"] for name in MUNICIPALITIES},
             "adopted_indicator_ids": [PREFIX + spec[0] for spec in specs],
             "adopted_observations": len(adopted),
             "unassessed_numbered_tables": [str(i) for i in range(1, 157) if str(i) not in TABLES],
             "warning": "Numeric inventory and selected-table arithmetic do not establish full semantic acceptance, current plans or geography polygons."}
    (project / "evidence/QAT_NPC2020_ADOPTION_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"territories": len(data["territories"]), "municipalities": 8,
                      "zones": 92, "zones_with_population": len(zones), "indicators": len(specs),
                      "observations": len(adopted), "national_population": municipality["Total"]["total"]},
                     ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
