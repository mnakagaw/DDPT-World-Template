#!/usr/bin/env python3
"""Build a private, historical 2011 India PCA candidate from saved ORGI originals.

The 2011 Census codes are not current LGD codes or present planning authorities.
ORGI portal reproduction permission has not been obtained; do not publish this
candidate or its raw workbook without a separate rights decision.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / "generated" / "india-areadata-20260927"
RAW = PROJECT / "raw" / "official-census"
EVIDENCE = PROJECT / "evidence"
DATA = PROJECT / "data" / "dashboard.json"
BOOTSTRAP = PROJECT / "reference" / "dashboard-bootstrap.json"
PCA = RAW / "2011-IndiaStateDistSbDistTwn-0000.xlsx"
LCD = RAW / "PC11_TV_DIR.xlsx"
GUIDE = PROJECT / "raw" / "official-planning" / "mopr-gpdp-booklet-2026-27.pdf"
SOURCE_ID = "orgi-pca-2011-national-state-district-subdistrict-town"
LCD_ID = "orgi-location-code-directory-2011"
GUIDE_ID = "mopr-panchayat-plan-booklet-2026-27"
PCA_URL = "https://censusindia.gov.in/nada/index.php/catalog/42559/download/46185/2011-IndiaStateDistSbDistTwn-0000.xlsx"
LCD_URL = "https://censusindia.gov.in/nada/index.php/catalog/42648/download/46323/PC11_TV_DIR.xlsx"
GUIDE_URL = "https://panchayat.gov.in/en/document/preparation-of-panchayat-development-plan-booklet-2026-27/"
GUIDE_PDF_URL = "https://cdnbbsr.s3waas.gov.in/s316026d60ff9b54410b3435b403afd226/uploads/2026/05/202607211776250416.pdf"
RIGHTS_URL = "https://censusindia.gov.in/census.website/en/node/286"
CAPTURED_AT = "2026-09-26T23:05:23.336237+00:00"
EXPECTED = {
    PCA.name: (16044425, "da487f7a1181fd9f43ccf354a8b5864ffca5920fc82b5183d644d483445cb380"),
    LCD.name: (23318601, "e5670123e836148cd4a869805333b519f45ce837d1145c8b5cece3a01dbf7dd0"),
    GUIDE.name: (1952128, "1fead200a14afc2c68b8aeddf3548f1f03eab4b36dbf44e105016309cc4bcb61"),
}
ADOPT = {
    "TOT_P": ("Census enumerated population", "people"),
    "TOT_M": ("Census enumerated males", "people"),
    "TOT_F": ("Census enumerated females", "people"),
    "No_HH": ("Census households", "households"),
    "P_06": ("Census children aged 0–6", "people"),
    "P_LIT": ("Census literate persons", "people"),
    "TOT_WORK_P": ("Census workers", "people"),
    "NON_WORK_P": ("Census non-workers", "people"),
}
LEVELS = {"India": "national", "STATE": "state_2011", "DISTRICT": "district_2011", "SUB-DISTRICT": "subdistrict_2011"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def census_key(row: tuple) -> tuple[str, str, str]:
    return tuple(str(row[index]).zfill(width) for index, width in ((0, 2), (1, 3), (2, 5)))


def territory_id(level: str, row: tuple) -> str:
    state, district, subdistrict = census_key(row)
    return {
        "India": "IND",
        "STATE": f"IND:census2011:state:{state}",
        "DISTRICT": f"IND:census2011:district:{district}",
        "SUB-DISTRICT": f"IND:census2011:subdistrict:{state}-{district}-{subdistrict}",
    }[level]


def main() -> None:
    for path in (PCA, LCD, GUIDE):
        size, digest = EXPECTED[path.name]
        if path.stat().st_size != size or sha(path) != digest:
            raise ValueError(f"Official original changed: {path.name}")
    if not BOOTSTRAP.exists():
        BOOTSTRAP.parent.mkdir(exist_ok=True)
        shutil.copy2(DATA, BOOTSTRAP)
    dataset = json.loads(BOOTSTRAP.read_text(encoding="utf-8"))
    workbook = load_workbook(PCA, read_only=True, data_only=True)
    sheet = workbook["Data"]
    rows = sheet.iter_rows(values_only=True)
    headers = tuple(next(rows))
    if len(headers) != 94 or tuple(headers[:11]) != (
        "State", "District", "Subdistt", "Town/Village", "Ward", "EB",
        "Level", "Name", "TRU", "No_HH", "TOT_P",
    ):
        raise ValueError("PCA structure changed")
    record_structure = {
        str(row[0]): str(row[3])
        for row in workbook["Record Structure"].iter_rows(min_row=3, values_only=True)
        if row[0] is not None and row[3] is not None
    }
    selected, level_counts, tru_counts = [], Counter(), Counter()
    field_counts = {header: Counter() for header in headers[9:]}
    for row_number, row in enumerate(rows, start=2):
        level_counts[str(row[6])] += 1
        tru_counts[str(row[8])] += 1
        for index, header in enumerate(headers[9:], start=9):
            value = row[index]
            field_counts[header]["numeric" if isinstance(value, (int, float)) and not isinstance(value, bool) else "blank_or_other"] += 1
            if value == 0:
                field_counts[header]["zero"] += 1
        if row[8] == "Total" and row[6] in LEVELS:
            selected.append((row_number, row))
    expected_levels = {"India": 1, "STATE": 35, "DISTRICT": 640, "SUB-DISTRICT": 5988}
    counts = Counter(row[6] for _, row in selected)
    if dict(counts) != expected_levels or len(selected) != 6664:
        raise ValueError(f"Selected historical hierarchy changed: {counts}")
    if len(headers[9:]) != 85 or set(ADOPT) - set(headers):
        raise ValueError("PCA numeric field catalogue changed")
    workbook.close()

    # A second independent ORGI workbook confirms the historical code keys.
    directory = load_workbook(LCD, read_only=True, data_only=True)
    code_rows = directory.active.iter_rows(values_only=True)
    if tuple(next(code_rows)) != ("State Code", "District Code", "Sub District Code", "Town-Village Code", "Town-Village Name"):
        raise ValueError("Location Code Directory structure changed")
    directory_codes = defaultdict(set)
    directory_count = 0
    for code_row in code_rows:
        directory_count += 1
        if str(code_row[3]).zfill(6) != "000000":
            continue
        key = census_key(code_row)
        directory_codes[key].add(str(code_row[4]))
    directory.close()
    # 99999 is an explicit PCA district residual, not a registered sub-district.
    residual_rows = [
        (row_number, row) for row_number, row in selected
        if row[6] == "SUB-DISTRICT" and str(row[2]) == "99999"
    ]
    if any(row[7] != "Area not under any Sub-district" for _, row in residual_rows):
        raise ValueError("Unregistered 99999 geography changed meaning")
    code_missing = [
        (row_number, census_key(row)) for row_number, row in selected
        if row[6] != "India" and census_key(row) not in directory_codes
        and not (row[6] == "SUB-DISTRICT" and str(row[2]) == "99999")
    ]
    if code_missing:
        raise ValueError(f"PCA codes absent from ORGI directory: {code_missing[:8]}")
    duplicate_directory_labels = [
        {"code": key, "labels": sorted(names)}
        for key, names in directory_codes.items() if len(names) > 1
    ]

    # Only directly published Total rows are used. Rural/urban and town rows
    # stay in the original and field inventory, never summed into a parent.
    territory_rows = {territory_id(row[6], row): (number, row) for number, row in selected}
    if len(territory_rows) != len(selected):
        raise ValueError("Historical territory ID collision")
    parents = {
        "STATE": lambda row: "IND",
        "DISTRICT": lambda row: f"IND:census2011:state:{str(row[0]).zfill(2)}",
        "SUB-DISTRICT": lambda row: f"IND:census2011:district:{str(row[1]).zfill(3)}",
    }
    national = dataset["territories"][0]
    territories = [national]
    for _, row in selected:
        if row[6] == "India":
            continue
        level = row[6]
        code = str(row[{"STATE": 0, "DISTRICT": 1, "SUB-DISTRICT": 2}[level]])
        parent_id = parents[level](row)
        if parent_id not in territory_rows and parent_id != "IND":
            raise ValueError(f"Missing 2011 parent for {code}")
        residual = level == "SUB-DISTRICT" and code == "99999"
        territories.append({
            "id": territory_id(level, row), "name": str(row[7]),
            "level": LEVELS[level], "type": "2011 Census unassigned district residual" if residual else f"2011 Census {level.lower()}",
            "parent_id": parent_id, "official_code": code,
            "code_system": "ORGI 2011 PCA special 99999 residual; absent from Location Code Directory" if residual else "ORGI 2011 PCA and Location Code Directory",
            "boundary_version": "ORGI-2011-statistical-geography; polygon-not-adopted",
            "source_id": SOURCE_ID,
        })
    dataset["territories"] = territories
    dataset["boundaries"] = {"type": "FeatureCollection", "features": []}
    dataset["country"]["geography_note"] = (
        "Historical ORGI Census 2011: 35 states/UTs, 640 districts and 5,988 "
        "sub-district-level source rows, including explicit district residuals coded 99999. "
        "The 2011 official code directory confirms ordinary historical keys but omits "
        "the 99999 residuals, which are not legal sub-districts. No Census-day or "
        "current legal polygons are adopted. "
        "The separately acquired geoBoundaries ADM1 provider shapes describe a different "
        "36-area configuration and are excluded from this edition. Census enumerated "
        "population and WDI midyear estimates are separate definitions and series."
    )
    for key in ("boundary_reconciliation", "subnational_statistics", "planning_documents"):
        dataset["gaps"] = [gap for gap in dataset["gaps"] if gap["category"] != key]
    dataset["gaps"].extend([
        {"category": "historical_to_current_geography", "status": "unverified", "detail": "2011 Census state, district and sub-district codes are not matched to current LGD local governments or current legal boundaries. The 36-area provider ADM1 shape is excluded from the 35-area 2011 hierarchy.", "next_action": "Acquire dated LGD change/crosswalk and official boundary editions; verify state/UT changes and legal Panchayat/Municipality identity before joining current plans."},
        {"category": "code_directory_labels", "status": "unverified", "detail": f"The official 2011 Location Code Directory repeats {len(duplicate_directory_labels)} geographic key(s) with different labels. Codes are matched as keys only; neither label is used to infer a current entity.", "next_action": "Review the duplicate code rows against the PCA and official errata before presenting directory labels as authoritative names."},
        {"category": "unassigned_census_areas", "status": "partial", "detail": f"The PCA contains {len(residual_rows)} district residual rows coded sub-district 99999 and named 'Area not under any Sub-district'. They retain observed counts for full district reconciliation but are not registered legal sub-districts or a current planning body.", "next_action": "Keep these residuals visibly distinct in comparisons and outputs; reconcile them against the official historical geography before polygon or current-authority adoption."},
        {"category": "census_source_coverage", "status": "partial", "detail": "One 2011 PCA workbook has 85 numeric fields and 28,389 data rows. Eight direct count fields from 6,664 national/state/district/sub-district Total rows are adopted. Rural/urban and town rows plus the other 77 fields remain in the source and field inventory; the separate 318 MB village workbook is catalogued but not acquired.", "next_action": "Assess all 85 PCA fields and obtain priority 2011 village/ward and other Census tables by theme; keep historical code/period and source rights visible."},
        {"category": "planning_documents", "status": "partial", "detail": "The national Ministry of Panchayati Raj 2026–27 planning booklet was acquired as guidance; no individual Gram, Block or District Panchayat plan, municipality plan, adopted budget, actual expenditure or official evaluation is attached to a verified present legal unit.", "next_action": "Choose representative current legal authorities; acquire the complete eGramSwaraj plan, budget, progress and accounts originals or document access failure, then verify current LGD codes and periods."},
        {"category": "source_reuse", "status": "unverified", "detail": "The ORGI website copyright policy requires prior permission to reproduce portal material. No such permission is recorded for this candidate.", "next_action": "Obtain and record ORGI publication/reuse permission before any public site or redistribution of Census-derived values or original workbooks."},
    ])
    indicators = []
    for field, (name, unit) in ADOPT.items():
        indicators.append({
            "id": f"IND_PCA_{field}", "name": name, "theme": "Population" if field not in ("TOT_WORK_P", "NON_WORK_P") else "Livelihoods",
            "unit": unit, "definition": record_structure[field],
            "population": "Persons or households enumerated in the 2011 Census as defined by ORGI PCA",
            "measurement_method": "2011 Primary Census Abstract direct published count",
            "source_id": SOURCE_ID, "aggregation": "none", "series_family": "census",
            "display_role": "primary", "period_policy": "same_period",
        })
    dataset["indicators"] = indicators + dataset["indicators"]
    observations = []
    parent_sums = defaultdict(lambda: Counter())
    parent_values = {}
    for row_number, row in selected:
        level, id_ = row[6], territory_id(row[6], row)
        if row[10] != row[11] + row[12]:
            raise ValueError(f"PCA sex sum differs at Data row {row_number}")
        for field in ADOPT:
            index = headers.index(field)
            value = row[index]
            if not isinstance(value, int) or value < 0:
                raise ValueError(f"Invalid adopted count at Data!{get_column_letter(index + 1)}{row_number}")
            parent_values[(id_, field)] = value
            if level != "India":
                parent_sums[("IND" if level == "STATE" else parents[level](row), field)]["sum"] += value
                parent_sums[("IND" if level == "STATE" else parents[level](row), field)]["children"] += 1
            observations.append({
                "territory_id": id_, "indicator_id": f"IND_PCA_{field}",
                "period": "2011", "value": value, "status": "observed",
                "source_id": SOURCE_ID,
                "boundary_version": None if level == "India" else "ORGI-2011-statistical-geography; polygon-not-adopted",
                "source_locator": f"Data!{get_column_letter(index + 1)}{row_number}; Level={level}; TRU=Total",
            })
    differences = [
        {"territory_id": id_, "field": field, "parent": parent_values[(id_, field)], "children": counter["sum"], "child_count": counter["children"]}
        for (id_, field), counter in parent_sums.items()
        if parent_values[(id_, field)] != counter["sum"]
    ]
    if differences:
        raise ValueError(f"2011 direct child count cover differs: {differences[:8]}")
    dataset["observations"] = observations + dataset["observations"]
    dataset["analysis"]["population_context"] = {"primary_indicator_id": "IND_PCA_TOT_P", "reference_indicator_id": "SP.POP.TOTL"}
    dataset["collection"]["adapters"].append("orgi-pca-2011-historical-offline")
    dataset["collection"]["notes"].extend([
        "2011 PCA population/household/work counts are direct source cells for the historical hierarchy; eight count fields adopted, 77 fields and town/rural/urban rows retained for further assessment.",
        "All adopted fields reconcile from sub-district to district to state to national in the 2011 statistical hierarchy; no current local-government value or boundary is inferred.",
        "ORGI portal reproduction permission is not recorded. The generated site is a private research candidate, not a public release.",
    ])
    for source in dataset["sources"]:
        if source["id"] == "geoboundaries-adm1":
            source["note"] = (source.get("note", "") + " Acquired provider source retained in raw only; its 36 features are not adopted with the 35-unit 2011 Census state/UT hierarchy.").strip()
    retrieved = CAPTURED_AT
    dataset["sources"].extend([
        {"id": SOURCE_ID, "name": "ORGI Census 2011 Primary Census Abstract: national/state/district/sub-district/town", "url": PCA_URL, "publisher": "Office of the Registrar General & Census Commissioner, India", "reference_period": "2011", "geographic_level": "historical_national_state_district_subdistrict_town", "status": "partial", "retrieved_at": retrieved, "sha256": EXPECTED[PCA.name][1], "raw_path": f"raw/official-census/{PCA.name}", "note": f"Eight direct count fields from historical Total rows adopted for private review. Official catalogue title says village but this downloaded file contains town rows; the separate 318 MB village file is not acquired. {RIGHTS_URL} requires prior permission to reproduce portal material; none is recorded."},
        {"id": LCD_ID, "name": "ORGI Census 2011 Location Code Directory", "url": LCD_URL, "publisher": "Office of the Registrar General & Census Commissioner, India", "reference_period": "2011", "geographic_level": "historical_code_register", "status": "ready", "retrieved_at": retrieved, "sha256": EXPECTED[LCD.name][1], "raw_path": f"raw/official-census/{LCD.name}", "note": f"Historical codes only; not a current LGD legal-authority register. Check reproduction permission at {RIGHTS_URL} before publication."},
        {"id": GUIDE_ID, "name": "Preparation of Panchayat Development Plan Booklet 2026–27", "url": GUIDE_URL, "publisher": "Ministry of Panchayati Raj, India", "reference_period": "financial year 2026–27", "geographic_level": "national_guidance", "status": "partial", "retrieved_at": retrieved, "sha256": EXPECTED[GUIDE.name][1], "raw_path": "raw/official-planning/mopr-gpdp-booklet-2026-27.pdf", "note": "National process guidance, not an approved plan, expenditure, or evaluation for a selected local authority."},
    ])
    DATA.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inventory = {
        "pca_source_url": PCA_URL, "pca_sha256": EXPECTED[PCA.name][1],
        "code_source_url": LCD_URL, "code_sha256": EXPECTED[LCD.name][1],
        "pca_data_rows": sum(level_counts.values()), "pca_levels": dict(level_counts),
        "pca_tru": dict(tru_counts), "selected_total_rows": len(selected),
        "selected_levels": dict(counts), "numeric_field_count": len(headers[9:]),
        "all_numeric_fields": [
            {"field": field, "column": get_column_letter(index + 1),
             "description": record_structure.get(field, ""),
             "numeric_cells": field_counts[field]["numeric"],
             "blank_or_other_cells": field_counts[field]["blank_or_other"],
             "source_zero_cells": field_counts[field]["zero"],
             "disposition": "adopted_historical_total_rows_only" if field in ADOPT else "retained_priority_unassessed"}
            for index, field in enumerate(headers) if index >= 9
        ],
        "code_directory_rows": directory_count,
        "matched_selected_subnational_codes": len(selected) - 1 - len(residual_rows),
        "pca_unassigned_subdistrict_rows": len(residual_rows),
        "directory_code_keys_with_multiple_labels": duplicate_directory_labels,
        "all_adopted_field_parent_child_sums_equal": True,
        "adopted_observations": len(observations),
        "dataset_sha256": sha(DATA),
        "rights": "ORGI portal reproduction permission required; no permission obtained; no publication",
    }
    EVIDENCE.mkdir(exist_ok=True)
    (EVIDENCE / "IND_ORGI_2011_PCA_FIELD_AUDIT.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (EVIDENCE / "IND_OFFICIAL_SOURCE_RECEIPTS.json").write_text(json.dumps({
        "capture_recorded_at_utc": retrieved,
        "sources": [
            {"url": url, "raw_path": str(path.relative_to(PROJECT)).replace("\\", "/"), "bytes": EXPECTED[path.name][0], "sha256": EXPECTED[path.name][1], "http_status": 200}
            for url, path in ((PCA_URL, PCA), (LCD_URL, LCD), (GUIDE_PDF_URL, GUIDE))
        ],
        "catalogue_pages": ["https://censusindia.gov.in/nada/index.php/catalog/42559", "https://censusindia.gov.in/nada/index.php/catalog/42648", "https://censusindia.gov.in/nada/index.php/catalog/42554", GUIDE_URL],
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"territories": len(territories), "census_observations": len(observations), "fields": len(headers[9:]), "dataset_sha256": sha(DATA)}))


if __name__ == "__main__":
    main()
