"""Build an unpublished Kazakhstan candidate from pinned official source bodies.

The July 2026 population roster is reconciled to the September 2026 KATO
classifier by its complete parent/child order and names. This does not assert
that an official polygon or a same-day code release has been acquired.
"""

import argparse
from collections import Counter, defaultdict
from decimal import Decimal
import difflib
import hashlib
import html
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/kazakhstan-2026-source-manifest.json"
CHECKED_AT = "2026-09-27"
PREFIX = "KAZ_2026_"
CODE_SYSTEM = "BNS KATO NK RK 11-2025, edition 2026-09-18"
GEOGRAPHY = "BNS 2026-07-01 reporting roster matched to KATO 2026-09-18; dated polygons pending"
PARENT_CODES = ("10", "11", "15", "19", "23", "27", "31", "33", "35", "39", "43",
                "47", "55", "59", "61", "62", "63", "71", "75", "79")
PARENT_NAMES = ("Abay", "Akmola", "Aktobe", "Almaty", "Atyrau", "Batys Kazakhstan",
                "Zhambyl", "Zhetisu", "Karaganda", "Kostanay", "Kyzylorda",
                "Mangystau", "Pavlodar", "Soltustik Kazakhstan", "Turkistan", "Ulytau",
                "Shygys Kazakhstan", "Astana city", "Almaty city", "Shymkent city")
CHILD_COUNTS = (12, 20, 13, 11, 8, 13, 11, 10, 13, 20, 9, 7, 13, 14, 17, 5, 13, 6, 8, 5)
JULY_KEYS = ("total", "men", "women", "urban", "urban_men", "urban_women",
             "rural", "rural_men", "rural_women")


def require(ok, message):
    if not ok:
        raise ValueError(message)


def pinned(project):
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    require(manifest["country_id"] == "KAZ" and len(manifest["sources"]) == 9,
            "Source manifest changed")
    sources = {}
    for source in manifest["sources"]:
        path = project / source["raw_path"]
        require(path.is_file(), f"Missing official original: {path}")
        body = path.read_bytes()
        require(len(body) == source["bytes"] and hashlib.sha256(body).hexdigest() == source["sha256"],
                f"Official original bytes differ: {path}")
        if path.suffix == ".xlsx":
            require(body.startswith(b"PK"), f"Invalid XLSX: {path}")
        if path.suffix == ".pdf":
            require(body.startswith(b"%PDF"), f"Invalid PDF: {path}")
        sources[source["id"]] = (source, path)
    return manifest, sources


def integer(value, context, signed=False):
    require(value is not None and str(value).strip() not in ("", "-", "..."),
            f"Missing value: {context}")
    number = Decimal(str(value).replace(" ", ""))
    require(number.is_finite() and number == int(number) and (signed or number >= 0),
            f"Non-integer or unexpected negative value: {context}={value!r}")
    return int(number)


def numeric_row(row, context):
    return tuple(integer(row[i], f"{context}/{JULY_KEYS[i-1]}") for i in range(1, 10))


def workbook_inventory(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    inventory = []
    for sheet in wb:
        number = zero = dash = 0
        numeric_columns = defaultdict(int)
        first_values = {}
        for row in sheet.iter_rows(values_only=True):
            for col, value in enumerate(row, start=1):
                if value is not None and str(value).strip() and col not in first_values:
                    first_values[col] = str(value).strip()[:100]
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    number += 1
                    zero += value == 0
                    numeric_columns[col] += 1
                elif isinstance(value, str):
                    dash += value.strip() == "-"
                    if re.fullmatch(r"-?\d+(?:\.\d+)?", value.strip()):
                        number += 1
                        zero += Decimal(value.strip()) == 0
                        numeric_columns[col] += 1
        adopted_columns = ({"2": set(range(2, 11)), "1": set(), "3": set()}
                           if path.name == "pop-sex-urban-2026-07-01.xlsx" else
                           {"1.": {3, 4, 5, 6}} if path.name == "pop-2026-08-01.xlsx" else {})
        inventory.append({"sheet": sheet.title, "title": str(sheet.cell(1, 1).value or "").strip(),
                          "rows": sheet.max_row, "columns": sheet.max_column,
                          "numeric_like_cells": number, "zero_cells": zero, "dash_cells": dash,
                          "numeric_columns": [{"column": get_column_letter(col),
                              "first_nonempty": first_values.get(col, ""),
                              "numeric_like_cells": count,
                              "assessment": "adopted_selected_rows" if col in adopted_columns.get(sheet.title, set())
                                  else "duplicate_crosscheck_selected_rows" if path.name == "pop-sex-urban-2026-07-01.xlsx" and sheet.title == "1" and col in range(2, 11)
                                  else "priority_unassessed"}
                              for col, count in sorted(numeric_columns.items())],
                          "assessment": "selected_columns_audited" if sheet.title in ("1", "2", "1.")
                          else "priority_unassessed"})
    wb.close()
    return inventory


TRANSLIT = dict(zip("абвгдезиклмнопрстуфыэ", "abvgdeziklmnoprstufye"))
TRANSLIT.update({"ж": "zh", "ч": "ch", "ш": "sh", "щ": "shch", "ё": "yo",
                "й": "y", "х": "kh", "ц": "ts", "ю": "yu", "я": "ya", "ә": "a",
                "ғ": "g", "қ": "k", "ң": "n", "ө": "o", "ұ": "u", "ү": "u",
                "һ": "h", "і": "i"})


def name_key(value):
    value = "".join(TRANSLIT.get(ch, ch) for ch in str(value).lower())
    value = re.sub(r"\b(c a|g a|district|distict|sity|city|area|oblast|region|audany|k a|ga|ca|c|q a|q|a)\b",
                   " ", value)
    value = re.sub(r"[^a-z0-9]+", "", value)
    for left, right in (("ai", "ay"), ("oy", "oi"), ("zh", "j"), ("kh", "h"),
                        ("sh", "s"), ("ch", "c")):
        value = value.replace(left, right)
    return value


def territory_id(code):
    return "KAZ" if code is None else f"KAZ:KATO:{code}"


def read_kato(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    require(wb.sheetnames == ["katonew1"], "KATO sheet changed")
    sheet = wb.active
    require(sheet.max_row == 15587 and tuple(c.value for c in sheet[1]) ==
            ("te", "ab", "cd", "ef", "hij", "k", "kaz_name", "rus_name", "nn"),
            "KATO edition/header changed")
    parents = {}
    children = defaultdict(list)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        code = str(row[0])
        require(re.fullmatch(r"\d{9}", code), f"Unexpected KATO code: {code}")
        if code[2:] == "0000000":
            require(code[:2] in PARENT_CODES and code not in parents, f"Unexpected KATO parent: {code}")
            parents[code] = row
        elif row[4] == "000" and row[2] != "00" and (
            row[3] == "10" if code[:2] in ("71", "75", "79") else row[3] == "00"
        ):
            children[code[:2]].append(row)
    require(set(parents) == {p + "0000000" for p in PARENT_CODES}, "KATO first-level roster differs")
    require(tuple(len(children[p]) for p in PARENT_CODES) == CHILD_COUNTS,
            "KATO district/city roster count differs")
    wb.close()
    return parents, children


def read_july(path, kato_parents, kato_children):
    wb = load_workbook(path, read_only=True, data_only=True)
    require("1" in wb and "2" in wb and "3" in wb, "BNS population sheet roster changed")
    groups = []
    last_numeric_row = 0
    for row_number, row in enumerate(wb["2"].iter_rows(values_only=True), 1):
        if not isinstance(row[1], (int, float)):
            continue
        if row_number - last_numeric_row > 1:
            groups.append([])
        groups[-1].append({"row": row_number, "name": str(row[0]).strip(),
                           "values": numeric_row(row, f"sheet 2 row {row_number}")})
        last_numeric_row = row_number
    require(len(groups) == 21 and len(groups[0]) == 1 and
            tuple(len(g) - 1 for g in groups[1:]) == CHILD_COUNTS,
            "BNS July national/region/child row roster changed")
    require(groups[0][0]["name"] == "Republic of Kazakhstan" and
            tuple(g[0]["name"] for g in groups[1:]) == PARENT_NAMES,
            "BNS July first-level names/order changed")
    records = [(None, None, groups[0][0], "country")]
    matches = []
    for parent, group in zip(PARENT_CODES, groups[1:]):
        parent_code = parent + "0000000"
        records.append((parent_code, None, group[0], "republican_city" if parent in ("71", "75", "79") else "region"))
        for population_row, kato_row in zip(group[1:], kato_children[parent]):
            code = str(kato_row[0])
            score = difflib.SequenceMatcher(None, name_key(population_row["name"]),
                                            name_key(kato_row[6])).ratio()
            alias = population_row["name"] == "Ust-Kamenogorsk c.a." and code == "631000000"
            require(score >= 0.64 or alias,
                    f"Unreconciled population/KATO row: {population_row['name']} / {code}")
            kind = "intracity_district" if parent in ("71", "75", "79") else (
                "district" if "district" in population_row["name"].lower() or
                "distict" in population_row["name"].lower() else "city_akimat")
            records.append((code, parent_code, population_row, kind))
            matches.append({"population_sheet_row": population_row["row"], "population_name": population_row["name"],
                            "kato_code": code, "kato_kaz_name": str(kato_row[6]),
                            "kato_rus_name": str(kato_row[7]), "name_similarity": round(score, 3),
                            "alias_checked": alias})
    require(len(records) == 249 and len(matches) == 228 and
            len({c for c, _, _, _ in records if c}) == 248, "Duplicate or missing 2026 KATO mapping")
    by_code = {code: row["values"] for code, _, row, _ in records}
    checks = Counter()
    for code, values in by_code.items():
        require(values[0] == values[1] + values[2] == values[3] + values[6] and
                values[3] == values[4] + values[5] and values[6] == values[7] + values[8] and
                values[1] == values[4] + values[7] and values[2] == values[5] + values[8],
                f"BNS July sex/locality cross-foot differs: {code}")
        checks["row_crossfoots"] += 5
    for parent in PARENT_CODES:
        code = parent + "0000000"
        child_values = [by_code[str(row[0])] for row in kato_children[parent]]
        for index in range(9):
            require(by_code[code][index] == sum(values[index] for values in child_values),
                    f"BNS July district subtotal differs: {code}/{index}")
            checks["region_child_subtotals"] += 1
    for index in range(9):
        require(by_code[None][index] == sum(by_code[p + "0000000"][index] for p in PARENT_CODES),
                f"BNS July national subtotal differs: {index}")
        checks["national_subtotals"] += 1
    first_sheet = [row for row in wb["1"].iter_rows(values_only=True)
                   if isinstance(row[1], (int, float))]
    require(len(first_sheet) == 21, "BNS July sheet 1 region count changed")
    for source_row, record in zip(first_sheet, [groups[0][0], *(g[0] for g in groups[1:])]):
        if record["name"] in ("Astana city", "Almaty city", "Shymkent city"):
            require(tuple(integer(source_row[i], f"sheet 1/{record['name']}/{i}") for i in range(1, 7)) ==
                    record["values"][:6] and tuple(source_row[7:10]) == ("-", "-", "-") and
                    record["values"][6:] == (0, 0, 0),
                    f"BNS July sheet 1 dash / sheet 2 zero distinction changed: {record['name']}")
            checks["sheet_1_2_region_equalities"] += 6
        else:
            require(numeric_row(source_row, "sheet 1") == record["values"],
                    f"BNS July sheet 1/2 differs: {record['name']}")
            checks["sheet_1_2_region_equalities"] += 9
    wb.close()
    require(by_code[None][0] == 20590589, "BNS July national anchor changed")
    return records, matches, dict(checks)


def read_august(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    require("1." in wb, "BNS August sheet changed")
    rows = list(wb["1."].iter_rows(values_only=True))
    groups = [rows[6:27], rows[28:49], rows[50:68]]
    require(tuple(len(g) for g in groups) == (21, 21, 18), "BNS August group count changed")
    for source_row, expected in zip(groups[0], ("Republic of Kazakhstan", *PARENT_NAMES)):
        require(difflib.SequenceMatcher(None, name_key(source_row[0]), name_key(expected)).ratio() >= 0.8,
                f"BNS August first-level roster differs: {source_row[0]} / {expected}")
    observations = []
    checks = Counter()
    for scope, group in zip(("total", "urban", "rural"), groups):
        for row in group:
            vals = [integer(row[i], f"August {scope}/{row[0]}/{i}", signed=(i in (2, 3, 4)))
                    for i in range(1, 6)]
            require(vals[0] + vals[1] == vals[4] and vals[2] + vals[3] == vals[1],
                    f"BNS August demographic balance differs: {scope}/{row[0]}")
            checks["row_demographic_balances"] += 2
            observations.append({"scope": scope, "name": str(row[0]).strip(),
                                 "start": vals[0], "growth": vals[1], "natural": vals[2],
                                 "migration": vals[3], "end": vals[4]})
    for key in ("start", "growth", "natural", "migration", "end"):
        require(observations[0][key] == sum(r[key] for r in observations[1:21]),
                f"BNS August national subtotal differs: {key}")
        checks["national_subtotals"] += 1
    wb.close()
    require(observations[0]["end"] == 20604819, "BNS August national anchor changed")
    return observations, dict(checks)


def census_audit(path):
    pdfinfo = subprocess.run(["pdfinfo", str(path)], capture_output=True, check=True, text=True)
    pages = int(re.search(r"^Pages:\s+(\d+)", pdfinfo.stdout, re.M).group(1))
    require(pages == 228, "2021 census PDF page count changed")
    output = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
                            capture_output=True, check=True)
    text = output.stdout.decode("utf-8", errors="replace")
    physical = text.split("\f")
    page5 = physical[5]
    page29 = physical[29]
    require(re.search(r"^2021\s+19 186 015\s+100\s+11 741 342\s+61,2\s+7 444 673", page5, re.M),
            "Census printed p.5 national locality row changed")
    require(re.search(r"^The Republic Of Kazakhstan\s+16 009 597.*19 186 015\s+9 324 840\s+9 861 175", page29, re.M),
            "Census printed p.29 national sex row changed")
    table_rows = []
    for line in physical[4].splitlines():
        match = re.match(r"\s*(\d+\.\d+(?:\.\d+)?)\s+(.+?)\.{3,}\s*(\d+)\s*$", line)
        if match:
            table_rows.append({"table": match.group(1), "title": match.group(2).strip(),
                               "printed_page": int(match.group(3)),
                               "assessment": "national_selected_columns_audited" if match.group(1) in ("1.1", "2.1")
                               else "priority_unassessed"})
    if not any(t["table"] == "1.2" for t in table_rows):
        table_rows.append({"table": "1.2", "title": "Population change (increase/decrease)",
                           "printed_page": 5, "assessment": "priority_unassessed"})
    require(len(table_rows) >= 17 and {"1.1", "1.2", "1.13", "2.1", "2.2", "3.1"} <=
            {t["table"] for t in table_rows}, "Census contents table inventory changed")
    return {"pdf_physical_pages": pages, "table_inventory": table_rows,
            "national_adopted": {"total": 19186015, "men": 9324840, "women": 9861175,
                                 "urban": 11741342, "rural": 7444673},
            "geography_note": "2021 census uses pre-2022 oblast boundaries; no 2021 subnational rows joined to the 2026 KATO roster."}


class TableInventory(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables = []
        self.table = None
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            self.row.append(" ".join(" ".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def html_table_inventory(path):
    parser = TableInventory()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    result = []
    for index, table in enumerate(parser.tables):
        numeric_columns = Counter(col for row in table for col, cell in enumerate(row, start=1)
                                  if re.search(r"\d", cell))
        result.append({"table_index": index + 1, "rows": len(table),
            "column_count_max": max((len(row) for row in table), default=0),
            "first_rows": [row[:12] for row in table[:2]],
            "numeric_like_cells": sum(numeric_columns.values()),
            "numeric_columns": [{"column": col, "cells_containing_digits": count,
                                 "assessment": "priority_unassessed"}
                                for col, count in sorted(numeric_columns.items())],
            "assessment": "priority_unassessed; target, baseline, forecast and budget amounts must be separated"})
    return result


def nsdi_audit(path, records):
    body = json.loads(path.read_text(encoding="utf-8"))
    features = body["features"]
    require(body.get("numberMatched") == 254 and len(features) == 254,
            "NSDI WFS feature count changed")
    props = [feature["properties"] for feature in features]
    codes = [row.get("kato") for row in props]
    counts = Counter(codes)
    official = {code for code, _, _, _ in records if code}
    present = {code for code in codes if code}
    result = {"wfs_layer": "geonode:border_districts", "features_returned": len(features),
        "null_kato_count": counts[None], "null_admin_level_count": sum(row.get("admin_level") is None for row in props),
        "unique_kato_codes": len(present), "matched_country_territory_codes": len(official & present),
        "missing_country_territory_codes": sorted(official - present),
        "duplicate_kato_feature_counts": {code: count for code, count in counts.items()
                                           if code is not None and count > 1},
        "assessment": "attributes_only; priority_unassessed geometry/date/rights; no WFS polygon joined or used for legal boundaries"}
    require(result["null_kato_count"] == 1 and result["null_admin_level_count"] == 40 and
            result["unique_kato_codes"] == 211 and result["matched_country_territory_codes"] == 211 and
            len(result["missing_country_territory_codes"]) == 37,
            "NSDI WFS code coverage changed")
    return result


def build(project, manifest, sources, records, matches, july_checks, august, august_checks,
          census, inventories, plan_tables, nsdi):
    path = project / "data/dashboard.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    require(data["schema_version"] == "0.2" and data["country"]["id"] == "KAZ", "Wrong dataset")
    data["generated_at"] = manifest["retrieved_at"]
    data["country"]["geography_note"] = (
        "BNS July 2026 population reporting roster has 17 oblasts, 3 republican cities and 228 child rows. "
        "Its hierarchy and names were reconciled to BNS KATO classifier edition 18 September 2026; "
        "the classification postdates the statistical reference date. No compatible dated polygons were acquired. "
        "The 2021 census predates the 2022 region split; only national census values are joined. "
        "July and August 2026 BNS estimates and WDI midyear estimates are separate series.")
    territories = []
    for code, parent, row, kind in records:
        territories.append({"id": territory_id(code), "name": row["name"],
                            "level": "national" if code is None else "adm1" if parent is None else "adm2",
                            "type": kind, "parent_id": territory_id(parent) if code else None,
                            "official_code": code, "code_system": CODE_SYSTEM if code else "ISO 3166-1 alpha-3",
                            "boundary_version": GEOGRAPHY if code else None,
                            "source_id": "kaz-bns-kato-2026-09-18" if code else "world-bank-countries",
                            "reconciliation_status": "2026 population roster order/name match; same-day code/polygon release pending"
                            if code else "country code"})
    data["territories"] = territories
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    source_info = {
        "kaz-bns-pop-2026-07-01-sex-locality": ("Population by sex and type of locality, 1 July 2026", "BNS", "2026-07-01", "nation, 20 first-level, 228 children"),
        "kaz-bns-pop-2026-08-01": ("Population from start of 2026 through 1 August 2026", "BNS", "2026-08-01 / year-to-date", "nation, 20 first-level"),
        "kaz-bns-administrative-2026-07-01": ("Administrative-territorial units, 1 July 2026", "BNS", "2026-07-01", "nation and local administrative hierarchy"),
        "kaz-bns-kato-2026-09-18": ("KATO NK RK 11-2025 classifier, update 18 September 2026", "BNS", "edition 2026-09-18", "official code roster"),
        "kaz-bns-census2021-volume1": ("2021 National Population Census, Population Volume I", "BNS", "census 2021-09-01", "national adopted; historical local tables retained separately"),
        "kaz-adilet-planning-system-2026": ("State Planning System decree 790, 2026 amendment", "Adilet", "amended 2026-07-14", "planning institutions"),
        "kaz-ulytau-plan-2026-2030": ("Ulytau oblast development plan 2026–2030", "Ulytau akimat", "2026–2030", "Ulytau oblast"),
        "kaz-burabay-plan-2026-2030": ("Burabay district development plan 2026–2030", "Burabay district akimat", "2026–2030", "Burabay district"),
        "kaz-nsdi-border-districts-attributes": ("NSDI Geoportal WFS district-boundary layer attributes", "Kazakhstan NSDI Geoportal", "layer edition unverified", "district layer; incomplete KATO coverage"),
    }
    sources_by_id = {s["id"]: s for s in data["sources"] if s["id"] != "geoboundaries-adm1"}
    for item in manifest["sources"]:
        sid = item["id"]
        name, publisher, period, geography = source_info[sid]
        sources_by_id[sid] = {"id": sid, "name": name, "url": item["url"],
            "catalogue_url": item["catalogue_url"], "publisher": publisher,
            "reference_period": period, "geographic_level": geography,
            "status": "partial" if sid == "kaz-nsdi-border-districts-attributes" else "ready",
            "retrieved_at": manifest["retrieved_at"], "raw_path": item["raw_path"],
            "sha256": item["sha256"],
            "license": "Official public source; source-specific redistribution terms not yet verified",
            "note": "Attribute-only WFS response, 254 features, 211 distinct KATO codes and no dated polygon acceptance."
                    if sid == "kaz-nsdi-border-districts-attributes" else
                    "Original body hash-pinned; only audited cells and statements are adopted."}
    for sid, name, url, publisher, note in [
        ("kaz-akmola-plan-location", "Akmola oblast 2026–2030 plan location",
         "https://www.gov.kz/memleket/entities/aqmola/documents/details/951803?lang=kk",
         "Akmola akimat", "Official listing found; the fetched page was a 2790-byte shell, so plan body and approval not verified."),
        ("kaz-almaty-oblast-plan-location", "Almaty oblast 2026–2030 plan location",
         "https://www.gov.kz/memleket/entities/almobl/documents/details/1015736?lang=ru",
         "Almaty oblast akimat", "Official listing found; the fetched page was a 2627-byte shell, so plan body and approval not verified."),
        ("kaz-akmola-budget-2026-07-location", "Akmola oblast budget execution report as at 1 July 2026",
         "https://www.gov.kz/memleket/entities/aqmola/documents/details/1039733?lang=ru",
         "Akmola akimat", "Official document listing verified; report body/amounts not acquired or adopted."),
    ]:
        sources_by_id[sid] = {"id": sid, "name": name, "url": url, "publisher": publisher,
            "reference_period": "checked 2026-09-27", "geographic_level": "named oblast",
            "status": "not_collected", "retrieved_at": manifest["retrieved_at"],
            "license": "Link location only; terms unverified", "note": note}
    data["sources"] = list(sources_by_id.values())
    indicators = [i for i in data["indicators"] if not i["id"].startswith(PREFIX)]
    observations = [o for o in data["observations"] if not o["indicator_id"].startswith(PREFIX)]
    july_labels = (
        "Total population", "Men", "Women", "Urban population", "Urban men", "Urban women",
        "Rural population", "Rural men", "Rural women")
    for key, label in zip(JULY_KEYS, july_labels):
        indicators.append({"id": PREFIX + "JUL_" + key.upper(),
            "name": f"{label}, 1 July 2026 (BNS estimate)", "theme": "Population",
            "unit": "people", "definition": f"BNS source-reported {label.lower()} under current population accounting as at 1 July 2026.",
            "population": "BNS current-accounting resident population", "aggregation": "none",
            "source_id": "kaz-bns-pop-2026-07-01-sex-locality",
            "measurement_method": "BNS published workbook table 2, direct row cell"})
    for key, label, definition in (
        ("AUG_TOTAL", "Total population, 1 August 2026 (BNS estimate)", "BNS current-accounting population as at 1 August 2026."),
        ("AUG_GROWTH", "Population change, Jan–Jul 2026", "BNS total population change from beginning of 2026 through 1 August 2026; signed persons."),
        ("AUG_NATURAL", "Natural population change, Jan–Jul 2026", "BNS natural increase/decrease for beginning of 2026 through 1 August 2026; signed persons."),
        ("AUG_MIGRATION", "Net migration, Jan–Jul 2026", "BNS net migration for beginning of 2026 through 1 August 2026; signed persons."),
    ):
        indicators.append({"id": PREFIX + key, "name": label, "theme": "Population",
                           "unit": "people", "definition": definition,
                           "population": "BNS current-accounting resident population", "aggregation": "none",
                           "source_id": "kaz-bns-pop-2026-08-01",
                           "measurement_method": "BNS published workbook table 1, total-population block direct row cell"})
    for key, label in (("TOTAL", "Total"), ("MEN", "Men"), ("WOMEN", "Women"),
                       ("URBAN", "Urban"), ("RURAL", "Rural")):
        indicators.append({"id": PREFIX + "CENSUS2021_" + key,
                           "name": f"2021 census {label.lower()} population (national)",
                           "theme": "Census", "unit": "people",
                           "definition": f"2021 census national {label.lower()} count; census geography predates the 2022 oblast division.",
                           "population": "2021 national census enumerated population", "aggregation": "none",
                           "source_id": "kaz-bns-census2021-volume1",
                           "measurement_method": "BNS 2021 census Volume I printed p.5 or p.29 direct cell"})
    data["indicators"] = indicators
    added = []
    for code, _, row, _ in records:
        for key, value in zip(JULY_KEYS, row["values"]):
            added.append({"territory_id": territory_id(code), "indicator_id": PREFIX + "JUL_" + key.upper(),
                "period": "2026", "value": value, "status": "observed",
                "source_id": "kaz-bns-pop-2026-07-01-sex-locality", "provenance": "source_reported",
                "population_scope": "BNS current-accounting residents, 2026-07-01",
                "measurement_method": "BNS published workbook table 2, direct row cell", "boundary_version": GEOGRAPHY,
                "source_locator": f"Population by sex/locality sheet 2 row {row['row']}, {key}",
                "footnote": "A reported zero is zero. The BNS July estimate is distinct from the August estimate, 2021 census and WDI midyear series."})
    for index, row in enumerate(august[:21]):
        code = None if index == 0 else PARENT_CODES[index - 1] + "0000000"
        expected = "Republic of Kazakhstan" if index == 0 else PARENT_NAMES[index - 1]
        require(difflib.SequenceMatcher(None, name_key(row["name"]), name_key(expected)).ratio() >= 0.8,
                f"BNS August name/order changed: {row['name']} / {expected}")
        for key, field, column in (("AUG_TOTAL", "end", "F"), ("AUG_GROWTH", "growth", "C"),
                                   ("AUG_NATURAL", "natural", "D"), ("AUG_MIGRATION", "migration", "E")):
            added.append({"territory_id": territory_id(code), "indicator_id": PREFIX + key,
                "period": "2026", "value": row[field], "status": "observed",
                "source_id": "kaz-bns-pop-2026-08-01", "provenance": "source_reported",
                "population_scope": "BNS current-accounting residents; 2026-08-01 stock or Jan–Jul flow as titled",
                "measurement_method": "BNS published workbook table 1, total-population block direct row cell", "boundary_version": GEOGRAPHY,
                "source_locator": f"Population 1 August workbook sheet 1 row {7 + index}, column {column}",
                "footnote": "Flow values may be negative. No August district/city estimate is available in this workbook."})
    for key, value in census["national_adopted"].items():
        added.append({"territory_id": "KAZ", "indicator_id": PREFIX + "CENSUS2021_" + key.upper(),
            "period": "2021", "value": value, "status": "observed",
            "source_id": "kaz-bns-census2021-volume1", "provenance": "source_reported",
            "population_scope": "2021 national census enumerated population",
            "measurement_method": "BNS 2021 census Volume I printed p.5 or p.29 direct cell",
            "source_locator": "Printed p.5 table 1.1 national row" if key in ("total", "urban", "rural")
                              else "Printed p.29 table 2.1 national row",
            "footnote": "2021 census subnational rows use boundaries before the 2022 oblast division and are not joined to 2026 units."})
    require(len(added) == 249 * 9 + 21 * 4 + 5, "Adopted Kazakhstan observation count differs")
    data["observations"] = observations + added
    children_by_parent = defaultdict(list)
    for code, parent, _, _ in records:
        if parent:
            children_by_parent[parent].append(territory_id(code))
    burabay = next(code for code, parent, row, _ in records
                   if parent == "110000000" and row["name"] == "Burabay district")
    data["documents"] = [
        {"id": "kaz-census2021-volume1", "territory_id": "KAZ", "category": "reference",
         "kind": "national_census_volume", "title": "2021 Population Census, Volume I",
         "url": sources_by_id["kaz-bns-census2021-volume1"]["catalogue_url"],
         "source_id": "kaz-bns-census2021-volume1", "period": "2021",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "Only five national counts from printed pp.5,29 adopted. Historical region/district tables use pre-2022 geography and remain unassessed."},
        {"id": "kaz-state-planning-system", "territory_id": "KAZ", "category": "reference",
         "kind": "planning_legal_framework", "title": "State Planning System decree 790, 2026 Chapter 7",
         "url": sources_by_id["kaz-adilet-planning-system-2026"]["url"],
         "source_id": "kaz-adilet-planning-system-2026", "period": "2026 amendment",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "Chapter 7 paragraphs 66–79 distinguish oblast/republican-city and district/regional-city plans, approval, investment plans and monitoring. It does not prove an individual plan was approved."},
        {"id": "kaz-ulytau-development-plan", "territory_id": territory_id("620000000"),
         "category": "plan", "kind": "oblast_development_plan",
         "title": "Ulytau oblast development plan 2026–2030",
         "url": sources_by_id["kaz-ulytau-plan-2026-2030"]["url"],
         "source_id": "kaz-ulytau-plan-2026-2030", "period": "2026–2030",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "Full official page body acquired, including current situation, targets and resources. Approval instrument and actual implementation remain separate verification steps; no numeric target is adopted as actual."},
        {"id": "kaz-burabay-development-plan", "territory_id": territory_id(burabay),
         "category": "plan", "kind": "district_development_plan",
         "title": "Burabay district development plan 2026–2030",
         "url": sources_by_id["kaz-burabay-plan-2026-2030"]["url"],
         "source_id": "kaz-burabay-plan-2026-2030", "period": "2026–2030",
         "availability": "body_acquired", "official_status": "approved",
         "official_evidence": {"source_id": "kaz-burabay-plan-2026-2030",
             "locator": "Official page opening: approved by Burabay district maslikhat decision 8C-39/3, 19 December 2025",
             "checked_at": CHECKED_AT, "authority": "Burabay district maslikhat"},
         "territory_match": {"territory_id": territory_id(burabay), "country_id": "KAZ",
             "type": "district", "code_system": CODE_SYSTEM, "official_code": burabay,
             "boundary_version": GEOGRAPHY,
             "method": "District name and Akmola parent match BNS July roster and September KATO row; code edition postdates plan decision",
             "source_id": "kaz-bns-kato-2026-09-18", "locator": f"KATO code {burabay}, district name and Akmola parent",
             "checked_at": CHECKED_AT},
         "note": "Official page identifies Burabay district maslikhat decision 8C-39/3 of 19 December 2025 as approval. Targets are not reported actuals."},
        {"id": "kaz-akmola-plan-location", "territory_id": territory_id("110000000"),
         "category": "plan", "kind": "oblast_development_plan_location",
         "title": "Akmola oblast development plan 2026–2030 location",
         "url": sources_by_id["kaz-akmola-plan-location"]["url"],
         "source_id": "kaz-akmola-plan-location", "period": "2026–2030",
         "availability": "link_verified", "official_status": "unverified",
         "note": "Official listing found; body and approval were not acquired from the accessible page shell."},
        {"id": "kaz-almaty-oblast-plan-location", "territory_id": territory_id("190000000"),
         "category": "plan", "kind": "oblast_development_plan_location",
         "title": "Almaty oblast development plan 2026–2030 location",
         "url": sources_by_id["kaz-almaty-oblast-plan-location"]["url"],
         "source_id": "kaz-almaty-oblast-plan-location", "period": "2026–2030",
         "availability": "link_verified", "official_status": "unverified",
         "note": "Official listing found; body and approval were not acquired from the accessible page shell."},
        {"id": "kaz-akmola-budget-location", "territory_id": territory_id("110000000"),
         "category": "budget", "kind": "budget_execution_report_location",
         "title": "Akmola oblast budget execution report, 1 July 2026",
         "url": sources_by_id["kaz-akmola-budget-2026-07-location"]["url"],
         "source_id": "kaz-akmola-budget-2026-07-location", "period": "2026-07-01",
         "availability": "link_verified", "official_status": "unverified",
         "note": "Report listing confirmed, but body, amounts, accounting scope and approval were not acquired. No budget execution number adopted."},
    ]
    analysis = data["analysis"]
    analysis["default_indicator_id"] = PREFIX + "JUL_TOTAL"
    analysis["latest_values_only"] = True
    analysis["terminal_territory_ids"] = [territory_id(code) for code, parent, _, _ in records if parent]
    analysis["comparisons"] = [{"parent_id": "KAZ",
        "member_ids": [territory_id(p + "0000000") for p in PARENT_CODES],
        "label": "20 BNS July 2026 oblasts and republican cities",
        "membership_note": "Complete 2026 BNS/KATO first-level roster; direct comparable 1 July population cells. August values are a separate date.",
        "source_ids": ["kaz-bns-pop-2026-07-01-sex-locality", "kaz-bns-kato-2026-09-18"]}]
    analysis["comparisons"].extend({"parent_id": territory_id(p + "0000000"),
        "member_ids": children_by_parent[p + "0000000"],
        "label": f"{PARENT_NAMES[index]} BNS July 2026 child units",
        "membership_note": "Complete source-reported child population roster matched to KATO by parent, order and name; no polygon join.",
        "source_ids": ["kaz-bns-pop-2026-07-01-sex-locality", "kaz-bns-kato-2026-09-18"]}
        for index, p in enumerate(PARENT_CODES))
    data["planning"] = {"title": "Territorial development plan and evidence",
        "purpose": "Inspect current local statistics, the responsible plan unit and verified plan, budget and monitoring sources for the selected territory.",
        "system": {"label": "Kazakhstan State Planning System, 2026 Chapter 7",
            "scope": "Oblast/republican-city and district/regional-city development plans are separate legal planning units. Intracity districts are analysis units in this candidate, not assumed plan authorities.",
            "cycle": "The July 2026 amendment says plans are prepared every three years for a five-year period, with approval by the relevant maslikhat.",
            "source_ids": ["kaz-adilet-planning-system-2026"]},
        "sections": [{"id": "plan", "label": "Development plans"},
                     {"id": "budget", "label": "Budget and allocations"},
                     {"id": "implementation", "label": "Implementation and execution"},
                     {"id": "evaluation", "label": "Official evaluation"},
                     {"id": "reference", "label": "Census, population and legal sources"}]}
    data["collection"]["status"] = "partial"
    data["collection"]["adapters"] = [a for a in data["collection"]["adapters"]
                                        if a not in ("geoboundaries", "bns-kato-july-august-2026-partial")] + ["bns-kato-july-august-2026-partial"]
    data["collection"]["notes"] = [note for note in data["collection"]["notes"]
        if not note.startswith(("Nine official bodies are SHA-256 pinned.",
                                "KATO classifier update 18 September 2026",
                                "A Burabay district plan approval statement"))] + [
        "Nine official bodies are SHA-256 pinned. July 2026 BNS sex/locality table 2, selected August 2026 population-change columns and five national 2021 census cells are adopted; other tables/years/columns remain unassessed.",
        "KATO classifier update 18 September 2026 reconciles all 20 first-level and 228 child population rows by parent, order and name but postdates July statistics; 2017 reference polygons are excluded.",
        "A Burabay district plan approval statement is verified in the official page. Ulytau plan body, two other plan locations and one budget-report location are separate acquisition states; no budget, target or evaluation values are adopted."]
    data["gaps"] = [g for g in data["gaps"] if g.get("category") not in (
        "boundary_reconciliation", "subnational_statistics", "official_geography",
        "census_table_coverage", "planning_documents")]
    data["gaps"].extend([
        {"category": "official_geography", "status": "partial",
         "detail": "2026 KATO code roster acquired and all July population reporting rows reconciled. KATO edition is 18 September, population date 1 July. NSDI WFS district-layer attribute response has 254 features but only 211 unique KATO codes, 37 candidate codes absent and no verified layer date/complete polygon match. No polygons joined; 2017 geoBoundaries omitted.",
         "next_action": "Audit NSDI layer metadata/date, duplicate geometries and missing ADM1/ADM2 IDs; obtain complete dated official polygons and confirm code history between July and September 2026."},
        {"category": "census_table_coverage", "status": "partial",
         "detail": "2021 census Volume I acquired; five national population cells adopted. Its oblast/district tables predate the 2022 division and remain priority_unassessed for a separate historical geography edition.",
         "next_action": "Inventory and assess remaining census table columns; build an explicit 2021 historical territory edition before adopting local census counts."},
        {"category": "planning_documents", "status": "partial",
         "detail": "2026 planning framework and Ulytau/Burabay plan bodies acquired; Burabay approval decision identified. Most local plans, investment annexes, budget reports and official evaluations not acquired.",
         "next_action": "Acquire approval instruments and content for each plan unit, investment annexes, dated budget execution and monitoring reports; distinguish target from actual."},
    ])
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit = {"source_manifest": "config/kazakhstan-2026-source-manifest.json",
        "checked_at": CHECKED_AT, "official_source_bodies": 9,
        "kato_classifier_rows": 15586, "first_level_units": 20, "child_units": 228,
        "kato_population_reconciliation": matches,
        "july_population_crosschecks": july_checks,
        "july_national_total": 20590589, "july_direct_observations": 249 * 9,
        "august_population_crosschecks": august_checks,
        "august_national_total": 20604819, "august_direct_observations": 21 * 4,
        "august_workbook_all_group_rows": len(august),
        "census": census, "workbook_inventories": inventories,
        "planning_body_table_inventories": plan_tables, "nsdi_wfs_attributes": nsdi,
        "adopted_direct_observations": len(added),
        "source_assessments": {
            "july_sheet_1": "nation and first-level nine numeric columns crosschecked against table 2; duplicate cells not double-counted",
            "july_sheet_2": "nine population count columns adopted for 249 distinct geographic rows",
            "july_sheet_3": "district-centre and settlement subrows may overlap table 2; priority_unassessed",
            "august_total_block": "end population, total change, natural change and migration adopted for nation plus 20 first-level rows",
            "august_urban_rural_blocks": "numeric fields checked for internal demographic balance but not adopted; priority_unassessed",
            "admin_units_workbook": "20 unit geography crosscheck; numeric administrative counts unassessed as dashboard indicators",
            "ulytau_plan": "body acquired, table columns inventoried; targets, baseline, resources, actual execution and approval instrument unassessed",
            "burabay_plan": "body and maslikhat approval statement acquired; numeric targets/baselines/resources unassessed",
            "akmola_and_almaty_plan_pages": "official link location only; fetched short shell was not accepted as source body",
            "akmola_budget_report": "official link location only; body and monetary values unassessed",
        },
        "warnings": ["KATO 2026-09-18 postdates July and August reporting dates; same-day code validity pending.",
                     "NSDI WFS attributes are incomplete/duplicated and layer date unresolved; no official polygon joined; 2017 provider shapes removed.",
                     "The 2021 census uses pre-2022 geography; local cells withheld.",
                     "Plan targets and budget/report listings are not implementation or expenditure actuals."]}
    audit_path = project / "evidence/KAZAKHSTAN_2026_SOURCE_AUDIT.json"
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return audit_path, len(added)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project.resolve()
    manifest, sources = pinned(project)
    kato_parents, kato_children = read_kato(sources["kaz-bns-kato-2026-09-18"][1])
    records, matches, july_checks = read_july(sources["kaz-bns-pop-2026-07-01-sex-locality"][1],
                                               kato_parents, kato_children)
    august, august_checks = read_august(sources["kaz-bns-pop-2026-08-01"][1])
    census = census_audit(sources["kaz-bns-census2021-volume1"][1])
    inventories = {name: workbook_inventory(sources[sid][1]) for name, sid in (
        ("july_population", "kaz-bns-pop-2026-07-01-sex-locality"),
        ("august_population", "kaz-bns-pop-2026-08-01"),
        ("administrative_units", "kaz-bns-administrative-2026-07-01"))}
    plan_tables = {name: html_table_inventory(sources[sid][1]) for name, sid in (
        ("ulytau", "kaz-ulytau-plan-2026-2030"),
        ("burabay", "kaz-burabay-plan-2026-2030"))}
    nsdi = nsdi_audit(sources["kaz-nsdi-border-districts-attributes"][1], records)
    audit_path, count = build(project, manifest, sources, records, matches, july_checks,
                              august, august_checks, census, inventories, plan_tables, nsdi)
    print(json.dumps({"project": str(project), "audit": str(audit_path),
                      "official_sources": len(sources), "territories": len(records),
                      "direct_observations": count, "july_checks": july_checks,
                      "august_checks": august_checks}, ensure_ascii=True))


if __name__ == "__main__":
    main()
