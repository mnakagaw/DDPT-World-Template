"""Import verified 2022 population rows from ArmStat Chapter 1 into Armenia.

Official HD 002-2021 classifier codes are read from a pinned 2025 ARLIS
incorporation. This candidate does not join the 2005 geoBoundaries polygons.
The source report and classifier agree on 70 marz communities and Yerevan's
12 districts; the narrative's stated 72 communities is separately audited.
"""

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PREFIX = "ARM_ARMSTAT2022_"
CLASSIFIER_SHA256 = "ef45a3e7741ec5948e2bcab4616947065448f20e146b6789b7a3f470fbc162b9"
DECISION_SHA256 = "f57da0dda1b8461dbb270918f48b3953e3f1df44c0b4da49d662661d710d0da1"
LAW_SHA256 = "60eed601ec26b8974972226d3fb9f291ec18eda890d935c03ba1118e3e4e564c"
MARZ_ROWS = (24, 42, 59, 77, 96, 125, 154, 169, 192, 207)
MARZ_NAMES = ("Aragatsotn", "Ararat", "Armavir", "Gegharkunik", "Lori",
              "Kotayk", "Shirak", "Syunik", "Vayots Dzor", "Tavush")
COMMUNITY_NAMES = {
    "02": ("Ashtarak", "Aparan", "Talin", "Alagyaz", "Metsadzor", "Arevut", "Tsaghkahovit", "Shamiram"),
    "03": ("Artashat", "Ararat", "Masis", "Vedi", "Verin Dvin"),
    "04": ("Armavir", "Vagharshapat", "Metsamor", "Baghramyan", "Araks", "Khoy", "Paraqar", "Ferik"),
    "05": ("Gavar", "Jambarak", "Martuni", "Sevan", "Vardenis"),
    "06": ("Vanadzor", "Alaverdi", "Tumanyan", "Spitak", "Stepanavan", "Tashir", "Gyulagarak",
           "Lermontovo", "Lori Berd", "Pambak", "Fioletovo"),
    "07": ("Hrazdan", "Abovyan", "Byureghavan", "Nairi", "Tsaghkadzor", "Nor Hachn", "Charentsavan",
           "Akunq", "Arzni", "Garni", "Jrvej"),
    "08": ("Gyumri", "Artik", "Ani", "Akhuryan", "Amasia", "Ashotsq"),
    "09": ("Kapan", "Goris", "Meghri", "Sisian", "Qajaran", "Tatev", "Tegh"),
    "10": ("Yeghegnadzor", "Jermuk", "Vayq", "Areni", "Yeghegis"),
    "11": ("Ijevan", "Berd", "Dilijan", "Noyemberyan"),
}
YEREVAN_DISTRICTS = ("Ajapnyak", "Avan", "Arabkir", "Davtashen", "Erebuni", "Kentron",
                     "Malatia-Sebastia", "Nor Nork", "Nork-Marash", "Nubarashen", "Shengavit", "Kanaker-Zeytun")


def check(condition, reason):
    if not condition:
        raise ValueError(reason)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag == "td" and self.row is not None:
            self.cell = []

    def handle_data(self, value):
        if self.cell is not None:
            self.cell.append(value)

    def handle_endtag(self, tag):
        if tag == "td" and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def read_classifier(path):
    check(digest(path) == CLASSIFIER_SHA256, "ARLIS 2025 classifier changed")
    parser = TableParser()
    parser.feed(path.read_text(encoding="utf-8"))
    result = []
    for row in parser.rows:
        if len(row) < 5 or not re.fullmatch(r"\d{2} \d{3} \d{3}", row[0]):
            continue
        code = row[0]
        level = "marz" if code[3:] == "000 000" else "community" if code.endswith("000") else "district_or_settlement"
        result.append({"code": code, "check_digit": row[1], "level": level,
                       "name_hy": row[2] or row[3] or row[4]})
    check(len([r for r in result if r["level"] == "marz"]) == 10, "Classifier marzes changed")
    check(len([r for r in result if r["level"] == "community"]) == 71,
          "Classifier communities changed")
    check(len({r["code"] for r in result}) == len(result), "Classifier duplicate code")
    return result


def census_count(sheet, row, column):
    value = sheet.cell(row, column).value
    check(isinstance(value, int) and not isinstance(value, bool) and value >= 0,
          f"Expected census count in {sheet.title}!{get_column_letter(column)}{row}: {value!r}")
    return value


def source(id_, name, url, raw, project, reference_period, geographic_level, note, checked_at):
    return {"id": id_, "name": name, "url": url, "publisher": "Statistical Committee of the Republic of Armenia"
            if id_.startswith("arm-armstat-") else "Armenian Legal Information System (ARLIS)",
            "reference_period": reference_period, "geographic_level": geographic_level,
            "status": "partial", "raw_path": raw.relative_to(project).as_posix(), "sha256": digest(raw),
            "retrieved_at": checked_at, "note": note}


def main(project):
    inventory = json.loads((project / "evidence/ARM_ARMSTAT2022_SOURCE_INVENTORY.json").read_text(encoding="utf-8"))
    root = project / "raw/armstat"
    by_filename = {Path(book["path"]).name: book for book in inventory["workbooks"]}
    file1 = project / by_filename["table 1.1.1.xlsx"]["path"]
    file2 = project / by_filename["table 1.1.2.xlsx"]["path"]
    for file in (file1, file2):
        check(digest(file) == by_filename[file.name]["sha256"], f"Census workbook changed: {file.name}")
    classifier_file = root / "arlis-territorial-classifier-2025.html"
    decision_file = root / "arlis-ashtarak-decision-194-N.html"
    law_file = root / "arlis-local-self-government-2026.html"
    check(digest(decision_file) == DECISION_SHA256 and digest(law_file) == LAW_SHA256,
          "ARLIS law or Ashtarak decision changed")
    decision_html = decision_file.read_text(encoding="utf-8")
    check("194-Ն" in decision_html and "2022-2026" in decision_html and "Հաստատել Աշտարակ" in decision_html,
          "Ashtarak approval decision content changed")
    classifier = read_classifier(classifier_file)
    by_code = {item["code"]: item for item in classifier}

    current_book = load_workbook(file1, read_only=False, data_only=True)
    permanent_book = load_workbook(file2, read_only=False, data_only=True)
    try:
        current, permanent = current_book.active, permanent_book.active
        check("Current" in current.cell(2, 1).value and "Permanent" in permanent.cell(2, 1).value,
              "Census population table definitions changed")
        for sheet in (current, permanent):
            check([sheet.cell(7, col).value for col in (4, 7, 10)] == [2022] * 3,
                  "2022 census columns changed")

        rows = [{"row": 8, "name": "Armenia", "level": "national", "type": "country",
                 "code": None, "parent_id": None, "name_hy": "Հայաստանի Հանրապետություն"}]
        yerevan_code = "01 001 000"
        check(by_code[yerevan_code]["name_hy"] == "Երևան", "Yerevan code changed")
        rows.append({"row": 11, "name": "Yerevan", "level": "community", "type": "Yerevan city community",
                     "code": yerevan_code, "parent_id": "ARM", "name_hy": "Երևան"})
        district_codes = [item for item in classifier if item["code"].startswith("01 001 ")
                          and item["code"] != yerevan_code]
        check(len(district_codes) == 12, "Yerevan district code count changed")
        for row_no, name, item in zip(range(12, 24), YEREVAN_DISTRICTS, district_codes):
            check(permanent.cell(row_no, 1).value.strip() == name, f"Yerevan district {row_no} changed")
            rows.append({"row": row_no, "name": name, "level": "administrative_district",
                         "type": "Yerevan administrative district", "code": item["code"],
                         "parent_id": f"ARM:HD002:{yerevan_code.replace(' ', '')}", "name_hy": item["name_hy"]})
        crosswalk = []
        for index, (marz_row, marz_name) in enumerate(zip(MARZ_ROWS, MARZ_NAMES), 2):
            marz_number = f"{index:02d}"
            marz_code = f"{marz_number} 000 000"
            marz_end = MARZ_ROWS[index - 1] if index < 11 else 224
            marz_label = str(permanent.cell(marz_row, 1).value).strip()
            check("Marz" in marz_label and marz_label.lower().startswith(marz_name[:4].lower()),
                  f"Census marz row {marz_row} changed: {marz_label}")
            check(marz_code in by_code and by_code[marz_code]["level"] == "marz", f"Classifier {marz_code}")
            rows.append({"row": marz_row, "name": marz_name, "level": "marz", "type": "marz",
                         "code": marz_code, "parent_id": "ARM", "name_hy": by_code[marz_code]["name_hy"]})
            census_community_rows = []
            for row_no in range(marz_row + 1, marz_end):
                cell = permanent.cell(row_no, 1)
                label = str(cell.value or "").strip()
                if (label and cell.alignment.indent == 0 and isinstance(permanent.cell(row_no, 4).value, int)
                        and label.lower() not in ("urban", "rural") and "town" not in label):
                    census_community_rows.append(row_no)
            classifier_communities = [item for item in classifier if item["level"] == "community"
                                      and item["code"].startswith(marz_number + " ")]
            expected = COMMUNITY_NAMES[marz_number]
            check(len(census_community_rows) == len(classifier_communities) == len(expected),
                  f"Community membership count changed in {marz_name}")
            for row_no, name, item in zip(census_community_rows, expected, classifier_communities):
                census_name = " ".join(str(permanent.cell(row_no, 1).value).replace("(rural)", "").split())
                check(census_name == name, f"Census community row {row_no} changed: {census_name}")
                entry = {"row": row_no, "name": name, "name_hy": item["name_hy"],
                         "level": "community", "type": "community", "code": item["code"],
                         "parent_id": f"ARM:HD002:{marz_code.replace(' ', '')}"}
                rows.append(entry)
                crosswalk.append({"marz": marz_name, "census_row": row_no,
                                  "census_name_en": name, "classifier_name_hy": item["name_hy"],
                                  "official_code": item["code"]})
        check(len(rows) == 94 and len(crosswalk) == 70, "Census territorial coverage changed")
        for row in rows:
            r = row["row"]
            check(current.cell(r, 1).value == permanent.cell(r, 1).value,
                  f"Current/permanent geographic label mismatch row {r}")
            for sheet in (current, permanent):
                check(census_count(sheet, r, 4) == census_count(sheet, r, 7) + census_count(sheet, r, 10),
                      f"Population sex subtotal row {r}")
        for parent in [row for row in rows if row["level"] in ("national", "marz")
                       or row["code"] == yerevan_code]:
            parent_id = "ARM" if parent["level"] == "national" else f"ARM:HD002:{parent['code'].replace(' ', '')}"
            children = [row for row in rows if row["parent_id"] == parent_id]
            check(children, f"No children for {parent_id}")
            for sheet in (current, permanent):
                for col in (4, 7, 10):
                    delta = census_count(sheet, parent["row"], col) - sum(
                        census_count(sheet, child["row"], col) for child in children)
                    expected_delta = ({7: 2, 10: -2}.get(col, 0)
                                      if sheet is current and parent["code"] == "08 000 000" else 0)
                    check(delta == expected_delta,
                          f"Unexpected subtotal difference {sheet.title} {parent_id} {get_column_letter(col)}: {delta}")
        check(census_count(permanent, 8, 4) == 2932731 and census_count(current, 8, 4) == 2689438,
              "2022 national census totals changed")
    finally:
        current_book.close()
        permanent_book.close()

    data_path = project / "data/dashboard.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    check(data["country"]["id"] == "ARM", "Expected Armenia country candidate")
    data["territories"] = [item for item in data["territories"] if item["id"] == "ARM"]
    data["boundaries"] = {"type": "FeatureCollection", "features": []}
    data["indicators"] = [item for item in data["indicators"] if not item["id"].startswith(PREFIX)]
    data["observations"] = [item for item in data["observations"] if not item["indicator_id"].startswith(PREFIX)]
    data["sources"] = [item for item in data["sources"] if not item["id"].startswith("arm-armstat-")
                       and not item["id"].startswith("arm-arlis-")
                       and item["id"] not in ("arm-ashtarak-plan-catalog-2022-2026",
                                              "arm-ashtarak-plan-2022-2026-pdf")]
    data["documents"] = [item for item in data["documents"] if not item["id"].startswith("arm-ashtarak-")]
    data["country"]["geography_note"] = (
        "ArmStat Census 2022 Chapter 1 reports direct values for Armenia, Yerevan, 12 Yerevan districts, "
        "10 marzes and 70 enlarged marz communities. Official HD 002-2021 codes are matched to the "
        "2025 ARLIS classifier incorporation by marz, ordered census rows and checked bilingual names. "
        "The chapter narrative says 72 communities, whereas its complete table and the classifier list 71 "
        "including Yerevan; this discrepancy is unresolved. The present-population Shirak male/female "
        "marz rows differ by +2/-2 from six community sums and remain as published. "
        "No compatible official polygon was acquired. "
        "The 2005 geoBoundaries reference shapes are unjoined. Present/current and permanent/usual resident "
        "population are separate census definitions, not WDI midyear estimates.")
    data["territories"][0]["source_id"] = "arm-armstat-census2022-table-1.1.2"
    data["territories"][0]["reconciliation_status"] = "Direct national census row; complete marz and Yerevan subtotal verified"
    code_system = "Armenia HD 002-2021 administrative territorial classifier, ARLIS 2025 incorporation"
    id_for = lambda row: "ARM" if row["level"] == "national" else f"ARM:HD002:{row['code'].replace(' ', '')}"
    for row in rows[1:]:
        data["territories"].append({"id": id_for(row), "name": row["name"],
            "name_local": row["name_hy"], "level": row["level"], "type": row["type"],
            "parent_id": row["parent_id"], "official_code": row["code"],
            "code_system": code_system, "boundary_version": None,
            "source_id": "arm-arlis-hd002-2025",
            "reconciliation_status": f"ArmStat Section 1 Table 1.1 row {row['row']} matched to ARLIS classifier code; polygon version unverified"})
    comparisons = []
    for parent in [row for row in rows if row["level"] in ("national", "marz") or row["code"] == yerevan_code]:
        parent_id = id_for(parent)
        children = [row for row in rows if row["parent_id"] == parent_id]
        discrepancy = parent["code"] == "08 000 000"
        comparisons.append({"parent_id": parent_id, "member_ids": [id_for(child) for child in children],
            "label": f"Census 2022 administrative members of {parent['name']}",
            "source_ids": ["arm-armstat-census2022-table-1.1.1",
                           "arm-armstat-census2022-table-1.1.2", "arm-arlis-hd002-2025"],
            "membership_note": ("Permanent counts and present total reconcile. Published present male/female "
                                "Shirak marz rows differ by +2/-2 from six community sums; no adjustment is made. "
                                if discrepancy else "All selected direct census counts reconcile to this parent. ")
                               + "The classifier supplies codes, not polygons."})
    data["analysis"]["comparisons"] = comparisons
    data["analysis"]["terminal_territory_ids"] = [id_for(row) for row in rows
        if row["level"] == "administrative_district" or
        (row["level"] == "community" and row["code"] != yerevan_code)]
    data["analysis"]["default_indicator_id"] = PREFIX + "PERMANENT_TOTAL"
    data["analysis"]["latest_values_only"] = True
    data["analysis"].pop("population_context", None)

    specs = (
        ("PERMANENT_TOTAL", "Census 2022 usual resident population", "permanent", 4, "All usual residents, present or temporarily absent"),
        ("PERMANENT_MALE", "Census 2022 usual resident male population", "permanent", 7, "Male usual residents"),
        ("PERMANENT_FEMALE", "Census 2022 usual resident female population", "permanent", 10, "Female usual residents"),
        ("PRESENT_TOTAL", "Census 2022 present population", "current", 4, "Persons present, including temporary visitors"),
        ("PRESENT_MALE", "Census 2022 present male population", "current", 7, "Male persons present"),
        ("PRESENT_FEMALE", "Census 2022 present female population", "current", 10, "Female persons present"),
    )
    for suffix, name, universe, _col, definition in specs:
        source_id = f"arm-armstat-census2022-table-1.1.{2 if universe == 'permanent' else 1}"
        indicator = {"id": PREFIX + suffix, "name": name, "theme": "Population", "unit": "people",
            "definition": definition + "; ArmStat 2022 Census Chapter 1 methodological definition. "
                          "Direct official rows, no interpolation or redistribution by AreaData.",
            "population": universe + " population of the reported administrative unit",
            "source_id": source_id, "aggregation": "none", "measurement_method": "source_reported",
            "display_decimals": 0}
        if suffix == "PERMANENT_TOTAL":
            indicator["series_family"] = "census"
        data["indicators"].append(indicator)
    workbooks = {"permanent": (file2, "arm-armstat-census2022-table-1.1.2"),
                 "current": (file1, "arm-armstat-census2022-table-1.1.1")}
    census_books = {key: load_workbook(file, read_only=True, data_only=True)
                    for key, (file, _sid) in workbooks.items()}
    try:
        for row in rows:
            for suffix, _name, universe, col, _definition in specs:
                file, source_id = workbooks[universe]
                value = census_count(census_books[universe].active, row["row"], col)
                observation = {"territory_id": id_for(row),
                    "indicator_id": PREFIX + suffix, "period": "2022", "value": value,
                    "status": "observed", "source_id": source_id,
                    "measurement_method": "source_reported",
                    "source_locator": f"ArmStat Census 2022 Section 1, {file.name}, sheet nor, "
                                      f"{get_column_letter(col)}{row['row']}"}
                if row["code"] == "08 000 000" and suffix in ("PRESENT_MALE", "PRESENT_FEMALE"):
                    observation["footnote"] = ("Published Shirak marz male/female count differs by "
                        "+2/-2 respectively from the six community counts; direct source value retained.")
                data["observations"].append(observation)
    finally:
        for book in census_books.values():
            book.close()

    now = datetime.now(timezone.utc).isoformat()
    for filename, table in (("table 1.1.1.xlsx", "1.1.1"), ("table 1.1.2.xlsx", "1.1.2")):
        book = by_filename[filename]
        raw = project / book["path"]
        entry = source(f"arm-armstat-census2022-table-{table}",
            f"ArmStat RA Population Census 2022, Section 1 Table {table}",
            "https://armstat.am/file/article/section_1.7z", raw, project,
            "2022 census; 2001 and 2011 columns unadopted",
            "national, Yerevan/districts, marzes/enlarged communities",
            "Only 2022 D/G/J cells of 94 direct administrative rows adopted; other columns and aggregate "
            "rows remain unassessed. Original archive and workbook hashes are in the private inventory.", now)
        entry["sha256"] = book["sha256"]
        data["sources"].append(entry)
    for sid, name, url, raw, period, scope, note in (
        ("arm-arlis-hd002-2025", "HD 002-2021 territorial classifier, 2025 ARLIS incorporation",
         "https://www.arlis.am/hy/acts/208945", classifier_file, "2025 incorporation", "marz, community, Yerevan district",
         "Official codes matched to 2022 census report table; administrative boundary polygons unacquired."),
        ("arm-arlis-lsg-2026", "Law on Local Self-Government, current ARLIS incorporation",
         "https://www.arlis.am/hy/acts/231070", law_file, "checked 2026-09-26", "Armenian communities",
         "Articles 82, 82.1 and 83 distinguish five-year development plan, annual work plan and budget."),
        ("arm-arlis-ashtarak-194-n", "Ashtarak Council Decision N 194-N, 28 December 2022",
         "https://www.arlis.am/hy/acts/173115", decision_file, "2022-2026 plan approval", "Ashtarak community",
         "Clause 1 approves the 2022-2026 plan; ARLIS reports the decision in force as checked. "
         "The separate full plan PDF was visible in web research but direct local retrieval failed.")):
        entry = source(sid, name, url, raw, project, period, scope, note, now)
        entry["sha256"] = digest(raw)
        data["sources"].append(entry)
    data["sources"].append({"id": "arm-armstat-census2022-results-catalog",
        "name": "ArmStat Population Census 2022 official results catalogue",
        "url": "https://armstat.am/en/?nid=82&id=2623",
        "publisher": "Statistical Committee of the Republic of Armenia",
        "reference_period": "2022 census, nine English chapters", "geographic_level": "national and published local tables",
        "status": "partial", "retrieved_at": now,
        "note": "Official chapter archive directory. AreaData acquired all nine chapter archives but adopted only selected Chapter 1 population fields."})
    data["sources"].append({"id": "arm-ashtarak-plan-catalog-2022-2026",
        "name": "Ashtarak Municipality development programme catalogue entry",
        "url": "https://ashtarak.am/Pages/DocFlow/Def.aspx?a=d&dt=Projects&g=3832faaa-dce7-460b-878f-6c5655668340",
        "publisher": "Ashtarak Municipality", "reference_period": "2022-2026 programme listing",
        "geographic_level": "Ashtarak community", "status": "partial", "retrieved_at": now,
        "note": "Municipal plan location; approval is separately evidenced by ARLIS decision. Full PDF body was not acquired locally."})
    data["sources"].append({"id": "arm-ashtarak-plan-2022-2026-pdf", "name": "Ashtarak five-year development plan 2022-2026 PDF",
        "url": "https://ashtarak.am/upload/DocFlow/Projects/Tu2321411511531797_2022-2026.pdf",
        "publisher": "Ashtarak Municipality", "reference_period": "2022-2026",
        "geographic_level": "Ashtarak community", "status": "partial", "retrieved_at": now,
        "note": "Municipal catalogue and 59-page PDF were inspected in web research; direct local PDF acquisition "
                "failed (connection timeout). No raw PDF or hash is claimed; plan content is not adopted."})
    ashtarak = next(row for row in rows if row["name"] == "Ashtarak" and row["level"] == "community")
    ashtarak_id = id_for(ashtarak)
    data["documents"].append({"id": "arm-ashtarak-development-plan-2022-2026",
        "territory_id": ashtarak_id, "category": "plan", "kind": "five-year-community-development-plan",
        "title": "Ashtarak Community Development Programme 2022-2026",
        "url": "https://ashtarak.am/upload/DocFlow/Projects/Tu2321411511531797_2022-2026.pdf",
        "source_id": "arm-ashtarak-plan-2022-2026-pdf", "period": "2022-2026",
        "target_period": {"label": "2022-2026", "kind": "multi_year",
                          "start": "2022-01-01", "end": "2026-12-31"},
        "availability": "link_verified", "official_status": "adopted",
        "official_evidence": {"source_id": "arm-arlis-ashtarak-194-n",
                              "locator": "Council Decision N 194-N, 28 Dec 2022, clause 1; ARLIS effective status",
                              "checked_at": "2026-09-26", "authority": "Ashtarak Community Council"},
        "territory_match": {"territory_id": ashtarak_id, "country_id": "ARM", "type": "community",
            "code_system": code_system, "official_code": ashtarak["code"], "boundary_version": None,
            "method": "Ashtarak name, community type and stable HD 002 code; exact 2022-to-2025 boundary equivalence remains unverified",
            "source_id": "arm-arlis-hd002-2025", "locator": "Classifier Table 1, row 02 001 000",
            "checked_at": "2026-09-26"},
        "note": "The council's approval decision was acquired and verified; the 59-page plan PDF was only "
                "inspected remotely because direct retrieval timed out. Its body, 2026 revisions and alignment "
                "to the census community boundary remain unverified."})
    data["planning"] = {"title": "Community planning sources", "purpose":
        "Use official community plans, annual work plans and budgets as evidence for local diagnosis and drafting.",
        "system": {"label": "Armenian community five-year development programme and annual work plan",
            "scope": "Community governments; Yerevan has separate local-government provisions",
            "cycle": "Five-year community development programme; annual work plan and budget",
            "source_ids": ["arm-arlis-lsg-2026"]},
        "sections": [{"id": "plan", "label": "Development plans"},
                     {"id": "budget", "label": "Budgets and annual work plans"},
                     {"id": "implementation", "label": "Implementation and execution"},
                     {"id": "evaluation", "label": "Official evaluations"},
                     {"id": "reference", "label": "Legal and methodological references"}]}
    data["gaps"] = [
        {"category": "administrative_boundaries", "status": "not_collected",
         "detail": "No official polygon version was matched to the ArmStat census/classifier communities. "
                   "The collected 2005 geoBoundaries ADM1 layer remains separate and unjoined.",
         "next_action": "Acquire the official dated marz, community and Yerevan district layers with codes and reuse terms."},
        {"category": "subnational_statistics", "status": "partial",
         "detail": "All 57 ArmStat chapter workbooks and 837 numeric columns were mechanically inventoried. "
                   "Only selected 2022 population cells from two Chapter 1 tables were semantically adopted.",
         "next_action": "Audit the remaining demographic, education, housing, migration and health tables by "
                        "geography, universe and year before adopting useful local fields."},
        {"category": "planning_documents", "status": "partial",
         "detail": "Ashtarak plan approval is documented, but direct plan PDF retrieval timed out and body "
                   "is unadopted. Other community plans, annual plans, budgets, execution and evaluations "
                   "are not collected.",
         "next_action": "Acquire the Ashtarak plan and check amendments; collect representative communities' "
                        "plans, budgets and performance records before scaling across 71 communities."},
        {"category": "administrative_register", "status": "partial",
         "detail": "Chapter 1 narrative says 72 communities, while its complete table and the 2025 classifier "
                   "show 70 marz communities plus Yerevan. Later 2026 legal changes have not been reconciled.",
         "next_action": "Confirm the source report's geography edition and 2026 law amendments, retaining historical IDs."},
        {"category": "source_data_quality", "status": "partial",
         "detail": "ArmStat Table 1.1.1 Shirak present-population male/female marz rows differ by +2/-2 "
                   "from the six published community sums; every direct row's total still equals male plus female.",
         "next_action": "Check official corrections or release notes. Do not harmonize the published direct counts."},
    ]
    data["collection"]["status"] = "partial"
    data["collection"]["adapters"] = list(dict.fromkeys(data["collection"]["adapters"] +
        ["armstat-census2022-chapter1-hd002-2025-crosswalk", "arlis-community-planning-sources"]))
    collection_note = ("ArmStat 2022 population and Ashtarak approval are a partial domestic edition; "
        "all other census columns, plan bodies, official polygons and independent ACCEPT remain pending.")
    if collection_note not in data["collection"]["notes"]:
        data["collection"]["notes"].append(collection_note)
    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    crosswalk_path = project / "evidence/ARM_ARMSTAT2022_CODE_CROSSWALK.json"
    crosswalk_path.write_text(json.dumps({"checked_at": now, "classifier_sha256": CLASSIFIER_SHA256,
        "census_workbooks": {filename: by_filename[filename]["sha256"] for filename in
                             ("table 1.1.1.xlsx", "table 1.1.2.xlsx")},
        "note": "English/Armenian pairs manually reviewed. One published Shirak present-sex subtotal "
                "difference is retained and documented; all other selected parent-child sums reconcile.",
        "community_crosswalk": crosswalk}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"territories": len(data["territories"]), "communities_outside_yerevan": 70,
        "districts": 12, "indicators": len(specs), "domestic_observations": len(rows) * len(specs),
        "national_permanent": 2932731, "national_present": 2689438}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
