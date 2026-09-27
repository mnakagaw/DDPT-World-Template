"""Audit official Singapore originals and build an unpublished 2026 partial site.

The 2020 Census and MP2019 shapes are retained as a separate historical edition.
The 2026 resident series uses the MP2025 planning-area roster and polygons only.
"""

import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config/singapore-population2026-source-manifest.json"
BOUNDARY = "URA Master Plan 2025, planning-area boundary, December 2025"
ZIP_SOURCE = "sgp-singstat-population-trends-2026-geospatial"
AREA_SOURCE = "sgp-ura-mp2025-planning-area"
INDICATORS = {
    "RESIDENT_TOTAL": ("Resident population, both sexes", "all residents", "Single year of age and sex; age Total, sex Total"),
    "RESIDENT_MALE": ("Male resident population", "male residents", "Single year of age and sex; age Total, sex Males"),
    "RESIDENT_FEMALE": ("Female resident population", "female residents", "Single year of age and sex; age Total, sex Females"),
    "RESIDENT_HDB": ("Residents in HDB flats", "residents in Total HDB^ dwelling category", "Age group, sex and dwelling; age Total, sex Total, Total HDB^"),
    "RESIDENT_CONDO": ("Residents in condominiums and other apartments", "residents in condominiums and other apartments", "Age group, sex and dwelling; age Total, sex Total, Condominiums and Other Apartments"),
    "RESIDENT_LANDED": ("Residents in landed properties", "residents in landed properties", "Age group, sex and dwelling; age Total, sex Total, Landed Properties"),
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def slug(value):
    return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def area_id(code):
    return f"SGP:URA:MP2025:PA:{code}"


def subzone_id(area_code, name):
    return f"SGP:SINGSTAT2026:SZ:{area_code}:{slug(name)}"


def pinned_sources(project, manifest):
    result = {}
    for item in manifest["sources"]:
        path = project / item["raw_path"]
        require(path.is_file(), f"Missing official original: {path}")
        require(path.stat().st_size == item["bytes"], f"Length changed: {path}")
        require(hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"],
                f"Hash changed: {path}")
        result[item["id"]] = path
    require(len(result) == 7, "Source manifest changed")
    return result


def selected_indicator(book, sheet, row):
    area, subzone, third, fourth, value = row
    if book.startswith("Singapore Residents by Planning Area, Subzone, Single"):
        if third != "Total":
            return None
        return {"Total": "RESIDENT_TOTAL", "Males": "RESIDENT_MALE",
                "Females": "RESIDENT_FEMALE"}.get(fourth)
    if "Type of Dwelling" in book and sheet == "2026(Total)" and third == "Total":
        return {"Total HDB^": "RESIDENT_HDB",
                "Condominiums and Other Apartments": "RESIDENT_CONDO",
                "Landed Properties": "RESIDENT_LANDED"}.get(fourth)
    return None


def inspect_workbooks(path):
    inventory = []
    selected = defaultdict(dict)
    crosscheck = {}
    with zipfile.ZipFile(path) as archive:
        members = [name for name in archive.namelist() if name.endswith(".xlsx")]
        require(len(members) == 3, "Expected all three 2026 geospatial workbooks")
        for member in members:
            with archive.open(member) as stream:
                workbook = load_workbook(stream, read_only=True, data_only=True)
                for sheet in workbook:
                    counts = Counter()
                    dimensions = {"third": Counter(), "fourth": Counter()}
                    footer = []
                    in_footer = False
                    for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
                        if len(row) != 5 or not row[0] or not row[1]:
                            in_footer = True
                            if row[0]:
                                footer.append(str(row[0]))
                            continue
                        require(not in_footer, f"Data after footer in {member}:{sheet.title}:{row_number}")
                        counts["rows"] += 1
                        dimensions["third"][str(row[2])] += 1
                        dimensions["fourth"][str(row[3])] += 1
                        value = row[4]
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            counts["numeric_cells"] += 1
                        elif value == "-":
                            counts["nil_or_negligible_cells"] += 1
                        elif value is None:
                            counts["blank_cells"] += 1
                        else:
                            counts["other_cells"] += 1
                        key = (str(row[0]), str(row[1]))
                        indicator = selected_indicator(member, sheet.title, row)
                        if indicator:
                            require(key not in selected[indicator], f"Duplicate {indicator} {key}")
                            require(isinstance(value, int) or value == "-",
                                    f"Invalid selected value {member}:{sheet.title}:{row_number}")
                            selected[indicator][key] = {"value": value, "book": member,
                                "sheet": sheet.title, "row": row_number}
                        if "Type of Dwelling" in member and sheet.title == "2026(Total)" \
                                and row[2] == "Total" and row[3] == "Total":
                            require(key not in crosscheck, f"Duplicate dwelling total {key}")
                            crosscheck[key] = value
                    require(counts["other_cells"] == 0 and footer,
                            f"Unclassified source cells in {member}:{sheet.title}")
                    inventory.append({"workbook": member, "sheet": sheet.title,
                        "header": [str(x) for x in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))],
                        "rows": counts["rows"], "numeric_cells": counts["numeric_cells"],
                        "nil_or_negligible_cells": counts["nil_or_negligible_cells"],
                        "blank_cells": counts["blank_cells"],
                        "footer_notes": footer,
                        "third_dimension": dict(dimensions["third"]),
                        "fourth_dimension": dict(dimensions["fourth"]),
                        "adopted": "selected cells only" if sheet.title in ("2026", "2026(Total)") else "none; priority_unassessed"})
                workbook.close()
    require(len(inventory) == 7, "Unexpected workbook sheet count")
    keys = set(selected["RESIDENT_TOTAL"])
    require(len(keys) == 388 and all(set(selected[i]) == keys for i in INDICATORS),
            "Selected geographic roster differs among indicators")
    require(set(crosscheck) == keys, "T1/T2 total geography differs")
    require(all(selected["RESIDENT_TOTAL"][k]["value"] == crosscheck[k] for k in keys),
            "T1/T2 resident totals differ")
    return inventory, selected


def historical_inventory(paths):
    with paths["sgp-singstat-cop2020-age-sex"].open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        require(len(reader.fieldnames) == 61, "2020 census column roster changed")
        counts = Counter()
        area_names = set()
        rows = []
        for row in reader:
            counts["rows"] += 1
            label = row["Number"]
            match = re.fullmatch(r"(.+?)\s*-\s*Total", label)
            if match:
                area_names.add(match.group(1).upper())
            rows.append(label)
            for field in reader.fieldnames[1:]:
                value = row[field]
                if value == "-":
                    counts["nil_or_negligible_cells"] += 1
                elif value.isdecimal():
                    counts["numeric_cells"] += 1
                else:
                    counts["other_cells"] += 1
    historical = json.loads(paths["sgp-ura-mp2019-subzone"].read_text(encoding="utf-8"))
    subzones = {(f["properties"]["PLN_AREA_N"].upper(),
                 f["properties"]["SUBZONE_N"].upper()) for f in historical["features"]}
    require(counts["rows"] == 388 and len(area_names) == 55 and len(subzones) == 332,
            "2020/2019 historical roster changed")
    require(counts["other_cells"] == 0, "Unclassified 2020 census cells")
    return {"csv_rows": counts["rows"], "columns": 61,
        "numeric_cells": counts["numeric_cells"],
        "nil_or_negligible_cells": counts["nil_or_negligible_cells"],
        "planning_area_rows": len(area_names), "mp2019_subzone_polygons": len(subzones),
        "adoption": "historical edition only; none transferred to MP2025 current areas"}


def source_record(item, retrieved_at):
    ident = item["id"]
    publisher = "Singapore Department of Statistics" if "singstat" in ident else "Urban Redevelopment Authority"
    return {"id": ident, "name": ident.replace("sgp-", "").replace("-", " ").title(),
        "url": item["url"], "catalogue_url": item["catalogue_url"],
        "publisher": publisher, "reference_period": "2020 census / MP2019" if "2020" in ident or "2019" in ident else "June 2026 / MP2025",
        "geographic_level": "Singapore, planning areas and subzones" if "singstat" in ident else "Singapore planning geography",
        "status": "ready", "retrieved_at": retrieved_at, "sha256": item["sha256"],
        "raw_path": item["raw_path"], "license": "Singapore Open Data Licence" if "ura" in ident or "cop2020" in ident else "Official public report; redistribution terms to confirm",
        "note": "Original hash pinned. Adopted cells and edition exclusions are recorded in SGP_POPULATION2026_AUDIT.json."}


def build(project, manifest, paths, inventory, selected, history):
    dashboard = project / "data/dashboard.json"
    data = json.loads(dashboard.read_text(encoding="utf-8"))
    require(data["country"]["id"] == "SGP", "Wrong country candidate")
    polygons = json.loads(paths[AREA_SOURCE].read_text(encoding="utf-8"))["features"]
    regions = json.loads(paths["sgp-ura-mp2025-region"].read_text(encoding="utf-8"))["features"]
    require(len(polygons) == 55 and len(regions) == 5, "MP2025 polygon roster changed")
    by_name = {f["properties"]["PLN_AREA_N"].upper(): f for f in polygons}
    by_code = {f["properties"]["PLN_AREA_C"]: f for f in polygons}
    require(len(by_name) == len(by_code) == 55, "Duplicate MP2025 area name/code")
    keys = set(selected["RESIDENT_TOTAL"])
    area_names = {area.upper() for area, subzone in keys if area != "Total" and subzone == "Total"}
    require(area_names == set(by_name), "SingStat 2026 areas do not exactly match URA MP2025")
    subzone_keys = {(area.upper(), subzone) for area, subzone in keys if subzone != "Total"}
    normalized_keys = {(area.upper() if area != "Total" else area, subzone) for area, subzone in keys}
    require(len(subzone_keys) == 332 and normalized_keys == {("Total", "Total")} |
            {(name, "Total") for name in area_names} | subzone_keys,
            "2026 reporting hierarchy changed")

    data["country"]["geography_note"] = (
        "SingStat June 2026 resident population reports 55 planning areas and 332 subzones "
        "under URA Master Plan 2025. Area names match all 55 URA MP2025 area polygons and "
        "PLN_AREA_C codes; polygons are indicative. Subzone codes and MP2025 subzone "
        "polygons are not acquired. MP2019/2020 census geography is a separate historical "
        "edition. Planning areas and subzones are statistical/planning geographies, not "
        "separate local governments. WDI national population is a different series.")
    national = next(t for t in data["territories"] if t["id"] == "SGP")
    national.update({"source_id": ZIP_SOURCE, "boundary_version": BOUNDARY})
    data["territories"] = [national]
    geometry = []
    for code in sorted(by_code):
        feature = by_code[code]
        props = feature["properties"]
        name = props["PLN_AREA_N"].title()
        tid = area_id(code)
        data["territories"].append({"id": tid, "name": name, "level": "adm1",
            "type": "Planning Area", "parent_id": "SGP", "official_code": code,
            "code_system": "URA MP2025 PLN_AREA_C", "boundary_version": BOUNDARY,
            "source_id": AREA_SOURCE, "region_name": props["REGION_N"],
            "region_code": props["REGION_C"],
            "reconciliation_status": "Exact 55-name SingStat June 2026 to URA MP2025 planning-area roster; URA code attached"})
        geometry.append({"type": "Feature", "geometry": feature["geometry"],
            "properties": {"territory_id": tid, "name": name, "official_code": code,
                "code_system": "URA MP2025 PLN_AREA_C", "boundary_version": BOUNDARY,
                "source_id": AREA_SOURCE}})
    ids = set(t["id"] for t in data["territories"])
    for area, subzone in sorted(subzone_keys):
        code = by_name[area]["properties"]["PLN_AREA_C"]
        tid = subzone_id(code, subzone)
        require(tid not in ids, f"Subzone slug collision: {area}/{subzone}")
        ids.add(tid)
        data["territories"].append({"id": tid, "name": subzone,
            "level": "adm2", "type": "Statistical Subzone",
            "parent_id": area_id(code), "official_code": None,
            "code_system": "SingStat 2026 planning-area plus subzone name; official subzone code pending",
            "boundary_version": BOUNDARY, "source_id": ZIP_SOURCE,
            "reconciliation_status": "Exact 2026 SingStat source row, MP2025 subzone polygon/code not acquired"})
    data["boundaries"] = {"type": "FeatureCollection", "features": geometry}

    sources = {s["id"]: s for s in data["sources"]}
    for item in manifest["sources"]:
        sources[item["id"]] = source_record(item, manifest["retrieved_at"])
    sources["sgp-ura-mp2025-gazette-location"] = {
        "id": "sgp-ura-mp2025-gazette-location", "name": "URA MP2025 gazette circular",
        "url": "https://www.ura.gov.sg/guidelines/circulars/ppg25-12/",
        "publisher": "Urban Redevelopment Authority", "reference_period": "1 December 2025",
        "geographic_level": "Singapore national statutory land-use plan", "status": "partial",
        "retrieved_at": manifest["retrieved_at"],
        "license": "Official public webpage; reuse terms to confirm",
        "note": "Official circular location and gazette date verified; HTML original not hash-pinned."}
    sources["sgp-ura-master-plan-overview-location"] = {
        "id": "sgp-ura-master-plan-overview-location", "name": "URA Master Plan overview and review cycle",
        "url": "https://www.ura.gov.sg/land-planning/master-plan/",
        "publisher": "Urban Redevelopment Authority", "reference_period": "web page checked 2026-09-27",
        "geographic_level": "Singapore national planning system", "status": "partial",
        "retrieved_at": manifest["retrieved_at"],
        "license": "Official public webpage; reuse terms to confirm",
        "note": "URA explains statutory Master Plan and five-year review; page body not hash-pinned and does not prescribe a separate plan for each statistical subzone."}
    sources["sgp-planning-act-location"] = {
        "id": "sgp-planning-act-location", "name": "Singapore Planning Act 1998 current text",
        "url": "https://sso.agc.gov.sg/Act/PA1998?ValidDate=20251001",
        "publisher": "Attorney-General's Chambers", "reference_period": "version checked 2026-09-27",
        "geographic_level": "Singapore", "status": "partial",
        "retrieved_at": manifest["retrieved_at"],
        "license": "Official public legal text; reuse terms to confirm",
        "note": "Part 2 Master Plan sections 6–11 located; full current legal applicability not separately audited."}
    data["sources"] = list(sources.values())

    data["indicators"] = [i for i in data["indicators"] if not i["id"].startswith("SGP_SINGSTAT_")]
    data["observations"] = [o for o in data["observations"] if not o["indicator_id"].startswith("SGP_SINGSTAT_")]
    for key, (name, scope, locator) in INDICATORS.items():
        data["indicators"].append({"id": "SGP_SINGSTAT_" + key,
            "name": "June 2026 " + name.lower(), "theme": "Housing" if key in ("RESIDENT_HDB", "RESIDENT_CONDO", "RESIDENT_LANDED") else "Population",
            "unit": "people", "display_decimals": 0,
            "definition": "SingStat June 2026 resident population, direct source cell; " + locator +
                          (". Total HDB includes non-privatised HUDC flats" if key == "RESIDENT_HDB" else "") +
                          ". Independently rounded to tens, with '-' meaning nil or negligible, not an exact zero.",
            "population": "Singapore citizens and permanent residents with local addresses, excluding those continuously away for 12 months or longer; " + scope,
            "source_id": ZIP_SOURCE, "aggregation": "none",
            "measurement_method": "SingStat June 2026 resident population, independently rounded to tens"})
    territory_for_key = {("Total", "Total"): "SGP"}
    for name in area_names:
        territory_for_key[(name, "Total")] = area_id(by_name[name]["properties"]["PLN_AREA_C"])
    for area, subzone in subzone_keys:
        territory_for_key[(area, subzone)] = subzone_id(by_name[area]["properties"]["PLN_AREA_C"], subzone)
    new_observations = []
    for indicator in INDICATORS:
        for key, source in sorted(selected[indicator].items()):
            area, subzone = key
            value = source["value"]
            new_observations.append({"territory_id": territory_for_key[(area.upper() if area != "Total" else area, subzone)],
                "indicator_id": "SGP_SINGSTAT_" + indicator, "period": "2026-06",
                "value": value if isinstance(value, int) else None,
                "status": "observed" if isinstance(value, int) else "missing",
                "source_id": ZIP_SOURCE, "provenance": "source_reported",
                "population_scope": "Singapore citizens and permanent residents with local addresses, excluding those continuously away for 12 months or longer",
                "measurement_method": "SingStat June 2026 resident population, independently rounded to tens",
                "boundary_version": BOUNDARY,
                "source_locator": f"{source['book']}:{source['sheet']}:E{source['row']} (area={area}, subzone={subzone})"})
    data["observations"].extend(new_observations)
    require(len(new_observations) == 2328, "Adopted observation count changed")

    data["documents"] = [d for d in data["documents"] if not d["id"].startswith("sgp-")]
    data["documents"].extend([
        {"id": "sgp-mp2025-national-written-statement", "territory_id": "SGP",
         "category": "plan", "kind": "statutory_master_plan",
         "title": "URA Master Plan 2025 Written Statement",
         "url": sources["sgp-ura-mp2025-written-statement"]["url"],
         "source_id": "sgp-ura-mp2025-written-statement", "period": "2025",
         "availability": "body_acquired", "official_status": "adopted_effective_2025",
         "official_evidence": {"source_id": "sgp-ura-mp2025-gazette-location",
             "locator": "URA circular PPG25-12; officially gazetted 1 December 2025",
             "checked_at": "2026-09-27", "authority": "URA"},
         "territory_match": {"territory_id": "SGP", "country_id": "SGP",
             "type": "country", "code_system": "World Bank economy code",
             "official_code": None, "boundary_version": BOUNDARY,
             "method": "National statutory land-use Master Plan; not a separate subzone local-government plan.",
             "source_id": "sgp-ura-mp2025-gazette-location",
             "locator": "URA circular PPG25-12", "checked_at": "2026-09-27"},
         "note": "26-page written statement acquired. Definitions and area list checked; full zoning/implementation content not semantically audited. Plan is national, not evidence of 55 independent local plans or of spending/evaluation."},
        {"id": "sgp-cop2020-historical-reference", "territory_id": "SGP",
         "category": "reference", "kind": "census_reference",
         "title": "Census of Population 2020 planning-area/subzone age-sex table (MP2019 edition)",
         "url": sources["sgp-singstat-cop2020-age-sex"]["url"],
         "source_id": "sgp-singstat-cop2020-age-sex", "period": "2020",
         "availability": "body_acquired", "official_status": "unverified",
         "note": "All 388 rows and 60 numeric columns inventoried. MP2019 geography is held as a historical edition and is not assigned to MP2025 dashboard territories; cross-edition population change is not calculated."}
    ])
    data["planning"] = {"title": "Master Plan and source evidence",
        "purpose": "Use a selected planning area/subzone for the June 2026 resident baseline; review URA's national statutory Master Plan and obtain location-specific implementation, budget and evaluation evidence separately.",
        "system": {"label": "Singapore statutory national land-use Master Plan",
            "scope": "Planning Act 1998 and URA's gazetted Master Plan 2025. URA planning areas/subzones are geographic units, not independent municipal planning authorities.",
            "cycle": "URA describes a five-year Master Plan review; MP2025 gives medium-term 10–15 year land-use direction, not a fixed local budget cycle.",
            "source_ids": ["sgp-planning-act-location", "sgp-ura-master-plan-overview-location", "sgp-ura-mp2025-gazette-location"]},
        "sections": [{"id": k, "label": v} for k, v in (
            ("plan", "Statutory Master Plan"), ("budget", "Budget evidence"),
            ("implementation", "Implementation and spending"),
            ("evaluation", "Official evaluation"),
            ("reference", "Census and legal references"))]}
    comparisons = [{"parent_id": "SGP", "member_ids": [area_id(c) for c in sorted(by_code)],
        "label": "55 URA MP2025 planning areas",
        "membership_note": "All 55 planning-area names match SingStat June 2026. URA MP2025 indicative area polygons are used for display; source cells are direct, independently rounded, and nil/negligible cells are missing. No sum is used as national population.",
        "source_ids": [ZIP_SOURCE, AREA_SOURCE]}]
    for code in sorted(by_code):
        members = [subzone_id(code, s) for a, s in sorted(subzone_keys)
                   if by_name[a]["properties"]["PLN_AREA_C"] == code]
        comparisons.append({"parent_id": area_id(code), "member_ids": members,
            "label": "SingStat June 2026 subzones of " + by_code[code]["properties"]["PLN_AREA_N"].title(),
            "membership_note": "Source planning-area/subzone name pairs only. MP2025 subzone official codes/polygons are pending; all values are direct cells, not area totals derived from subzones.",
            "source_ids": [ZIP_SOURCE]})
    data["analysis"]["comparisons"] = comparisons
    data["analysis"]["terminal_territory_ids"] = [subzone_id(by_name[a]["properties"]["PLN_AREA_C"], s) for a, s in sorted(subzone_keys)]
    data["analysis"]["default_indicator_id"] = "SGP_SINGSTAT_RESIDENT_TOTAL"
    data["analysis"]["population_context"] = {"primary_indicator_id": "SGP_SINGSTAT_RESIDENT_TOTAL",
        "reference_indicator_id": "SP.POP.TOTL"}
    data["analysis"]["latest_values_only"] = True
    data["collection"]["status"] = "partial"
    data["collection"]["adapters"] = list(dict.fromkeys([
        *data["collection"].get("adapters", []), "singstat-population-trends2026-mp2025-2026-09-27"]))
    data["collection"]["notes"] = [
        "Seven official originals hash-pinned: SingStat 2026 ZIP/report, URA MP2025 planning-area/region GeoJSON and written statement, 2020 census CSV and MP2019 subzone GeoJSON. All 3 ZIP workbooks/7 sheets and the 2020 60 value columns inventoried; only six direct 2026 indicators are adopted.",
        "SingStat 2026 uses MP2025 planning areas. The 2020 census table uses MP2019. The five URA MP2025 regions are retained in the raw inventory but are not shown as population nodes because no direct 2026 regional cells were acquired.",
        "The 55 MP2025 planning-area polygons are indicative display geometry. Subzone codes/MP2025 geometry and cross-edition reconciliation are pending. Never infer zero from '-' (nil/negligible), sum independently rounded cells into a parent, or compare WDI total population with SingStat residents as a same-population change.",
        "URA national Master Plan 2025 was gazetted 1 December 2025. Area/subzone-specific plan content, budget, actual spending and official evaluation have not been acquired or assessed."]
    data["gaps"] = [g for g in data["gaps"] if g["category"] not in
                    {"boundary_reconciliation", "subnational_statistics", "planning_documents"}]
    data["gaps"].extend([
        {"category": "boundary_reconciliation", "status": "partial",
         "detail": "55 MP2025 planning-area source names/code polygons match. No MP2025 subzone code/polygon or MP2019-to-MP2025 correspondence has been established; prior 2016 provider polygons were removed.",
         "next_action": "Acquire URA MP2025 subzone roster/code/geometry and a documented cross-edition correspondence before mapping subzones or measuring 2020–2026 area change."},
        {"category": "subnational_statistics", "status": "partial",
         "detail": "Six direct June 2026 resident-count indicators adopted; remaining age, dwelling, floor-area and 2020 census cells remain priority_unassessed. Nil/negligible cells are explicit missing.",
         "next_action": "Semantically audit remaining table dimensions, footnotes, denominators and subzone crosswalk; add only validated indicators."},
        {"category": "planning_documents", "status": "partial",
         "detail": "National URA MP2025 Written Statement and gazette date established. Planning-area implementation, budget, actual expenditure and official evaluation evidence not acquired.",
         "next_action": "Identify applicable URA/agency plans and location-linked budget/implementation/evaluation evidence without making planning areas into local authorities."}
    ])
    dashboard.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return len(new_observations)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    paths = pinned_sources(args.project, manifest)
    zip_path = paths[ZIP_SOURCE]
    inventory, selected = inspect_workbooks(zip_path)
    history = historical_inventory(paths)
    pdf_pages = {ident: len(PdfReader(str(paths[ident])).pages) for ident in
                 ("sgp-singstat-population-trends-2026-report", "sgp-ura-mp2025-written-statement")}
    count = build(args.project, manifest, paths, inventory, selected, history)
    audit = {"country_id": "SGP", "candidate_status": "partial_unpublished",
        "retrieved_at": manifest["retrieved_at"],
        "originals": [{k: item[k] for k in ("id", "raw_path", "bytes", "sha256")} for item in manifest["sources"]],
        "pdf_pages": pdf_pages, "workbooks": inventory, "historical_2020": history,
        "adopted_indicators": list(INDICATORS), "adopted_slots": count,
        "adopted_by_indicator": {indicator: {
            "observed": sum(isinstance(row["value"], int) for row in selected[indicator].values()),
            "missing_nil_or_negligible": sum(row["value"] == "-" for row in selected[indicator].values())}
            for indicator in INDICATORS},
        "observed_slots": sum(isinstance(x["value"], int) for rows in selected.values() for x in rows.values()),
        "missing_nil_or_negligible_slots": sum(x["value"] == "-" for rows in selected.values() for x in rows.values()),
        "total_population_geography": {
            "planning_area_observed": sum(isinstance(row["value"], int) for (area,subzone),row in selected["RESIDENT_TOTAL"].items() if area != "Total" and subzone == "Total"),
            "planning_area_missing": sum(row["value"] == "-" for (area,subzone),row in selected["RESIDENT_TOTAL"].items() if area != "Total" and subzone == "Total"),
            "subzone_observed": sum(isinstance(row["value"], int) for (area,subzone),row in selected["RESIDENT_TOTAL"].items() if subzone != "Total"),
            "subzone_missing": sum(row["value"] == "-" for (area,subzone),row in selected["RESIDENT_TOTAL"].items() if subzone != "Total"),
            "numeric_planning_area_sum_not_adopted": sum(row["value"] for (area,subzone),row in selected["RESIDENT_TOTAL"].items() if area != "Total" and subzone == "Total" and isinstance(row["value"], int))},
        "source_tieouts": {"national_residents_total": selected["RESIDENT_TOTAL"][("Total", "Total")]["value"],
            "national_residents_male": selected["RESIDENT_MALE"][("Total", "Total")]["value"],
            "national_residents_female": selected["RESIDENT_FEMALE"][("Total", "Total")]["value"],
            "tampines_residents": selected["RESIDENT_TOTAL"][("Tampines", "Total")]["value"]},
        "adoption_rules": ["2026 SingStat planning-area names exactly match all 55 URA MP2025 PLN_AREA_N values.",
            "2020 MP2019 values and polygons are separate historical evidence; no 2025 territory assignment.",
            "MP2025 subzones have direct 2026 source rows but no verified code or polygon.",
            "Dash means nil/negligible and is missing, not zero; rounded child sums are not parent estimates.",
            "Unselected age/dwelling/floor area dimensions are priority_unassessed, not rejected or zero."],
        "independent_audit": "not_done", "hosting": "not_deployed", "public": "not_deployed"}
    path = args.project / "evidence/SGP_POPULATION2026_AUDIT.json"
    path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported 388 territories, 55 MP2025 polygons, 6 indicators, {count} source slots; audit {path}")


if __name__ == "__main__":
    main()
