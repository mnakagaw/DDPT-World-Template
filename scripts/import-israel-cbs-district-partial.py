"""Adopt a checked 2022 CBS nation/six-district subset; keep other fields unassessed."""

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FILE = "raw/israel-cbs-census-2022/broad-geographical-units.xlsx"
SHA256 = "79a967d1e72127e4ea01296d77a32c2bd7d6e9fca132ac2b2029b1ac7c7d866d"
SHEET = "District Sub-District Natural  "
SOURCE_ID = "isr-cbs-census-2022-broad-geographical"
NAMES = {
    1: ("Jerusalem District", "Jerusalem District"),
    2: ("Northern District", "Northern District"),
    3: ("Haifa District", "Haifa"),
    4: ("Central District", "Central District"),
    5: ("Tel Aviv District", "Tel Aviv"),
    6: ("Southern District", "Southern District"),
    7: ("Judea and Samaria Area", None),
}
# Excel column, CBS machine field, ID suffix, theme, unit, denominator/meaning.
FIELDS = (
    (8, "pop_approx", "POP", "2022 census estimate population", "Population", "people", "CBS 2022 census population estimate, rounded by the source."),
    (9, "DependencyRatio", "DEPENDENCY", "Dependency ratio per 1,000", "Demography", "per 1,000", "Source-defined demographic dependency ratio per 1,000 persons."),
    (10, "Foreign_pcnt", "FOREIGN", "Foreign population", "Demography", "percent", "Percent of population classified by CBS as foreign."),
    (200, "Acadm1Cert_pcnt", "FIRST_DEGREE", "Age 15+ with a first academic degree", "Education", "percent", "Percent of persons aged 15 and over with a first academic degree."),
    (219, "walk5_pcnt", "WALK_DIFFICULTY", "Age 5+ with great walking or stair difficulty", "Functioning", "percent", "Percent of persons aged 5 and over reporting great difficulty walking around the house or using stairs."),
    (235, "WrkY_pcnt", "WORKED_LAST_YEAR", "Age 15+ who worked in the last 12 months", "Employment", "percent", "Percent of persons aged 15 and over who worked during the last 12 months; not an unemployment rate."),
    (385, "employeesAnnual_medWage", "EMPLOYEE_MEDIAN_WAGE", "Annual median employee wage", "Wages", "NIS/year", "Median annual wage for employees, as reported by CBS; not household income."),
    (407, "hh_total_approx", "HOUSEHOLDS", "Households, approximate", "Households", "households", "Approximate number of households; do not equate it with dwellings."),
    (438, "HousingDens2_pcnt", "CROWDED_HOUSING", "Households with two or more persons per room", "Housing", "percent", "Percent of households residing at a housing density of two or more persons per room."),
    (440, "rent_pcnt", "RENTED_HOUSING", "Households in rented dwellings", "Housing", "percent", "Percent of households residing in rented dwellings."),
    (450, "Vehicle1up_pcnt", "HOUSEHOLD_CAR", "Households with access to at least one car", "Transport access", "percent", "Percent of households with at least one car at their disposal."),
    (454, "Computer_avg", "HOUSEHOLD_COMPUTERS", "Personal computers per household", "Digital access", "computers/household", "Average number of personal computers per household."),
)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True, type=Path)
    args = parser.parse_args()
    project = args.project.resolve(strict=True)
    dataset_path = project / "data/dashboard.json"
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    if dataset["country"]["id"] != "ISR" or any(s["id"] == SOURCE_ID for s in dataset["sources"]):
        raise ValueError("Expected unimproved Israel candidate")
    target = project / FILE
    receipt = json.loads((project / (FILE + ".receipt.json")).read_text(encoding="utf-8"))
    actual = hashlib.sha256(target.read_bytes()).hexdigest()
    if actual != SHA256 or receipt["sha256"] != actual or receipt["path"] != FILE or receipt["status"] != "acquired":
        raise ValueError("CBS original/receipt does not match audited edition")
    workbook = load_workbook(target, read_only=True, data_only=True)
    sheet = workbook[SHEET]
    headers = list(sheet.iter_rows(min_row=2, max_row=3, values_only=True))
    for col, code, *_ in FIELDS:
        if headers[1][col - 1] != code or not headers[0][col - 1]:
            raise ValueError(f"CBS column {col} changed")
    nation = None
    source_rows = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
        if row[0] == "Nationwide":
            nation = (row_number, row)
        elif isinstance(row[0], int) and row[1] is None and row[2] is None and row[3]:
            code = row[0]
            if code not in NAMES or row[3] != NAMES[code][0] or code in source_rows:
                raise ValueError(f"Unexpected CBS top-level area at row {row_number}")
            source_rows[code] = (row_number, row)
    if nation is None or set(source_rows) != set(NAMES):
        raise ValueError("CBS top-level coverage changed")
    if nation[1][7] != 9_601_720 or source_rows[7][1][7] != 481_940 or \
            sum(record[1][7] for record in source_rows.values()) != nation[1][7]:
        raise ValueError("CBS national/area population controls changed")
    workbook.close()

    regions = {area["name"]: area for area in dataset["territories"] if area["parent_id"] == "ISR"}
    if set(regions) != {name for _, name in NAMES.values() if name}:
        raise ValueError("Reference district coverage changed")
    now = datetime.now(timezone.utc).isoformat()
    dataset["sources"].append({"id": SOURCE_ID, "name": "CBS 2022 Census summarized broad geographical units",
        "url": receipt["source_url"], "publisher": "Central Bureau of Statistics (Israel)",
        "reference_period": "April 2022 Census estimate; Excel linked from 2025 source catalogue",
        "geographic_level": "Nationwide, six districts, and separately reported Judea and Samaria Area",
        "status": "ready", "retrieved_at": receipt["retrieved_at"], "sha256": actual,
        "raw_path": FILE, "license": "terms_review_required",
        "note": "Nationwide includes the separately reported Judea and Samaria Area. Only six district rows are matched by name to 2017 geoBoundaries provider polygons. The seventh area row is retained in the source audit and is not adopted as a mapped administrative district. CBS census estimates differ from WDI population."})
    crosswalk = []
    for code, (source_name, provider_name) in NAMES.items():
        area = regions.get(provider_name) if provider_name else None
        crosswalk.append((code, source_name, provider_name or "", area["id"] if area else "",
                          "name_correspondence_only; legal_boundary_equivalence_unverified" if area else
                          "source_only; not_one_of_six_districts; no_matched_polygon"))
        if area:
            area["source_name_en"] = source_name
            area["reconciliation_status"] = "cbs_2022_code_and_name_known_provider_2017_boundary_unverified"
    observations_added = 0
    selected_rows = [("ISR", *nation)] + [(regions[NAMES[code][1]]["id"], *source_rows[code]) for code in range(1, 7)]
    for col, field_code, suffix, title, theme, unit, meaning in FIELDS:
        indicator_id = "ISR_CBS_2022_" + suffix
        dataset["indicators"].append({"id": indicator_id, "name": title, "theme": theme,
            "unit": unit, "definition": meaning + " Source table: " + headers[0][col - 1] +
                " The published nationwide value includes the separately reported Judea and Samaria Area (CBS code 7); the six mapped district rows do not exhaust the nationwide value.",
            "population": "CBS 2022 Census reported population/households as specified by the source column",
            "source_id": SOURCE_ID, "aggregation": "none", "measurement_method": "source_reported",
            "series_family": "census", "display_role": "primary", "period_policy": "latest_available_per_indicator",
            "display_decimals": 0 if unit in ("people", "households", "NIS/year") else 1})
        for territory_id, row_number, row in selected_rows:
            value = row[col - 1]
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise ValueError(f"Nonnumeric adopted field: {field_code}, Excel row {row_number}")
            dataset["observations"].append({"territory_id": territory_id, "indicator_id": indicator_id,
                "period": "2022", "value": value, "status": "observed", "measurement_method": "source_reported",
                "source_id": SOURCE_ID, "source_locator": f"{SHEET.strip()}!{get_column_letter(col)}{row_number}; {field_code}; CBS district code {row[0] if territory_id != 'ISR' else 'Nationwide'}"})
            observations_added += 1
    dataset["country"]["geography_note"] = (
        "CBS 2022 nationwide estimate includes six districts plus its separately reported Judea and Samaria Area (481,940). Only six districts are shown with 2017 geoBoundaries provider reference shapes, matched by name; current legal boundary equivalence is unverified. The seventh CBS reporting area is retained in source evidence, not shown as an administrative district or mapped polygon. WDI population is a separate series with its own definition.")
    for gap in dataset["gaps"]:
        if gap["category"] == "boundary_reconciliation":
            gap.update(status="partial", detail="Six CBS 2022 district codes/names matched by name to six 2017 provider polygons. The separately reported Judea and Samaria Area (CBS code 7) has no matched polygon and is not one of the six districts.",
                       next_action="Acquire dated official district, local-authority and statistical-area boundaries and reconcile territorial scope before adding lower-unit polygons.")
        elif gap["category"] == "subnational_statistics":
            gap.update(status="partial", detail="Twelve checked CBS 2022 fields were adopted for nation and six districts. The separately reported Judea and Samaria Area and other workbook fields remain source-only/unassessed; locality and statistical-area products were acquired but not yet adopted.",
                       next_action="Review all 14 acquired sheets and each numeric column; join official locality codes, classifications and boundary editions; expand priority themes with comparable definitions.")
        elif gap["category"] == "planning_documents":
            gap.update(status="not_collected", detail="Planning law, XPLAN and local-authority audited finance locations identified, but no plan/budget/execution/evaluation document has been linked to a selected district or local planning authority.",
                       next_action="Inspect the current Planning and Building Law, XPLAN records and local committee geography; acquire actual plans and municipal financial records by code.")
    dataset["collection"]["adapters"] = sorted(set(dataset["collection"]["adapters"] + ["israel-cbs-2022-six-district-partial"]))
    dataset["analysis"]["latest_values_only"] = True
    dataset["generated_at"] = now
    evidence = project / "evidence"
    with (evidence / "ISR_CBS_2022_DISTRICT_CROSSWALK.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("cbs_code", "cbs_name", "provider_name", "provider_territory_id", "decision"))
        writer.writerows(crosswalk)
    audit = {"status": "partial_candidate_not_accepted", "checked_at": now, "raw_sha256": actual,
             "adopted_columns": [{"excel_column": col, "field_code": field_code, "indicator_id": "ISR_CBS_2022_" + suffix}
                                 for col, field_code, suffix, *_ in FIELDS],
             "adopted_rows": ["Nationwide", *(NAMES[code][0] for code in range(1, 7))],
             "observations_added": observations_added,
             "source_only_area": {"cbs_code": 7, "name": NAMES[7][0], "population_2022": source_rows[7][1][7]},
             "six_district_population_sum": sum(source_rows[code][1][7] for code in range(1, 7)),
             "nationwide_population": nation[1][7],
             "all_14_sheets_and_numeric_columns_decided": False,
             "unresolved": ["Official legal boundary equivalence", "Seventh CBS area display and scope policy",
                            "Locality, municipal and statistical-area products", "Remaining fields across six acquired workbooks",
                            "Planning institutions, plan/finance originals and independent acceptance"]}
    (evidence / "ISR_CBS_2022_IMPORT_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    dataset_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"fields": len(FIELDS), "districts": 6, "observations_added": observations_added,
                      "six_district_population_sum": audit["six_district_population_sum"],
                      "nationwide_population": audit["nationwide_population"],
                      "source_only_area_population": audit["source_only_area"]["population_2022"]}))


if __name__ == "__main__":
    main()
