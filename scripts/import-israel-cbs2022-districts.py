"""Import a narrow, direct-value Israel CBS 2022 Census reporting-area edition.

The source has seven first-level rows, including a separate Judea and Samaria
reporting area. It is not a six-district administrative polygon join. Other
acquired CBS workbooks and most fields of the broad-geography file await audit.
"""

import argparse
import json
import runpy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


_fetch = runpy.run_path(str(Path(__file__).with_name("fetch-israel-cbs-census.py")))
FILES, EXPECTED_SHA256, hash_file, BASE = (_fetch[key] for key in ("FILES", "EXPECTED_SHA256", "hash_file", "BASE"))
SOURCE = "isr-cbs-census2022-broad-geographical-units"
OTHER_SOURCES = {
    "selected-data-localities-statistical-areas.xlsx": "isr-cbs-census2022-selected-localities-statistical-areas",
    "population-group-religion-age-sex.xlsx": "isr-cbs-census2022-age-sex-national",
    "population-households-locality.xlsx": "isr-cbs-census2022-population-households-locality",
    "population-selected-characteristics.xlsx": "isr-cbs-census2022-population-selected-characteristics",
    "households-selected-characteristics.xlsx": "isr-cbs-census2022-households-selected-characteristics",
}
OTHER_REFERENCE_PERIODS = {
    "population-selected-characteristics.xlsx": "2022 Census estimate, end-2021 population estimate and 2008 Census estimate; fields unassessed",
    "households-selected-characteristics.xlsx": "2022 Census estimate and 2008 Census estimate; fields unassessed",
}
PLAN_CATALOG = "isr-planning-administration-xplan"
PLAN_LAW_INDEX = "isr-planning-building-law-index"
PLAN_GUIDANCE = "isr-planning-mavat-guidance"
INDICATORS = {
    8: ("ISR_CBS22_POP_APPROX", "2022 Census population estimate", "Population", "people, rounded to 10", "CBS 2022 Census reported population estimate, rounded to tens. The reporting universe includes CBS area 7 for Israeli localities in Judea and Samaria; not all residents of that area.", "CBS reporting population", 0),
    49: ("ISR_CBS22_MEDIAN_AGE", "Median age, 2022 Census", "Demography", "years", "Median age for the population covered by the CBS 2022 Census reporting row.", "CBS reporting population", 0),
    200: ("ISR_CBS22_FIRST_DEGREE_15P_PCT", "First academic degree, age 15+, 2022", "Education", "percent", "Source percentage of persons aged 15 and over with a first academic degree; not all tertiary education and not a school enrolment rate.", "Persons aged 15 and over", 1),
    219: ("ISR_CBS22_WALK_DIFFICULTY_5P_PCT", "Great difficulty walking, age 5+, 2022", "Daily functioning", "percent", "Source percentage of persons aged 5 and over with great difficulty walking around the house or going up/down stairs; this is not a general disability rate.", "Persons aged 5 and over", 1),
    407: ("ISR_CBS22_HOUSEHOLDS_APPROX", "Households, 2022 Census estimate", "Households", "households, rounded to 10", "CBS 2022 Census approximate household total; the definition is retained separately from population and all dwelling counts.", "Households in CBS reporting population", 0),
    438: ("ISR_CBS22_HOUSING_DENSITY_2PLUS_PCT", "Households with at least 2 persons per room, 2022", "Housing", "percent", "Source percentage of households residing at a housing density of 2 persons or more per room.", "Households", 1),
    439: ("ISR_CBS22_OWNER_OCCUPIED_PCT", "Households in self-owned dwellings, 2022", "Housing", "percent", "Source percentage of households residing in self-owned dwellings; this is a household share, not a dwelling count.", "Households", 1),
    450: ("ISR_CBS22_CAR_ACCESS_PCT", "Households with at least one car, 2022", "Mobility", "percent", "Source percentage of households with at least one car at their disposal; this does not measure public transport access.", "Households", 1),
}
SOURCE_ANCHORS = [
    (5, 1, "Jerusalem District", 1252940),
    (9, 2, "Northern District", 1519660),
    (38, 3, "Haifa District", 1104670),
    (46, 4, "Central District", 2319540),
    (59, 5, "Tel Aviv District", 1502610),
    (64, 6, "Southern District", 1420360),
    (78, 7, "Judea and Samaria Area", 481940),
]


def main(project):
    raw = project / "raw" / "official"
    evidence = project / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    for name, expected in EXPECTED_SHA256.items():
        if hash_file(raw / name) != expected:
            raise ValueError(f"Changed CBS original: {name}")
    book = load_workbook(raw / "broad-geographical-units.xlsx", read_only=True, data_only=True)
    if len(book.sheetnames) != 9:
        raise ValueError("CBS broad workbook sheet list changed")
    sheet = book["District Sub-District Natural  "]
    if sheet.max_row != 80 or sheet.max_column != 454:
        raise ValueError("CBS district sheet dimensions changed")
    headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    field_ids = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    expected_ids = {8: "pop_approx", 49: "age_median", 200: "Acadm1Cert_pcnt", 219: "walk5_pcnt",
                    407: "hh_total_approx", 438: "HousingDens2_pcnt", 439: "own_pcnt", 450: "Vehicle1up_pcnt"}
    if any(field_ids[col-1] != expected for col, expected in expected_ids.items()):
        raise ValueError("CBS selected field IDs changed")
    source_rows = [4] + [row for row, _, _, _ in SOURCE_ANCHORS]
    values = {row: next(sheet.iter_rows(min_row=row, max_row=row, values_only=True)) for row in source_rows}
    if values[4][0] != "Nationwide" or values[4][7] != 9601720:
        raise ValueError("CBS nationwide source anchor changed")
    for row, code, name, population in SOURCE_ANCHORS:
        data = values[row]
        if (data[0], data[1], data[2], data[3], data[7]) != (code, None, None, name, population):
            raise ValueError(f"CBS reporting row changed: {row}")
    if sum(population for _, _, _, population in SOURCE_ANCHORS) != values[4][7]:
        raise ValueError("Seven source rows do not reconcile with CBS nationwide value")
    output = project / "data" / "dashboard.json"
    dashboard = json.loads(output.read_text(encoding="utf-8"))
    if dashboard["country"]["id"] != "ISR":
        raise ValueError("Expected Israel project")
    own_sources = {SOURCE, *OTHER_SOURCES.values(), PLAN_CATALOG, PLAN_LAW_INDEX, PLAN_GUIDANCE}
    own_indicators = {value[0] for value in INDICATORS.values()}
    dashboard["territories"] = [area for area in dashboard["territories"] if area["id"] == "ISR"]
    if len(dashboard["territories"]) != 1:
        raise ValueError("Unexpected bootstrap territory registry")
    dashboard["observations"] = [row for row in dashboard["observations"] if row["indicator_id"] not in own_indicators]
    dashboard["indicators"] = [item for item in dashboard["indicators"] if item["id"] not in own_indicators]
    dashboard["sources"] = [item for item in dashboard["sources"] if item["id"] not in own_sources]
    dashboard["boundaries"] = {"type": "FeatureCollection", "features": []}
    district_ids = []
    for _, code, name, _ in SOURCE_ANCHORS:
        area_id = f"ISR:CBS22:{'AREA' if code == 7 else 'DIST'}:{code}"
        district_ids.append(area_id)
        dashboard["territories"].append({
            "id": area_id, "name": name, "level": "cbs_reporting_area", "type": "CBS first-level reporting unit",
            "source_area_subtype": "Israeli localities in Judea and Samaria only" if code == 7 else "district",
            "parent_id": "ISR", "official_code": str(code), "code_system": "CBS 2022 Census DistrictCode reporting classification",
            "boundary_version": None, "source_id": SOURCE,
            "reconciliation_status": "CBS area 7 covers Israeli localities only; no legal status or 2022 polygon inferred" if code == 7 else "CBS district code and row only; no verified 2022 polygon or legal planning-unit correspondence",
        })
    for col, (indicator_id, name, theme, unit, definition, population, decimals) in INDICATORS.items():
        dashboard["indicators"].append({"id": indicator_id, "name": name, "theme": theme, "unit": unit,
            "definition": definition, "population": population, "source_id": SOURCE, "aggregation": "none",
            "measurement_method": "source_reported", "display_decimals": decimals})
        for idx, row in enumerate(source_rows):
            source_value = values[row][col-1]
            if not isinstance(source_value, (int, float)):
                raise ValueError(f"Selected CBS cell not numeric: {row} {get_column_letter(col)}")
            area_id = "ISR" if idx == 0 else district_ids[idx-1]
            value = float(source_value)
            if col in (8, 407) and value != int(value):
                raise ValueError("CBS rounded count is not an integer")
            if col in (200, 219, 438, 439, 450) and not 0 <= value <= 100:
                raise ValueError("CBS share outside 0–100")
            dashboard["observations"].append({"territory_id": area_id, "indicator_id": indicator_id,
                "period": "2022", "value": int(value) if col in (8, 407) else value,
                "status": "observed", "measurement_method": "source_reported", "source_id": SOURCE,
                "source_locator": f"CBS 2022 broad geographical units, sheet District Sub-District Natural, Excel row {row}, column {get_column_letter(col)} ({field_ids[col-1]})"})
    if len(district_ids) != 7:
        raise ValueError("Incomplete CBS first-level reporting set")
    dashboard["analysis"]["comparisons"] = [{"parent_id": "ISR", "member_ids": district_ids,
        "label": "Seven CBS 2022 Census first-level reporting rows",
        "membership_note": "The official CBS worksheet lists these seven direct rows; area 7 reports Israeli localities in Judea and Samaria only. Population rows reconcile with its nationwide CBS value. Percentages and medians remain direct source observations, never summed or averaged. This reporting set is not the six-shape 2006 geoBoundaries edition or a legal planning-unit classification.",
        "source_ids": [SOURCE]}]
    dashboard["analysis"]["default_indicator_id"] = INDICATORS[8][0]
    dashboard["analysis"]["latest_values_only"] = True
    dashboard["analysis"].pop("population_context", None)
    dashboard["country"]["geography_note"] = "CBS 2022 Census reports six districts plus Judea and Samaria Area (Israeli localities only). The 2006 six-part geoBoundaries reference is not joined to these official source rows; no 2022 local polygon or legal planning-unit equivalence is verified. WDI national values have their own date and geography definition."
    for filename, tail in FILES.items():
        url_id = SOURCE if filename == "broad-geographical-units.xlsx" else OTHER_SOURCES[filename]
        dashboard["sources"].append({"id": url_id,
            "name": "CBS 2022 Census broad geographical units" if url_id == SOURCE else f"CBS 2022 Census {filename.replace('.xlsx','').replace('-', ' ')}",
            "url": BASE + tail, "publisher": "Israel Central Bureau of Statistics",
            "reference_period": OTHER_REFERENCE_PERIODS.get(filename, "2022 Census, April 2022 reference"),
            "geographic_level": "varies by sheet" if url_id == SOURCE else "varies; unassessed",
            "status": "partial" if url_id == SOURCE else "partial", "license": "terms_review_required",
            "raw_path": "raw/official/" + filename, "sha256": EXPECTED_SHA256[filename],
            "retrieved_at": datetime.fromtimestamp((raw / filename).stat().st_mtime, timezone.utc).isoformat(),
            "note": "Eight fields adopted only from the 8 nationwide/first-level rows of one broad-geography sheet; remaining 446 columns and other eight sheets unassessed." if url_id == SOURCE else "Official file acquired and sheet structure inventoried; no numeric fields adopted pending table-level audit."})
    dashboard["sources"] += [
        {"id": PLAN_CATALOG, "name": "Planning Administration XPLAN online plan search",
         "url": "https://www.gov.il/en/service/searching-plans-submitted-planning-institutes-in-xplan-site",
         "publisher": "Israel Planning Administration", "reference_period": "service page checked 2026-09-26",
         "geographic_level": "plans and plots; selected-area match unverified", "status": "not_collected",
         "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required", "note": "Official plan-search service location only. No individual plan body, statutory status or local budget adopted."},
        {"id": PLAN_LAW_INDEX, "name": "Planning Administration planning and building laws index",
         "url": "https://www.gov.il/he/pages/planning_building_laws?chapterindex=0",
         "publisher": "Israel Planning Administration", "reference_period": "index checked 2026-09-26",
         "geographic_level": "national law and regulations", "status": "not_collected",
         "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required", "note": "Law index lead. Current consolidated duties and legal planning unit have not been verified."},
        {"id": PLAN_GUIDANCE, "name": "Planning Administration Mavat plan submission guidance",
         "url": "https://www.gov.il/he/pages/mavat_2020",
         "publisher": "Israel Planning Administration", "reference_period": "guide page last updated 2025-07-01",
         "geographic_level": "planning submissions; applicable unit to verify", "status": "not_collected",
         "retrieved_at": datetime.now(timezone.utc).isoformat(),
         "license": "terms_review_required", "note": "Official procedural guidance lead. A plan template or approved selected-area plan has not been acquired."},
    ]
    dashboard["gaps"] = [g for g in dashboard["gaps"] if g["category"] not in ("subnational_statistics", "boundary_reconciliation", "planning_documents")]
    dashboard["gaps"] += [
        {"category": "subnational_statistics", "status": "partial",
         "detail": "Eight direct 2022 Census indicators in 7 CBS first-level rows and nationwide only. Six acquired workbooks have 14 sheets; only 8 of 454 columns in one sheet are adopted. Locality/statistical-area and other census fields remain unassessed.",
         "next_action": "Audit every table and numeric field, crosswalk locality codes and municipal status, then adopt compatible local themes with full coverage and explicit denominators."},
        {"category": "boundary_reconciliation", "status": "partial",
         "detail": "CBS reports 7 first-level rows including Israeli localities in Judea and Samaria; 2006 geoBoundaries has 6 shapes. No official current polygons or legal planning geography joined.",
         "next_action": "Obtain official 2022 district/local-authority/statistical-area geometry and code histories; verify exact scopes before any map join."},
        {"category": "planning_documents", "status": "not_collected",
         "detail": "XPLAN, planning-law index and Mavat guidance locations found. Current law, local plan bodies, approval, budgets, implementation and evaluation not acquired.",
         "next_action": "Determine the applicable local planning authority and retrieve actual selected-area plan, budget and official status evidence."},
    ]
    dashboard["collection"]["adapters"] = sorted(set(dashboard["collection"]["adapters"] + ["israel-cbs2022-first-level-partial-v1"]))
    dashboard["generated_at"] = datetime.now(timezone.utc).isoformat()
    output.write_text(json.dumps(dashboard, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    other_sheets = []
    for filename in FILES:
        workbook = book if filename == "broad-geographical-units.xlsx" else load_workbook(raw / filename, read_only=True, data_only=True)
        for tab in workbook:
            other_sheets.append({"file": filename, "sheet": tab.title, "rows": tab.max_row,
                "columns": tab.max_column, "decision": "eight_selected_fields_partial" if filename == "broad-geographical-units.xlsx" and tab.title == sheet.title else "unassessed"})
    audit = {"source_files": {name: {"sha256": digest, "bytes": (raw/name).stat().st_size} for name,digest in EXPECTED_SHA256.items()},
        "sheet_inventory": other_sheets,
        "district_sheet_fields": [{"excel_column": get_column_letter(i), "field_id": field_ids[i-1], "source_label": headers[i-1],
            "decision": "adopted_8_direct_rows" if i in INDICATORS else "unassessed_not_adopted"} for i in range(1,455)],
        "row_inventory": [{"excel_row": row,"source_code": code,"name": name,"population": pop} for row,code,name,pop in SOURCE_ANCHORS],
        "reference_shape_gap": "Six 2006 provider shapes are not joined to seven CBS 2022 source rows",
        "historical_finer_rows": "16 sub-district and 52 natural-area rows are identified but not adopted; published child values do not completely cover every parent",
        "scope_note": "CBS area 7 is Israeli localities in Judea and Samaria, not all territory residents; CBS yearbook table 2.17 footnote 8."}
    (evidence / "ISR_CBS2022_FIELD_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"reporting_areas": 7, "local_indicators": len(INDICATORS), "domestic_observations": 8*len(INDICATORS),
        "acquired_workbook_sheets": len(other_sheets), "first_level_population_sum": sum(x[3] for x in SOURCE_ANCHORS)}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True)
    main(Path(parser.parse_args().project))
