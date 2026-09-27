"""Inventory the DOP demographic appendix and cross-check Table A-1 against the Union Report."""

import argparse
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
NAMES = ("Kachin", "Kayah", "Kayin", "Chin", "Sagaing", "Tanintharyi", "Bago",
         "Magway", "Mandalay", "Mon", "Rakhine", "Yangon", "Shan", "Ayeyawady",
         "Nay Pyi Taw")
FIELDS = ("conventional_both", "conventional_male", "conventional_female",
          "institution_both", "institution_male", "institution_female",
          "estimated_both", "estimated_male", "estimated_female",
          "total_both", "total_male", "total_female")
NUMBER = re.compile(r"[0-9][0-9,]*(?:\.[0-9]+)?|-")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def value(raw, locator):
    if raw == "-":
        return 0
    require(isinstance(raw, int) and raw >= 0, f"Non-count at {locator}: {raw!r}")
    return raw


def row_values(row, row_number):
    return {field: value(row[col], f"Table A-1!{get_column_letter(col+1)}{row_number}")
            for col, field in enumerate(FIELDS, 1)}


def pdf_numbers(page_text, name, expected_count):
    lines = [line.strip() for line in page_text.splitlines()
             if line.strip().startswith(name + " ")]
    require(len(lines) == 1, f"Union Report row ambiguous: {name}")
    tokens = NUMBER.findall(lines[0][len(name):])
    require(len(tokens) == expected_count, f"Union Report field count changed: {name}: {tokens}")
    return [0 if token == "-" else float(token.replace(",", ""))
            for token in tokens]


def main(project):
    manifest = json.loads((ROOT / "config/myanmar-census2024-source-manifest.json").read_text(encoding="utf-8"))
    for source in manifest["sources"]:
        path = project / source["raw_path"]
        require(path.is_file() and path.stat().st_size == source["bytes"] and
                sha(path) == source["sha256"], f"Pinned original differs: {source['id']}")
    workbook_path = project / manifest["sources"][0]["raw_path"]
    report_path = project / manifest["sources"][1]["raw_path"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    require(len(workbook.sheetnames) == 19 and
            workbook.sheetnames[0] == "List of Tables" and
            workbook.sheetnames[1] == "Table A-1" and
            workbook.sheetnames[-1] == "Table A-18", "Workbook table list changed")
    inventory = []
    for sheet in workbook:
        numeric = defaultdict(int)
        headers = {}
        field_labels = set()
        for row_number, row in enumerate(sheet.values, 1):
            for col, item in enumerate(row, 1):
                if row_number <= 4 and isinstance(item, str) and item.strip():
                    headers.setdefault(col, item.strip()[:120])
                if isinstance(item, (int, float)) and not isinstance(item, bool):
                    numeric[col] += 1
            if isinstance(row[0], str) and any(isinstance(x, (int, float)) for x in row[1:]):
                field_labels.add(row[0].strip())
        inventory.append({"sheet": sheet.title, "rows": sheet.max_row,
                          "columns": sheet.max_column,
                          "numeric_cells": sum(numeric.values()),
                          "numeric_columns": [{"column": get_column_letter(col),
                                               "numeric_cells": count,
                                               "top_header": headers.get(col)}
                                              for col, count in sorted(numeric.items())],
                          "numeric_row_labels": sorted(field_labels),
                          "decision": "A-1 state/region and UNION twelve source columns adopted; urban/rural used for reconciliation"
                          if sheet.title == "Table A-1" else
                          "metadata_only" if sheet.title == "List of Tables" else
                          "priority_unassessed: table, column population/definition and geography need separate semantic audit"})
    rows = list(workbook["Table A-1"].values)
    require(len(rows) == 52 and rows[4][0] == "UNION", "Table A-1 layout changed")
    report = PdfReader(report_path)
    require(len(report.pages) == 146, "Union Report pagination changed")
    page_22 = report.pages[40].extract_text()
    page_31 = report.pages[43].extract_text()
    require("Table 2.2 Percentage of enumeration" in page_22 and
            "Table 3.1 Population distribution" in page_31 and
            "51,375,327" in page_22, "Union Report cross-check tables changed")
    report_rows = ["Union", *NAMES]
    xlsx_rows = [5, *range(8, 53, 3)]
    areas = []
    for report_name, row_number in zip(report_rows, xlsx_rows, strict=True):
        row = rows[row_number - 1]
        require(row[0].upper() == report_name.upper(), f"A-1 row changed at {row_number}")
        fields = row_values(row, row_number)
        for prefix in ("conventional", "institution", "estimated", "total"):
            require(fields[prefix + "_both"] == fields[prefix + "_male"] + fields[prefix + "_female"],
                    f"Sex total mismatch: {report_name}/{prefix}")
        for suffix in ("both", "male", "female"):
            require(fields["total_" + suffix] ==
                    sum(fields[prefix + "_" + suffix]
                        for prefix in ("conventional", "institution", "estimated")),
                    f"Component mismatch: {report_name}/{suffix}")
        enum = fields["conventional_both"] + fields["institution_both"]
        pdf22 = pdf_numbers(page_22, report_name, 6)
        pdf31 = pdf_numbers(page_31, report_name, 9)
        require(pdf22[:3] == [enum, fields["estimated_both"], fields["total_both"]] and
                pdf31 == [enum,
                          fields["conventional_male"] + fields["institution_male"],
                          fields["conventional_female"] + fields["institution_female"],
                          fields["estimated_both"], fields["estimated_male"], fields["estimated_female"],
                          fields["total_both"], fields["total_male"], fields["total_female"]],
                f"Union Report differs from workbook: {report_name}")
        require(abs(pdf22[3] - round(100 * enum / fields["total_both"], 1)) <= .11 and
                abs(pdf22[4] - round(100 * fields["estimated_both"] / fields["total_both"], 1)) <= .11 and
                pdf22[5] == 100,
                f"Union Report percentages do not reconcile: {report_name}")
        if report_name != "Union":
            for offset in (1, 2):
                sub = row_values(rows[row_number - 1 + offset], row_number + offset)
                require(rows[row_number - 1 + offset][0] in ("Urban", "Rural"),
                        f"A-1 urban/rural label changed: {report_name}")
            for field in FIELDS:
                require(fields[field] == sum(row_values(rows[row_number - 1 + offset], row_number + offset)[field]
                                             for offset in (1, 2)),
                        f"A-1 urban/rural mismatch: {report_name}/{field}")
        areas.append({"name": "Myanmar" if report_name == "Union" else report_name,
                      "source_row": row_number, "fields": fields,
                      "enumerated_both": enum,
                      "enumerated_pct": pdf22[3], "estimated_pct": pdf22[4],
                      "source_dash_is_zero": row[7] == "-",
                      "report_table_2_2_page": 15,
                      "report_table_3_1_page": 18})
    for field in FIELDS:
        require(areas[0]["fields"][field] == sum(area["fields"][field] for area in areas[1:]),
                f"State/region rows do not cover the Union: {field}")
    require([a["name"] for a in areas[1:]] == list(NAMES), "Fifteen reporting areas changed")
    require(areas[0]["enumerated_both"] == 32183599 and
            areas[0]["fields"]["estimated_both"] == 19191728 and
            areas[0]["fields"]["total_both"] == 51375327,
            "Union published population changed")
    result = {"source_sha256": manifest["sources"][0]["sha256"],
              "report_sha256": manifest["sources"][1]["sha256"],
              "sheet_inventory": inventory, "areas": areas,
              "geography": "Source reporting rows without a published official area code in Table A-1; no polygon join",
              "source_dash_decision": "Three areas print a dash for estimated persons and Table 2.2 prints 0.0%; zero is inferred from reconciled total-minus-enumerated, not a numeric A-1 cell",
              "adopted_indicators": 15,
              "adopted_observations": 240,
              "unassessed": "Tables A-2 through A-18, including 2024 district/township and demographic detail; provisional 2024 PDF is a separate edition"}
    out = project / "evidence/MMR_CENSUS2024_AUDIT.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Audited {len(inventory)} workbook sheets, 15 state/region rows + Union, 240 observation slots; {out}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    main(args.project)
