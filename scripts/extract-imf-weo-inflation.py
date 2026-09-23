"""Extract annual consumer-price inflation from one fixed IMF WEO vintage.

IMF country groups are not silently mapped to UN M49 regions. Only the IMF
World row and exact ISO3-matched country rows are retained.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


INDICATOR_ID = "IMF_WEO_PCPI_ANNUAL_CHANGE"
SOURCE_URL = "https://data.imf.org/Datasets/WEO"
FILE_URL = "https://data.imf.org/-/media/iData/External-Storage/Documents/2F78EE59F79143A7921E5E203D3AAA80/en/WEOApr2026all.xlsx"


def sha256(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--registry", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    registry = json.loads(args.registry.read_text(encoding="utf-8"))
    iso_to_id = {area["id"]: area["id"] for area in registry["territories"] if area.get("type") == "country" and len(area["id"]) == 3}
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    observations = []
    counts = {}
    for sheet_name in ("Countries", "Country Groups"):
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = next(rows)
        if header[4] != "INDICATOR.ID" or header[9] != "UNIT" or 2024 not in header:
            raise ValueError(f"Unexpected IMF WEO sheet {sheet_name}")
        year_columns = [(index, str(value)) for index, value in enumerate(header) if isinstance(value, int) and 2015 <= value <= 2024]
        adopted = 0
        for row in rows:
            if row[4] != "PCPIPCH":
                continue
            if row[7] != "Annual" or row[9] != "Percent":
                raise ValueError(f"Unexpected IMF WEO inflation dimensions: {row[1]}")
            area_id = iso_to_id.get(row[2]) if sheet_name == "Countries" else "WLD" if row[2] == "G001" else None
            if not area_id:
                continue
            for index, period in year_columns:
                value = row[index]
                if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value):
                    continue
                observations.append({"territory_id": area_id, "indicator_id": INDICATOR_ID, "period": period, "value": float(value), "status": "observed", "source_id": "imf-weo-2026-04", "nature": "IMF WEO published value; original report or IMF estimate according to WEO country notes", "latest_actual_annual_data": row[15]})
                adopted += 1
        counts[sheet_name] = adopted
    workbook.close()
    unique = set()
    for row in observations:
        key = row["territory_id"], row["indicator_id"], row["period"]
        if key in unique:
            raise ValueError(f"Duplicate IMF WEO observation: {key}")
        unique.add(key)
    if ("WLD", INDICATOR_ID, "2024") not in unique:
        raise ValueError("Missing published World 2024 inflation")
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    source = {"id": "imf-weo-2026-04", "name": "IMF World Economic Outlook, April 2026 — annual consumer price inflation", "publisher": "International Monetary Fund", "url": SOURCE_URL, "status": "ready", "retrieved_at": now, "reference_period": "2015–2024 (published historical WEO cells)", "geographic_level": "world_and_country_series", "license": "IMF data terms; raw workbook is not published by AreaData", "raw_redistribution_status": "not_in_public_site", "source_files": [{"url": FILE_URL, "sha256": sha256(args.workbook), "bytes": args.workbook.stat().st_size}], "note": "PCPIPCH annual percent change of average consumer prices. Only the IMF World group and exact ISO3 countries are adopted. IMF groups differ from UN M49 regions; the latter have no fabricated overall inflation rate. WEO country values can contain IMF estimates; 2025–2026 projection cells were deliberately excluded."}
    indicator = {"id": INDICATOR_ID, "name": "Consumer price inflation, annual", "theme": "Economy", "unit": "%", "definition": "IMF WEO PCPIPCH: annual average consumer price inflation, percent change. World is the IMF-published World group. Country observations may include IMF estimates; 2025 and later WEO projections are excluded. UN M49 region observations are unavailable from this source.", "definition_id": "imf-weo-2026-04-pcpipch", "population": "Published IMF World group or matched country", "measurement_method": "IMF WEO, annual average CPI change", "aggregation": "official_only", "series_family": "international_reference", "display_role": "supplementary", "period_policy": "same_period", "display_decimals": 2, "visualization": "trend", "source_id": source["id"]}
    result = {"schema_version": "1.0", "generated_at": now, "source": source, "indicators": [indicator], "observations": sorted(observations, key=lambda row: (row["territory_id"], row["period"])), "summary": {"observations": len(observations), "by_sheet": counts, "geographic_ids": len({row["territory_id"] for row in observations}), "periods": [str(year) for year in range(2015, 2025)]}}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
