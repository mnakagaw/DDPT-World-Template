"""Count every populated numeric source column in acquired Geostat XLSX files.

The inventory is mechanical, and leaves semantic/adoption decisions explicit.
It does not infer zero from blanks or a table title from its filename.
"""

import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SELECTED_COLUMNS = {
    "909-02": {"B", "C", "D", "E", "F", "G", "H", "I", "J"},
    "910-02": {"B"},
    "912-01": {"B", "C", "I"},
    "913-01": {"B", "C", "D", "K"},
    "915-01": {"B", "C"},
}


def main(project: Path) -> None:
    manifest = json.loads((project / "evidence/GEO_GEOSTAT2024_CATALOG_INVENTORY.json").read_text(encoding="utf-8"))
    result = {"scope": "all acquired 2024 census XLSX files listed by eight Geostat catalog sections",
              "audited_at": datetime.now(timezone.utc).isoformat(), "tables": []}
    for table in manifest["tables"]:
        if table.get("status") != "acquired_field_audit_pending":
            continue
        path = project / table["path"]
        raw = path.read_bytes()
        if hashlib.sha256(raw).hexdigest() != table["sha256"]:
            raise ValueError(f"Original changed: {table['id']}")
        wb = load_workbook(path, read_only=True, data_only=True)
        item = {"id": table["id"], "title": table["title"], "url": table["url"],
                "sha256": table["sha256"], "status": ("selected_arithmetic_and_scope_audited" if table["id"] in SELECTED_COLUMNS
                                                      else "acquired_mechanical_field_inventory_semantic_review_pending"),
                "worksheets": []}
        for ws in wb.worksheets:
            counts = Counter()
            nonempty = Counter()
            text_tokens = {}
            sample_numeric = {}
            headers = []
            for row_num, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_num <= 10:
                    headers.append({"row": row_num, "cells": [value for value in row]})
                for col_num, value in enumerate(row, 1):
                    if value is None or value == "":
                        continue
                    nonempty[col_num] += 1
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        counts[col_num] += 1
                        if len(sample_numeric.get(col_num, [])) < 5:
                            sample_numeric.setdefault(col_num, []).append({"row": row_num, "value": value})
                    elif isinstance(value, str) and row_num > 10:
                        tokens = text_tokens.setdefault(col_num, Counter())
                        tokens[value.strip()] += 1
            fields = []
            for col_num in range(1, ws.max_column + 1):
                letter = get_column_letter(col_num)
                disposition = ("adopted_direct" if table["id"] != "910-02" else "adopted_calculation_input") if letter in SELECTED_COLUMNS.get(table["id"], set()) else (
                    "checked_retained_unadopted" if table["id"] in SELECTED_COLUMNS else "unassessed")
                fields.append({"column": letter, "nonempty_cells": nonempty[col_num],
                               "numeric_cells": counts[col_num], "numeric_samples": sample_numeric.get(col_num, []),
                               "text_tokens_after_row10": text_tokens.get(col_num, Counter()).most_common(8),
                               "semantic_disposition": disposition})
            item["worksheets"].append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                                       "numeric_cells": sum(counts.values()), "headers_rows1_to10": headers,
                                       "fields": fields})
        wb.close()
        result["tables"].append(item)
        print(f"{item['id']}: {sum(s['numeric_cells'] for s in item['worksheets'])} numeric cells")
    result["total_numeric_cells"] = sum(s["numeric_cells"] for t in result["tables"] for s in t["worksheets"])
    output = project / "evidence/GEO_GEOSTAT2024_FIELD_INVENTORY.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{len(result['tables'])} originals; {result['total_numeric_cells']} numeric cells mechanically inventoried")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    main(parser.parse_args().project)
