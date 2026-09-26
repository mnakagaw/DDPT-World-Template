"""Import the official 2025 TÜİK ADNKS province and district population tables.

The workbook contains seven sheets. This importer adopts the first two only;
the remaining tables need separate denominator and geography review. A dash in
the workbook means information unavailable and is never converted to zero.
"""

import argparse
import csv
import hashlib
import json
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


SOURCE_URL = "https://www.tuik.gov.tr/media/announcements/2025ADNKS_FavoriTablolar.xlsx"
SOURCE_FILE = "2025ADNKS_FavoriTablolar.xlsx"
SOURCE_HASH = "72e36cf8f1eeb2e8c12480e14148b42448c6aad05d19932d94add37f360980c1"
SOURCE_ID = "tur-tuik-adnks-2025-favorite-tables"
SHEET_PROVINCE = "İL NÜFUSU"
SHEET_DISTRICT = "İLÇE NÜFUSU"
FIELDS = ("POP_TOTAL", "POP_MALE", "POP_FEMALE", "POP_CENTER_TOTAL",
          "POP_CENTER_MALE", "POP_CENTER_FEMALE", "POP_TOWN_VILLAGE_TOTAL",
          "POP_TOWN_VILLAGE_MALE", "POP_TOWN_VILLAGE_FEMALE")
LABELS = ("Population", "Male population", "Female population",
          "Province and district centre population", "Male centre population",
          "Female centre population", "Town and village population",
          "Male town and village population", "Female town and village population")


def fold(value):
    value = str(value).replace("ı", "i").replace("I", "i")
    return "".join(character for character in unicodedata.normalize("NFKD", value).casefold()
                   if not unicodedata.combining(character) and character.isalnum())


def value_at(value, label):
    if value in (None, "-"):
        return None
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"Expected nonnegative integer or unavailable at {label}: {value!r}")
    return value


def values(row, first, sheet, row_number):
    return tuple(value_at(row[first + index], f"{sheet}!{openpyxl.utils.get_column_letter(first + index + 1)}{row_number}")
                 for index in range(9))


def balance(row, label):
    for total, male, female in ((0, 1, 2), (3, 4, 5), (6, 7, 8)):
        if all(row[index] is not None for index in (total, male, female)):
            if row[total] != row[male] + row[female]:
                raise ValueError(f"Sex mismatch: {label}: {total}")
    for column in range(3):
        if all(row[index] is not None for index in (column, column + 3, column + 6)):
            if row[column] != row[column + 3] + row[column + 6]:
                raise ValueError(f"Centre + town/village mismatch: {label}: {column}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", required=True)
    args = parser.parse_args()
    project = Path(args.project)
    source_path = project / "raw" / SOURCE_FILE
    actual_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    if actual_hash != SOURCE_HASH:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    book = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    if book.sheetnames[:2] != [SHEET_PROVINCE, SHEET_DISTRICT] or len(book.sheetnames) != 7:
        raise ValueError(f"Unexpected workbook tabs: {book.sheetnames}")
    province_rows = list(book[SHEET_PROVINCE].iter_rows(min_row=8, values_only=True))
    district_rows = list(book[SHEET_DISTRICT].iter_rows(min_row=8, values_only=True))
    if len(province_rows) != 82 or len(district_rows) != 974:
        raise ValueError(f"Unexpected row counts: {len(province_rows)}, {len(district_rows)}")

    dataset_path = project / "data" / "dashboard.json"
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    geometry = json.loads((project / "raw" / "geoboundaries-adm1-simplified.json").read_text(encoding="utf-8"))
    geo_by_iso = {feature["properties"]["shapeISO"]: feature["properties"]
                  for feature in geometry["features"]}
    if len(geo_by_iso) != 81:
        raise ValueError("Reference ADM1 shapeISO codes are not unique and complete")
    territory_by_id = {territory["id"]: territory for territory in dataset["territories"]}
    boundary_by_id = {feature["properties"]["territory_id"]: feature for feature in dataset["boundaries"]["features"]}
    if len(boundary_by_id) != 81:
        raise ValueError("Expected 81 reference province boundaries")

    national = values(province_rows[0], 2, SHEET_PROVINCE, 8)
    if national[:3] != (86092168, 43059434, 43032734):
        raise ValueError("2025 published national totals changed")
    if national != values(district_rows[0], 4, SHEET_DISTRICT, 8):
        raise ValueError("National values differ between workbook tabs")
    balance(national, "national")

    provinces = {}
    crosswalk = []
    for row_number, row in enumerate(province_rows[1:], 9):
        code = int(row[0])
        if code in provinces or code not in range(1, 82):
            raise ValueError(f"Duplicate/invalid province code {code}")
        iso = f"TR-{code:02d}"
        geo = geo_by_iso.get(iso)
        if geo is None or fold(row[1]) != fold(geo["shapeName"]):
            raise ValueError(f"Province code/name geometry mismatch: {code}, {row[1]}, {geo}")
        territory_id = next((x["id"] for x in dataset["territories"]
                             if x.get("provider_code") == geo["shapeID"]), None)
        if territory_id is None or territory_id not in boundary_by_id:
            raise ValueError(f"No registered province/boundary: {code}")
        territory = territory_by_id[territory_id]
        territory.update(official_code=f"{code:02d}", code_system="TÜİK il kodu",
                         reconciliation_status="code_and_name_matched_to_2021_reference_geometry")
        boundary_by_id[territory_id]["properties"]["official_code"] = f"{code:02d}"
        boundary_by_id[territory_id]["properties"]["code_system"] = "TÜİK il kodu"
        extracted = values(row, 2, SHEET_PROVINCE, row_number)
        balance(extracted, f"province {code}")
        provinces[code] = (territory, extracted, row_number)
        crosswalk.append((code, row[1], territory_id, iso, geo["shapeName"], "code_and_name_matched"))
    if len(provinces) != 81 or {f"TR-{i:02d}" for i in provinces} != set(geo_by_iso):
        raise ValueError("Province workbook/boundary coverage mismatch")

    districts = []
    district_codes = set()
    district_by_province = defaultdict(list)
    for row_number, row in enumerate(district_rows[1:], 9):
        code, registry = row[:2]
        if code not in provinces or not isinstance(registry, int) or registry in district_codes:
            raise ValueError(f"Invalid district identifier at row {row_number}")
        district_codes.add(registry)
        if fold(row[2]) != fold(provinces[code][0]["name"]):
            raise ValueError(f"District province label mismatch at row {row_number}")
        extracted = values(row, 4, SHEET_DISTRICT, row_number)
        balance(extracted, f"district {registry}")
        territory_id = f"TUR-ADNKS2025-ILCE-{registry}"
        territory = {"id": territory_id, "name": row[3], "level": "adm2", "type": "district",
                     "parent_id": provinces[code][0]["id"], "official_code": str(registry),
                     "code_system": "TÜİK ilçe kayıt no", "boundary_version": None,
                     "source_id": SOURCE_ID, "reconciliation_status": "official_hierarchy_no_boundary"}
        districts.append((territory, extracted, row_number))
        district_by_province[code].append(extracted)
    if len(districts) != 973:
        raise ValueError("Expected 973 districts")
    for code, (_, province_values, _) in provinces.items():
        children = district_by_province[code]
        if not children or sum(item[0] for item in children) != province_values[0]:
            raise ValueError(f"District totals do not equal province {code}")
        for index in (1, 2):
            if sum(item[index] for item in children) != province_values[index]:
                raise ValueError(f"District sex totals do not equal province {code}: {index}")
    for index in (0, 1, 2):
        if sum(item[1][index] for item in provinces.values()) != national[index]:
            raise ValueError(f"Province totals do not equal national: {index}")

    def observation(territory_id, extracted, sheet, row_number, first):
        result = []
        for index, numeric in enumerate(extracted):
            if numeric is None:
                continue
            column = openpyxl.utils.get_column_letter(first + index + 1)
            result.append({"territory_id": territory_id,
                           "indicator_id": f"TUR_ADNKS_2025_{FIELDS[index]}",
                           "period": "2025", "value": numeric, "status": "observed",
                           "measurement_method": "source_reported", "source_id": SOURCE_ID,
                           "source_locator": f"{sheet}!{column}{row_number}"})
        return result

    observations = observation("TUR", national, SHEET_PROVINCE, 8, 2)
    for territory, extracted, row_number in provinces.values():
        observations.extend(observation(territory["id"], extracted, SHEET_PROVINCE, row_number, 2))
    for territory, extracted, row_number in districts:
        observations.extend(observation(territory["id"], extracted, SHEET_DISTRICT, row_number, 4))
    indicators = []
    for key, name in zip(FIELDS, LABELS):
        indicators.append({"id": f"TUR_ADNKS_2025_{key}", "name": name,
                           "theme": "Population", "unit": "people",
                           "definition": f"Address Based Population Registration System (ADNKS), 31 December 2025: {name.lower()}.",
                           "population": "Persons recorded in the ADNKS administrative register",
                           "source_id": SOURCE_ID, "aggregation": "sum",
                           "measurement_method": "source_reported", "series_family": "administrative",
                           "period_policy": "latest_available_per_indicator", "display_decimals": 0})

    dataset["territories"] = [item for item in dataset["territories"]
                              if not item["id"].startswith("TUR-ADNKS2025-ILCE-")] + [item[0] for item in districts]
    dataset["indicators"] = [item for item in dataset["indicators"]
                             if not item["id"].startswith("TUR_ADNKS_2025_")] + indicators
    dataset["observations"] = [item for item in dataset["observations"]
                               if not item["indicator_id"].startswith("TUR_ADNKS_2025_")] + observations
    dataset["sources"] = [item for item in dataset["sources"] if item["id"] != SOURCE_ID]
    stamp = datetime.now(timezone.utc).isoformat()
    dataset["generated_at"] = stamp
    dataset["sources"].append({"id": SOURCE_ID, "name": "TÜİK ADNKS 2025 — favorite tables",
                               "url": SOURCE_URL, "publisher": "Turkish Statistical Institute (TÜİK)",
                               "reference_period": "31 December 2025", "geographic_level": "country, il, ilçe",
                               "status": "ready", "retrieved_at": stamp, "sha256": SOURCE_HASH,
                               "raw_path": f"raw/{SOURCE_FILE}", "license": "terms_review_required",
                               "note": "Annual address-based administrative register, not a decennial census. First two of seven sheets adopted; dash values remain missing. 2025 statistics are matched by province code to 2021 reference geometry for navigation only."})
    dataset["country"]["geography_note"] = (
        "TÜİK 2025 ADNKS provides 81 provinces and 973 districts; the annual register is separate "
        "from the World Bank national series. Province code/name matches the geoBoundaries 2021 "
        "reference shapes, which are not certified as 2025 legal boundaries. District geometry is pending.")
    dataset["gaps"] = [item for item in dataset["gaps"] if item["category"] != "subnational_statistics"]
    dataset["gaps"].append({"category": "subnational_statistics", "status": "partial",
                            "detail": "TÜİK 2025 ADNKS population and sex for all 81 provinces and 973 districts imported. Centre/town/village columns retain source dashes as missing. Five other workbook tabs, district boundaries and census-specific domains remain to review.",
                            "next_action": "Audit remaining tables, planning evidence, 2025 administrative boundary validity, and redistribution terms."})
    dataset["collection"]["adapters"] = sorted(set(dataset["collection"]["adapters"] + ["turkey-tuik-adnks-2025"]))
    dataset["analysis"]["terminal_territory_ids"] = sorted(item[0]["id"] for item in districts)
    dataset["analysis"]["latest_values_only"] = True
    dataset_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    evidence = project / "evidence"
    evidence.mkdir(exist_ok=True)
    receipt = {"url": SOURCE_URL, "publisher": "Turkish Statistical Institute (TÜİK)",
               "retrieved_at_utc": stamp, "sha256": SOURCE_HASH,
               "size_bytes": source_path.stat().st_size, "redistribution_terms": "review_required"}
    (project / "raw" / (SOURCE_FILE + ".receipt.json")).write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit = {"country": "TUR", "reference_year": 2025, "status": "partial_import_not_accepted",
             "source_url": SOURCE_URL, "sha256": SOURCE_HASH,
             "counts": {"provinces": 81, "districts": 973, "indicators": len(indicators),
                        "observations": len(observations)},
             "national_values": {"population": national[0], "male": national[1], "female": national[2]},
             "reconciliation": "All district total/male/female sums match each province; all province sums match the national row; each available sex and centre/town/village decomposition balances.",
             "unresolved": ["2025 legal boundary geometry", "District geometry", "Planning law and documents",
                            "Five other workbook sheets", "Census-specific domains", "Redistribution terms"]}
    (evidence / "TUR_ADNKS_IMPORT_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inventory = []
    for sheet, columns in ((SHEET_PROVINCE, ("İL KODU", "İL ADI") + FIELDS),
                           (SHEET_DISTRICT, ("İL KODU", "İLÇE KAYIT NO", "İL ADI", "İLÇE ADI") + FIELDS)):
        for index, label in enumerate(columns, 1):
            inventory.append({"sheet": sheet, "column": openpyxl.utils.get_column_letter(index),
                              "source_field": label, "adoption": "identifier" if label not in FIELDS else "adopted",
                              "missing_marker": "dash_or_blank" if label in FIELDS else None})
    for sheet in book.sheetnames[2:]:
        inventory.append({"sheet": sheet, "column": "all", "source_field": "all columns",
                          "adoption": "not_adopted_yet", "reason": "Geography/denominator audit pending"})
    (evidence / "SOURCE_TABLE_INVENTORY.json").write_text(
        json.dumps({"country": "TUR", "reference_year": 2025, "fields": inventory},
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (evidence / "CODE_CROSSWALK.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("tuik_il_code", "tuik_il_name", "areadata_territory_id",
                         "geoboundaries_shape_iso", "geoboundaries_shape_name", "status"))
        writer.writerows(crosswalk)
    print(json.dumps(audit["counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
