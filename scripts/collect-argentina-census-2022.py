#!/usr/bin/env python3
"""Collect an evidence-grounded Argentina country-depth bundle.

The adapter uses INDEC 2022 Census tables for the themes the Census publishes
in comparable form, the official Georef registry for province/department codes
and reference geometry, the official 2018 disability study for its stated
national universe, and clearly labelled national WDI supplements.  Values from
different years and universes remain separate observations.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook
from shapely.geometry import mapping, shape

INDEC_PAGE = "https://www.indec.gob.ar/indec/web/Nivel4-Tema-2-41-165?lang=es"
INDEC_FILES = "https://www.indec.gob.ar/ftp/cuadros/poblacion/"
GEOREF = "https://apis.datos.gob.ar/georef/api"
DISABILITY = "https://www.indec.gob.ar/ftp/cuadros/poblacion/estudio_discapacidad_12_18.pdf"
FILES = [
    "c2022_tp_c_resumen.xlsx",
    "c2022_tp_vivienda_c1.xlsx",
    "c2022_tp_hogares_c2.xlsx",
    "c2022_tp_hogares_c3.xlsx",
    "c2022_tp_educacion_c3.xlsx",
    "c2022_tp_actividad_economica_c1.xlsx",
    "c2022_tp_migraciones_c1.xlsx",
    "c2022_tp_salud_c1.xlsx",
    "c2022_tp_poblacion_indigena_c10.xlsx",
]
WDI = {
    "SP.URB.TOTL.IN.ZS": ("Urban population", "Settlement", "% of total population", "urban_rural"),
    "SN.ITK.DEFC.ZS": ("Prevalence of undernourishment", "Nutrition", "% of population", "nutrition"),
    "SI.POV.NAHC": ("National poverty headcount ratio", "Poverty", "% of population", "poverty"),
}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def fetch(url: str, target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    response = requests.get(url, timeout=240, headers={"User-Agent": "AreaData/0.10.2 source collector"})
    response.raise_for_status()
    target.write_bytes(response.content)
    return target


def number(value):
    if value in (None, "", "-", "///"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def pct(n, d):
    n, d = number(n), number(d)
    return None if n is None or d in (None, 0) else round(n / d * 100, 4)


def norm(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def data_sheet(path: Path, preferred: str | None = None):
    wb = load_workbook(path, read_only=True, data_only=True)
    if preferred:
        for ws in wb.worksheets:
            if preferred.lower() in ws.title.lower():
                return ws
    return max((ws for ws in wb.worksheets if ws.max_column > 2), key=lambda ws: ws.max_row * ws.max_column)


def coded_rows(path: Path, start: int, code_col=0, name_col=1):
    ws = data_sheet(path)
    out = {}
    for row in ws.iter_rows(min_row=start, values_only=True):
        code = re.sub(r"\D", "", str(row[code_col] or ""))
        name = norm(row[name_col])
        if code and name and len(code) <= 3:
            out[code.zfill(2)[-2:]] = row
        elif name in ("total", "total del pais"):
            out["ARG"] = row
    return out


def load_ndjson(path: Path):
    with path.open("r", encoding="utf-8") as stream:
        next(stream)
        for line in stream:
            if line.strip():
                yield json.loads(line)


def simplified(geometry):
    geom = shape(geometry)
    result = geom.simplify(0.02, preserve_topology=True)
    output = mapping(result)
    def clamp(value):
        if isinstance(value, (list, tuple)) and len(value) >= 2 and all(isinstance(v, (int, float)) for v in value[:2]):
            return [max(-180.0, min(180.0, value[0])), max(-90.0, min(90.0, value[1])), *value[2:]]
        if isinstance(value, (list, tuple)):
            return [clamp(item) for item in value]
        return value
    output["coordinates"] = clamp(output["coordinates"])
    return output


def inventory(files: list[Path]):
    records = []
    for file in files:
        wb = load_workbook(file, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if ws.max_column < 2:
                continue
            title = next((str(row[0]) for row in ws.iter_rows(min_row=1, max_row=min(4, ws.max_row), values_only=True) if row and row[0]), ws.title)
            counts = [0] * ws.max_column
            for row in ws.iter_rows(values_only=True):
                for i, value in enumerate(row):
                    if number(value) is not None:
                        counts[i] += 1
            for i, count in enumerate(counts):
                if count:
                    records.append({"country_area_id": "ARG", "source_id": "arg-indec-census-2022",
                                    "source_path": f"raw/argentina-census-2022/{file.name}", "source_url": INDEC_PAGE,
                                    "table_id": f"{file.stem}:{ws.title}", "table_title": title,
                                    "field_id": f"col-{i + 1}", "field_label": f"Numeric column {i + 1}",
                                    "numeric_cell_count": count, "theme": "source_inventory",
                                    "disposition": "not_adopted",
                                    "reason": "Official numeric field inventoried; retained for traceability but not selected as a dashboard indicator in this edition.",
                                    "coverage_complete": True, "country_edition_eligible": False})
    return records


def wdi_latest(code: str, raw: Path):
    url = f"https://api.worldbank.org/v2/country/ARG/indicator/{code}?format=json&per_page=100&source=2"
    target = fetch(url, raw / f"wb-{code}.json")
    payload = json.loads(target.read_text(encoding="utf-8"))
    row = next(item for item in payload[1] if item.get("value") is not None)
    return {"code": code, "period": str(row["date"]), "value": float(row["value"]), "url": url,
            "sha256": sha(target), "raw_path": target.name}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    out = Path(args.out).resolve(); raw = out / "raw"; raw.mkdir(parents=True, exist_ok=True)
    acquired = [fetch(INDEC_FILES + name, raw / name) for name in FILES]
    province_nd = fetch(GEOREF + "/provincias.ndjson", raw / "provincias.ndjson")
    department_nd = fetch(GEOREF + "/departamentos.ndjson", raw / "departamentos.ndjson")
    disability_pdf = fetch(DISABILITY, raw / "estudio_discapacidad_12_18.pdf")

    provinces = list(load_ndjson(province_nd)); departments = list(load_ndjson(department_nd))
    if len(provinces) != 24 or len(departments) != 529:
        raise RuntimeError(f"Unexpected Georef counts: {len(provinces)} provinces, {len(departments)} departments")
    province_ids = {p["id"]: f"ARG:C2022:PROV:{p['id']}" for p in provinces}
    territories, features = [], []
    geo_source = "arg-georef-2026"
    for p in provinces:
        tid = province_ids[p["id"]]
        territories.append({"id": tid, "country_id": "ARG", "name": p["nombre"], "level": "province", "type": "province",
                            "parent_id": "ARG", "official_code": p["id"], "code_system": "Argentina Georef province ID v13",
                            "boundary_version": "Georef 13.0.0 / 2026-04-09", "valid_from": "2026-04-09"})
        features.append({"type": "Feature", "properties": {"territory_id": tid, "source_id": geo_source,
                         "official_code": p["id"], "code_system": "Argentina Georef province ID v13",
                         "geometry_edition": "Georef 13.0.0", "join_method": "Exact Georef province ID", "reference_only": True},
                         "geometry": simplified(p["geometria"])})
    department_ids = {}
    for d in departments:
        tid = f"ARG:C2022:DEPT:{d['id']}"; department_ids[d["id"]] = tid; parent = province_ids[d["provincia"]["id"]]
        territories.append({"id": tid, "country_id": "ARG", "name": d["nombre_completo"], "level": "department", "type": "department",
                            "parent_id": parent, "official_code": d["id"], "code_system": "Argentina Georef department ID v13",
                            "boundary_version": "Georef 13.0.0 / 2026-04-09", "valid_from": "2026-04-09"})
        features.append({"type": "Feature", "properties": {"territory_id": tid, "source_id": geo_source,
                         "official_code": d["id"], "code_system": "Argentina Georef department ID v13",
                         "geometry_edition": "Georef 13.0.0", "join_method": "Exact Georef department ID", "reference_only": True},
                         "geometry": simplified(d["geometria"])})

    # Source rows.  The summary has no code column, so join its 24 jurisdiction
    # names against the official province registry and fail on any ambiguity.
    summary_ws = data_sheet(raw / "c2022_tp_c_resumen.xlsx")
    by_name = {norm(p["nombre"]): p["id"] for p in provinces}
    aliases = {"tierra del fuego antartida e islas del atlantico sur": "94"}
    summary = {}
    for row in summary_ws.iter_rows(min_row=5, values_only=True):
        name = norm(row[0])
        if name == "total del pais": summary["ARG"] = row
        elif name in by_name or name in aliases: summary[(by_name | aliases)[name]] = row
    if set(summary) != {"ARG", *province_ids}:
        raise RuntimeError(f"Census summary join incomplete: {sorted(set(summary) ^ {'ARG', *province_ids})}")
    housing = coded_rows(raw / "c2022_tp_vivienda_c1.xlsx", 6)
    employment = coded_rows(raw / "c2022_tp_actividad_economica_c1.xlsx", 6)
    migration = coded_rows(raw / "c2022_tp_migraciones_c1.xlsx", 5)
    health = coded_rows(raw / "c2022_tp_salud_c1.xlsx", 5)
    indigenous = coded_rows(raw / "c2022_tp_poblacion_indigena_c10.xlsx", 5)

    national = {
        "water": list(data_sheet(raw / "c2022_tp_hogares_c2.xlsx").iter_rows(min_row=5, values_only=True)),
        "sanitation": list(data_sheet(raw / "c2022_tp_hogares_c3.xlsx").iter_rows(min_row=5, values_only=True)),
        "education": list(data_sheet(raw / "c2022_tp_educacion_c3.xlsx").iter_rows(min_row=6, values_only=True)),
    }
    wdi = [wdi_latest(code, raw) for code in WDI]

    source = "arg-indec-census-2022"
    specs = [
        ("ARG_C2022_POP_TOTAL", "Census population", "Population", "people", lambda key: number(summary[key][1]), "sum", "Enumerated population in the 2022 Census."),
        ("ARG_C2022_FEMALE_PCT", "Female population", "Population", "%", lambda key: pct(summary[key][9], summary[key][1]), "none", "Female population as a share of the 2022 Census population."),
        ("ARG_C2022_PRIVATE_DWELLINGS", "Private dwellings", "Households and housing", "dwellings", lambda key: number(housing[key][3]) if key in housing else None, "sum", "Private dwellings recorded by the 2022 Census."),
        ("ARG_C2022_PUBLIC_WATER_PCT", "Households using public-network water", "Water", "% of households", lambda key: pct(national['water'][1][1], national['water'][0][1]) if key == "ARG" else None, "none", "Households whose water source is the public network."),
        ("ARG_C2022_PUBLIC_SEWER_PCT", "Households connected to public sewer", "Sanitation", "% of households", lambda key: pct(national['sanitation'][1][2], national['sanitation'][0][2]) if key == "ARG" else None, "none", "Households whose toilet drainage is connected to the public sewer network."),
        ("ARG_C2022_NO_INSTRUCTION_PCT", "Population with no instruction", "Education", "% of population age 5+", lambda key: pct(national['education'][0][4], national['education'][0][2]) if key == "ARG" else None, "none", "Population age 5 or older with no instruction among the stated Census education universe."),
        ("ARG_C2022_LABOUR_FORCE_PCT", "Labour-force participation", "Employment", "% of population age 14+", lambda key: pct(employment[key][3], employment[key][2]) if key in employment else None, "none", "Economically active population as a share of people age 14 or older in private dwellings."),
        ("ARG_C2022_FOREIGN_BORN_PCT", "Foreign-born population", "Migration", "% of private-dwelling population", lambda key: pct(migration[key][5], migration[key][2]) if key in migration else None, "none", "Population born in another country as a share of residents in private dwellings."),
        ("ARG_C2022_NO_HEALTH_COVERAGE_PCT", "Population without listed health coverage", "Health", "% of private-dwelling population", lambda key: pct(health[key][5], health[key][2]) if key in health else None, "none", "Population without obra social, prepaid cover or a state health plan."),
        ("ARG_C2022_INDIGENOUS_PCT", "Indigenous or Indigenous-descendant population", "Ethnicity", "% of private-dwelling population", lambda key: pct(indigenous[key][3], indigenous[key][2]) if key in indigenous else None, "none", "Population self-identifying as Indigenous or descendant of Indigenous peoples."),
    ]
    indicators, observations = [], []
    keys = ["ARG", *sorted(province_ids)]
    for iid, name, theme, unit, calc, aggregation, definition in specs:
        indicators.append({"id": iid, "name": name, "theme": theme, "unit": unit, "definition": definition,
                           "definition_id": "ARG-INDEC-CENSUS-2022", "population": "Published Census universe stated per indicator",
                           "measurement_method": "source_reported" if aggregation == "sum" else "derived_from_source_counts",
                           "aggregation": aggregation, "period_policy": "fixed_source_period", "series_family": "census",
                           "display_role": "primary", "source_id": source, "source_locator": "Selected official INDEC 2022 Census XLSX tables"})
        for key in keys:
            tid = "ARG" if key == "ARG" else province_ids[key]; value = calc(key)
            observations.append({"territory_id": tid, "indicator_id": iid, "period": "2022", "value": value,
                                 "status": "observed" if value is not None else "missing", "source_id": source,
                                 "definition": definition, "definition_id": "ARG-INDEC-CENSUS-2022", "unit": unit,
                                 "population": "Published Census universe stated per indicator",
                                 "measurement_method": "source_reported" if aggregation == "sum" else "derived_from_source_counts",
                                 "source_locator": "Selected official INDEC 2022 Census XLSX tables"})

    # The official 2018 study applies to urban localities of 5,000+ people; it
    # remains a national study observation and is never copied to provinces.
    disability_id = "ARG_INDEC_2018_DIFFICULTY_PCT"
    indicators.append({"id": disability_id, "name": "Population with difficulty", "theme": "Disability", "unit": "% of population age 6+",
                       "definition": "People age 6 or older with at least one difficulty in urban localities of 5,000 or more inhabitants.",
                       "definition_id": "ARG-INDEC-ENPD-2018", "population": "Population age 6+ in urban localities of 5,000+ inhabitants",
                       "measurement_method": "source_reported", "aggregation": "none", "period_policy": "fixed_source_period",
                       "series_family": "survey", "display_role": "supplementary", "source_id": "arg-indec-disability-2018", "source_locator": "Executive summary, prevalence result"})
    observations.append({"territory_id": "ARG", "indicator_id": disability_id, "period": "2018", "value": 10.2, "status": "observed",
                         "source_id": "arg-indec-disability-2018", "definition": indicators[-1]["definition"], "definition_id": "ARG-INDEC-ENPD-2018",
                         "unit": indicators[-1]["unit"], "population": indicators[-1]["population"], "measurement_method": "source_reported",
                         "source_locator": "Executive summary, prevalence result"})
    for row in wdi:
        name, theme, unit, _ = WDI[row["code"]]; iid = f"ARG_WDI_{row['code'].replace('.', '_')}"
        definition = f"{name} under the provider's published national series definition."
        indicators.append({"id": iid, "name": name, "theme": theme, "unit": unit, "definition": definition,
                           "definition_id": f"WDI-{row['code']}", "population": "National population defined by the provider series",
                           "measurement_method": "source_reported", "aggregation": "none", "period_policy": "latest_available",
                           "series_family": "international_reference", "display_role": "supplementary", "source_id": f"arg-wdi-{row['code']}",
                           "source_locator": f"World Bank WDI {row['code']}"})
        observations.append({"territory_id": "ARG", "indicator_id": iid, "period": row["period"], "value": row["value"], "status": "observed",
                             "source_id": f"arg-wdi-{row['code']}", "definition": definition, "definition_id": f"WDI-{row['code']}",
                             "unit": unit, "population": "National population defined by the provider series", "measurement_method": "source_reported",
                             "source_locator": f"World Bank WDI {row['code']}"})

    comparisons = [{"parent_id": "ARG", "member_ids": [province_ids[x] for x in sorted(province_ids)], "label": "Provinces",
                    "membership_note": "Official Georef v13 province registry.", "source_ids": [geo_source]}]
    for code, tid in sorted(province_ids.items()):
        members = [department_ids[x] for x in sorted(department_ids) if x.startswith(code)]
        comparisons.append({"parent_id": tid, "member_ids": members, "label": "Departments / partidos / communes",
                            "membership_note": "Official Georef v13 first-order subdivisions; category names remain in territory labels.", "source_ids": [geo_source]})

    today = datetime.now(timezone.utc).date().isoformat()
    sources = [
        {"id": source, "name": "Censo Nacional de Poblacion, Hogares y Viviendas 2022 - definitive tables", "publisher": "INDEC",
         "url": INDEC_PAGE, "status": "ready", "retrieved_at": today, "reference_period": "2022", "geographic_level": "country and province",
         "license": "Official public statistical tables; source attribution retained.", "note": "Nine selected official XLSX products are acquired and hashed. Province coverage varies by theme and missing local cells remain missing."},
        {"id": geo_source, "name": "Argentina Georef complete province and department registry v13", "publisher": "Argentina Datos / Georef",
         "url": GEOREF, "status": "ready", "retrieved_at": today, "reference_period": "2026-04-09", "geographic_level": "province and department",
         "license": "Official open government geographic service; source attribution retained.", "note": "24 provinces and 529 departments with exact official IDs and simplified reference geometry."},
        {"id": "arg-indec-disability-2018", "name": "Estudio Nacional sobre el Perfil de las Personas con Discapacidad 2018", "publisher": "INDEC",
         "url": DISABILITY, "status": "ready", "retrieved_at": today, "reference_period": "2018", "geographic_level": "country study universe",
         "raw_path": "raw/argentina-census-2022/estudio_discapacidad_12_18.pdf", "sha256": sha(disability_pdf),
         "license": "Official public report; source attribution retained.", "note": "Survey universe is urban localities of 5,000+ people; not a 2022 Census result."},
    ]
    for row in wdi:
        name = WDI[row["code"]][0]
        sources.append({"id": f"arg-wdi-{row['code']}", "name": f"World Bank WDI - {name}", "publisher": "World Bank", "url": row["url"],
                        "indicator_url": f"https://data.worldbank.org/indicator/{row['code']}", "status": "ready", "retrieved_at": today,
                        "reference_period": row["period"], "geographic_level": "country", "raw_path": f"raw/argentina-census-2022/{row['raw_path']}",
                        "sha256": row["sha256"], "license": "World Bank dataset terms apply.",
                        "note": "National supplementary series; never assigned to provinces or departments."})
    bundle = {"schema_version": "1.0", "country_area_id": "ARG", "period": "2022", "replace_country_branch": True,
              "territories": territories, "terminal_territory_ids": list(department_ids.values()), "comparisons": comparisons,
              "indicators": indicators, "observations": observations, "boundaries": {"type": "FeatureCollection", "features": features}, "sources": sources}
    (out / "argentina-country-depth-bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = inventory(acquired)
    (out / "argentina-workbook-field-inventory.json").write_text(json.dumps({"schema_version": "1.0", "country_area_id": "ARG", "records": fields}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    receipts = [{"file": p.name, "bytes": p.stat().st_size, "sha256": sha(p)} for p in sorted(raw.iterdir()) if p.is_file()]
    audit = {"schema_version": "1.0", "country_area_id": "ARG", "status": "complete", "counts": {"provinces": 24, "departments": 529,
             "boundary_features": len(features), "indicators": len(indicators), "observations": len(observations), "workbook_numeric_fields": len(fields)},
             "mixed_periods": sorted({o["period"] for o in observations}), "source_receipts": receipts}
    (out / "ARGENTINA_COLLECTION_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
